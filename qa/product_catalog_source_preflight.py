"""产品来源表只读预检。输出统计与单元格位置，不输出整表、不写任何文件。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


HEADERS = ["类型", "产品编码", "类别", "品名", "规格", "单位", "货号", "仓库号", "仓库名称", "规模"]
XML_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def duplicate_summary(rows, columns, *, ignore_blank=False):
    groups = defaultdict(list)
    for row_number, values in rows:
        key = tuple(values[index] for index in columns)
        if ignore_blank and any(value in (None, "") for value in key):
            continue
        groups[key].append(row_number)
    repeats = [numbers for numbers in groups.values() if len(numbers) > 1]
    return {
        "distinct_keys": len(groups),
        "duplicate_groups": len(repeats),
        "rows_in_duplicate_groups": sum(map(len, repeats)),
        "extra_rows_if_deduplicated": sum(len(group) - 1 for group in repeats),
        "duplicate_row_groups": repeats,
    }


def audit(path: Path):
    source_bytes = path.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    report = {
        "source_file": path.name,
        "source_bytes": len(source_bytes),
        "source_sha256": source_hash,
        "sheets": [],
        "xlsx_structure": [],
    }
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook:
            sheet_rows = list(sheet.iter_rows())
            headers = [cell.value for cell in sheet_rows[0]] if sheet_rows else []
            if headers != HEADERS:
                raise ValueError(f"{sheet.title}: 表头不符合已确认的 10 列契约：{headers!r}")
            rows = [
                (index, [cell.value for cell in row])
                for index, row in enumerate(sheet_rows[1:], start=2)
                if any(cell.value is not None for cell in row)
            ]
            fields = []
            for index, header in enumerate(headers):
                populated = [(number, values[index]) for number, values in rows if values[index] is not None]
                max_length = max((len(str(value)) for _, value in populated), default=0)
                fields.append({
                    "header": header,
                    "types": dict(Counter(type(values[index]).__name__ for _, values in rows)),
                    "blank_rows": [number for number, values in rows if values[index] in (None, "")],
                    "single_slash_rows": [number for number, values in rows if values[index] == "/"],
                    "distinct_nonblank": len({value for _, value in populated if value != ""}),
                    "max_characters": max_length,
                    "max_utf8_bytes": max((len(str(value).encode("utf-8")) for _, value in populated), default=0),
                    "max_length_rows_first_10": [number for number, value in populated if len(str(value)) == max_length][:10],
                    "number_formats": dict(Counter(row[index].number_format for row in sheet_rows[1:])),
                })
            cells = [cell for row in sheet_rows for cell in row if cell.value is not None]
            string_cells = [cell for cell in cells if isinstance(cell.value, str)]
            normalized = [(number, [None if index in (4, 5) and value == "/" else value for index, value in enumerate(values)]) for number, values in rows]
            report["sheets"].append({
                "sheet": sheet.title,
                "state": sheet.sheet_state,
                "headers": headers,
                "data_rows": len(rows),
                "actual_row_numbers": [number for number, _ in rows],
                "blank_rows": [index for index, row in enumerate(sheet_rows, start=1) if all(cell.value is None for cell in row)],
                "actual_max_columns": max(map(len, sheet_rows), default=0),
                "declared_max_row": sheet.max_row,
                "declared_max_column": sheet.max_column,
                "fields": fields,
                "type_counts": dict(Counter(values[0] for _, values in rows)),
                "unit_counts": dict(Counter(values[5] for _, values in rows)),
                "scale_counts": dict(Counter(values[9] for _, values in rows)),
                "formula_cells": [cell.coordinate for cell in cells if cell.data_type == "f"],
                "error_cells": [cell.coordinate for cell in cells if cell.data_type == "e"],
                "trim_changes": [cell.coordinate for cell in string_cells if cell.value != cell.value.strip()],
                "internal_whitespace_cells": [cell.coordinate for cell in string_cells if any(character.isspace() for character in cell.value)],
                "control_characters": [cell.coordinate for cell in string_cells if re.search(r"[\x00-\x1f\x7f]", cell.value)],
                "formula_like_text": [cell.coordinate for cell in string_cells if cell.data_type != "f" and cell.value.startswith(("=", "+", "-", "@"))],
                "fullwidth_slash_cells": [cell.coordinate for cell in string_cells if "／" in cell.value],
                "numeric_identifier_cells": [cell.coordinate for cell in string_cells if cell.column in (2, 7) and cell.value.isdecimal()],
                "leading_zero_identifiers": [cell.coordinate for cell in string_cells if cell.column in (2, 7) and cell.value.startswith("0")],
                "embedded_slash_specification_rows": [number for number, values in rows if isinstance(values[4], str) and "/" in values[4] and values[4] != "/"],
                "duplicates": {
                    "product_code_nonblank": duplicate_summary(rows, [1], ignore_blank=True),
                    "article_number": duplicate_summary(rows, [6]),
                    "category_and_name": duplicate_summary(rows, [2, 3]),
                    "category_name_specification": duplicate_summary(rows, [2, 3, 4]),
                    "category_name_specification_unit_article_scale": duplicate_summary(normalized, [2, 3, 4, 5, 6, 9]),
                    "all_raw_fields": duplicate_summary(rows, list(range(10))),
                },
            })
    finally:
        workbook.close()
    with ZipFile(path) as archive:
        for filename in archive.namelist():
            if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", filename):
                continue
            root = ET.fromstring(archive.read(filename))
            dimension = root.find("s:dimension", XML_NS)
            report["xlsx_structure"].append({
                "part": filename,
                "dimension": dimension.attrib if dimension is not None else None,
                "xml_row_numbers": [int(row.attrib["r"]) for row in root.findall("s:sheetData/s:row", XML_NS)],
                "merged_cells": [node.attrib for node in root.findall("s:mergeCells/s:mergeCell", XML_NS)],
                "hidden_rows": [row.attrib["r"] for row in root.findall("s:sheetData/s:row", XML_NS) if row.attrib.get("hidden") in ("1", "true")],
                "hidden_columns": [node.attrib for node in root.findall("s:cols/s:col", XML_NS) if node.attrib.get("hidden") in ("1", "true")],
                "hyperlink_count": len(root.findall("s:hyperlinks/s:hyperlink", XML_NS)),
            })
        report["external_link_parts"] = [name for name in archive.namelist() if name.startswith("xl/externalLinks/")]
    report["source_unchanged"] = hashlib.sha256(path.read_bytes()).hexdigest() == source_hash
    if not report["source_unchanged"]:
        raise RuntimeError("预检期间来源文件发生变化，统计结果不可用于导入。")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="仅供读取的产品信息 .xlsx 文件")
    args = parser.parse_args()
    if not args.source.is_file() or args.source.suffix.lower() != ".xlsx":
        parser.error("来源必须是已经存在的 .xlsx 文件。")
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(audit(args.source), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
