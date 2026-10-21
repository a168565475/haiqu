from copy import deepcopy
import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIClient, APITestCase

from apps.audit.models import AuditLog
from apps.customers.models import CustomerAccount, CustomerAssignment
from apps.files.models import FileScanStatus
from apps.imports.security import InfectedFile, ScannerUnavailable
from apps.organizations.models import Employee
from apps.products.models import ProductInformation, RawMaterialInformation
from apps.core.api.concurrency import etag_for_instance
from apps.core.api.exceptions import Conflict
from apps.production_tasks.models import (
    DeliveryMethod,
    OrderSource,
    ProductionDeliveryBatch,
    ProductionTask,
    ProductionTaskProduct,
    ProductionTaskStatus,
)
from apps.production_tasks.services import replace_unbound_product_line
from apps.reference_data.models import (
    AccessoryColor,
    AccessoryMaterial,
    AccessoryProductName,
    AccessorySpecification,
    Currency,
    FabricMaterial,
    FabricSpecification,
    FabricWidth,
    ProductColor,
)

from ..maintenance import match_maintenance_configuration
from ..models import Bom, BomLine, BomSection
from ..services import PREVIEW_LINE_LIMIT, _safe_preview_sections
from ..workbook_import import BomWorkbookError, parse_bom_workbook


@override_settings(FILE_VIRUS_SCAN_ENABLED=False)
class BomApiTests(APITestCase):
    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.settings_override = self.settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_override.enable()

        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="bom_user")
        self.employee = Employee.objects.create(
            user=self.user,
            employee_code="BOM_USER",
            display_name="BOM 测试员工",
        )
        self.user.user_permissions.set(
            Permission.objects.filter(content_type__app_label__in=["boms", "production_tasks"])
        )
        self.user.user_permissions.add(Permission.objects.get(
            content_type__app_label="products", codename="add_rawmaterialinformation",
        ))
        self.customer = CustomerAccount.objects.create(
            account_code="BOM_CUSTOMER",
            account_name="BOM 测试客户",
            primary_country_code="CN",
        )
        CustomerAssignment.objects.create(
            customer_account=self.customer,
            employee=self.employee,
            assignment_role="PRIMARY_OWNER",
            effective_from=timezone.now(),
        )
        self.product = (
            ProductInformation.objects.filter(status="ACTIVE", archived_at__isnull=True)
            .exclude(product_code="")
            .select_related("unit")
            .first()
        )
        self.assertIsNotNone(self.product)
        self.currency = Currency.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        today = timezone.localdate()
        self.task = ProductionTask.objects.create(
            task_number="HT-TEST-001",
            contract_number="HTC-TEST-001",
            order_source=OrderSource.NEW,
            order_number="ORDER-BOM-001",
            customer_account=self.customer,
            order_date=today,
            delivery_date=today + timedelta(days=30),
            salesperson=self.employee,
            currency=self.currency,
            includes_tax=True,
            order_tolerance_percent=Decimal("3"),
            is_repeat_order=False,
            preproduction_sample_required=False,
            delivery_method=DeliveryMethod.SPLIT,
            status=ProductionTaskStatus.READY,
            created_by_employee=self.employee,
        )
        self.product_line = ProductionTaskProduct.objects.create(
            production_task=self.task,
            line_number=1,
            product=self.product,
            product_number_snapshot=self.product.product_number,
            product_code_snapshot=self.product.product_code,
            item_number_snapshot=self.product.product_item.item_number,
            product_name_snapshot=self.product.product_name,
            specification_snapshot=self.product.specification,
            quantity=10,
            unit=self.product.unit,
            currency=self.currency,
            unit_price=Decimal("1"),
            total_amount=Decimal("10"),
            exchange_rate=Decimal("1"),
        )
        ProductionDeliveryBatch.objects.create(
            production_task=self.task,
            line_number=1,
            quantity=Decimal("5"),
            delivery_date=today + timedelta(days=30),
        )
        ProductionDeliveryBatch.objects.create(
            production_task=self.task,
            line_number=2,
            quantity=Decimal("7"),
            delivery_date=today + timedelta(days=30),
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def tearDown(self):
        self.settings_override.disable()
        self.media_dir.cleanup()

    def complete_pipe_sources(self, bom_id):
        # S4 确认前必须补齐真实根制原料及每根切割支数；保留原测试的权限/ETag目的。
        from apps.reference_data.models import ProductionUnit, ProductCategory
        unit, _ = ProductionUnit.objects.get_or_create(unit_name="根", defaults={"unit_code": "S4LEGROOT"})
        source, _ = RawMaterialInformation.objects.get_or_create(material_number="S4LEGPIPE", defaults={
            "material_name": "测试管材", "material": "铁", "specification": "Φ25 × 1.2 mm", "unit": unit,
            "item_type": "MATERIAL", "category": ProductCategory.objects.get(category_name="管材", item_type="MATERIAL")})
        for line in BomLine.objects.filter(section__bom_id=bom_id, section__section_type="PROFILE"):
            line.raw_values = {**line.raw_values, "source_inventory_material_id": str(source.pk), "tube_cut_piece_count": "10"}
            line.specification = source.specification
            line.save()
        meter, _ = ProductionUnit.objects.get_or_create(unit_name="米", defaults={"unit_code": "S4LEGMETER"})
        fabric, _ = RawMaterialInformation.objects.get_or_create(material_number="S4LEGFABRIC", defaults={
            "material_name": "测试面料", "unit": meter, "item_type": "MATERIAL",
            "category": ProductCategory.objects.get(category_name="面料", item_type="MATERIAL")})
        for line in BomLine.objects.filter(section__bom_id=bom_id, section__section_type="FABRIC"):
            if "面" in line.section.title:
                line.raw_values = {**line.raw_values, "source_inventory_material_id": str(fabric.pk)}
                line.unit_consumption = line.unit_consumption or line.purchase_consumption
                line.save()
        return source

    def review_product_information(self, bom_id, etag):
        color = ProductColor.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        self.assertIsNotNone(color)
        response = self.client.post(
            f"/api/v1/boms/{bom_id}/confirm-product/",
            {"product_id": str(self.product.pk), "product_color_id": str(color.pk), "color_code": color.color_code},
            format="json",
            HTTP_IF_MATCH=etag,
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["data"]["product_information_confirmed"])
        return response

    def save_confirmed_material_sections(self, bom_id, sections, etag):
        body = {"sections": sections}
        preview = self.client.post(
            f"/api/v1/boms/{bom_id}/material-save-preview/", body,
            format="json", HTTP_IF_MATCH=etag,
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        token = preview.data["data"]["confirmation_token"]
        if token:
            body["confirmation_token"] = token
        return self.client.patch(
            f"/api/v1/boms/{bom_id}/", body,
            format="json", HTTP_IF_MATCH=etag,
        )

    @staticmethod
    def fixed_sections(source_sections):
        definitions = (
            ("管材类", "PROFILE"),
            ("五金件", "ACCESSORY"),
            ("塑料件类", "ACCESSORY"),
            ("辅料类", "FABRIC"),
            ("面料类", "FABRIC"),
            ("包材类", "ACCESSORY"),
            ("裁剪明细", "FABRIC"),
        )
        expected_types = dict(definitions)
        used_titles = set()
        sections = []
        keyword_groups = (
            ("裁剪明细", ("裁剪", "裁片")),
            ("塑料件类", ("塑料",)),
            ("包材类", ("包材", "包装")),
            ("五金件", ("五金", "配件")),
            ("面料类", ("面辅料", "面料")),
            ("辅料类", ("辅料",)),
        )
        for section in source_sections:
            if section["section_type"] == "PROFILE":
                fixed_title = "管材类"
            else:
                compact_title = "".join(section["title"].split())
                fixed_title = next(
                    (
                        title
                        for title, keywords in keyword_groups
                        if title not in used_titles and any(keyword in compact_title for keyword in keywords)
                    ),
                    "",
                )
                if not fixed_title:
                    candidates = (
                        ("面料类", "辅料类", "裁剪明细")
                        if section["section_type"] == "FABRIC"
                        else ("五金件", "塑料件类", "包材类")
                    )
                    fixed_title = next(title for title in candidates if title not in used_titles)
            section["title"] = fixed_title
            used_titles.add(fixed_title)
            sections.append(section)

        for title, section_type in definitions:
            if title not in used_titles:
                sections.append({"title": title, "section_type": section_type, "lines": []})
        return sorted(sections, key=lambda section: list(expected_types).index(section["title"]))

    @staticmethod
    def section_by_title(sections, title):
        return next(section for section in sections if section["title"] == title)

    def workbook_bytes(
        self,
        *,
        task_number=None,
        file_name="bom.xlsx",
        quantity_label="产品数量",
        bom_sheet_index=1,
        product_name=None,
        product=None,
    ):
        workbook_product = product or self.product
        workbook = Workbook()
        workbook.active.title = "生产任务单"
        while len(workbook.worksheets) < bom_sheet_index:
            workbook.create_sheet(f"产品任务单{len(workbook.worksheets)}")
        sheet = workbook.create_sheet("BOM数据")
        sheet["A2"] = "产品编号"
        sheet["C2"] = workbook_product.product_code
        sheet["A3"] = "产品名称"
        sheet["C3"] = product_name or workbook_product.product_name
        sheet["A4"] = "产品规格"
        sheet["C4"] = workbook_product.specification
        sheet["E4"] = "生产单号"
        sheet["H4"] = task_number or self.task.task_number
        sheet["A5"] = quantity_label
        sheet["C5"] = 11

        sheet["A8"] = "型材清单"
        for column, value in {"A9": "序号", "B9": "类别", "C9": "材质", "D9": "配件名称", "E9": "规格/厚度（MM）", "H9": "长度(MM)", "I9": "单耗", "J9": "总耗", "K9": "单位", "O9": "物料编码"}.items():
            sheet[column] = value
        for column, value in {"A10": 1, "B10": "型材", "C10": "铁", "D10": "交叉管", "E10": "Φ", "F10": 16, "G10": 0.6, "H10": 580, "I10": 6, "J10": "=C$5*I10", "K10": "PCS", "O10": "MT-STEEL-001"}.items():
            sheet[column] = value

        sheet["A12"] = "配件清单"
        for column, value in {"A13": "序号", "B13": "类别", "C13": "材质", "D13": "配件名称", "E13": "规格(MM)", "I13": "单耗", "J13": "订单量", "K13": "单位"}.items():
            sheet[column] = value
        for column, value in {"A14": 1, "B14": "五金件", "C14": "标准件", "D14": "铆钉", "E14": "Φ4×35", "I14": 4, "J14": 11, "K14": "PCS"}.items():
            sheet[column] = value
        sheet["D15"] = "型号"
        sheet["A16"] = "Q235铆钉"
        sheet["D16"] = "Φ4"

        sheet["A18"] = "绿色面辅料清单"
        for column, value in {"A19": "序号", "B19": "配件名称", "D19": "裁片名称", "G19": "裁片量", "H19": "原材料", "K19": "单位", "L19": "单采购量", "M19": "订单量", "N19": "总采购量"}.items():
            sheet[column] = value
        for column, value in {"A20": 1, "B20": "椅面", "D20": "椅面大片", "G20": 1, "H20": "300D/1.5M门幅", "K20": "M", "L20": 0.5, "M20": 11, "N20": 5.5}.items():
            sheet[column] = value
        sheet["A21"] = "总共"
        sheet["N21"] = 5.5

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return SimpleUploadedFile(
            file_name,
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_parser_reads_second_sheet_and_excludes_calculation_and_total_rows(self):
        upload = self.workbook_bytes()
        with TemporaryDirectory() as folder:
            path = f"{folder}/bom.xlsx"
            with open(path, "wb") as target:
                target.write(upload.read())
            plan = parse_bom_workbook(path, "bom.xlsx")

        self.assertEqual(plan["sheet_name"], "BOM数据")
        self.assertEqual(plan["sheet_number"], 2)
        self.assertEqual(plan["section_count"], 3)
        self.assertEqual(plan["line_count"], 3)
        self.assertEqual([item["line_count"] for item in plan["sections"]], [1, 1, 1])
        self.assertEqual(plan["sections"][1]["calculation_rows"][0]["source_row_number"], 15)
        self.assertEqual(plan["sections"][2]["calculation_rows"][0]["source_row_number"], 21)
        self.assertIn("J", plan["sections"][0]["lines"][0]["source_formulas"])

    def test_safe_preview_is_capped_and_does_not_expose_raw_workbook_data(self):
        source_lines = [
            {
                "source_row_number": index + 10,
                "category": "型材",
                "part_name": f"物料 {index + 1}",
                "specification": "Φ16 × 0.6 mm",
                "unit": "PCS",
                "unit_consumption": "2",
                "notes": "",
                "raw_values": {"O": f"MT-{index + 1:03d}"},
                "source_formulas": {"J": "=A1"},
            }
            for index in range(PREVIEW_LINE_LIMIT + 2)
        ]
        plan = {
            "sections": [{
                "section_type": "PROFILE",
                "title": "型材清单",
                "line_count": len(source_lines),
                "headers": {"O": "物料编码"},
                "lines": source_lines,
            }],
        }

        sections = _safe_preview_sections(plan)
        rows = sections[0]["preview_rows"]

        self.assertEqual(len(rows), PREVIEW_LINE_LIMIT)
        self.assertEqual(rows[-1]["line_number"], PREVIEW_LINE_LIMIT)
        self.assertEqual(rows[0]["material_code"], "MT-001")
        self.assertNotIn("raw_values", rows[0])
        self.assertNotIn("source_formulas", rows[0])

    def test_sales_order_bound_product_line_cannot_be_replaced_from_bom_flow(self):
        self.product_line.sales_order_delivery_batch_line_id = uuid4()

        with self.assertRaisesMessage(Conflict, "来源于销售订单"):
            replace_unbound_product_line(
                task=self.task,
                product_line=self.product_line,
                product_data={},
                files={},
                actor_employee=self.employee,
                request=None,
            )

    def test_parser_accepts_order_quantity_and_falls_back_to_named_bom_sheet(self):
        upload = self.workbook_bytes(quantity_label="订单数量", bom_sheet_index=3)
        with TemporaryDirectory() as folder:
            path = f"{folder}/bom.xlsx"
            with open(path, "wb") as target:
                target.write(upload.read())
            plan = parse_bom_workbook(path, "bom.xlsx")

        self.assertEqual(plan["sheet_name"], "BOM数据")
        self.assertEqual(plan["sheet_number"], 4)
        self.assertEqual(plan["source_quantity"], Decimal("11"))
        self.assertEqual(plan["section_count"], 3)
        self.assertEqual(plan["line_count"], 3)

    @patch("apps.boms.uploads._clamav_scan", side_effect=ScannerUnavailable())
    def test_preview_accepts_order_number_reference_and_product_name_variant(self, _scan):
        response = self.client.post(
            "/api/v1/boms/preview-import/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "workbook": self.workbook_bytes(
                    task_number=self.task.order_number,
                    product_name=f"{self.product.product_name}（表格名称）",
                ),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["identity_errors"], [])
        _scan.assert_not_called()

    @override_settings(FILE_VIRUS_SCAN_ENABLED=True)
    @patch("apps.boms.uploads._clamav_scan", return_value="stream: OK")
    def test_preview_scans_when_enabled_and_cleans_temporary_file(self, scanner):
        response = self.client.post(
            "/api/v1/boms/preview-import/",
            {"production_task_product_id": str(self.product_line.pk), "workbook": self.workbook_bytes()},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)
        scanner.assert_called_once()
        temporary_path = scanner.call_args.args[0]
        self.assertFalse(temporary_path.exists())
        self.assertFalse(temporary_path.parent.exists())

    @override_settings(FILE_VIRUS_SCAN_ENABLED=True)
    def test_preview_scan_failure_blocks_parsing_and_cleans_temporary_file(self):
        for error in (ScannerUnavailable(), InfectedFile()):
            with self.subTest(error=type(error).__name__):
                with patch("apps.boms.uploads._clamav_scan", side_effect=error) as scanner, patch("apps.boms.services.parse_bom_workbook") as parser:
                    response = self.client.post(
                        "/api/v1/boms/preview-import/",
                        {"production_task_product_id": str(self.product_line.pk), "workbook": self.workbook_bytes()},
                        format="multipart",
                    )

                self.assertEqual(response.status_code, error.status_code, response.data)
                scanner.assert_called_once()
                parser.assert_not_called()
                self.assertEqual([path for path in Path(self.media_dir.name).rglob("*") if path.is_file()], [])
                temporary_path = scanner.call_args.args[0]
                self.assertFalse(temporary_path.parent.exists())

    def replacement_product(self):
        return ProductInformation.objects.create(
            product_number="BOM-REPLACE-PRODUCT",
            product_code="BOM-REPLACE-CODE",
            category=self.product.category,
            product_name="BOM 更换产品",
            specification="100×50×20",
            unit=self.product.unit,
            product_item=self.product.product_item,
            warehouse_number="",
            warehouse_name="",
            scale="CUSTOM",
            status="ACTIVE",
        )

    def replacement_payload(self, product):
        color = ProductColor.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        self.assertIsNotNone(color)
        return {
            "client_key": str(self.product_line.pk),
            "product_id": str(product.pk),
            "product_color_id": str(color.pk),
            "color_code": color.color_code,
            "quantity": "25",
            "unit_id": str(product.unit_id),
            "currency_id": str(self.currency.pk),
            "unit_price": "2.5000",
            "exchange_rate": "1.000" if self.currency.code == "CNY" else "7.123",
            "special_requirements": "更换产品后的特殊要求",
            "packaging_method": "独立纸箱",
            "retained_image_ids": [],
            "details": [{
                "client_key": "bom_replace_detail_1",
                "template_key": "material_tube",
                "detail_type": "材料",
                "category_label": "管材",
                "is_selected": True,
                "description": "使用新产品管材",
            }],
        }

    def test_preview_replacement_uses_draft_without_writing_task(self):
        self.task.status = ProductionTaskStatus.DRAFT
        self.task.save(update_fields=["status", "updated_at"])
        replacement = self.replacement_product()
        payload = self.replacement_payload(replacement)

        response = self.client.post(
            "/api/v1/boms/preview-import/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "production_task_product_update": json.dumps(payload),
                "workbook": self.workbook_bytes(product=replacement),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["identity_errors"], [])
        self.assertEqual(response.data["data"]["quantity"]["order_quantity"], Decimal("25"))
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)
        self.assertEqual(self.product_line.quantity, 10)
        self.assertEqual(Bom.objects.count(), 0)

    def test_create_replacement_updates_task_and_bom_atomically(self):
        replacement = self.replacement_product()
        payload = self.replacement_payload(replacement)
        before_audits = AuditLog.objects.count()

        response = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "production_task_product_update": json.dumps(payload),
            },
            format="multipart",
            HTTP_IF_MATCH=etag_for_instance(self.task),
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, replacement.pk)
        self.assertEqual(self.product_line.quantity, 25)
        self.assertEqual(self.product_line.details.count(), 1)
        bom = Bom.objects.get()
        self.assertEqual(bom.product_line_id, self.product_line.pk)
        self.assertEqual(bom.product_code_snapshot, replacement.product_code)
        self.assertEqual(bom.order_quantity, Decimal("25"))
        self.assertEqual(AuditLog.objects.count(), before_audits + 2)

    def test_review_then_verify_updates_unbound_task_line_and_bom(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        bom = Bom.objects.get(pk=created.data["data"]["id"])
        color = ProductColor.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        self.assertIsNotNone(color)

        from apps.reference_data.models import ProductCategory, ProductionUnit
        accessory = RawMaterialInformation.objects.create(
            material_number="BOM-REVIEW-ACCESSORY", material_name="核对测试配件", item_type="MATERIAL",
            category=ProductCategory.objects.get(category_name="五金", item_type="MATERIAL"),
            unit=ProductionUnit.objects.get(unit_name="PCS"),
        )
        section = BomSection.objects.create(bom=bom, line_number=2, section_type="ACCESSORY", title="五金件")
        BomLine.objects.create(
            section=section, line_number=1, part_name=accessory.material_name,
            unit=accessory.unit.unit_name, unit_consumption="1", total_consumption="10",
            raw_values={"source_inventory_material_id": str(accessory.pk)},
        )
        Bom.objects.filter(pk=bom.pk).update(line_count=1, section_count=2)

        # 模拟销售/BOM之前只保存客户手工输入的产品；正式确认时再绑定主档。
        self.product_line.product = None
        self.product_line.product_name_snapshot = "客户手工产品"
        self.product_line.specification_snapshot = "客户手工规格"
        self.product_line.customer_product_name_snapshot = "客户手工产品"
        self.product_line.customer_specification_snapshot = "客户手工规格"
        self.product_line.customer_color_snapshot = "客户手工颜色"
        self.product_line.color_name_snapshot = "客户手工颜色"
        self.product_line.product_code_snapshot = ""
        self.product_line.product_number_snapshot = "CUSTOMER-001"
        self.product_line.save(update_fields=[
            "product", "product_name_snapshot", "specification_snapshot",
            "customer_product_name_snapshot", "customer_specification_snapshot",
            "customer_color_snapshot", "color_name_snapshot", "product_code_snapshot",
            "product_number_snapshot", "updated_at",
        ])

        detail = self.client.get(f"/api/v1/boms/{bom.pk}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        response = self.client.post(
            f"/api/v1/boms/{bom.pk}/confirm-product/",
            {
                "product_id": str(self.product.pk),
                "product_color_id": str(color.pk),
                "color_code": color.color_code,
            },
            format="json",
            HTTP_IF_MATCH=detail.headers["ETag"],
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.product_line.refresh_from_db()
        bom.refresh_from_db()
        self.assertIsNone(self.product_line.product_id)
        self.assertEqual(self.product_line.customer_product_name_snapshot, "客户手工产品")
        self.assertEqual(self.product_line.customer_color_snapshot, "客户手工颜色")
        self.assertEqual(bom.product_name_snapshot, self.product.product_name)
        self.assertTrue(response.data["data"]["product_information_confirmed"])

        verified = self.client.post(
            f"/api/v1/boms/{bom.pk}/verify/",
            {"quantity_confirmation_reason": "按当前任务产品数量核对，批次数量差异不改变本次产品回写范围。"},
            format="json",
            HTTP_IF_MATCH=response.headers["ETag"],
        )
        self.assertEqual(verified.status_code, 200, verified.data)
        self.assertEqual(verified.data["data"]["status"], "VERIFIED")
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)
        self.assertEqual(self.product_line.product_color_id, color.pk)
        self.assertEqual(self.product_line.customer_product_name_snapshot, self.product.product_name)
        self.assertEqual(self.product_line.customer_color_snapshot, color.color_name)

    def test_create_saves_quantity_difference_reason_for_later_verification(self):
        reason = "Excel原表多100件，按当前任务数量生产。"
        response = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "quantity_difference_reason": reason,
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 201, response.data)
        bom = Bom.objects.get()
        self.assertEqual(bom.quantity_difference_reason, reason)
        self.assertEqual(response.data["data"]["quantity_difference_reason"], reason)

    def test_verify_uses_quantity_difference_reason_saved_with_draft(self):
        reason = "Excel数量仅供参考，按当前生产任务数量执行。"
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(),
                "quantity_difference_reason": reason,
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.complete_pipe_sources(created.data["data"]["id"])
        reviewed = self.review_product_information(created.data["data"]["id"], created.headers["ETag"])

        verified = self.client.post(
            f"/api/v1/boms/{created.data['data']['id']}/verify/",
            {},
            format="json",
            HTTP_IF_MATCH=reviewed.headers["ETag"],
        )

        self.assertEqual(verified.status_code, 200, verified.data)
        self.assertEqual(verified.data["data"]["status"], "VERIFIED")
        self.assertEqual(verified.data["data"]["confirmation_reason"], reason)

    def test_create_replacement_requires_current_task_etag(self):
        replacement = self.replacement_product()
        payload = self.replacement_payload(replacement)

        response = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "production_task_product_update": json.dumps(payload),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 428, response.data)
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)
        self.assertEqual(Bom.objects.count(), 0)

        stale = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "production_task_product_update": json.dumps(payload),
            },
            format="multipart",
            HTTP_IF_MATCH='"stale-task-etag"',
        )
        self.assertEqual(stale.status_code, 412, stale.data)
        self.assertEqual(Bom.objects.count(), 0)

    def test_create_replacement_rejects_existing_bom_and_missing_permission(self):
        created = self.client.post(
            "/api/v1/boms/",
            {"production_task_product_id": str(self.product_line.pk), "creation_method": "BLANK"},
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.task.status = ProductionTaskStatus.DRAFT
        self.task.save(update_fields=["status", "updated_at"])
        replacement = self.replacement_product()
        payload = self.replacement_payload(replacement)

        conflict = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "production_task_product_update": json.dumps(payload),
            },
            format="multipart",
            HTTP_IF_MATCH=etag_for_instance(self.task),
        )
        self.assertEqual(conflict.status_code, 409, conflict.data)
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)

        Bom.objects.all().delete()
        change_permission = Permission.objects.get(
            content_type__app_label="production_tasks",
            codename="change_productiontask",
        )
        self.user.user_permissions.remove(change_permission)
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(self.user)
        denied = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
                "production_task_product_update": json.dumps(payload),
            },
            format="multipart",
            HTTP_IF_MATCH=etag_for_instance(self.task),
        )
        self.assertEqual(denied.status_code, 403, denied.data)
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)
        self.assertEqual(Bom.objects.count(), 0)

    def test_create_replacement_rolls_back_when_bom_import_fails(self):
        replacement = self.replacement_product()
        payload = self.replacement_payload(replacement)
        before_audits = AuditLog.objects.count()

        response = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "production_task_product_update": json.dumps(payload),
                "workbook": self.workbook_bytes(task_number="WRONG-ORDER", product=replacement),
            },
            format="multipart",
            HTTP_IF_MATCH=etag_for_instance(self.task),
        )

        self.assertEqual(response.status_code, 400, response.data)
        self.product_line.refresh_from_db()
        self.assertEqual(self.product_line.product_id, self.product.pk)
        self.assertEqual(self.product_line.quantity, 10)
        self.assertEqual(Bom.objects.count(), 0)
        self.assertEqual(AuditLog.objects.count(), before_audits)

    def test_preview_create_list_and_stats_use_real_persisted_data(self):
        preview = self.client.post(
            "/api/v1/boms/preview-import/",
            {"production_task_product_id": str(self.product_line.pk), "workbook": self.workbook_bytes()},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["data"]["line_count"], 3)
        self.assertEqual(preview.data["data"]["preview_limit"], PREVIEW_LINE_LIMIT)
        self.assertEqual(preview.data["data"]["preview_row_count"], 3)
        self.assertFalse(preview.data["data"]["has_more_preview_rows"])
        preview_rows = [
            row
            for section in preview.data["data"]["sections"]
            for row in section["preview_rows"]
        ]
        self.assertEqual(len(preview_rows), 3)
        self.assertEqual(preview_rows[0]["material_code"], "MT-STEEL-001")
        self.assertEqual(preview_rows[0]["material_name"], "交叉管")
        self.assertEqual(preview_rows[0]["specification"], "Φ16 × 0.6 mm")
        self.assertEqual(preview_rows[0]["unit_consumption"], "6")
        self.assertNotIn("raw_values", preview_rows[0])
        self.assertNotIn("source_formulas", preview_rows[0])
        self.assertEqual(preview.data["data"]["quantity"]["order_quantity"], Decimal("10"))
        self.assertEqual(preview.data["data"]["quantity"]["source_quantity"], Decimal("11"))
        self.assertEqual(preview.data["data"]["quantity"]["batch_quantity"], Decimal("12"))
        self.assertTrue(preview.data["data"]["quantity"]["requires_confirmation"])

        options = self.client.get("/api/v1/boms/options/")
        self.assertEqual(options.status_code, 200, options.data)
        self.assertEqual(options.data["data"]["production_task_products"][0]["order_number"], "ORDER-BOM-001")
        self.assertEqual(str(options.data["data"]["production_task_products"][0]["product_id"]), str(self.product.pk))
        self.assertTrue(options.data["data"]["production_task_products"][0]["task_etag"])
        self.assertTrue(any(item["id"] == str(self.product.pk) for item in options.data["data"]["products"]))
        self.assertTrue(options.data["data"]["currencies"])
        self.assertTrue(options.data["data"]["units"])
        self.assertIn("detail_template", options.data["data"])
        matched_options = self.client.get("/api/v1/boms/options/", {"task_q": "ORDER-BOM-001"})
        unmatched_options = self.client.get("/api/v1/boms/options/", {"task_q": "ORDER-NOT-FOUND"})
        self.assertEqual(len(matched_options.data["data"]["production_task_products"]), 1)
        self.assertEqual(unmatched_options.data["data"]["production_task_products"], [])

        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(),
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["data"]["order_number"], "ORDER-BOM-001")
        self.assertRegex(created.data["data"]["bom_number"], r"^BOM[0-9]{6}[0-9]{3,}$")
        self.assertEqual(created.data["data"]["line_count"], 3)
        self.assertEqual(BomSection.objects.count(), 3)
        self.assertEqual(BomLine.objects.count(), 3)
        source_file = Bom.objects.get().source_file
        self.assertEqual(source_file.scan_status, FileScanStatus.SKIPPED)
        self.assertIsNone(source_file.scanned_at)
        self.assertEqual(source_file.scan_message, "BOM 上传已关闭病毒扫描。")

        listed = self.client.get("/api/v1/boms/", {"status": "PENDING_CONFIRMATION"})
        stats = self.client.get("/api/v1/boms/stats/")
        self.assertEqual(listed.status_code, 200, listed.data)
        self.assertEqual(listed.data["meta"]["total"], 1)
        self.assertEqual(listed.data["data"][0]["order_number"], "ORDER-BOM-001")
        self.assertEqual(listed.data["data"][0]["version_number"], 1)
        self.assertEqual(listed.data["data"][0]["batch_quantity_breakdown"], [Decimal("5"), Decimal("7")])
        searched = self.client.get("/api/v1/boms/", {"q": "ORDER-BOM-001"})
        self.assertEqual(searched.status_code, 200, searched.data)
        self.assertEqual(searched.data["meta"]["total"], 1)
        detail = self.client.get(f"/api/v1/boms/{created.data['data']['id']}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertEqual(detail.data["data"]["order_number"], "ORDER-BOM-001")
        self.assertEqual(detail.data["data"]["customer_name"], "BOM 测试客户")
        self.assertIsNone(detail.data["data"]["planned_start_date"])
        self.assertIsNone(detail.data["data"]["planned_end_date"])
        self.assertEqual(stats.data["data"]["total"], 1)
        self.assertEqual(stats.data["data"]["pending_confirmation"], 1)
        filtered_stats = self.client.get(
            "/api/v1/boms/stats/",
            {"q": "ORDER-BOM-001", "status": "VERIFIED"},
        )
        self.assertEqual(filtered_stats.status_code, 200, filtered_stats.data)
        self.assertEqual(filtered_stats.data["data"]["total"], 1)
        bom = Bom.objects.get(pk=created.data["data"]["id"])
        bom.product_name_snapshot = "@危险公式"
        bom.save(update_fields=["product_name_snapshot", "updated_at"])
        exported = self.client.post(
            "/api/v1/boms/export/",
            {"ids": [str(bom.pk)]},
            format="json",
            QUERY_STRING="status=PENDING_CONFIRMATION",
        )
        self.assertEqual(exported.status_code, 200)
        self.assertIn("text/csv", exported["Content-Type"])
        self.assertEqual(exported["X-Export-Count"], "1")
        exported_text = exported.content.decode("utf-8-sig")
        self.assertIn("'@危险公式", exported_text)
        self.assertIn(bom.bom_number, exported_text)
        invalid_stats_status = self.client.get("/api/v1/boms/stats/", {"status": "UNKNOWN"})
        self.assertEqual(invalid_stats_status.status_code, 400, invalid_stats_status.data)
        oversized_search = self.client.get("/api/v1/boms/", {"q": "测" * 201})
        self.assertEqual(oversized_search.status_code, 400, oversized_search.data)
        invalid_range = self.client.get(
            "/api/v1/boms/",
            {"updated_after": "2026-09-11", "updated_before": "2026-09-10"},
        )
        self.assertEqual(invalid_range.status_code, 400, invalid_range.data)
        self.assertIn("updated_before", invalid_range.data["error"]["field_errors"])

    def test_create_rejects_workbook_for_another_task(self):
        response = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(task_number="HT-OTHER-001"),
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Bom.objects.count(), 0)

    def test_options_hide_only_product_lines_that_already_have_a_bom(self):
        available_product_line = ProductionTaskProduct.objects.create(
            production_task=self.task,
            line_number=2,
            product=self.product,
            product_number_snapshot=self.product.product_number,
            product_code_snapshot=self.product.product_code,
            item_number_snapshot=self.product.product_item.item_number,
            product_name_snapshot="同订单待建 BOM 产品",
            specification_snapshot=self.product.specification,
            quantity=8,
            unit=self.product.unit,
            currency=self.currency,
            unit_price=Decimal("1"),
            total_amount=Decimal("8"),
            exchange_rate=Decimal("1"),
        )

        before = self.client.get("/api/v1/boms/options/")
        self.assertEqual(before.status_code, 200, before.data)
        self.assertEqual(
            {str(item["id"]) for item in before.data["data"]["production_task_products"]},
            {str(self.product_line.pk), str(available_product_line.pk)},
        )

        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)

        after = self.client.get("/api/v1/boms/options/")
        self.assertEqual(after.status_code, 200, after.data)
        candidates = after.data["data"]["production_task_products"]
        self.assertEqual([str(item["id"]) for item in candidates], [str(available_product_line.pk)])
        self.assertEqual(candidates[0]["existing_bom_count"], 0)
        self.assertTrue(any(item["id"] == created.data["data"]["id"] for item in after.data["data"]["history_boms"]))

    def test_create_and_options_require_pmc_confirmed_task(self):
        for task_status in (ProductionTaskStatus.DRAFT, ProductionTaskStatus.PENDING_VALIDATION):
            with self.subTest(task_status=task_status):
                self.task.status = task_status
                self.task.save(update_fields=["status", "updated_at"])

                options = self.client.get("/api/v1/boms/options/")
                self.assertEqual(options.status_code, 200, options.data)
                self.assertEqual(options.data["data"]["production_task_products"], [])

                created = self.client.post(
                    "/api/v1/boms/",
                    {
                        "production_task_product_id": str(self.product_line.pk),
                        "creation_method": "BLANK",
                    },
                    format="json",
                )
                self.assertEqual(created.status_code, 409, created.data)
                self.assertIn("PMC", str(created.data))
                self.assertFalse(Bom.objects.exists())

    def test_copy_and_blank_creation_keep_real_source_rules(self):
        imported = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(),
            },
            format="multipart",
        )
        self.assertEqual(imported.status_code, 201, imported.data)
        imported_number = imported.data["data"]["bom_number"]
        self.assertRegex(imported_number, r"^BOM[0-9]{6}[0-9]{3,}$")

        copied = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "COPY",
                "source_bom_id": imported.data["data"]["id"],
            },
            format="json",
        )
        self.assertEqual(copied.status_code, 201, copied.data)
        copied_number = copied.data["data"]["bom_number"]
        self.assertRegex(copied_number, r"^BOM[0-9]{6}[0-9]{3,}$")
        self.assertEqual(copied_number[3:9], imported_number[3:9])
        self.assertEqual(int(copied_number[9:]), int(imported_number[9:]) + 1)
        self.assertEqual(copied.data["data"]["copied_from_number"], imported_number)
        self.assertEqual(copied.data["data"]["source_sheet_name"], "BOM数据")
        self.assertEqual(copied.data["data"]["line_count"], 3)
        self.assertEqual(copied.data["data"]["sections"][1]["calculation_rows"][0]["source_row_number"], 15)

        blank = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(blank.status_code, 201, blank.data)
        blank_number = blank.data["data"]["bom_number"]
        self.assertRegex(blank_number, r"^BOM[0-9]{6}[0-9]{3,}$")
        self.assertEqual(blank_number[3:9], copied_number[3:9])
        self.assertEqual(int(blank_number[9:]), int(copied_number[9:]) + 1)
        self.assertEqual(blank.data["data"]["line_count"], 0)
        self.assertEqual(blank.data["data"]["section_count"], 1)
        self.assertEqual(len(blank.data["data"]["sections"]), 1)
        self.assertEqual(blank.data["data"]["sections"][0]["section_type"], "PROFILE")
        self.assertEqual(blank.data["data"]["sections"][0]["title"], "管材类")

    def test_update_requires_fixed_seven_sections_and_preserves_existing_identity(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        reversed_sections = list(reversed(self.fixed_sections(created.data["data"]["sections"])))
        normalized = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": reversed_sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(normalized.status_code, 200, normalized.data)
        sections = normalized.data["data"]["sections"]
        self.assertEqual(
            [section["title"] for section in sections],
            ["管材类", "五金件", "塑料件类", "辅料类", "面料类", "包材类", "裁剪明细"],
        )

        missing = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": deepcopy(sections[:-1])},
            format="json",
            HTTP_IF_MATCH=normalized.headers["ETag"],
        )
        self.assertEqual(missing.status_code, 400, missing.data)
        self.assertIn("固定七个分区", str(missing.data))

        extra_sections = deepcopy(sections)
        extra_sections.append({"section_type": "FABRIC", "title": "自定义分区", "lines": []})
        extra = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": extra_sections},
            format="json",
            HTTP_IF_MATCH=normalized.headers["ETag"],
        )
        self.assertEqual(extra.status_code, 400, extra.data)
        self.assertIn("固定七个分区", str(extra.data))

        renamed_sections = deepcopy(sections)
        renamed_sections[1]["title"], renamed_sections[2]["title"] = (
            renamed_sections[2]["title"],
            renamed_sections[1]["title"],
        )
        renamed = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": renamed_sections},
            format="json",
            HTTP_IF_MATCH=normalized.headers["ETag"],
        )
        self.assertEqual(renamed.status_code, 400, renamed.data)
        self.assertIn("名称不能修改", str(renamed.data))

        retyped_sections = deepcopy(sections)
        retyped_sections[0]["title"], retyped_sections[1]["title"] = (
            retyped_sections[1]["title"],
            retyped_sections[0]["title"],
        )
        retyped_sections[0]["section_type"], retyped_sections[1]["section_type"] = (
            retyped_sections[1]["section_type"],
            retyped_sections[0]["section_type"],
        )
        retyped = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": retyped_sections},
            format="json",
            HTTP_IF_MATCH=normalized.headers["ETag"],
        )
        self.assertEqual(retyped.status_code, 400, retyped.data)
        self.assertIn("类型不能修改", str(retyped.data))

        replaced_existing = deepcopy(sections)
        replaced_existing[2].pop("id")
        replaced = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": replaced_existing},
            format="json",
            HTTP_IF_MATCH=normalized.headers["ETag"],
        )
        self.assertEqual(replaced.status_code, 400, replaced.data)
        self.assertIn("不能删除", str(replaced.data))

    def test_hardware_material_picker_is_bom_scoped_and_only_returns_hardware(self):
        response = self.client.get("/api/v1/boms/hardware-materials/?q=铆钉&page_size=5")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertGreater(response.data["meta"]["total"], 0)
        self.assertLessEqual(len(response.data["data"]), 5)
        for item in response.data["data"]:
            self.assertEqual(item["category_name"], "五金")
            self.assertIn("material_number", item)
            self.assertIn("material", item)
            self.assertIn("color_name", item)
            self.assertIn("weight_kg", item)
            self.assertIn("flow_unit", item)

        source = (
            RawMaterialInformation.objects.filter(
                item_type="MATERIAL",
                category__category_name="五金",
                category__status="ACTIVE",
                category__archived_at__isnull=True,
                archived_at__isnull=True,
            )
            .select_related("category", "unit", "color")
            .first()
        )
        self.assertIsNotNone(source)

        def create_hardware(material_number, material_name, **extra_fields):
            return RawMaterialInformation.objects.create(
                material_number=material_number,
                item_type=source.item_type,
                category=source.category,
                material_name=material_name,
                specification="BOM 分组测试规格",
                unit=source.unit,
                color=source.color,
                color_snapshot=source.color_snapshot,
                material=source.material,
                warehouse_number=source.warehouse_number,
                warehouse_name=source.warehouse_name,
                **extra_fields,
            )

        group_name = "BOM分组专测QX9"
        first_group_item = create_hardware("BOM-GROUP-QX9-001", group_name)
        second_group_item = create_hardware("BOM-GROUP-QX9-002", group_name)
        fuzzy_only_item = create_hardware("BOM-GROUP-QX9-003", f"{group_name}扩展")
        archived_group_item = create_hardware(
            "BOM-GROUP-QX9-004",
            group_name,
            archived_at=timezone.now(),
        )

        fuzzy_response = self.client.get(
            "/api/v1/boms/hardware-materials/",
            {"q": group_name, "page_size": 100},
        )
        self.assertEqual(fuzzy_response.status_code, 200, fuzzy_response.data)
        self.assertEqual(fuzzy_response.data["meta"]["total"], 3)
        self.assertEqual(
            {item["id"] for item in fuzzy_response.data["data"]},
            {str(first_group_item.pk), str(second_group_item.pk), str(fuzzy_only_item.pk)},
        )
        self.assertNotIn(
            str(archived_group_item.pk),
            {item["id"] for item in fuzzy_response.data["data"]},
        )

        exact_response = self.client.get(
            "/api/v1/boms/hardware-materials/",
            {"material_name": group_name, "page_size": 100},
        )
        self.assertEqual(exact_response.status_code, 200, exact_response.data)
        self.assertEqual(exact_response.data["meta"]["total"], 2)
        self.assertEqual(
            {item["id"] for item in exact_response.data["data"]},
            {str(first_group_item.pk), str(second_group_item.pk)},
        )

        groups_response = self.client.get("/api/v1/boms/hardware-material-groups/")
        self.assertEqual(groups_response.status_code, 200, groups_response.data)
        group_counts = {
            item["material_name"]: item["count"]
            for item in groups_response.data["data"]
        }
        self.assertEqual(group_counts[group_name], 2)
        self.assertEqual(group_counts[f"{group_name}扩展"], 1)

        change_permission = Permission.objects.get(codename="change_bom", content_type__app_label="boms")
        self.user.user_permissions.remove(change_permission)
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(self.user)
        forbidden = self.client.get("/api/v1/boms/hardware-materials/")
        self.assertEqual(forbidden.status_code, 403, forbidden.data)
        groups_forbidden = self.client.get("/api/v1/boms/hardware-material-groups/")
        self.assertEqual(groups_forbidden.status_code, 403, groups_forbidden.data)

    def test_plastic_picker_and_editor_reference_options_are_bom_scoped(self):
        plastic_response = self.client.get("/api/v1/boms/plastic-materials/?page_size=5")
        self.assertEqual(plastic_response.status_code, 200, plastic_response.data)
        self.assertGreater(plastic_response.data["meta"]["total"], 0)
        for item in plastic_response.data["data"]:
            self.assertEqual(item["category_name"], "塑料件")
            self.assertEqual(item["item_type"], "MATERIAL")
            self.assertEqual(item["item_type_display"], "材料")
            self.assertEqual(
                set(item),
                {
                    "id",
                    "material_number",
                    "item_type",
                    "item_type_display",
                    "category_name",
                    "material_name",
                    "material",
                    "color_name",
                    "specification",
                    "unit_code",
                    "unit_name",
                    "weight_kg",
                    "flow_unit",
                    "number_value",
                },
            )

        options_response = self.client.get("/api/v1/boms/editor-reference-options/")
        self.assertEqual(options_response.status_code, 200, options_response.data)
        options = options_response.data["data"]
        self.assertEqual(
            set(options),
            {
                "auxiliary_products",
                "accessory_materials",
                "fabric_materials",
                "fabric_specifications",
                "fabric_widths",
                "product_colors",
            },
        )
        self.assertGreater(len(options["auxiliary_products"]), 0)
        self.assertEqual(
            set(options["auxiliary_products"][0]),
            {"id", "product_name", "pinyin_search", "unit_name", "specifications", "colors"},
        )
        if options["auxiliary_products"][0]["specifications"]:
            self.assertEqual(
                set(options["auxiliary_products"][0]["specifications"][0]),
                {"id", "specification_value"},
            )
        if options["auxiliary_products"][0]["colors"]:
            self.assertEqual(set(options["auxiliary_products"][0]["colors"][0]), {"id", "color_name"})
        self.assertGreater(len(options["accessory_materials"]), 0)
        self.assertGreater(len(options["fabric_materials"]), 0)
        self.assertGreater(len(options["fabric_specifications"]), 0)
        self.assertGreater(len(options["fabric_widths"]), 0)
        self.assertGreater(len(options["product_colors"]), 0)

        change_permission = Permission.objects.get(codename="change_bom", content_type__app_label="boms")
        self.user.user_permissions.remove(change_permission)
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(self.user)
        self.assertEqual(self.client.get("/api/v1/boms/plastic-materials/").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/boms/editor-reference-options/").status_code, 403)

    def test_plastic_source_snapshot_is_authoritative_and_cannot_be_cleared(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        source = (
            RawMaterialInformation.objects.filter(
                item_type="MATERIAL",
                category__category_name="塑料件",
                category__status="ACTIVE",
                category__archived_at__isnull=True,
                archived_at__isnull=True,
            )
            .select_related("category", "unit", "color")
            .first()
        )
        self.assertIsNotNone(source)
        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "塑料件类")["lines"] = [
                    {
                        "part_name": "客户端伪造品名",
                        "category": "客户端伪造类别",
                        "material": "客户端伪造材质",
                        "specification": "客户端伪造规格",
                        "raw_material": "客户端伪造品号",
                        "unit": "客户端伪造单位",
                        "source_raw_material_id": str(source.pk),
                        "plastic_material_number": "客户端伪造品号",
                        "plastic_color": "客户端伪造颜色",
                        "unit_consumption": "2",
                        "notes": "塑料测试行",
                    }
                ]
        updated = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        line = updated.data["data"]["sections"][2]["lines"][0]
        expected_color = source.color.color_name if source.color_id else source.color_snapshot
        self.assertEqual(line["category"], "塑料件")
        self.assertEqual(line["part_name"], source.material_name)
        self.assertEqual(line["material"], source.material)
        self.assertEqual(line["specification"], source.specification)
        self.assertEqual(line["raw_material"], source.material_number)
        self.assertEqual(line["unit"], source.unit.unit_name)
        self.assertEqual(line["source_raw_material_id"], str(source.pk))
        self.assertEqual(line["plastic_material_number"], source.material_number)
        self.assertEqual(line["plastic_color"], expected_color)
        self.assertEqual(line["raw_values"]["source_raw_material_id"], str(source.pk))

        line["part_name"] = "同来源时也不能篡改快照"
        source.archived_at = timezone.now()
        source.save(update_fields=["archived_at", "updated_at"])
        historical = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": updated.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=updated.headers["ETag"],
        )
        self.assertEqual(historical.status_code, 200, historical.data)
        historical_line = historical.data["data"]["sections"][2]["lines"][0]
        self.assertEqual(historical_line["part_name"], source.material_name)

        historical_line["source_raw_material_id"] = None
        cleared = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": historical.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=historical.headers["ETag"],
        )
        self.assertEqual(cleared.status_code, 400, cleared.data)
        self.assertIn("不能清空数据来源", str(cleared.data))

    def test_auxiliary_line_can_freeze_exact_inventory_material_identity(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        source = (
            RawMaterialInformation.objects.filter(
                item_type="MATERIAL",
                category__category_name="辅料",
                category__status="ACTIVE",
                category__archived_at__isnull=True,
                archived_at__isnull=True,
            )
            .select_related("category", "unit")
            .first()
        )
        self.assertIsNotNone(source)
        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "辅料类")["lines"] = [
            {
                "part_name": "客户端伪造辅料",
                "source_inventory_material_id": str(source.pk),
                "purchase_consumption": "0.5",
            }
        ]

        updated = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )

        self.assertEqual(updated.status_code, 200, updated.data)
        line = self.section_by_title(updated.data["data"]["sections"], "辅料类")["lines"][0]
        self.assertEqual(line["source_inventory_material_id"], str(source.pk))
        self.assertEqual(line["category"], "辅料")
        self.assertEqual(line["part_name"], source.material_name)
        self.assertEqual(line["material"], source.material)
        self.assertEqual(line["specification"], source.specification)
        self.assertEqual(line["raw_material"], source.material_number)
        self.assertEqual(line["unit"], source.unit.unit_name)

    def test_auxiliary_sources_validate_relations_and_round_trip(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        product = (
            AccessoryProductName.objects.filter(
                status="ACTIVE",
                archived_at__isnull=True,
                specifications__status="ACTIVE",
                specifications__archived_at__isnull=True,
                colors__status="ACTIVE",
                colors__archived_at__isnull=True,
            )
            .distinct()
            .first()
        )
        self.assertIsNotNone(product)
        specification = product.specifications.filter(status="ACTIVE", archived_at__isnull=True).first()
        color = product.colors.filter(status="ACTIVE", archived_at__isnull=True).first()
        material = AccessoryMaterial.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        other_specification = (
            AccessorySpecification.objects.filter(status="ACTIVE", archived_at__isnull=True)
            .exclude(product_name=product)
            .first()
        )
        self.assertIsNotNone(other_specification)
        self.assertIsNotNone(material)
        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "辅料类")["lines"] = [
                    {
                        "part_name": "客户端伪造品名",
                        "category": "客户端伪造类别",
                        "material": "客户端伪造材质",
                        "specification": "客户端伪造规格",
                        "unit": "客户端伪造单位",
                        "source_accessory_product_name_id": str(product.pk),
                        "source_accessory_specification_id": str(other_specification.pk),
                        "source_accessory_color_id": str(color.pk),
                        "source_accessory_material_id": str(material.pk),
                        "accessory_color": "客户端伪造颜色",
                        "unit_consumption": "1.5",
                        "notes": "辅料测试行",
                    }
                ]
        wrong_relation = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(wrong_relation.status_code, 400, wrong_relation.data)
        self.assertIn("规格不属于当前辅料品名", str(wrong_relation.data))

        self.section_by_title(sections, "辅料类")["lines"][0]["source_accessory_specification_id"] = str(
            specification.pk
        )
        updated = self.save_confirmed_material_sections(
            created.data["data"]["id"], sections, created.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        line = updated.data["data"]["sections"][3]["lines"][0]
        self.assertEqual(line["category"], "辅料")
        self.assertEqual(line["part_name"], product.product_name)
        self.assertEqual(line["specification"], specification.specification_value)
        self.assertEqual(line["accessory_color"], color.color_name)
        self.assertEqual(line["material"], material.material_name)
        self.assertEqual(line["unit"], product.unit_name)
        self.assertEqual(line["source_accessory_product_name_id"], str(product.pk))
        self.assertEqual(line["source_accessory_specification_id"], str(specification.pk))
        self.assertEqual(line["source_accessory_color_id"], str(color.pk))
        self.assertEqual(line["source_accessory_material_id"], str(material.pk))
        self.assertEqual(line["raw_values"]["accessory_color"], color.color_name)

        line["part_name"] = "同来源时不能篡改"
        line["material"] = "同来源时不能篡改"
        preserved = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": updated.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=updated.headers["ETag"],
        )
        self.assertEqual(preserved.status_code, 200, preserved.data)
        preserved_line = preserved.data["data"]["sections"][3]["lines"][0]
        self.assertEqual(preserved_line["part_name"], product.product_name)
        self.assertEqual(preserved_line["material"], material.material_name)

        BomLine.objects.filter(pk=preserved_line["id"]).update(unit="")
        preserved_line["unit"] = ""
        repaired = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": preserved.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=preserved.headers["ETag"],
        )
        self.assertEqual(repaired.status_code, 200, repaired.data)
        repaired_line = repaired.data["data"]["sections"][3]["lines"][0]
        self.assertEqual(repaired_line["unit"], product.unit_name)

        repaired_line["source_accessory_product_name_id"] = None
        cleared = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": repaired.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=repaired.headers["ETag"],
        )
        self.assertEqual(cleared.status_code, 400, cleared.data)
        self.assertIn("不能清空数据来源", str(cleared.data))

    def test_legacy_plastic_and_auxiliary_lines_without_sources_remain_editable(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        bom = Bom.objects.get(pk=created.data["data"]["id"])
        plastic_section = BomSection.objects.create(
            bom=bom,
            line_number=2,
            section_type="ACCESSORY",
            title="塑料件类",
        )
        BomLine.objects.create(section=plastic_section, line_number=1, part_name="历史塑料件")
        auxiliary_section = BomSection.objects.create(
            bom=bom,
            line_number=3,
            section_type="FABRIC",
            title="辅料类",
        )
        BomLine.objects.create(section=auxiliary_section, line_number=1, part_name="历史辅料")

        detail = self.client.get(f"/api/v1/boms/{bom.pk}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        sections = self.fixed_sections(detail.data["data"]["sections"])
        self.section_by_title(sections, "塑料件类")["lines"][0]["notes"] = "历史塑料件无来源仍可编辑"
        self.section_by_title(sections, "辅料类")["lines"][0]["notes"] = "历史辅料无来源仍可编辑"

        updated = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=detail.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        plastic_line = updated.data["data"]["sections"][2]["lines"][0]
        auxiliary_line = updated.data["data"]["sections"][3]["lines"][0]
        self.assertIsNone(plastic_line["source_raw_material_id"])
        self.assertIsNone(auxiliary_line["source_accessory_product_name_id"])
        self.assertEqual(plastic_line["notes"], "历史塑料件无来源仍可编辑")
        self.assertEqual(auxiliary_line["notes"], "历史辅料无来源仍可编辑")

    def test_fabric_sources_preserve_when_omitted_and_clear_only_on_explicit_null(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        material = FabricMaterial.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        specification = FabricSpecification.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        width = FabricWidth.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        color = ProductColor.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        self.assertTrue(all([material, specification, width, color]))
        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "面料类")["lines"] = [
                    {
                        "part_name": "面料测试行",
                        "material": "客户端伪造材质",
                        "specification": "客户端伪造规格",
                        "unit": "码",
                        "source_fabric_material_id": str(material.pk),
                        "source_fabric_specification_id": str(specification.pk),
                        "source_fabric_width_id": str(width.pk),
                        "source_product_color_id": str(color.pk),
                        "fabric_width": "客户端伪造门幅",
                        "fabric_color_card": "客户端伪造色卡",
                        "fabric_color_code": "客户端伪造色号",
                        "fabric_color_name": "客户端伪造颜色",
                        "unit_consumption": "2.5",
                        "cut_piece_quantity": "3",
                        "notes": "面料测试行",
                    }
                ]
        preview = self.client.post(
            f"/api/v1/boms/{created.data['data']['id']}/material-save-preview/",
            {"sections": sections}, format="json", HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        updated = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections, "confirmation_token": preview.data["data"]["confirmation_token"]},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        line = updated.data["data"]["sections"][4]["lines"][0]
        self.assertEqual(line["material"], material.material_name)
        self.assertEqual(line["specification"], specification.specification_name)
        self.assertEqual(line["unit"], "米")
        self.assertEqual(line["fabric_width"], width.width_value)
        self.assertEqual(line["fabric_color_card"], color.color_card)
        self.assertEqual(line["fabric_color_code"], color.color_code)
        self.assertEqual(line["fabric_color_name"], color.color_name)

        for field in (
            "source_fabric_material_id",
            "source_fabric_specification_id",
            "source_fabric_width_id",
            "source_product_color_id",
        ):
            line.pop(field)
        line.update(
            {
                "material": "省略来源时不能篡改材质",
                "specification": "省略来源时不能篡改规格",
                "fabric_width": "省略来源时不能篡改门幅",
                "fabric_color_card": "省略来源时不能篡改色卡",
                "fabric_color_code": "省略来源时不能篡改色号",
                "fabric_color_name": "省略来源时不能篡改颜色",
            }
        )
        preserved = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": updated.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=updated.headers["ETag"],
        )
        self.assertEqual(preserved.status_code, 200, preserved.data)
        preserved_line = preserved.data["data"]["sections"][4]["lines"][0]
        self.assertEqual(preserved_line["source_fabric_material_id"], str(material.pk))
        self.assertEqual(preserved_line["source_fabric_specification_id"], str(specification.pk))
        self.assertEqual(preserved_line["source_fabric_width_id"], str(width.pk))
        self.assertEqual(preserved_line["source_product_color_id"], str(color.pk))
        self.assertEqual(preserved_line["material"], material.material_name)
        self.assertEqual(preserved_line["specification"], specification.specification_name)
        self.assertEqual(preserved_line["fabric_width"], width.width_value)
        self.assertEqual(preserved_line["fabric_color_card"], color.color_card)
        self.assertEqual(preserved_line["fabric_color_code"], color.color_code)
        self.assertEqual(preserved_line["fabric_color_name"], color.color_name)

        for field in (
            "source_fabric_material_id",
            "source_fabric_specification_id",
            "source_fabric_width_id",
            "source_product_color_id",
        ):
            preserved_line[field] = None
        cleared = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": preserved.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=preserved.headers["ETag"],
        )
        self.assertEqual(cleared.status_code, 200, cleared.data)
        cleared_line = cleared.data["data"]["sections"][4]["lines"][0]
        self.assertIsNone(cleared_line["source_fabric_material_id"])
        self.assertIsNone(cleared_line["source_fabric_specification_id"])
        self.assertIsNone(cleared_line["source_fabric_width_id"])
        self.assertIsNone(cleared_line["source_product_color_id"])
        self.assertEqual(cleared_line["material"], "")
        self.assertEqual(cleared_line["specification"], "")
        self.assertEqual(cleared_line["fabric_width"], "")
        self.assertEqual(cleared_line["fabric_color_card"], "")
        self.assertEqual(cleared_line["fabric_color_code"], "")
        self.assertEqual(cleared_line["fabric_color_name"], "")
        self.assertEqual(cleared_line["unit"], "米")

    def test_auxiliary_maintenance_matching_uses_accessory_references(self):
        color = (
            AccessoryColor.objects.filter(
                status="ACTIVE",
                archived_at__isnull=True,
                product_name__status="ACTIVE",
                product_name__archived_at__isnull=True,
                product_name__specifications__status="ACTIVE",
                product_name__specifications__archived_at__isnull=True,
            )
            .exclude(color_name__in=["", "/"])
            .select_related("product_name")
            .distinct()
            .first()
        )
        product = color.product_name
        specification = product.specifications.filter(status="ACTIVE", archived_at__isnull=True).first()
        material = AccessoryMaterial.objects.filter(status="ACTIVE", archived_at__isnull=True).first()
        result = match_maintenance_configuration(
            {
                "sections": [
                    {
                        "section_type": "FABRIC",
                        "title": "辅料类",
                        "lines": [
                            {
                                "part_name": product.product_name,
                                "material": material.material_name,
                                "specification": specification.specification_value,
                                "raw_material": "",
                                "unit": "",
                                "raw_values": {"accessory_color": color.color_name},
                            }
                        ],
                    }
                ]
            }
        )
        self.assertEqual(result["matched_count"], 4)
        self.assertEqual(result["unmatched_count"], 0)
        self.assertEqual(result["unmatched_values"], [])

    def test_hardware_weight_and_powder_coating_fields_are_validated(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        source = RawMaterialInformation.objects.filter(
            item_type="MATERIAL",
            category__category_name="五金",
            flow_unit="KG",
        ).exclude(weight_kg="").first()
        self.assertIsNotNone(source)
        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "五金件")["lines"] = [
                    {
                        "part_name": source.material_name,
                        "source_raw_material_id": str(source.pk),
                        # 即使客户端清空隐藏的流转单位，也必须按来源原料的 KG 规则校验。
                        "hardware_flow_unit": "",
                        "unit_material_weight_kg": "",
                    }
                ]
        source_id = sections[1]["lines"][0].pop("source_raw_material_id")
        missing_source = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(missing_source.status_code, 400, missing_source.data)
        self.assertIn("选择数据", str(missing_source.data))
        sections[1]["lines"][0]["source_raw_material_id"] = source_id

        source.archived_at = timezone.now()
        source.save(update_fields=["archived_at", "updated_at"])
        archived_source = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(archived_source.status_code, 400, archived_source.data)
        self.assertIn("已停用、已归档", str(archived_source.data))
        source.archived_at = None
        source.save(update_fields=["archived_at", "updated_at"])

        missing_weight = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(missing_weight.status_code, 400, missing_weight.data)
        self.assertIn("按重量流转", str(missing_weight.data))

        sections[1]["lines"][0]["unit_material_weight_kg"] = source.weight_kg
        sections[1]["lines"][0]["powder_coating_color"] = "黑色磨砂"
        color_only = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(color_only.status_code, 200, color_only.data)
        saved_hardware = self.section_by_title(color_only.data["data"]["sections"], "五金件")["lines"][0]
        self.assertEqual(saved_hardware["powder_coating_color"], "黑色磨砂")
        self.assertEqual(saved_hardware["powder_coating_unit_area_sqm"], "")

        source.archived_at = timezone.now()
        source.save(update_fields=["archived_at", "updated_at"])
        sections = self.fixed_sections(color_only.data["data"]["sections"])
        hardware_line = self.section_by_title(sections, "五金件")["lines"][0]
        hardware_line["powder_coating_color"] = ""
        hardware_line["powder_coating_unit_area_sqm"] = "0.0065"
        area_only = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=color_only.headers["ETag"],
        )
        self.assertEqual(area_only.status_code, 200, area_only.data)
        saved_hardware = self.section_by_title(area_only.data["data"]["sections"], "五金件")["lines"][0]
        self.assertEqual(saved_hardware["powder_coating_color"], "")
        self.assertEqual(saved_hardware["powder_coating_unit_area_sqm"], "0.0065")

        sections = self.fixed_sections(area_only.data["data"]["sections"])
        hardware_line = self.section_by_title(sections, "五金件")["lines"][0]
        hardware_line["powder_coating_color"] = "不喷塑"
        hardware_line["powder_coating_unit_area_sqm"] = ""
        no_coating = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=area_only.headers["ETag"],
        )
        self.assertEqual(no_coating.status_code, 200, no_coating.data)
        self.assertEqual(no_coating.data["data"]["sections"][1]["lines"][0]["hardware_flow_unit"], "KG")

    def test_edit_then_verify_requires_etag_and_quantity_reason(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(),
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        bom_id = created.data["data"]["id"]
        pipe_source = self.complete_pipe_sources(bom_id)
        self.review_product_information(bom_id, created.headers["ETag"])
        detail = self.client.get(f"/api/v1/boms/{bom_id}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        etag = detail.headers["ETag"]
        sections = self.fixed_sections(detail.data["data"]["sections"])
        hardware_source = RawMaterialInformation.objects.filter(
            item_type="MATERIAL",
            category__category_name="五金",
            flow_unit="KG",
        ).exclude(weight_kg="").select_related("unit", "color").first()
        self.assertIsNotNone(hardware_source)
        sections[0]["lines"][0]["part_name"] = "修改后的交叉管"
        sections[0]["lines"][0].update(
            {
                "tube_shape": "圆管",
                "tube_diameter": "25",
                "tube_thickness": "1.2",
                "tube_coating_color": "黑色磨砂",
                "coating_unit_area_sqm": "0.18",
                "purchase_unit_weight_kg": "0.56",
                "tube_cut_piece_count": "10",
            }
        )
        sections[0]["lines"].append(
            {
                "part_name": "新增连接管",
                "material": "铁-硬管",
                "tube_shape": "圆管",
                "tube_diameter": "25",
                "tube_thickness": "1.2",
                "source_inventory_material_id": str(pipe_source.pk),
                "length_value": "581",
                "unit_consumption": "2",
                "original_material_length": "6000",
                "tube_cut_piece_count": "10",
                "tube_coating_color": "黑色磨砂",
                "coating_unit_area_sqm": "0.18",
                "purchase_unit_weight_kg": "0.56",
            }
        )
        sections[1]["title"] = "五金件"
        sections[1]["lines"][0].update(
            {
                "category": "伪造类别",
                "part_name": "伪造品名",
                "material": "伪造材质",
                "specification": "伪造规格",
                "raw_material": "伪造品号",
                "unit": "伪造单位",
                "source_raw_material_id": str(hardware_source.pk),
                "hardware_material_number": "伪造品号",
                "hardware_color": "伪造颜色",
                # 服务端必须忽略客户端伪造值并保存来源原料的真实流转单位。
                "hardware_flow_unit": "PCS",
                "unit_material_weight_kg": hardware_source.weight_kg,
                "powder_coating_color": "黑色磨砂",
                "powder_coating_unit_area_sqm": "0.0065",
                "unit_consumption": "4",
                "notes": "五金测试行",
            }
        )
        sections[0]["lines"].append(
            {
                "part_name": "舍入边界管",
                "source_inventory_material_id": str(pipe_source.pk),
                "length_value": "2005",
                "unit_consumption": "1",
                "tube_cut_piece_count": "5",
            }
        )

        from apps.reference_data.models import ProductionUnit, ProductCategory
        meter_unit, _ = ProductionUnit.objects.get_or_create(unit_name="米", defaults={"unit_code": "S4LEGMETER"})
        fabric_source = RawMaterialInformation.objects.create(material_number="S4EDITFABRIC", material_name="测试面料", item_type="MATERIAL",
            category=ProductCategory.objects.get(category_name="面料", item_type="MATERIAL"), unit=meter_unit)
        for section in sections:
            if section["title"] == "面料类":
                for line in section["lines"]:
                    line.update({"source_inventory_material_id": str(fabric_source.pk), "unit_consumption": "0.5"})

        missing_etag = self.client.patch(
            f"/api/v1/boms/{bom_id}/",
            {"sections": sections},
            format="json",
        )
        self.assertEqual(missing_etag.status_code, 428, missing_etag.data)

        updated = self.save_confirmed_material_sections(bom_id, sections, etag)
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][0]["part_name"], "修改后的交叉管")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["part_name"], "新增连接管")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][0]["specification"], "Φ25 × 1.2 mm")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][0]["tube_coating_color"], "黑色磨砂")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["purchase_unit_weight_kg"], "0.56")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["unit"], "根")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["total_consumption"], "20")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["tube_cut_piece_count"], "10")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["tube_cut_length_m"], "5.86")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["branch_count"], "3")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["long_material_unit_weight_kg"], "2.719")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["tube_unit_weight_kg"], "0.5392")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][1]["long_material_total_weight_kg"], "8.1571")
        self.assertEqual(
            updated.data["data"]["sections"][0]["lines"][1]["raw_values"]["coating_unit_area_sqm"],
            "0.18",
        )
        self.assertEqual(
            updated.data["data"]["sections"][0]["lines"][1]["raw_values"]["tube_cut_piece_count"],
            "10",
        )
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][2]["tube_cut_length_m"], "10.08")
        self.assertEqual(updated.data["data"]["sections"][0]["lines"][2]["branch_count"], "3")
        hardware_line = updated.data["data"]["sections"][1]["lines"][0]
        self.assertEqual(hardware_line["category"], "五金")
        self.assertEqual(hardware_line["part_name"], hardware_source.material_name)
        self.assertEqual(hardware_line["material"], hardware_source.material)
        self.assertEqual(hardware_line["specification"], hardware_source.specification)
        self.assertEqual(hardware_line["raw_material"], hardware_source.material_number)
        self.assertEqual(hardware_line["unit"], hardware_source.unit.unit_name)
        self.assertEqual(hardware_line["source_raw_material_id"], str(hardware_source.pk))
        self.assertEqual(hardware_line["hardware_material_number"], hardware_source.material_number)
        self.assertEqual(hardware_line["hardware_color"], hardware_source.color.color_name if hardware_source.color_id else hardware_source.color_snapshot)
        self.assertEqual(hardware_line["hardware_flow_unit"], "KG")
        self.assertEqual(hardware_line["unit_material_weight_kg"], hardware_source.weight_kg)
        self.assertEqual(hardware_line["powder_coating_color"], "黑色磨砂")
        self.assertEqual(hardware_line["powder_coating_unit_area_sqm"], "0.0065")
        self.assertEqual(hardware_line["raw_values"]["source_raw_material_id"], str(hardware_source.pk))
        self.assertEqual(updated.data["data"]["line_count"], 5)
        self.assertEqual(updated.data["data"]["status"], "PENDING_CONFIRMATION")
        updated_etag = updated.headers["ETag"]

        reloaded = self.client.get(f"/api/v1/boms/{bom_id}/")
        self.assertEqual(reloaded.status_code, 200, reloaded.data)
        reloaded_hardware_line = reloaded.data["data"]["sections"][1]["lines"][0]
        for field in [
            "source_raw_material_id",
            "hardware_material_number",
            "hardware_color",
            "hardware_flow_unit",
            "unit_material_weight_kg",
            "powder_coating_color",
            "powder_coating_unit_area_sqm",
        ]:
            self.assertEqual(reloaded_hardware_line[field], hardware_line[field])

        reloaded_hardware_line["source_raw_material_id"] = None
        reloaded_hardware_line["part_name"] = "尝试解除来源后篡改"
        cleared_source = self.client.patch(
            f"/api/v1/boms/{bom_id}/",
            {"sections": reloaded.data["data"]["sections"]},
            format="json",
            HTTP_IF_MATCH=updated_etag,
        )
        self.assertEqual(cleared_source.status_code, 400, cleared_source.data)
        self.assertIn("不能清空数据来源", str(cleared_source.data))

        stale = self.client.patch(
            f"/api/v1/boms/{bom_id}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=etag,
        )
        self.assertEqual(stale.status_code, 412, stale.data)

        missing_reason = self.client.post(
            f"/api/v1/boms/{bom_id}/verify/",
            {},
            format="json",
            HTTP_IF_MATCH=updated_etag,
        )
        self.assertEqual(missing_reason.status_code, 400, missing_reason.data)
        self.assertIn("quantity_confirmation_reason", missing_reason.data["error"]["field_errors"])
        quantity_error = str(missing_reason.data["error"]["field_errors"]["quantity_confirmation_reason"])
        self.assertIn("本 BOM 保存的订单数量为 10", quantity_error)
        self.assertIn("Excel 原表数量为 11", quantity_error)
        self.assertIn("生产批次合计数量为 12", quantity_error)
        self.assertIn("相差 +1", quantity_error)
        self.assertIn("相差 +2", quantity_error)

        verified = self.client.post(
            f"/api/v1/boms/{bom_id}/verify/",
            {"quantity_confirmation_reason": "按当前生产任务订单数量执行。"},
            format="json",
            HTTP_IF_MATCH=updated_etag,
        )
        self.assertEqual(verified.status_code, 200, verified.data)
        self.assertEqual(verified.data["data"]["status"], "VERIFIED")
        self.assertTrue(AuditLog.objects.filter(object_type="BOM", object_id=bom_id, action="UPDATE").exists())
        self.assertTrue(AuditLog.objects.filter(object_type="BOM", object_id=bom_id, action="VERIFY").exists())

        submitted_sections = self.fixed_sections(verified.data["data"]["sections"])
        submitted_sections[0]["lines"][0]["notes"] = "提交后按实际生产情况调整"
        reopened = self.client.patch(
            f"/api/v1/boms/{bom_id}/",
            {"sections": submitted_sections},
            format="json",
            HTTP_IF_MATCH=verified.headers["ETag"],
        )
        self.assertEqual(reopened.status_code, 200, reopened.data)
        self.assertEqual(reopened.data["data"]["status"], "PENDING_CONFIRMATION")
        self.assertEqual(reopened.data["data"]["sections"][0]["lines"][0]["notes"], "提交后按实际生产情况调整")
        self.assertIsNone(reopened.data["data"]["confirmed_by_name"])
        self.assertIsNone(reopened.data["data"]["confirmed_at"])
        self.assertEqual(reopened.data["data"]["confirmation_reason"], "")

        self.user.user_permissions.set(
            Permission.objects.filter(codename="view_bom", content_type__app_label="boms")
        )
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(self.user)
        forbidden = self.client.patch(
            f"/api/v1/boms/{bom_id}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=reopened.headers["ETag"],
        )
        self.assertEqual(forbidden.status_code, 403, forbidden.data)

    def test_verify_allows_any_production_task_status(self):
        formerly_blocked_statuses = (
            ProductionTaskStatus.DRAFT,
            ProductionTaskStatus.PENDING_VALIDATION,
            ProductionTaskStatus.CANCELLED,
        )
        for task_status in formerly_blocked_statuses:
            with self.subTest(task_status=task_status):
                self.task.status = ProductionTaskStatus.READY
                self.task.save(update_fields=["status", "updated_at"])
                created = self.client.post(
                    "/api/v1/boms/",
                    {
                        "production_task_product_id": str(self.product_line.pk),
                        "creation_method": "IMPORT",
                        "workbook": self.workbook_bytes(),
                    },
                    format="multipart",
                )
                self.assertEqual(created.status_code, 201, created.data)
                self.complete_pipe_sources(created.data["data"]["id"])
                reviewed = self.review_product_information(created.data["data"]["id"], created.headers["ETag"])

                self.task.status = task_status
                self.task.submitted_at = None
                self.task.save(update_fields=["status", "submitted_at", "updated_at"])
                verified = self.client.post(
                    f"/api/v1/boms/{created.data['data']['id']}/verify/",
                    {"quantity_confirmation_reason": "测试生产任务任意状态均可确认 BOM。"},
                    format="json",
                    HTTP_IF_MATCH=reviewed.headers["ETag"],
                )
                self.assertEqual(verified.status_code, 200, verified.data)
                self.assertEqual(verified.data["data"]["status"], "VERIFIED")

    def test_editor_can_sync_changed_order_quantity(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "IMPORT",
                "workbook": self.workbook_bytes(),
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.data)
        bom_id = created.data["data"]["id"]

        self.product_line.quantity = Decimal("14")
        self.product_line.save(update_fields=["quantity", "updated_at"])
        detail = self.client.get(f"/api/v1/boms/{bom_id}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertTrue(detail.data["data"]["is_order_quantity_outdated"])
        self.assertEqual(Decimal(detail.data["data"]["order_quantity"]), Decimal("10"))
        self.assertEqual(Decimal(detail.data["data"]["current_order_quantity"]), Decimal("14"))

        missing_etag = self.client.post(f"/api/v1/boms/{bom_id}/sync-order-quantity/", {}, format="json")
        self.assertEqual(missing_etag.status_code, 428, missing_etag.data)

        synced = self.client.post(
            f"/api/v1/boms/{bom_id}/sync-order-quantity/",
            {},
            format="json",
            HTTP_IF_MATCH=detail.headers["ETag"],
        )
        self.assertEqual(synced.status_code, 200, synced.data)
        self.assertEqual(synced.data["data"]["id"], bom_id)
        self.assertEqual(Decimal(synced.data["data"]["order_quantity"]), Decimal("14"))
        self.assertEqual(Decimal(synced.data["data"]["current_order_quantity"]), Decimal("14"))
        self.assertFalse(synced.data["data"]["is_order_quantity_outdated"])
        self.assertEqual(synced.data["data"]["sections"][0]["lines"][0]["total_consumption"], "84")
        self.assertTrue(
            AuditLog.objects.filter(
                object_type="BOM",
                object_id=bom_id,
                action="SYNC_FROM_PRODUCTION_TASK",
            ).exists()
        )

    def test_temporary_lock_file_is_rejected(self):
        response = self.client.post(
            "/api/v1/boms/preview-import/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "workbook": self.workbook_bytes(file_name="~$bom.xlsx"),
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("临时锁定文件", str(response.data))

    def test_list_requires_bom_permission(self):
        self.user.user_permissions.clear()
        response = self.client.get("/api/v1/boms/")
        self.assertEqual(response.status_code, 403, response.data)

    def test_view_only_user_can_load_list_options_without_task_choices(self):
        self.user.user_permissions.set(
            Permission.objects.filter(codename="view_bom", content_type__app_label="boms")
        )
        response = self.client.get("/api/v1/boms/options/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["production_task_products"], [])
        self.assertIn("product_categories", response.data["data"])

    def test_packaging_and_cutting_sources_are_authoritative_and_round_trip(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        packaging_source = (
            RawMaterialInformation.objects.filter(
                item_type="MATERIAL",
                category__category_name="包材",
                category__status="ACTIVE",
                category__archived_at__isnull=True,
                archived_at__isnull=True,
            )
            .select_related("category", "unit", "color")
            .first()
        )
        cutting_source = (
            RawMaterialInformation.objects.filter(
                item_type="SEMI_FINISHED",
                category__status="ACTIVE",
                category__archived_at__isnull=True,
                archived_at__isnull=True,
            )
            .select_related("category", "unit", "color")
            .first()
        )
        self.assertIsNotNone(packaging_source)
        self.assertIsNotNone(cutting_source)

        sections = self.fixed_sections(created.data["data"]["sections"])
        self.section_by_title(sections, "包材类")["lines"] = [
            {
                "part_name": "伪造包材品名",
                "material": "伪造材质",
                "specification": "伪造规格",
                "unit": "伪造单位",
                "source_packaging_material_id": str(packaging_source.pk),
                "packaging_material_number": "伪造品号",
                "packaging_color": "伪造颜色",
                "units_per_box": "",
                "unit_consumption": "0.5",
            }
        ]
        self.section_by_title(sections, "裁剪明细")["lines"] = [
            {
                "part_name": "伪造裁剪名称",
                "material": "BOM 自填半成品材质",
                "specification": "BOM 自填半成品规格",
                "unit": "BOM 自填单位",
                "source_cutting_material_id": str(cutting_source.pk),
                "cutting_material_number": "伪造半成品号",
                "cutting_color": "自定义红色",
                "cut_piece_quantity": "2",
                "notes": "裁剪备注",
            }
        ]

        missing_box_quantity = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(missing_box_quantity.status_code, 400, missing_box_quantity.data)
        self.assertIn("多少一箱", str(missing_box_quantity.data))

        packaging_line = self.section_by_title(sections, "包材类")["lines"][0]
        packaging_line["units_per_box"] = "12"
        cutting_line = self.section_by_title(sections, "裁剪明细")["lines"][0]
        cutting_line["cut_piece_quantity"] = "-1"
        invalid_cut_piece_quantity = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(invalid_cut_piece_quantity.status_code, 400, invalid_cut_piece_quantity.data)
        self.assertIn("裁片量不能小于 0", str(invalid_cut_piece_quantity.data))

        cutting_line["cut_piece_quantity"] = "2"
        saved = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        saved_packaging = self.section_by_title(saved.data["data"]["sections"], "包材类")["lines"][0]
        self.assertEqual(saved_packaging["part_name"], packaging_source.material_name)
        self.assertEqual(saved_packaging["material"], packaging_source.material)
        self.assertEqual(saved_packaging["specification"], packaging_source.specification)
        self.assertEqual(saved_packaging["unit"], packaging_source.unit.unit_name)
        self.assertEqual(saved_packaging["packaging_material_number"], packaging_source.material_number)
        self.assertEqual(
            saved_packaging["packaging_color"],
            packaging_source.color.color_name if packaging_source.color_id else packaging_source.color_snapshot,
        )
        self.assertEqual(saved_packaging["units_per_box"], "12")
        self.assertEqual(saved_packaging["unit_consumption"], "0.5")

        saved_cutting = self.section_by_title(saved.data["data"]["sections"], "裁剪明细")["lines"][0]
        self.assertEqual(saved_cutting["part_name"], cutting_source.material_name)
        self.assertEqual(saved_cutting["material"], "BOM 自填半成品材质")
        self.assertEqual(saved_cutting["specification"], "BOM 自填半成品规格")
        self.assertEqual(saved_cutting["unit"], "BOM 自填单位")
        self.assertEqual(saved_cutting["cutting_material_number"], cutting_source.material_number)
        self.assertEqual(saved_cutting["cutting_item_type"], cutting_source.item_type)
        self.assertEqual(saved_cutting["cutting_item_type_display"], cutting_source.get_item_type_display())
        self.assertEqual(saved_cutting["cutting_category"], cutting_source.category.category_name)
        self.assertEqual(
            saved_cutting["cutting_color"],
            cutting_source.color.color_name if cutting_source.color_id else cutting_source.color_snapshot,
        )
        self.assertEqual(saved_cutting["cut_piece_quantity"], "2")
        self.assertEqual(saved_cutting["notes"], "裁剪备注")
        self.assertEqual(saved_cutting["images"], [])

        persisted_sections = self.fixed_sections(saved.data["data"]["sections"])
        edited_cutting = self.section_by_title(persisted_sections, "裁剪明细")["lines"][0]
        edited_cutting["material"] = "编辑后的半成品材质"
        edited_cutting["specification"] = "编辑后的半成品规格"
        edited_cutting["unit"] = "编辑单位"
        edited = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": persisted_sections},
            format="json",
            HTTP_IF_MATCH=saved.headers["ETag"],
        )
        self.assertEqual(edited.status_code, 200, edited.data)
        edited_cutting = self.section_by_title(edited.data["data"]["sections"], "裁剪明细")["lines"][0]
        self.assertEqual(edited_cutting["material"], "编辑后的半成品材质")
        self.assertEqual(edited_cutting["specification"], "编辑后的半成品规格")
        self.assertEqual(edited_cutting["unit"], "编辑单位")
        saved = edited
        persisted_sections = self.fixed_sections(saved.data["data"]["sections"])
        cleared_packaging = deepcopy(persisted_sections)
        self.section_by_title(cleared_packaging, "包材类")["lines"][0]["source_packaging_material_id"] = None
        clear_packaging_response = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": cleared_packaging},
            format="json",
            HTTP_IF_MATCH=saved.headers["ETag"],
        )
        self.assertEqual(clear_packaging_response.status_code, 400, clear_packaging_response.data)
        self.assertIn("不能清空数据来源", str(clear_packaging_response.data))

        cleared_cutting = deepcopy(persisted_sections)
        self.section_by_title(cleared_cutting, "裁剪明细")["lines"][0]["source_cutting_material_id"] = None
        clear_cutting_response = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": cleared_cutting},
            format="json",
            HTTP_IF_MATCH=saved.headers["ETag"],
        )
        self.assertEqual(clear_cutting_response.status_code, 400, clear_cutting_response.data)
        self.assertIn("不能清空数据来源", str(clear_cutting_response.data))

        packaging_source.archived_at = timezone.now()
        packaging_source.save(update_fields=["archived_at", "updated_at"])
        cutting_source.archived_at = timezone.now()
        cutting_source.save(update_fields=["archived_at", "updated_at"])
        historical_sources = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": persisted_sections},
            format="json",
            HTTP_IF_MATCH=saved.headers["ETag"],
        )
        self.assertEqual(historical_sources.status_code, 200, historical_sources.data)
        historical_packaging = self.section_by_title(historical_sources.data["data"]["sections"], "包材类")["lines"][0]
        historical_cutting = self.section_by_title(historical_sources.data["data"]["sections"], "裁剪明细")["lines"][0]
        self.assertEqual(historical_packaging["part_name"], saved_packaging["part_name"])
        self.assertEqual(historical_packaging["packaging_material_number"], saved_packaging["packaging_material_number"])
        self.assertEqual(historical_cutting["part_name"], saved_cutting["part_name"])
        self.assertEqual(historical_cutting["cutting_material_number"], saved_cutting["cutting_material_number"])

    def test_packaging_requires_positive_integer_box_quantity_and_positive_consumption(self):
        created = self.client.post(
            "/api/v1/boms/",
            {
                "production_task_product_id": str(self.product_line.pk),
                "creation_method": "BLANK",
            },
            format="json",
        )
        source = RawMaterialInformation.objects.filter(
            item_type="MATERIAL",
            category__category_name="包材",
            category__status="ACTIVE",
            category__archived_at__isnull=True,
            archived_at__isnull=True,
        ).first()
        sections = self.fixed_sections(created.data["data"]["sections"])
        packaging_section = self.section_by_title(sections, "包材类")
        packaging_section["lines"] = [
            {
                "part_name": source.material_name,
                "source_packaging_material_id": str(source.pk),
                "units_per_box": "1.5",
                "unit_consumption": "0",
            }
        ]
        invalid_box_quantity = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(invalid_box_quantity.status_code, 400, invalid_box_quantity.data)
        self.assertIn("正整数", str(invalid_box_quantity.data))

        packaging_section["lines"][0]["units_per_box"] = "10"
        invalid_consumption = self.client.patch(
            f"/api/v1/boms/{created.data['data']['id']}/",
            {"sections": sections},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(invalid_consumption.status_code, 400, invalid_consumption.data)
        self.assertIn("单个用量必须大于 0", str(invalid_consumption.data))

    def test_archive_removes_bom_from_list_and_writes_audit(self):
        created = self.client.post(
            "/api/v1/boms/",
            {"production_task_product_id": str(self.product_line.pk), "creation_method": "BLANK"},
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        bom_id = created.data["data"]["id"]

        archived = self.client.post(
            f"/api/v1/boms/{bom_id}/archive/",
            {},
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )

        self.assertEqual(archived.status_code, 200, archived.data)
        self.assertTrue(Bom.objects.get(pk=bom_id).archived_at)
        self.assertFalse(any(item["id"] == bom_id for item in self.client.get("/api/v1/boms/").data["data"]))
        self.assertTrue(AuditLog.objects.filter(object_type="BOM", object_id=bom_id, action="ARCHIVE").exists())

    def test_create_auxiliary_reference_is_atomic_and_immediately_selectable(self):
        self.user.user_permissions.add(
            *Permission.objects.filter(
                content_type__app_label="reference_data",
                codename__in=[
                    "add_accessoryproductname",
                    "add_accessoryspecification",
                    "add_accessorycolor",
                    "add_accessorymaterial",
                ],
            )
        )
        response = self.client.post(
            "/api/v1/boms/create-auxiliary-reference/",
            {
                "product_name": "QA即时新增辅料",
                "unit_name": "PCS",
                "specification": "QA规格",
                "color": "QA颜色",
                "material": "QA材质",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["data"]["specifications"][0]["specification_value"], "QA规格")
        self.assertEqual(response.data["data"]["colors"][0]["color_name"], "QA颜色")
        self.assertTrue(AccessoryMaterial.objects.filter(material_name="QA材质").exists())
        options = self.client.get("/api/v1/boms/editor-reference-options/")
        self.assertTrue(any(item["id"] == response.data["data"]["id"] for item in options.data["data"]["auxiliary_products"]))
