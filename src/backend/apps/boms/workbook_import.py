from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class BomWorkbookError(ValueError):
    pass


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _value_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def _decimal_value(value: str) -> Decimal | None:
    normalized = value.replace(",", "").strip()
    if not normalized:
        return None
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


class _XlsxGrid:
    def __init__(self, path: Path):
        try:
            self.value_book = load_workbook(path, data_only=True, read_only=False)
            self.formula_book = load_workbook(path, data_only=False, read_only=False)
        except Exception as exc:
            raise BomWorkbookError("XLSX 文件无法打开，请确认文件未损坏且未加密。") from exc
        if len(self.value_book.worksheets) < 2:
            self.close()
            raise BomWorkbookError("BOM 文件必须包含第二个 Sheet。")
        self.sheet_names = [sheet.title for sheet in self.value_book.worksheets]
        self.sheet_count = len(self.sheet_names)
        self.select_sheet(1)

    def select_sheet(self, index: int):
        self.sheet_index = index
        self.value_sheet = self.value_book.worksheets[index]
        self.formula_sheet = self.formula_book.worksheets[index]
        self.sheet_name = self.value_sheet.title
        self.max_row = self.value_sheet.max_row
        self.max_column = self.value_sheet.max_column

    def value(self, row: int, column: int) -> str:
        return _value_text(self.value_sheet.cell(row=row, column=column).value)

    def formula(self, row: int, column: int) -> str:
        value = self.formula_sheet.cell(row=row, column=column).value
        return value if isinstance(value, str) and value.startswith("=") else ""

    def close(self):
        self.value_book.close()
        self.formula_book.close()


class _XlsGrid:
    def __init__(self, path: Path):
        try:
            import xlrd

            self.xlrd = xlrd
            self.book = xlrd.open_workbook(path, on_demand=True)
        except ModuleNotFoundError as exc:
            raise BomWorkbookError("服务器缺少 XLS 解析依赖，请联系管理员安装 xlrd。") from exc
        except Exception as exc:
            raise BomWorkbookError("XLS 文件无法打开，请确认文件未损坏且未加密。") from exc
        if self.book.nsheets < 2:
            self.close()
            raise BomWorkbookError("BOM 文件必须包含第二个 Sheet。")
        self.sheet_names = self.book.sheet_names()
        self.sheet_count = len(self.sheet_names)
        self.select_sheet(1)

    def select_sheet(self, index: int):
        self.sheet_index = index
        self.sheet = self.book.sheet_by_index(index)
        self.sheet_name = self.sheet.name
        self.max_row = self.sheet.nrows
        self.max_column = self.sheet.ncols

    def value(self, row: int, column: int) -> str:
        if row < 1 or column < 1 or row > self.max_row or column > self.max_column:
            return ""
        cell = self.sheet.cell(row - 1, column - 1)
        if cell.ctype == self.xlrd.XL_CELL_DATE:
            value = self.xlrd.xldate_as_datetime(cell.value, self.book.datemode)
            return value.date().isoformat()
        return _value_text(cell.value)

    def formula(self, row: int, column: int) -> str:
        # xlrd 只提供旧式 XLS 公式的缓存结果；原值仍会保存到 raw_values。
        return ""

    def close(self):
        self.book.release_resources()


def _open_grid(path: Path, original_name: str):
    suffix = Path(original_name).suffix.lower()
    if suffix == ".xlsx":
        return _XlsxGrid(path)
    if suffix == ".xls":
        return _XlsGrid(path)
    raise BomWorkbookError("BOM 文件仅支持 XLS 或 XLSX 格式。")


def _raw_row(grid, row: int) -> tuple[dict[str, str], dict[str, str]]:
    values: dict[str, str] = {}
    formulas: dict[str, str] = {}
    for column in range(1, grid.max_column + 1):
        value = grid.value(row, column)
        formula = grid.formula(row, column)
        letter = _column_letter(column)
        if value:
            values[letter] = value
        if formula:
            formulas[letter] = formula
    return values, formulas


def _metadata(grid) -> dict[str, str]:
    target_labels = {
        "产品编号": "product_code",
        "产品名称": "product_name",
        "产品规格": "specification",
        "产品数量": "product_quantity",
        "订单数量": "product_quantity",
        "生产单号": "production_task_number",
        "开单日期": "order_date",
        "交货日期": "delivery_date",
    }
    result: dict[str, str] = {}
    for row in range(1, min(8, grid.max_row) + 1):
        for column in range(1, min(12, grid.max_column) + 1):
            label = grid.value(row, column).replace("\n", "").replace(" ", "").strip("：:")
            key = target_labels.get(label)
            if not key:
                continue
            for value_column in range(column + 1, min(column + 4, grid.max_column) + 1):
                value = grid.value(row, value_column).strip()
                if value:
                    result[key] = value
                    break
    missing = [label for label, key in target_labels.items() if key not in result and key in {"product_code", "product_name", "production_task_number", "product_quantity"}]
    if missing:
        raise BomWorkbookError(f"第二个 Sheet 缺少必要字段：{'、'.join(missing)}。")
    return result


def _section_type(title: str) -> str | None:
    if "型材" in title:
        return "PROFILE"
    if any(keyword in title for keyword in ("面辅料", "面料", "布料")):
        return "FABRIC"
    if "配件" in title:
        return "ACCESSORY"
    return None


def _section_starts(grid) -> list[tuple[int, str, str]]:
    result = []
    for row in range(1, grid.max_row + 1):
        title = grid.value(row, 1)
        section_type = _section_type(title)
        if title and "清单" in title and section_type:
            result.append((row, title, section_type))
    return result


def _select_bom_sheet(grid) -> tuple[dict[str, str], list[tuple[int, str, str]]]:
    candidate_indexes = [1]
    candidate_indexes.extend(
        index
        for index, name in enumerate(grid.sheet_names)
        if index != 1 and "BOM" in name.upper()
    )
    candidate_indexes.extend(index for index in range(grid.sheet_count) if index not in candidate_indexes)

    second_sheet_error = None
    for index in candidate_indexes:
        grid.select_sheet(index)
        section_starts = _section_starts(grid)
        if not section_starts:
            continue
        try:
            metadata = _metadata(grid)
        except BomWorkbookError as exc:
            if index == 1:
                second_sheet_error = exc
            continue
        return metadata, section_starts

    grid.select_sheet(1)
    if second_sheet_error:
        raise second_sheet_error
    raise BomWorkbookError(
        f"第二个 Sheet“{grid.sheet_name}”不是可识别的 BOM，其他 Sheet 也未找到型材、配件或面辅料清单。"
    )


def _normalized_header(value: str) -> str:
    return value.replace("\n", "").replace(" ", "").replace("（", "(").replace("）", ")").lower()


def _header_column(grid, header_row: int, aliases: tuple[str, ...], default: int) -> int:
    normalized_aliases = tuple(_normalized_header(alias) for alias in aliases)
    for column in range(1, grid.max_column + 1):
        header = _normalized_header(grid.value(header_row, column))
        if header and any(alias in header for alias in normalized_aliases):
            return column
    return default


def _section_columns(grid, header_row: int, section_type: str) -> dict[str, int]:
    columns = {
        "source_serial": _header_column(grid, header_row, ("序号",), 1),
        "category": _header_column(grid, header_row, ("类别",), 2),
        "material": _header_column(grid, header_row, ("材质",), 3),
        "unit": _header_column(grid, header_row, ("单位",), 11),
        "notes": _header_column(grid, header_row, ("备注",), 17),
    }
    if section_type == "PROFILE":
        columns.update(
            part_name=_header_column(grid, header_row, ("配件名称",), 4),
            specification=_header_column(grid, header_row, ("规格/厚度", "规格厚度"), 5),
            length_value=_header_column(grid, header_row, ("切割长度(mm)", "切割长度（mm）", "长度(mm)", "长度"), 8),
            tube_cut_piece_count=_header_column(grid, header_row, ("每根切割支数", "切割支数"), 0),
            unit_consumption=_header_column(grid, header_row, ("单耗",), 9),
            total_consumption=_header_column(grid, header_row, ("总耗",), 10),
            hardness=_header_column(grid, header_row, ("材料硬度", "硬度"), 12),
            original_material_length=_header_column(grid, header_row, ("原材料长度",), 13),
            branch_count=_header_column(grid, header_row, ("采购根数", "支数"), 17),
        )
    elif section_type == "ACCESSORY":
        columns.update(
            part_name=_header_column(grid, header_row, ("配件名称",), 4),
            specification=_header_column(grid, header_row, ("规格",), 5),
            unit_consumption=_header_column(grid, header_row, ("单耗",), 9),
            source_order_quantity=_header_column(grid, header_row, ("订单量", "订单数量"), 10),
            budget_quantity=_header_column(grid, header_row, ("预算量",), 12),
            inventory_quantity=_header_column(grid, header_row, ("库存",), 13),
            weight_grams=_header_column(grid, header_row, ("重量/g", "重量(g)"), 14),
        )
    else:
        columns.update(
            part_name=_header_column(grid, header_row, ("裁片名称",), 4),
            raw_material=_header_column(grid, header_row, ("原材料",), 8),
            cut_piece_quantity=_header_column(grid, header_row, ("裁片量",), 7),
            purchase_consumption=_header_column(grid, header_row, ("单采购量", "单套用量"), 12),
            source_order_quantity=_header_column(grid, header_row, ("订单量", "订单数量"), 13),
            total_purchase_quantity=_header_column(grid, header_row, ("总采购量",), 14),
        )
    return columns


def _profile_specification(grid, row: int, specification_column: int) -> str:
    shape = grid.value(row, specification_column)
    diameter = grid.value(row, specification_column + 1)
    thickness = grid.value(row, specification_column + 2)
    parts = [item for item in (shape, diameter, thickness) if item]
    if shape in {"Φ", "Ф", "φ"} and diameter:
        base = f"Φ{diameter}"
        return f"{base} × {thickness} mm" if thickness else base
    return " × ".join(parts)


def _line_payload(
    grid,
    row: int,
    section_type: str,
    state: dict[str, str],
    columns: dict[str, int],
) -> dict[str, Any]:
    source_serial = grid.value(row, columns["source_serial"])
    category_value = grid.value(row, columns["category"])
    material_value = grid.value(row, columns["material"])
    part_name = grid.value(row, columns["part_name"])
    raw_material_value = grid.value(row, columns.get("raw_material", 8))

    if category_value:
        state["category"] = category_value
    if material_value:
        state["material"] = material_value
    if raw_material_value:
        state["raw_material"] = raw_material_value

    material = material_value or state.get("material", "")
    if section_type == "PROFILE" and part_name == "塑粉":
        material = ""

    raw_values, source_formulas = _raw_row(grid, row)
    common = {
        "source_row_number": row,
        "source_serial": source_serial,
        "category": category_value or state.get("category", ""),
        "material": material,
        "part_name": part_name,
        "specification": "",
        "raw_material": "",
        "length_value": "",
        "unit_consumption": "",
        "total_consumption": "",
        "unit": grid.value(row, columns["unit"]),
        "hardness": "",
        "original_material_length": "",
        "branch_count": "",
        "budget_quantity": "",
        "inventory_quantity": "",
        "weight_grams": "",
        "cut_piece_quantity": "",
        "purchase_consumption": "",
        "source_order_quantity": "",
        "total_purchase_quantity": "",
        "notes": grid.value(row, columns["notes"]),
        "raw_values": raw_values,
        "source_formulas": source_formulas,
    }
    if section_type == "PROFILE":
        raw_values.update(
            {
                "tube_shape": grid.value(row, columns["specification"]),
                "tube_diameter": grid.value(row, columns["specification"] + 1),
                "tube_thickness": grid.value(row, columns["specification"] + 2),
            }
        )
        if columns.get("tube_cut_piece_count"):
            raw_values["tube_cut_piece_count"] = grid.value(row, columns["tube_cut_piece_count"])
        common.update(
            specification=_profile_specification(grid, row, columns["specification"]),
            length_value=grid.value(row, columns["length_value"]),
            unit_consumption=grid.value(row, columns["unit_consumption"]),
            total_consumption=grid.value(row, columns["total_consumption"]),
            hardness=grid.value(row, columns["hardness"]),
            original_material_length=grid.value(row, columns["original_material_length"]),
            branch_count=grid.value(row, columns["branch_count"]),
        )
    elif section_type == "ACCESSORY":
        common.update(
            specification=grid.value(row, columns["specification"]),
            unit_consumption=grid.value(row, columns["unit_consumption"]),
            source_order_quantity=grid.value(row, columns["source_order_quantity"]),
            budget_quantity=grid.value(row, columns["budget_quantity"]),
            inventory_quantity=grid.value(row, columns["inventory_quantity"]),
            weight_grams=grid.value(row, columns["weight_grams"]),
        )
    else:
        common.update(
            raw_material=raw_material_value or state.get("raw_material", ""),
            cut_piece_quantity=grid.value(row, columns["cut_piece_quantity"]),
            purchase_consumption=grid.value(row, columns["purchase_consumption"]),
            source_order_quantity=grid.value(row, columns["source_order_quantity"]),
            total_purchase_quantity=grid.value(row, columns["total_purchase_quantity"]),
        )
        if not common["unit"]:
            common["unit"] = state.get("unit", "")
        else:
            state["unit"] = common["unit"]
    return common


def _calculation_row(grid, row: int) -> dict[str, Any]:
    values, formulas = _raw_row(grid, row)
    return {"source_row_number": row, "raw_values": values, "source_formulas": formulas}


def parse_bom_workbook(path: str | Path, original_name: str) -> dict[str, Any]:
    safe_name = Path(original_name).name
    if safe_name.startswith("~$"):
        raise BomWorkbookError("“~$”开头的是 Excel 临时锁定文件，不能作为 BOM 数据源。")
    source_path = Path(path)
    if not source_path.is_file():
        raise BomWorkbookError("BOM 文件不存在或已失效。")

    grid = _open_grid(source_path, safe_name)
    try:
        metadata, section_starts = _select_bom_sheet(grid)

        sections = []
        all_formula_columns: set[str] = set()
        for index, (title_row, title, section_type) in enumerate(section_starts):
            boundary = section_starts[index + 1][0] - 1 if index + 1 < len(section_starts) else grid.max_row
            header_row = None
            for candidate in range(title_row + 1, min(title_row + 3, boundary) + 1):
                if grid.value(candidate, 1) == "序号":
                    header_row = candidate
                    break
            if header_row is None:
                raise BomWorkbookError(f"“{title}”后未找到“序号”表头。")

            headers = {}
            for column in range(1, grid.max_column + 1):
                value = grid.value(header_row, column)
                if value:
                    headers[_column_letter(column)] = value

            lines = []
            calculation_rows = []
            calculation_started = False
            state: dict[str, str] = {}
            columns = _section_columns(grid, header_row, section_type)
            for row in range(header_row + 1, boundary + 1):
                first = grid.value(row, 1)
                second = grid.value(row, 2)
                name = grid.value(row, columns["part_name"])
                if section_type == "ACCESSORY" and name == "型号":
                    calculation_started = True
                if section_type == "FABRIC" and (first in {"总共", "合计"} or second in {"总共", "合计"}):
                    calculation_started = True
                if calculation_started:
                    if any(grid.value(row, column) for column in range(1, grid.max_column + 1)):
                        calculation_rows.append(_calculation_row(grid, row))
                    continue
                if not name:
                    continue
                line = _line_payload(grid, row, section_type, state, columns)
                lines.append(line)
                all_formula_columns.update(line["source_formulas"].keys())

            if not lines:
                raise BomWorkbookError(f"“{title}”未识别到 BOM 明细。")
            sections.append(
                {
                    "section_type": section_type,
                    "title": title,
                    "source_title_row": title_row,
                    "source_header_row": header_row,
                    "headers": headers,
                    "calculation_rows": calculation_rows,
                    "lines": lines,
                    "line_count": len(lines),
                }
            )

        return {
            "original_name": safe_name,
            "sheet_name": grid.sheet_name,
            "sheet_number": grid.sheet_index + 1,
            "metadata": metadata,
            "source_quantity": _decimal_value(metadata.get("product_quantity", "")),
            "section_count": len(sections),
            "line_count": sum(section["line_count"] for section in sections),
            "calculation_summary": {
                "calculation_row_count": sum(len(section["calculation_rows"]) for section in sections),
                "formula_column_count": len(all_formula_columns),
                "formula_columns": sorted(all_formula_columns),
            },
            "sections": sections,
        }
    finally:
        grid.close()


def validate_workbook_identity(plan: dict[str, Any], product_line) -> list[str]:
    metadata = plan["metadata"]
    errors = []

    def normalized_identifier(value):
        return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())

    actual_task_text = str(metadata.get("production_task_number") or "").strip().upper()
    actual_task_number = normalized_identifier(actual_task_text)
    task = product_line.production_task
    task_references = {
        normalized_identifier(task.task_number),
        normalized_identifier(task.order_number),
        normalized_identifier(task.contract_number),
    }
    task_references.discard("")
    order_reference = normalized_identifier(task.order_number)
    task_matches = actual_task_number in task_references
    if not task_matches and order_reference:
        product_suffix = re.fullmatch(r"(.+)-(\d{1,2})", actual_task_text)
        task_matches = bool(product_suffix and normalized_identifier(product_suffix.group(1)) == order_reference)
    if not task_matches:
        errors.append("生产单号与所选生产任务不一致。")

    expected_code = normalized_identifier(product_line.product_code_snapshot)
    actual_code = normalized_identifier(metadata.get("product_code"))
    if expected_code and actual_code != expected_code:
        errors.append("产品编号与所选生产任务产品不一致。")
    if not expected_code and metadata.get("product_name", "").strip() != product_line.product_name_snapshot.strip():
        errors.append("产品名称与所选生产任务产品不一致。")
    return errors
