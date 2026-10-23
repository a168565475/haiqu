import hashlib
from io import BytesIO
from zipfile import BadZipFile

from django.db import connection, transaction
from django.db.models import Q
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

from apps.reference_data.models import Currency
from .exceptions import FinanceConflict
from .models import FinanceAccount
from .serializers import AccountWriteSerializer
from .services import insert_account, validate_currency

HEADERS = {"账户类型": "account_type", "账户编码": "code", "账户类别": "category", "账户名称": "name",
           "账户账号": "account_number", "账户状态": "status", "币种": "currency"}


def read_account_workbook(content, expected_hash):
    if hashlib.sha256(content).hexdigest().lower() != expected_hash.lower():
        raise ValidationError("文件 SHA256 不符，停止导入。")
    if len(content) > 10 * 1024 * 1024:
        raise ValidationError("账户文件超过 10 MB。")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValidationError("无法读取有效的 XLSX 账户文件。") from exc
    try:
        if "账户信息" not in workbook.sheetnames:
            raise ValidationError("缺少账户信息工作表。")
        rows = iter(workbook["账户信息"].iter_rows())
        headers = [cell.value for cell in next(rows, [])]
        if any(headers.count(header) != 1 for header in HEADERS):
            raise ValidationError("账户表必须包含且不能重复七个业务表头。")
        result, seen = [], set()
        for index, row in enumerate(rows, start=2):
            if not any(cell.value is not None for cell in row):
                continue
            data = {}
            for header, field in HEADERS.items():
                cell = row[headers.index(header)]
                if cell.data_type == "f" or not isinstance(cell.value, str) or not cell.value.strip():
                    raise ValidationError(f"第 {index} 行 {header} 必须是非空原始文本，不接受公式或数值转换。")
                data[field] = cell.value
            if data["code"] != data["code"].strip():
                raise ValidationError(f"第 {index} 行账户编码含首尾空白，请核对原文件。")
            if data["code"] in seen:
                raise ValidationError(f"第 {index} 行账户编码在原表重复。")
            seen.add(data["code"])
            if data["status"] not in ("启用", "停用", "ACTIVE", "INACTIVE"):
                raise ValidationError(f"第 {index} 行账户状态无法识别。")
            data["status"] = {"启用": "ACTIVE", "停用": "INACTIVE"}.get(data["status"], data["status"])
            # 原表使用中文币名；只匹配字典的确切代码/中文名，不推断汇率。
            currency = Currency.objects.filter(status="ACTIVE", archived_at__isnull=True).filter(
                Q(code=data["currency"]) | Q(chinese_name=data["currency"]))
            if currency.count() != 1:
                raise ValidationError(f"第 {index} 行币种未唯一匹配有效字典。")
            data["currency_id"] = str(currency.get().pk)
            del data["currency"]
            serializer = AccountWriteSerializer(data=data)
            if not serializer.is_valid():
                raise ValidationError(f"第 {index} 行字段校验失败：{', '.join(serializer.errors)}。")
            if any(serializer.validated_data[field] != data[field] for field in ("code", "account_type", "category", "name", "account_number")):
                raise ValidationError(f"第 {index} 行文本包含需核实的首尾空白，停止而不自动修改原字段。")
            result.append((index, serializer.validated_data))
        if not result:
            raise ValidationError("原表没有有效账户行。")
        return result
    finally:
        workbook.close()


def account_import_plan(rows, *, table_exists=True):
    result = []
    for index, data in rows:
        account = FinanceAccount.objects.filter(code=data["code"]).first() if table_exists else None
        different = [] if account is None else [key for key, value in data.items() if getattr(account, key) != value]
        result.append({"row": index, "action": "CREATE" if account is None else "CONFLICT" if different else "UNCHANGED",
            "different_fields": different, "currency": data["currency"].code, "status": data["status"],
            "code_length": len(data["code"]), "account_number_length": len(data["account_number"]),
            "account_number_is_text": True, "leading_zero": data["account_number"].startswith("0")})
    return result


@transaction.atomic
def apply_account_import(rows, actor):
    # 与普通新增/编辑统一先币种、后账户。多币种按主键排序，避免反向等待。
    currencies = {data["currency"].pk: data["currency"] for _, data in rows}
    for pk in sorted(currencies):
        validate_currency(currencies[pk])
    # 两行原表的受控批次；阻止导入预检与并行CRUD之间出现差异窗口。
    with connection.cursor() as cursor:
        cursor.execute("LOCK TABLE finance_account IN SHARE ROW EXCLUSIVE MODE")
    plan = account_import_plan(rows)
    if any(item["action"] == "CONFLICT" for item in plan):
        raise FinanceConflict("DUPLICATE_CODE", "存在同编码不同内容，整批停止；请查看脱敏预检。")
    for (_, data), item in zip(rows, plan):
        if item["action"] == "CREATE":
            insert_account(data, actor)
    if any(item["action"] != "UNCHANGED" for item in account_import_plan(rows)):
        raise RuntimeError("账户导入后七字段精确比对失败，事务回滚。")
    return plan
