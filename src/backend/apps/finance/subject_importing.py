import hashlib
from io import BytesIO
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

from .exceptions import FinanceConflict
from .models import AccountingSubject
from .subject_services import insert_subject, lock_subjects

HEADERS = ("费用类型", "一级科目", "二级科目")


def read_subject_workbook(content, expected_hash):
    if len(content) > 10 * 1024 * 1024 or hashlib.sha256(content).hexdigest().lower() != expected_hash.lower():
        raise ValidationError("科目文件大小超过10MB或SHA256不符。")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            if "会计科目" not in workbook.sheetnames:
                raise ValidationError("缺少会计科目工作表。")
            rows = iter(workbook["会计科目"].iter_rows())
            headers = [cell.value for cell in next(rows, [])]
            if any(headers.count(header) != 1 for header in HEADERS):
                raise ValidationError("必须包含且不能重复费用类型、一级科目、二级科目表头。")
            paths, seen = [], set()
            for index, row in enumerate(rows, 2):
                if not any(cell.value is not None for cell in row):
                    continue
                parts = []
                for header in HEADERS:
                    cell = row[headers.index(header)]
                    if cell.data_type == "f" or not isinstance(cell.value, str) or not 1 <= len(cell.value.strip()) <= 100:
                        raise ValidationError(f"第{index}行{header}必须是1—100字非空原始文本，不接受公式。")
                    parts.append(cell.value.strip())
                path = tuple(parts)
                if path in seen:
                    raise ValidationError(f"第{index}行完整路径重复。")
                seen.add(path)
                paths.append(path)
            if not paths:
                raise ValidationError("原表没有有效完整路径。")
            return paths
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, OSError, ValueError, ParseError, KeyError) as exc:
        raise ValidationError("无法读取有效XLSX科目文件。") from exc


def subject_import_plan(paths, *, table_exists=True):
    nodes = sorted({path[:level] for path in paths for level in (1, 2, 3)}, key=lambda path: (len(path), path))
    existing, plan = {}, []
    if table_exists:
        for subject in AccountingSubject.objects.select_related("parent__parent"):
            chain = [subject.name]
            parent = subject.parent
            while parent:
                chain.insert(0, parent.name)
                parent = parent.parent
            existing[tuple(chain)] = subject
    for path in nodes:
        node = existing.get(path)
        conflict = node is not None and (node.status != "ACTIVE" or node.level != len(path))
        plan.append({"path": list(path), "level": len(path), "action": "CONFLICT" if conflict else "UNCHANGED" if node else "CREATE",
                     "different_fields": ["status/level"] if conflict else []})
    return {"path_count": len(paths), "node_count": len(nodes),
            "level_counts": {str(level): sum(len(path) == level for path in nodes) for level in (1, 2, 3)},
            "paths": [list(path) for path in paths], "nodes": plan}


@transaction.atomic
def apply_subject_import(paths, actor):
    lock_subjects()
    plan = subject_import_plan(paths)
    if any(item["action"] == "CONFLICT" for item in plan["nodes"]):
        raise FinanceConflict("DUPLICATE_SUBJECT_PATH", "同路径现有状态或层级不一致，整批停止。")
    ids = {}
    for item in plan["nodes"]:
        path = tuple(item["path"])
        parent_id = ids[path[:-1]] if len(path) > 1 else None
        if item["action"] == "CREATE":
            subject = insert_subject({"parent_id": parent_id, "level": len(path), "name": path[-1], "status": "ACTIVE"}, actor)
        else:
            subject = AccountingSubject.objects.get(parent_id=parent_id, name=path[-1])
        ids[path] = subject.pk
    if any(item["action"] != "UNCHANGED" for item in subject_import_plan(paths)["nodes"]):
        raise RuntimeError("科目导入后完整路径比对失败，整批回滚。")
    return plan
