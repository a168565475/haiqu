import csv
import io
import zipfile
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from apps.customers.contact_services import create_contact
from apps.customers.models import (
    ContactMethodType,
    ContactRoleType,
    ContactStatus,
    CustomerAccount,
    CustomerContactMethod,
)

from .models import ImportBatch, ImportBatchStatus, ImportRecord, ImportRecordStatus


CONTACT_COLUMNS = [
    "customer_account_code", "display_name", "position_title", "status", "is_key_contact",
    "email", "mobile", "phone", "wechat", "whatsapp", "role_types", "notes",
]
METHOD_COLUMNS = {
    "email": ContactMethodType.EMAIL,
    "mobile": ContactMethodType.MOBILE,
    "phone": ContactMethodType.PHONE,
    "wechat": ContactMethodType.WECHAT,
    "whatsapp": ContactMethodType.WHATSAPP,
}


def _check_xlsx_archive(path):
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 1000:
                raise ValidationError("XLSX 文件内部项目过多。")
            total_size = sum(item.file_size for item in infos)
            if total_size > 20 * 1024 * 1024:
                raise ValidationError("XLSX 解压后内容超过 20MB。")
            lowered = {item.filename.lower() for item in infos}
            if any(name.endswith("vbaproject.bin") for name in lowered):
                raise ValidationError("不允许导入包含宏的 XLSX 文件。")
            if any("externallinks/" in name for name in lowered):
                raise ValidationError("不允许导入包含外部链接的 XLSX 文件。")
    except zipfile.BadZipFile as exc:
        raise ValidationError("XLSX 文件结构损坏。") from exc


def read_contact_rows(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        try:
            with Path(path).open("r", encoding="utf-8-sig", newline="") as source:
                reader = csv.DictReader(source)
                headers = reader.fieldnames or []
                rows = [(index, dict(row)) for index, row in enumerate(reader, start=2)]
        except UnicodeDecodeError as exc:
            raise ValidationError("CSV 必须使用 UTF-8 或 UTF-8-SIG 编码。") from exc
    else:
        _check_xlsx_archive(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=False)
        header_cells = next(iterator, ())
        headers = [str(cell.value or "").strip() for cell in header_cells]
        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                workbook.close()
                raise ValidationError(f"第 {row_number} 行包含公式，不允许导入。")
            values = [cell.value for cell in cells]
            if not any(value not in (None, "") for value in values):
                continue
            rows.append((row_number, dict(zip(headers, values, strict=False))))
        workbook.close()
    missing = {"display_name"} - set(headers)
    if missing:
        raise ValidationError(f"缺少必填列：{', '.join(sorted(missing))}。")
    if len(rows) > settings.CONTACT_IMPORT_MAX_ROWS:
        raise ValidationError(f"导入数据不能超过 {settings.CONTACT_IMPORT_MAX_ROWS} 行。")
    return rows


def _text(value):
    return "" if value is None else str(value).strip()


def _boolean(value):
    normalized = _text(value).lower()
    if normalized in {"1", "true", "yes", "是", "关键"}:
        return True
    if normalized in {"", "0", "false", "no", "否"}:
        return False
    raise ValidationError("is_key_contact 只能填写是/否、true/false 或 1/0。")


def _row_payload(row):
    display_name = _text(row.get("display_name"))
    if not display_name:
        raise ValidationError("display_name 不能为空。")
    methods = []
    for column, method_type in METHOD_COLUMNS.items():
        value = _text(row.get(column))
        if value:
            methods.append(
                {
                    "method_type": method_type,
                    "raw_value": value,
                    "is_primary": True,
                    "is_preferred": True,
                    "verification_status": "UNVERIFIED",
                }
            )
    if not methods:
        raise ValidationError("邮箱、手机、电话、微信或 WhatsApp 至少填写一项。")
    status = _text(row.get("status")).upper() or ContactStatus.ACTIVE
    if status not in {ContactStatus.ACTIVE, ContactStatus.INACTIVE}:
        raise ValidationError("status 只能填写 ACTIVE 或 INACTIVE。")
    role_types = [item.strip().upper() for item in _text(row.get("role_types")).split("|") if item.strip()]
    invalid_roles = set(role_types) - set(ContactRoleType.values)
    if invalid_roles:
        raise ValidationError(f"不支持的角色：{', '.join(sorted(invalid_roles))}。")
    return {
        "display_name": display_name,
        "position_title": _text(row.get("position_title")),
        "status": status,
        "is_key_contact": _boolean(row.get("is_key_contact")),
        "notes": _text(row.get("notes")),
        "methods": methods,
        "role_types": role_types,
    }


def _domain_warnings(customer, methods, request):
    warnings = []
    for method in methods:
        if method["method_type"] != ContactMethodType.EMAIL or "@" not in method["raw_value"]:
            continue
        domain = method["raw_value"].lower().rsplit("@", 1)[1]
        from apps.customers.selectors import accessible_customers as visible_customers
        if CustomerContactMethod.objects.filter(
            customer_contact__customer_account__in=visible_customers(request, "view_customercontact"),
            method_type=ContactMethodType.EMAIL,
            normalized_value__iendswith=f"@{domain}",
            archived_at__isnull=True,
            customer_contact__archived_at__isnull=True,
        ).exclude(customer_contact__customer_account=customer).exists():
            warnings.append(f"邮箱域名 {domain} 也出现在其他客户中，请核对客户归属。")
    return warnings


def process_contact_import(*, business_file, path, actor_employee, request, accessible_customers, customer_id=None, allow_duplicates=False):
    fixed_customer = None
    if customer_id:
        fixed_customer = accessible_customers.filter(pk=customer_id, archived_at__isnull=True).first()
        if fixed_customer is None:
            raise ValidationError({"customer_id": "客户不存在或无权导入。"})

    started_at = timezone.now()
    batch = ImportBatch.objects.create(
        source_file=business_file,
        created_by_employee=actor_employee,
        import_type="CONTACT",
        status=ImportBatchStatus.PROCESSING,
        started_at=started_at,
        config_data={"customer_id": str(customer_id) if customer_id else None, "allow_duplicates": allow_duplicates},
    )
    try:
        rows = read_contact_rows(path)
    except Exception as exc:
        batch.status = ImportBatchStatus.FAILED
        batch.error_summary = str(getattr(exc, "detail", exc))[:2000]
        batch.finished_at = timezone.now()
        batch.save(update_fields=["status", "error_summary", "finished_at", "updated_at"])
        raise

    accessible_by_code = {
        customer.account_code: customer
        for customer in accessible_customers.filter(archived_at__isnull=True)
    }

    success_count = 0
    failure_count = 0
    for row_number, row in rows:
        raw_data = {key: _text(value) for key, value in row.items() if key}
        record = ImportRecord.objects.create(
            import_batch=batch,
            row_number=row_number,
            raw_data=raw_data,
        )
        try:
            with transaction.atomic():
                customer = fixed_customer
                row_customer_code = _text(row.get("customer_account_code"))
                if customer is None:
                    if not row_customer_code:
                        raise ValidationError("全局导入必须填写 customer_account_code。")
                    customer = accessible_by_code.get(row_customer_code)
                    if customer is None:
                        raise ValidationError("客户编号不存在或无权导入。")
                elif row_customer_code and row_customer_code != customer.account_code:
                    raise ValidationError("文件中的客户编号与当前客户不一致。")
                payload = _row_payload(row)
                warnings = _domain_warnings(customer, payload["methods"], request)
                payload["duplicate_acknowledged"] = allow_duplicates
                contact = create_contact(
                    customer=customer,
                    validated_data=payload,
                    actor_employee=actor_employee,
                    request=request,
                )
                record.status = ImportRecordStatus.SUCCEEDED
                record.warning_messages = warnings
                record.target_type = "CUSTOMER_CONTACT"
                record.target_id = contact.pk
                record.save(update_fields=[
                    "status", "warning_messages", "target_type", "target_id", "updated_at"
                ])
            success_count += 1
        except Exception as exc:
            failure_count += 1
            detail = getattr(exc, "detail", exc)
            record.status = ImportRecordStatus.FAILED
            record.error_message = str(detail)[:2000]
            record.save(update_fields=["status", "error_message", "updated_at"])

    batch.total_count = len(rows)
    batch.success_count = success_count
    batch.failure_count = failure_count
    batch.status = (
        ImportBatchStatus.COMPLETED if failure_count == 0 else
        ImportBatchStatus.FAILED if success_count == 0 else
        ImportBatchStatus.PARTIAL_FAILED
    )
    batch.finished_at = timezone.now()
    batch.save(update_fields=[
        "total_count", "success_count", "failure_count", "status", "finished_at", "updated_at"
    ])
    return batch
