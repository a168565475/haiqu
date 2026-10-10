import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from types import SimpleNamespace

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from apps.customers.contact_services import create_contact
from apps.customers.models import (
    AssignmentRole,
    ContactMethodType,
    ContactStatus,
    CustomerAccount,
    CustomerAssignment,
    CustomerCompanyContactMethod,
    CustomerCompanyContactMethodType,
    CustomerInterestedProduct,
    CustomerLifecycleStage,
    CustomerRegion,
)
from apps.customers.services import (
    create_customer,
    set_initial_activity_to_follow_up_from_imported_contact,
)
from apps.imports.models import ImportBatch, ImportBatchStatus, ImportRecord, ImportRecordStatus
from apps.imports.security import save_and_scan_upload
from apps.organizations.models import Employee, EmployeeStatus
from apps.reference_data.models import ReferenceDictionary, ReferenceStatus


SOURCE_SYSTEM = "ENTE_ERP"
IMPORT_TYPE = "ENTE_CUSTOMER"

LIFECYCLE_MAP = {
    "潜在客户": CustomerLifecycleStage.LEAD,
    "正式客户": CustomerLifecycleStage.CUSTOMER,
    "暂无需求": CustomerLifecycleStage.LEAD,
}
CUSTOMER_TYPE_MAP = {
    "潜在客户": "POTENTIAL",
    "正式客户": "FORMAL",
    "货代": "FORWARDER",
    # 恩特旧值没有单独的“暂无需求”类型，按已确认规则直接归入潜在客户。
    "暂无需求": "POTENTIAL",
}
GRADE_MAP = {
    "A": "A", "A级": "A",
    "B": "B", "B级": "B",
    "C": "C", "C级": "C",
    "D": "D", "D级": "D",
    "E": "E", "E级": "E",
    "★询价": "F",
}
SOURCE_MAP = {
    "阿里诚信通": "ALIBABA_CHINA",
    "阿里国际": "ALIBABA_INTERNATIONAL",
    "中国制造网": "MADE_IN_CHINA",
    "企业邮箱": "EMAIL",
    "广交会": "CANTON_FAIR",
    "德国科隆展": "COLOGNE_FAIR",
    "拉斯维加斯展": "LAS_VEGAS_FAIR",
    "巴西展": "BRAZIL_FAIR",
    "其他": "OTHER",
    "香港展": "HONG_KONG_FAIR",
}
INDUSTRY_MAP = {
    "国外进出口公司": "FOREIGN_IMPORT_EXPORT",
    "港台外贸公司": "HONG_KONG_TAIWAN_FOREIGN_TRADE",
    "国内外贸公司": "DOMESTIC_FOREIGN_TRADE",
    "国外促销礼品公司": "FOREIGN_PROMOTIONAL_GIFTS",
    "国外大型连锁超市": "FOREIGN_CHAIN_SUPERMARKET",
    "国外连锁日用购物": "FOREIGN_DAILY_CHAIN",
    "国外最终客户": "FOREIGN_END_CUSTOMER",
    "国外批发商": "FOREIGN_WHOLESALER",
    "国外零售商": "FOREIGN_RETAILER",
    "国内礼品公司": "DOMESTIC_GIFTS",
    "国内最终客户": "DOMESTIC_END_CUSTOMER",
    "国内零售商": "DOMESTIC_RETAILER",
    "国内批发公司": "DOMESTIC_WHOLESALER",
    "其他": "OTHER",
}
INTERESTED_PRODUCT_MAP = {
    "折叠床/桌/椅": "FOLDING_FURNITURE",
    "购物车/篮/袋/包": "CART_BASKET_BAG",
    "购物车/篮+行李车": "CART_BASKET_BAG",
    "车棚/凉棚/遮阳伞": "CANOPY_AWNING_UMBRELLA",
    "帐篷/睡袋": "TENT_SLEEPING_BAG",
    "帐蓬": "TENT_SLEEPING_BAG",
    "编藤套装桌椅": "RATTAN_FURNITURE",
    "吊篮": "HANGING_CHAIR",
    "秋千": "SWING",
    "杯子，酒具，柠檬夹，饭盒等五金类": "HARDWARE",
    "健身器材，哑铃，跑步机类": "FITNESS",
    "防疫物资，口罩，试剂类": "PANDEMIC_SUPPLIES",
    "车类": "VEHICLE",
    "反光背心": "REFLECTIVE_VEST",
    "其他": "OTHER",
}
CUSTOMER_DIRECT_FIELDS = {
    "序号",
    "系统唯一标识",
    "客户名称",
    "客户简称",
    "客户类型",
    "客户等级",
    "客户来源",
    "所属行业",
    "感兴趣产品",
    "询盘关键词",
    "建档人",
    "建档人代码",
    "建档时间",
    "最后联系时间",
    "所属业务员",
    "业务员代码",
    "原分管人",
    "区域代码",
    "国家/地区",
    "电话",
    "电子邮箱",
    "其他",
    "网站",
    "读取完整性",
    "未完整原因",
    "联系人数量",
}
CONTACT_DIRECT_FIELDS = {
    "客户代码",
    "客户名称",
    "系统唯一标识",
    "联系人序号",
    "联系人",
    "职位",
    "邮箱",
    "手机",
    "电话",
    "办公电话",
    "宅电",
    "WhatsApp",
    "微信",
    "是否主联系人",
    "部门",
    "备注",
    "读取完整性",
    "未完整原因",
}


class ImportValidationError(ValueError):
    pass


class _SystemUser:
    @staticmethod
    def has_perm(permission):
        return True


SYSTEM_REQUEST = SimpleNamespace(user=_SystemUser(), META={})


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _nonempty_row(row):
    return {key: _json_value(value) for key, value in row.items() if value is not None and value != ""}


def _legacy_data(row, *, excluded_fields, source_row):
    fields = {
        key: _json_value(value)
        for key, value in row.items()
        if key not in excluded_fields and value is not None and value != ""
    }
    return {"source_system": SOURCE_SYSTEM, "source_row": source_row, "fields": fields}


def _read_sheet(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise ImportValidationError(f"缺少工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    headers = [_text(cell.value) for cell in sheet[1]]
    if not headers or any(not header for header in headers):
        raise ImportValidationError(f"工作表 {sheet_name} 的表头不能为空。")
    if len(headers) != len(set(headers)):
        raise ImportValidationError(f"工作表 {sheet_name} 存在重复表头。")
    rows = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or value == "" for value in values):
            continue
        rows.append((source_row, dict(zip(headers, values, strict=True))))
    return headers, rows


def _required_headers(actual_headers, required_headers, sheet_name):
    missing = sorted(set(required_headers) - set(actual_headers))
    if missing:
        raise ImportValidationError(f"工作表 {sheet_name} 缺少字段：{', '.join(missing)}")


def _parse_owners(customer_row, source_row):
    names = [_text(value) for value in _text(customer_row.get("所属业务员")).split(";") if _text(value)]
    codes = [_text(value) for value in _text(customer_row.get("业务员代码")).split(";") if _text(value)]
    if not names or not codes or len(names) != len(codes):
        raise ImportValidationError(f"客户信息第 {source_row} 行业务员姓名与代码数量不一致。")
    return list(zip(codes, names, strict=True))


def _split_values(value):
    return [_text(item) for item in _text(value).replace("；", ";").split(";") if _text(item)]


def _parse_managers(customer_row):
    # 原分管人没有旧员工编码；空编码是原始数据的事实，不擅自与同名人员合并。
    return [("MANAGER", name) for name in _split_values(customer_row.get("原分管人"))]


def _parse_recorded_by(customer_row):
    name = _text(customer_row.get("建档人"))
    code = _text(customer_row.get("建档人代码"))
    return (code, name) if name else None


def _parse_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return timezone.make_aware(value) if timezone.is_naive(value) else value
    if isinstance(value, date):
        return timezone.make_aware(datetime.combine(value, time.min))
    try:
        parsed = datetime.fromisoformat(_text(value).replace("Z", "+00:00"))
    except ValueError as exc:
        raise ImportValidationError(f"无法识别日期时间：{_text(value)}") from exc
    return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed


def _map_interested_products(value, source_row):
    result = []
    for raw_value in _split_values(value):
        item_code = INTERESTED_PRODUCT_MAP.get(raw_value)
        if item_code is None:
            raise ImportValidationError(f"客户信息第 {source_row} 行感兴趣产品无法映射：{raw_value}")
        if item_code not in result:
            result.append(item_code)
    return result


def _company_contact_methods(row):
    field_types = (
        ("电子邮箱", CustomerCompanyContactMethodType.EMAIL),
        ("电话", CustomerCompanyContactMethodType.PHONE),
        ("其他", CustomerCompanyContactMethodType.SKYPE),
    )
    result = []
    seen = set()
    for field_name, method_type in field_types:
        raw_value = _text(row.get(field_name))
        normalized_value = raw_value.casefold()
        key = (method_type, normalized_value)
        if not raw_value or key in seen:
            continue
        seen.add(key)
        result.append(
            {
                "method_type": method_type,
                "raw_value": raw_value,
                "normalized_value": normalized_value,
                "is_primary": not any(item["method_type"] == method_type for item in result),
            }
        )
    return result


def _normalize_website(value):
    website = _text(value)
    if not website:
        return ""
    if not website.lower().startswith(("http://", "https://")):
        website = f"https://{website}"
    return website


def _contact_methods(row):
    field_types = (
        ("邮箱", ContactMethodType.EMAIL),
        ("手机", ContactMethodType.MOBILE),
        ("电话", ContactMethodType.PHONE),
        ("办公电话", ContactMethodType.PHONE),
        ("宅电", ContactMethodType.PHONE),
        ("WhatsApp", ContactMethodType.WHATSAPP),
        ("微信", ContactMethodType.WECHAT),
    )
    seen = set()
    seen_types = set()
    methods = []
    for field_name, method_type in field_types:
        raw_value = _text(row.get(field_name))
        key = (method_type, raw_value.casefold())
        if not raw_value or key in seen:
            continue
        seen.add(key)
        first_of_type = method_type not in seen_types
        seen_types.add(method_type)
        methods.append(
            {
                "method_type": method_type,
                "raw_value": raw_value,
                "is_primary": first_of_type,
                "is_preferred": first_of_type,
                "verification_status": "UNVERIFIED",
            }
        )
    return methods


def build_import_plan(source_path):
    source_path = Path(source_path)
    if not source_path.is_file():
        raise ImportValidationError(f"源文件不存在：{source_path}")
    if source_path.suffix.lower() != ".xlsx":
        raise ImportValidationError("恩特客户导入只支持 XLSX 文件。")
    if source_path.stat().st_size <= 0:
        raise ImportValidationError("源文件为空。")

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        customer_headers, customer_rows = _read_sheet(workbook, "客户信息")
        contact_headers, contact_rows = _read_sheet(workbook, "联系人信息")
    finally:
        workbook.close()

    _required_headers(
        customer_headers,
        {
            "客户代码", "系统唯一标识", "客户名称", "客户类型", "客户等级", "客户来源",
            "所属业务员", "业务员代码", "区域代码", "网站", "读取完整性",
        },
        "客户信息",
    )
    _required_headers(
        contact_headers,
        {"客户代码", "系统唯一标识", "联系人", "是否主联系人", "读取完整性"},
        "联系人信息",
    )

    customers = []
    customer_codes = set()
    external_ids = set()
    owner_pairs = []
    person_pairs = []
    region_without_country = []
    for source_row, row in customer_rows:
        old_code = _text(row.get("客户代码"))
        external_id = _text(row.get("系统唯一标识"))
        account_name = _text(row.get("客户名称"))
        lifecycle_raw = _text(row.get("客户类型"))
        grade_raw = _text(row.get("客户等级"))
        source_raw = _text(row.get("客户来源"))
        industry_raw = _text(row.get("所属行业"))
        completeness = _text(row.get("读取完整性"))
        if not old_code or not external_id or not account_name:
            raise ImportValidationError(f"客户信息第 {source_row} 行缺少客户代码、系统唯一标识或客户名称。")
        if completeness and completeness != "完整":
            raise ImportValidationError(f"客户信息第 {source_row} 行读取不完整：{_text(row.get('未完整原因'))}")
        if old_code in customer_codes:
            raise ImportValidationError(f"客户代码重复：{old_code}")
        if external_id in external_ids:
            raise ImportValidationError(f"系统唯一标识重复：{external_id}")
        if lifecycle_raw not in LIFECYCLE_MAP:
            raise ImportValidationError(f"客户信息第 {source_row} 行客户类型无法映射：{lifecycle_raw}")
        if lifecycle_raw not in CUSTOMER_TYPE_MAP:
            raise ImportValidationError(f"客户信息第 {source_row} 行客户类型无法映射：{lifecycle_raw}")
        if grade_raw not in GRADE_MAP:
            raise ImportValidationError(f"客户信息第 {source_row} 行客户等级无法映射：{grade_raw}")
        if source_raw not in SOURCE_MAP:
            raise ImportValidationError(f"客户信息第 {source_row} 行客户来源无法映射：{source_raw}")
        if industry_raw and industry_raw not in INDUSTRY_MAP:
            raise ImportValidationError(f"客户信息第 {source_row} 行所属行业无法映射：{industry_raw}")
        owners = _parse_owners(row, source_row)
        for pair in owners:
            if pair not in owner_pairs:
                owner_pairs.append(pair)
            if pair not in person_pairs:
                person_pairs.append(pair)
        managers = _parse_managers(row)
        for pair in managers:
            if pair not in person_pairs:
                person_pairs.append(pair)
        recorded_by = _parse_recorded_by(row)
        if recorded_by and recorded_by not in person_pairs:
            person_pairs.append(recorded_by)
        country_code = _text(row.get("区域代码")).upper()
        valid_country_code = country_code if len(country_code) == 2 and country_code.isalpha() else ""
        if country_code and not valid_country_code:
            region_without_country.append(
                {"source_row": source_row, "customer_code": old_code, "raw_country_code": country_code}
            )
        customer_codes.add(old_code)
        external_ids.add(external_id)
        legacy_source_data = _legacy_data(
            row, excluded_fields=CUSTOMER_DIRECT_FIELDS, source_row=source_row
        )
        legacy_source_data["mapped_values"] = {
            key: _json_value(row.get(key))
            for key in CUSTOMER_DIRECT_FIELDS
            if row.get(key) is not None and row.get(key) != ""
        }
        customers.append(
            {
                "source_row": source_row,
                "old_code": old_code,
                "external_id": external_id,
                "account_name": account_name,
                "short_name": _text(row.get("客户简称")),
                "lifecycle_stage": LIFECYCLE_MAP[lifecycle_raw],
                "customer_type_code": CUSTOMER_TYPE_MAP[lifecycle_raw],
                "grade_code": GRADE_MAP[grade_raw],
                "source_code": SOURCE_MAP[source_raw],
                "industry_code": INDUSTRY_MAP.get(industry_raw, ""),
                "interested_product_codes": _map_interested_products(row.get("感兴趣产品"), source_row),
                "country_code": valid_country_code,
                "region_selector_code": country_code,
                "website": _normalize_website(row.get("网站")),
                "owners": owners,
                "managers": managers,
                "recorded_by": recorded_by,
                "recorded_at": _parse_datetime(row.get("建档时间")),
                "last_contact_at": _parse_datetime(row.get("最后联系时间")),
                "inquiry_keywords": _text(row.get("询盘关键词")),
                "company_contact_methods": _company_contact_methods(row),
                "legacy_source_data": legacy_source_data,
                "raw_data": _nonempty_row(row),
                "contacts": [],
            }
        )

    customer_by_code = {customer["old_code"]: customer for customer in customers}
    primary_seen = set()
    contact_identity = set()
    for source_row, row in contact_rows:
        old_code = _text(row.get("客户代码"))
        external_id = _text(row.get("系统唯一标识"))
        display_name = _text(row.get("联系人"))
        completeness = _text(row.get("读取完整性"))
        if old_code not in customer_by_code:
            raise ImportValidationError(f"联系人信息第 {source_row} 行引用了不存在的客户代码：{old_code}")
        if external_id != customer_by_code[old_code]["external_id"]:
            raise ImportValidationError(f"联系人信息第 {source_row} 行系统唯一标识与客户不一致。")
        if not display_name:
            raise ImportValidationError(f"联系人信息第 {source_row} 行联系人姓名为空。")
        if completeness and completeness != "完整":
            raise ImportValidationError(f"联系人信息第 {source_row} 行读取不完整：{_text(row.get('未完整原因'))}")
        identity = (old_code, _text(row.get("联系人序号")))
        if identity in contact_identity:
            raise ImportValidationError(f"联系人信息第 {source_row} 行联系人序号重复。")
        contact_identity.add(identity)
        is_primary = _text(row.get("是否主联系人")).upper() == "Y"
        if is_primary and old_code in primary_seen:
            raise ImportValidationError(f"客户 {old_code} 存在多个主联系人。")
        if is_primary:
            primary_seen.add(old_code)
        contact = {
            "source_row": source_row,
            "display_name": display_name,
            "position_title": _text(row.get("职位")),
            "department_name": _text(row.get("部门")),
            "is_primary_contact": is_primary,
            "notes": _text(row.get("备注")),
            "methods": _contact_methods(row),
            "legacy_source_data": _legacy_data(
                row, excluded_fields=CONTACT_DIRECT_FIELDS, source_row=source_row
            ),
            "raw_data": _nonempty_row(row),
        }
        customer_by_code[old_code]["contacts"].append(contact)

    owner_names_by_code = defaultdict(list)
    for old_code, person_name in owner_pairs:
        if person_name not in owner_names_by_code[old_code]:
            owner_names_by_code[old_code].append(person_name)
    extra_names_by_code = defaultdict(list)
    for old_code, person_name in person_pairs:
        if person_name not in owner_names_by_code[old_code] and person_name not in extra_names_by_code[old_code]:
            extra_names_by_code[old_code].append(person_name)
    person_mapping = []
    for old_code, person_name in person_pairs:
        names = [*owner_names_by_code[old_code], *extra_names_by_code[old_code]]
        suffix = f"_{names.index(person_name) + 1:02d}" if len(names) > 1 else ""
        person_mapping.append(
            {
                "legacy_employee_code": old_code,
                "legacy_employee_name": person_name,
                "employee_code": f"ENTE_{old_code}{suffix}",
            }
        )

    digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    return {
        "source_path": str(source_path),
        "source_name": source_path.name,
        "source_size": source_path.stat().st_size,
        "source_hash": digest,
        "customers": customers,
        "customer_count": len(customers),
        "contact_count": sum(len(customer["contacts"]) for customer in customers),
        "owner_mapping": [item for item in person_mapping if (item["legacy_employee_code"], item["legacy_employee_name"]) in owner_pairs],
        "person_mapping": person_mapping,
        "invalid_countries": region_without_country,
    }


def plan_summary(plan):
    existing_ids = set(
        CustomerAccount.objects.filter(
            source_system=SOURCE_SYSTEM,
            external_id__in=[customer["external_id"] for customer in plan["customers"]],
        ).values_list("external_id", flat=True)
    )
    return {
        "source_name": plan["source_name"],
        "source_hash": plan["source_hash"],
        "customer_count": plan["customer_count"],
        "contact_count": plan["contact_count"],
        "owner_count": len(plan["owner_mapping"]),
        "person_count": len(plan["person_mapping"]),
        "invalid_country_count": len(plan["invalid_countries"]),
        "existing_customer_count": len(existing_ids),
        "existing_external_ids": sorted(existing_ids),
        "invalid_countries": plan["invalid_countries"],
        "owner_mapping": plan["owner_mapping"],
    }


def _reference_item(category_code, item_code):
    try:
        return ReferenceDictionary.objects.get(
            dictionary_type__code=category_code,
            item_code=item_code,
            status=ReferenceStatus.ACTIVE,
            archived_at__isnull=True,
        )
    except ReferenceDictionary.DoesNotExist as exc:
        raise ImportValidationError(f"ERP 缺少启用的字典项：{category_code}.{item_code}") from exc


def _employee_mapping(plan):
    result = {}
    for item in plan["person_mapping"]:
        employee, created = Employee.objects.get_or_create(
            employee_code=item["employee_code"],
            defaults={"display_name": item["legacy_employee_name"], "status": EmployeeStatus.ACTIVE},
        )
        if not created and employee.display_name != item["legacy_employee_name"]:
            raise ImportValidationError(
                f"员工编码 {employee.employee_code} 已存在，但姓名为 {employee.display_name}，"
                f"与源数据 {item['legacy_employee_name']} 不一致。"
            )
        if employee.archived_at is not None or employee.status != EmployeeStatus.ACTIVE:
            raise ImportValidationError(f"员工 {employee.employee_code} 不是启用状态。")
        result[(item["legacy_employee_code"], item["legacy_employee_name"])] = employee
    return result


def _region_mapping(customer_data):
    selector_code = customer_data["region_selector_code"]
    if not selector_code:
        return None, None, ""
    try:
        region = CustomerRegion.objects.select_related("parent", "country").get(
            legacy_selector_code=selector_code,
            is_active=True,
            archived_at__isnull=True,
        )
    except CustomerRegion.DoesNotExist as exc:
        raise ImportValidationError(
            f"客户信息第 {customer_data['source_row']} 行区域代码无法匹配系统配置：{selector_code}"
        ) from exc
    if region.level == CustomerRegion.Level.PRIMARY:
        return region, None, ""
    return region.parent, region, region.country.code if region.country_id else ""


def _reference_items(category_code, item_codes):
    return {item_code: _reference_item(category_code, item_code) for item_code in sorted(set(item_codes)) if item_code}


def _ensure_history_contact(*, customer, customer_data, created_by_employee):
    occurred_at = customer_data["last_contact_at"]
    if occurred_at is None:
        return False
    from apps.crm.models import CustomerCommunication

    _communication, created = CustomerCommunication.objects.get_or_create(
        source_system=SOURCE_SYSTEM,
        external_id=f"{customer_data['external_id']}:LAST_CONTACT",
        defaults={
            "customer_account": customer,
            "created_by_employee": created_by_employee,
            "communication_type": "OTHER",
            "direction": "INTERNAL",
            "subject": "恩特 ERP 历史联系时间",
            "content": "恩特 ERP 历史联系时间（仅有时间，无明细）。",
            "occurred_at": occurred_at,
        },
    )
    return created


def execute_import(plan):
    existing = plan_summary(plan)["existing_external_ids"]
    if existing:
        raise ImportValidationError(f"ERP 已存在 {len(existing)} 个相同恩特系统标识，已停止导入以防重复。")

    with Path(plan["source_path"]).open("rb") as source:
        uploaded_file = SimpleUploadedFile(
            plan["source_name"],
            source.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    business_file, _ = save_and_scan_upload(uploaded_file, None)

    with transaction.atomic():
        import_operator, _ = Employee.objects.get_or_create(
            employee_code="SYSTEM_ENTE_IMPORT",
            defaults={"display_name": "恩特 ERP 导入服务", "status": EmployeeStatus.ACTIVE},
        )
        if import_operator.archived_at is not None or import_operator.status != EmployeeStatus.ACTIVE:
            raise ImportValidationError("恩特 ERP 导入服务员工不是启用状态。")
        business_file.uploaded_by_employee = import_operator
        business_file.save(update_fields=["uploaded_by_employee", "updated_at"])

        employee_mapping = _employee_mapping(plan)
        grade_items = _reference_items("CUSTOMER_GRADE", GRADE_MAP.values())
        source_items = _reference_items("CUSTOMER_SOURCE", SOURCE_MAP.values())
        type_items = _reference_items("CUSTOMER_TYPE", CUSTOMER_TYPE_MAP.values())
        industry_items = _reference_items("CUSTOMER_INDUSTRY", INDUSTRY_MAP.values())
        interested_product_items = _reference_items("CUSTOMER_INTERESTED_PRODUCT", INTERESTED_PRODUCT_MAP.values())
        started_at = timezone.now()
        batch = ImportBatch.objects.create(
            source_file=business_file,
            created_by_employee=import_operator,
            import_type=IMPORT_TYPE,
            status=ImportBatchStatus.PROCESSING,
            total_count=plan["customer_count"],
            started_at=started_at,
            config_data={
                "source_system": SOURCE_SYSTEM,
                "source_hash": plan["source_hash"],
                "customer_count": plan["customer_count"],
                "contact_count": plan["contact_count"],
                "owner_mapping": plan["owner_mapping"],
                "person_mapping": plan["person_mapping"],
                "mapping": {
                    "lifecycle": LIFECYCLE_MAP,
                    "customer_type": CUSTOMER_TYPE_MAP,
                    "grade": GRADE_MAP,
                    "source": SOURCE_MAP,
                    "industry": INDUSTRY_MAP,
                    "interested_product": INTERESTED_PRODUCT_MAP,
                    "owner_rule": "所属业务员第一位为负责人，其余为协同人；原分管人单独写入分管人",
                    "region_rule": "区域代码通过客户地区配置映射；一级地区代码不填写国家",
                },
            },
        )

        historical_contact_count = 0
        for customer_data in plan["customers"]:
            owner_employees = [employee_mapping[pair] for pair in customer_data["owners"]]
            manager_employees = [employee_mapping[pair] for pair in customer_data["managers"]]
            recorded_by_employee = (
                employee_mapping[customer_data["recorded_by"]]
                if customer_data["recorded_by"] is not None
                else None
            )
            primary_region, secondary_region, country_code = _region_mapping(customer_data)
            customer, primary_assignment = create_customer(
                validated_data={
                    "account_name": customer_data["account_name"],
                    "short_name": customer_data["short_name"],
                    "lifecycle_stage": customer_data["lifecycle_stage"],
                    "primary_country_code": country_code,
                    "primary_region": primary_region,
                    "secondary_region": secondary_region,
                    "primary_website_url": customer_data["website"],
                    "industry_item": industry_items.get(customer_data["industry_code"]),
                    "inquiry_keywords": customer_data["inquiry_keywords"],
                    "recorded_by_employee": recorded_by_employee,
                    "recorded_at": customer_data["recorded_at"],
                    "source_system": SOURCE_SYSTEM,
                    "external_id": customer_data["external_id"],
                    "notes": "",
                    "legacy_source_data": customer_data["legacy_source_data"],
                    "primary_owner_employee": owner_employees[0],
                    "customer_source_item": source_items[customer_data["source_code"]],
                    "customer_grade_item": grade_items[customer_data["grade_code"]],
                    "customer_type_items": [type_items[customer_data["customer_type_code"]]],
                    "customer_tags": [],
                    "interested_product_items": [
                        interested_product_items[code]
                        for code in customer_data["interested_product_codes"]
                    ],
                },
                actor_employee=import_operator,
                request=SYSTEM_REQUEST,
            )
            primary_assignment.reason = "恩特 ERP 历史负责人导入（所属业务员第一位）"
            primary_assignment.save(update_fields=["reason", "updated_at"])
            for collaborator in owner_employees[1:]:
                CustomerAssignment.objects.create(
                    customer_account=customer,
                    employee=collaborator,
                    assignment_role=AssignmentRole.COLLABORATOR,
                    effective_from=started_at,
                    assigned_by_employee=import_operator,
                    reason="恩特 ERP 历史协同人导入（所属业务员非第一位）",
                )
            for manager in manager_employees:
                CustomerAssignment.objects.create(
                    customer_account=customer,
                    employee=manager,
                    assignment_role=AssignmentRole.MANAGER,
                    effective_from=started_at,
                    assigned_by_employee=import_operator,
                    reason="恩特 ERP 历史分管人导入",
                )
            CustomerCompanyContactMethod.objects.bulk_create(
                [
                    CustomerCompanyContactMethod(customer_account=customer, **method_data)
                    for method_data in customer_data["company_contact_methods"]
                ]
            )
            historical_contact_count += int(
                _ensure_history_contact(
                    customer=customer,
                    customer_data=customer_data,
                    created_by_employee=recorded_by_employee or import_operator,
                )
            )
            if customer_data["last_contact_at"] is not None:
                set_initial_activity_to_follow_up_from_imported_contact(customer=customer)
            for contact_data in customer_data["contacts"]:
                create_contact(
                    customer=customer,
                    validated_data={
                        "display_name": contact_data["display_name"],
                        "position_title": contact_data["position_title"],
                        "department_name": contact_data["department_name"],
                        "is_primary_contact": contact_data["is_primary_contact"],
                        "status": ContactStatus.ACTIVE,
                        "notes": contact_data["notes"],
                        "legacy_source_data": contact_data["legacy_source_data"],
                        "methods": contact_data["methods"],
                        "role_types": [],
                        "duplicate_acknowledged": False,
                    },
                    actor_employee=import_operator,
                    request=SYSTEM_REQUEST,
                )
            ImportRecord.objects.create(
                import_batch=batch,
                row_number=customer_data["source_row"],
                status=ImportRecordStatus.SUCCEEDED,
                raw_data={
                    "customer": customer_data["raw_data"],
                    "contacts": [contact["raw_data"] for contact in customer_data["contacts"]],
                },
                warning_messages=(
                    ["区域代码只对应一级地区，国家字段按规则留空。"]
                    if customer_data["region_selector_code"] in {"1", "2", "3", "4", "5", "6", "7"}
                    else []
                ),
                target_type="CUSTOMER_ACCOUNT",
                target_id=customer.pk,
            )

        finished_at = timezone.now()
        batch.status = ImportBatchStatus.COMPLETED
        batch.success_count = plan["customer_count"]
        batch.failure_count = 0
        batch.finished_at = finished_at
        batch.save(
            update_fields=["status", "success_count", "failure_count", "finished_at", "updated_at"]
        )
        return {
            "batch_id": str(batch.pk),
            "status": batch.status,
            "customer_count": plan["customer_count"],
            "contact_count": plan["contact_count"],
            "employee_count": len(plan["owner_mapping"]),
            "person_count": len(plan["person_mapping"]),
            "historical_contact_count": historical_contact_count,
            "invalid_country_count": len(plan["invalid_countries"]),
            "source_hash": plan["source_hash"],
        }


def execute_sync(plan):
    """按已确认规则重读并重映射已导入的恩特客户，不创建重复客户。"""
    existing_customers = {
        customer.external_id: customer
        for customer in CustomerAccount.objects.filter(
            source_system=SOURCE_SYSTEM,
            external_id__in=[item["external_id"] for item in plan["customers"]],
        )
    }
    missing_external_ids = sorted(
        item["external_id"] for item in plan["customers"] if item["external_id"] not in existing_customers
    )
    if missing_external_ids:
        raise ImportValidationError(
            f"有 {len(missing_external_ids)} 条恩特客户尚未导入，不能执行仅重映射同步："
            f"{', '.join(missing_external_ids[:5])}"
        )

    with Path(plan["source_path"]).open("rb") as source:
        uploaded_file = SimpleUploadedFile(
            plan["source_name"],
            source.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    business_file, _ = save_and_scan_upload(uploaded_file, None)

    from apps.customers.models import CustomerSourceAttribution, CustomerTypeAssignment

    with transaction.atomic():
        import_operator, _ = Employee.objects.get_or_create(
            employee_code="SYSTEM_ENTE_IMPORT",
            defaults={"display_name": "恩特 ERP 导入服务", "status": EmployeeStatus.ACTIVE},
        )
        if import_operator.archived_at is not None or import_operator.status != EmployeeStatus.ACTIVE:
            raise ImportValidationError("恩特 ERP 导入服务员工不是启用状态。")
        business_file.uploaded_by_employee = import_operator
        business_file.save(update_fields=["uploaded_by_employee", "updated_at"])

        employee_mapping = _employee_mapping(plan)
        source_items = _reference_items("CUSTOMER_SOURCE", SOURCE_MAP.values())
        type_items = _reference_items("CUSTOMER_TYPE", CUSTOMER_TYPE_MAP.values())
        industry_items = _reference_items("CUSTOMER_INDUSTRY", INDUSTRY_MAP.values())
        interested_product_items = _reference_items("CUSTOMER_INTERESTED_PRODUCT", INTERESTED_PRODUCT_MAP.values())
        started_at = timezone.now()
        batch = ImportBatch.objects.create(
            source_file=business_file,
            created_by_employee=import_operator,
            import_type=IMPORT_TYPE,
            status=ImportBatchStatus.PROCESSING,
            total_count=plan["customer_count"],
            started_at=started_at,
            config_data={
                "source_system": SOURCE_SYSTEM,
                "source_hash": plan["source_hash"],
                "mode": "REMAP_EXISTING",
                "customer_count": plan["customer_count"],
                "person_mapping": plan["person_mapping"],
                "mapping": {
                    "customer_type": CUSTOMER_TYPE_MAP,
                    "source": SOURCE_MAP,
                    "industry": INDUSTRY_MAP,
                    "interested_product": INTERESTED_PRODUCT_MAP,
                    "assignment_rule": "所属业务员映射负责人和协同人；原分管人映射分管人",
                },
            },
        )

        historical_contact_count = 0
        for customer_data in plan["customers"]:
            customer = CustomerAccount.objects.select_for_update().get(
                pk=existing_customers[customer_data["external_id"]].pk
            )
            primary_region, secondary_region, country_code = _region_mapping(customer_data)
            recorded_by_employee = (
                employee_mapping[customer_data["recorded_by"]]
                if customer_data["recorded_by"] is not None
                else None
            )
            customer.account_name = customer_data["account_name"]
            customer.lifecycle_stage = customer_data["lifecycle_stage"]
            customer.primary_region = primary_region
            customer.secondary_region = secondary_region
            customer.primary_country_code = country_code
            customer.legacy_source_data = customer_data["legacy_source_data"]
            if customer_data["short_name"]:
                customer.short_name = customer_data["short_name"]
            if customer_data["website"]:
                customer.primary_website_url = customer_data["website"]
            if customer_data["industry_code"]:
                customer.industry_item = industry_items[customer_data["industry_code"]]
            if customer_data["inquiry_keywords"]:
                customer.inquiry_keywords = customer_data["inquiry_keywords"]
            if recorded_by_employee is not None:
                customer.recorded_by_employee = recorded_by_employee
            if customer_data["recorded_at"] is not None:
                customer.recorded_at = customer_data["recorded_at"]
            customer.save()

            CustomerSourceAttribution.objects.select_for_update().filter(
                customer_account=customer,
                effective_to__isnull=True,
            ).update(effective_to=started_at, updated_at=started_at)
            CustomerSourceAttribution.objects.create(
                customer_account=customer,
                source_item=source_items[customer_data["source_code"]],
                recorded_by_employee=recorded_by_employee or import_operator,
                notes="按恩特 ERP 客户明细重新映射客户来源",
                effective_from=started_at,
            )
            CustomerTypeAssignment.objects.select_for_update().filter(
                customer_account=customer,
                effective_to__isnull=True,
            ).update(effective_to=started_at, updated_at=started_at)
            CustomerTypeAssignment.objects.create(
                customer_account=customer,
                type_item=type_items[customer_data["customer_type_code"]],
                assigned_by_employee=import_operator,
                reason="按恩特 ERP 客户类型重新映射",
                effective_from=started_at,
            )

            CustomerInterestedProduct.objects.filter(customer_account=customer).delete()
            CustomerInterestedProduct.objects.bulk_create(
                [
                    CustomerInterestedProduct(
                        customer_account=customer,
                        product_item=interested_product_items[item_code],
                        assigned_by_employee=import_operator,
                    )
                    for item_code in customer_data["interested_product_codes"]
                ]
            )
            for method_data in customer_data["company_contact_methods"]:
                CustomerCompanyContactMethod.objects.update_or_create(
                    customer_account=customer,
                    method_type=method_data["method_type"],
                    normalized_value=method_data["normalized_value"],
                    defaults={
                        "raw_value": method_data["raw_value"],
                        "is_primary": method_data["is_primary"],
                    },
                )

            CustomerAssignment.objects.select_for_update().filter(
                customer_account=customer,
                assignment_role__in=[
                    AssignmentRole.PRIMARY_OWNER,
                    AssignmentRole.COLLABORATOR,
                    AssignmentRole.MANAGER,
                ],
                effective_to__isnull=True,
            ).update(effective_to=started_at, updated_at=started_at)
            owner_employees = [employee_mapping[pair] for pair in customer_data["owners"]]
            CustomerAssignment.objects.create(
                customer_account=customer,
                employee=owner_employees[0],
                assignment_role=AssignmentRole.PRIMARY_OWNER,
                effective_from=started_at,
                assigned_by_employee=import_operator,
                reason="恩特 ERP 历史负责人重映射（所属业务员第一位）",
            )
            CustomerAssignment.objects.bulk_create(
                [
                    CustomerAssignment(
                        customer_account=customer,
                        employee=employee,
                        assignment_role=AssignmentRole.COLLABORATOR,
                        effective_from=started_at,
                        assigned_by_employee=import_operator,
                        reason="恩特 ERP 历史协同人重映射（所属业务员非第一位）",
                    )
                    for employee in owner_employees[1:]
                ]
            )
            CustomerAssignment.objects.bulk_create(
                [
                    CustomerAssignment(
                        customer_account=customer,
                        employee=employee_mapping[pair],
                        assignment_role=AssignmentRole.MANAGER,
                        effective_from=started_at,
                        assigned_by_employee=import_operator,
                        reason="恩特 ERP 历史分管人重映射",
                    )
                    for pair in customer_data["managers"]
                ]
            )
            historical_contact_count += int(
                _ensure_history_contact(
                    customer=customer,
                    customer_data=customer_data,
                    created_by_employee=recorded_by_employee or import_operator,
                )
            )
            if customer_data["last_contact_at"] is not None:
                set_initial_activity_to_follow_up_from_imported_contact(customer=customer)
            ImportRecord.objects.create(
                import_batch=batch,
                row_number=customer_data["source_row"],
                status=ImportRecordStatus.SUCCEEDED,
                raw_data={
                    "customer": customer_data["raw_data"],
                    "contacts": [contact["raw_data"] for contact in customer_data["contacts"]],
                },
                warning_messages=(
                    ["区域代码只对应一级地区，国家字段按规则留空。"]
                    if customer_data["region_selector_code"] in {"1", "2", "3", "4", "5", "6", "7"}
                    else []
                ),
                target_type="CUSTOMER_ACCOUNT",
                target_id=customer.pk,
            )

        finished_at = timezone.now()
        batch.status = ImportBatchStatus.COMPLETED
        batch.success_count = plan["customer_count"]
        batch.failure_count = 0
        batch.finished_at = finished_at
        batch.save(update_fields=["status", "success_count", "failure_count", "finished_at", "updated_at"])
        return {
            "batch_id": str(batch.pk),
            "status": batch.status,
            "customer_count": plan["customer_count"],
            "person_count": len(plan["person_mapping"]),
            "historical_contact_count": historical_contact_count,
            "source_hash": plan["source_hash"],
        }


def dump_summary(data):
    return json.dumps(data, ensure_ascii=False, indent=2, default=str)
