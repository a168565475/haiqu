from dataclasses import dataclass
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from .models import RecordStatus, Warehouse, WarehouseLocation, WarehouseType, WarehouseZone


class WarehouseWorkbookImportError(ValueError):
    pass


TYPE_MAP = {"材料仓": WarehouseType.MATERIAL, "半成品仓": WarehouseType.SEMI_FINISHED, "成品仓": WarehouseType.FINISHED, "不良品仓": WarehouseType.DEFECTIVE}
STATUS_MAP = {"正常": RecordStatus.ACTIVE, "停用": RecordStatus.DISABLED}
WAREHOUSE_HEADERS = ["仓库号", "仓库名称", "仓库类型", "仓库地址", "仓库状态", "仓库负责人", "负责人联系电话"]
LOCATION_HEADERS = ["仓库号", "仓库名称", "库区", "库位号", "仓库库容", "库位号唯一校验", "仓库类型"]


@dataclass
class WarehouseImportPlan:
    warehouses: list[dict]
    zones: list[dict]
    locations: list[dict]
    duplicate_location_keys: list[str]
    source_location_rows: int


def _text(value):
    return "" if value is None else str(value).strip()


def _rows(path, sheet_name, headers):
    source = Path(path)
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise WarehouseWorkbookImportError(f"源文件不存在或不是 XLSX：{source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise WarehouseWorkbookImportError(f"缺少工作表：{sheet_name}")
        sheet = workbook[sheet_name]
        actual = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(headers)))]
        if actual != headers:
            raise WarehouseWorkbookImportError(f"{sheet_name} 表头与已确认模板不一致。")
        return [tuple(row) for row in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True) if any(value is not None for value in row)]
    finally:
        workbook.close()


def build_warehouse_import_plan(warehouse_path, location_path):
    warehouse_rows = _rows(warehouse_path, "仓库信息", WAREHOUSE_HEADERS)
    location_rows = _rows(location_path, "库位维护", LOCATION_HEADERS)
    warehouses = []
    known = set()
    warehouse_meta = {}
    for row_number, row in enumerate(warehouse_rows, start=2):
        number, name, type_label, address, status_label, responsible, phone = [_text(value) for value in row]
        if not number or not name or type_label not in TYPE_MAP or status_label not in STATUS_MAP:
            raise WarehouseWorkbookImportError(f"仓库信息第 {row_number} 行必填值或枚举值不正确。")
        if number in known:
            raise WarehouseWorkbookImportError(f"仓库号重复：{number}")
        known.add(number)
        warehouse_meta[number] = (name, TYPE_MAP[type_label])
        warehouses.append({"warehouse_number": number.upper(), "warehouse_name": name, "warehouse_type": TYPE_MAP[type_label], "address": address, "status": STATUS_MAP[status_label], "responsible_person_name": responsible, "contact_phone": phone})

    zones_by_key = {}
    locations = []
    seen_location_keys = set()
    duplicates = set()
    for row_number, row in enumerate(location_rows, start=2):
        number, name, zone_name, location_number, capacity, unique_key, type_label = [_text(value) for value in row]
        if number not in known or not zone_name or type_label not in TYPE_MAP:
            raise WarehouseWorkbookImportError(f"库位维护第 {row_number} 行仓库、库区或类型不正确。")
        if warehouse_meta[number] != (name, TYPE_MAP[type_label]):
            raise WarehouseWorkbookImportError(f"库位维护第 {row_number} 行仓库名称或类型与仓库信息不一致。")
        calculated = f"{number}{zone_name}{location_number}"
        if unique_key != calculated:
            raise WarehouseWorkbookImportError(f"库位维护第 {row_number} 行唯一校验值不一致。")
        zones_by_key[(number, zone_name)] = {"warehouse_number": number, "zone_name": zone_name}
        if location_number == "-":
            continue
        key = (number, zone_name, location_number.upper())
        if key in seen_location_keys:
            duplicates.add(calculated)
            continue
        seen_location_keys.add(key)
        try:
            parsed_capacity = None if not capacity else int(capacity)
        except ValueError as exc:
            raise WarehouseWorkbookImportError(f"库位维护第 {row_number} 行库容不是正整数。") from exc
        if parsed_capacity is not None and parsed_capacity <= 0:
            raise WarehouseWorkbookImportError(f"库位维护第 {row_number} 行库容必须为正整数或空。")
        locations.append({"warehouse_number": number, "zone_name": zone_name, "location_number": location_number.upper(), "capacity_box_slots": parsed_capacity})
    return WarehouseImportPlan(warehouses, list(zones_by_key.values()), locations, sorted(duplicates), len(location_rows))


@transaction.atomic
def execute_warehouse_import(plan, *, expected_existing_warehouses, expected_existing_locations):
    if plan.duplicate_location_keys:
        raise WarehouseWorkbookImportError("源文件存在重复库位：" + "、".join(plan.duplicate_location_keys))
    if Warehouse.objects.select_for_update().count() != expected_existing_warehouses:
        raise WarehouseWorkbookImportError("现有仓库数量与保护参数不一致，已停止导入。")
    if WarehouseLocation.objects.select_for_update().count() != expected_existing_locations:
        raise WarehouseWorkbookImportError("现有库位数量与保护参数不一致，已停止导入。")
    warehouse_map = {}
    for values in plan.warehouses:
        warehouse, _ = Warehouse.objects.update_or_create(warehouse_number=values["warehouse_number"], defaults=values)
        warehouse_map[values["warehouse_number"]] = warehouse
    zone_map = {}
    for values in plan.zones:
        zone, _ = WarehouseZone.objects.update_or_create(warehouse=warehouse_map[values["warehouse_number"]], zone_name=values["zone_name"], defaults={"status": RecordStatus.ACTIVE})
        zone_map[(values["warehouse_number"], values["zone_name"])] = zone
    for values in plan.locations:
        WarehouseLocation.objects.update_or_create(zone=zone_map[(values["warehouse_number"], values["zone_name"])], location_number=values["location_number"], defaults={"capacity_box_slots": values["capacity_box_slots"], "status": RecordStatus.ACTIVE})
    return {"warehouses": len(plan.warehouses), "zones": len(plan.zones), "locations": len(plan.locations)}
