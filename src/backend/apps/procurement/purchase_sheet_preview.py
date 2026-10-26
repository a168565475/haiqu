"""从下载所用的 XLSX 读取预览；不另建一套合同排版，不保存文件。"""

import base64
import re
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter

from .purchase_sheet_xlsx import build_supplier_purchase_sheet_xlsx
from .purchase_sheet_template_edits import original_template_terms
from .fabric_purchase_sheet_template import FabricPurchaseSheetTemplateError, uses_fabric_contract_template
from .original_purchase_sheet_templates import uses_original_purchase_contract_template


def _color(color, default="#000000"):
    if color is not None and color.type == "rgb":
        return f"#{color.rgb[-6:]}"
    return default


def _font(font):
    return {
        "fontFamily": f'"{getattr(font, "name", None) or getattr(font, "rFont", None) or "宋体"}", serif',
        "fontSize": f"{float(font.sz or 11) * 4 / 3}px",
        "fontWeight": "700" if font.b else "400",
        "fontStyle": "italic" if font.i else "normal",
        "color": _color(font.color),
    }


def _border(side):
    if not side or not side.style:
        return "none"
    width = 3 if side.style == "thick" else 2 if side.style in {"medium", "double"} else 1
    style = "double" if side.style == "double" else "dotted" if side.style == "dotted" else "dashed" if "dash" in side.style.lower() else "solid"
    return f"{width}px {style} {_color(side.color)}"


def _display(value, number_format):
    if value is None:
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        text = format(Decimal(str(value)), "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        # 原版导出用引号保存混合单位的后缀。
        return text + "".join(re.findall(r'"([^"]*)"', number_format))
    return str(value)


def build_purchase_sheet_preview(sheet):
    if not (uses_fabric_contract_template(sheet) or uses_original_purchase_contract_template(sheet)):
        raise FabricPurchaseSheetTemplateError("该采购表未配置原版 Excel，继续使用现有编辑视图。")
    data = build_supplier_purchase_sheet_xlsx(sheet)
    workbook = load_workbook(BytesIO(data), rich_text=True)
    try:
        ws = workbook.active
        # 原件在打印区外留有格式单元格；不能把这些空列算进合同宽度。
        last_column = max([cell.column for row in ws for cell in row if cell.value is not None] + [area.max_col for area in ws.merged_cells.ranges])
        with ZipFile(BytesIO(data)) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        cached = {cell.get("r"): cell.findtext("m:v", namespaces=ns) for cell in root.findall(".//m:c[m:f]", ns)}
        columns = []
        for index in range(1, last_column + 1):
            dimension = next((dim for dim in ws.column_dimensions.values() if dim.min <= index <= dim.max), None)
            width = dimension.width if dimension else ws.sheet_format.defaultColWidth or 8.43
            columns.append(0 if dimension and dimension.hidden else int(width * 7 + 5))
        heights = [0 if ws.row_dimensions[index].hidden else float(ws.row_dimensions[index].height or ws.sheet_format.defaultRowHeight or 15) * 4 / 3 for index in range(1, ws.max_row + 1)]
        merges, covered = {}, set()
        for merged in ws.merged_cells.ranges:
            merges[(merged.min_row, merged.min_col)] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
            covered.update((r, c) for r in range(merged.min_row, merged.max_row + 1) for c in range(merged.min_col, merged.max_col + 1) if (r, c) != (merged.min_row, merged.min_col))
        rows = []
        for r, height in enumerate(heights, 1):
            cells = []
            for c in range(1, len(columns) + 1):
                if (r, c) in covered:
                    continue
                cell = ws.cell(r, c)
                row_span, col_span = merges.get((r, c), (1, 1))
                value = cell.value
                if cell.data_type == "f":
                    value = Decimal(cached[cell.coordinate]) if cached.get(cell.coordinate) else ""
                runs = ([{"text": str(part.text) if isinstance(part, TextBlock) else str(part), "style": _font(part.font) if isinstance(part, TextBlock) else {}} for part in value]
                        if isinstance(value, CellRichText) else [{"text": _display(value, cell.number_format), "style": {}}])
                style = {
                    **_font(cell.font),
                    "textAlign": cell.alignment.horizontal if cell.alignment.horizontal in {"left", "center", "right", "justify"} else "right" if isinstance(value, (int, float, Decimal)) else "left",
                    "verticalAlign": {"center": "middle"}.get(cell.alignment.vertical, cell.alignment.vertical or "bottom"),
                    "whiteSpace": "pre-wrap" if cell.alignment.wrap_text else "pre",
                    "borderTop": _border(cell.border.top), "borderRight": _border(cell.border.right),
                    "borderBottom": _border(cell.border.bottom), "borderLeft": _border(cell.border.left),
                    "backgroundColor": _color(cell.fill.fgColor, "transparent") if cell.fill.patternType == "solid" else "transparent",
                }
                cells.append({"address": f"{get_column_letter(c)}{r}", "rowspan": row_span, "colspan": col_span, "runs": runs, "style": style})
            rows.append({"height": height, "cells": cells})
        images = []
        for picture in ws._images:
            anchor = picture.anchor
            start = anchor._from
            left = sum(columns[:start.col]) + start.colOff / 9525
            top = sum(heights[:start.row]) + start.rowOff / 9525
            if hasattr(anchor, "ext"):
                width, height = anchor.ext.cx / 9525, anchor.ext.cy / 9525
            else:
                end = anchor.to
                width = sum(columns[:end.col]) + end.colOff / 9525 - left
                height = sum(heights[:end.row]) + end.rowOff / 9525 - top
            images.append({"src": f"data:image/{picture.format};base64,{base64.b64encode(picture._data()).decode('ascii')}", "left": left, "top": top, "width": width, "height": height})
        return {"columns": columns, "rows": rows, "images": images,
                "width": sum(columns), "height": sum(heights),
                "terms": original_template_terms([str(ws.cell(r, 1).value or "") for r in range(1, ws.max_row + 1)]),
                "margins": {key: float(getattr(ws.page_margins, key) or 0) * 96 for key in ("top", "right", "bottom", "left")}}
    finally:
        workbook.close()
