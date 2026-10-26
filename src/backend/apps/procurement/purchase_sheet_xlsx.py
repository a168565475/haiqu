"""按确认的原版模板或通用七列表格导出可编辑采购合同。"""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import unicodedata

from django.conf import settings
from apps.files.models import is_file_scan_usable
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .fabric_purchase_sheet_template import build_fabric_purchase_sheet_xlsx, uses_fabric_contract_template
from .original_purchase_sheet_templates import build_original_purchase_contract_xlsx, uses_original_purchase_contract_template
from .purchase_sheet_pdf import _contract_date, _structured_sections
from .purchase_sheet_line_display import purchase_sheet_line_display
from .services import MONEY_QUANT


THIN = Side(style="thin", color="000000")
TABLE_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUMBER_ALIGNMENT = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
TITLE_FONT = Font(name="宋体", size=18, bold=True)
BODY_FONT = Font(name="宋体", size=10)
BOLD_FONT = Font(name="宋体", size=10, bold=True)


def _text(value, fallback="—"):
    return str(value if value is not None else "").strip() or fallback


def _number(value):
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            return None
        return int(number) if number == number.to_integral_value() else float(number)
    except (InvalidOperation, TypeError, ValueError):
        return None


def _number_format(value):
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return "General"
    if not number.is_finite():
        return "General"
    fractional = format(number, "f").partition(".")[2].rstrip("0")
    return "0." + "#" * len(fractional) if fractional else "0"


def _media_image_path(business_file):
    if business_file is None or not str(business_file.content_type or "").startswith("image/") or not is_file_scan_usable(business_file.scan_status):
        return None
    media_root = Path(settings.MEDIA_ROOT).resolve()
    file_path = (media_root / business_file.storage_key).resolve()
    try:
        file_path.relative_to(media_root)
    except ValueError:
        return None
    return file_path if file_path.is_file() else None


def _style_range(ws, row_start, row_end, column_start, column_end, *, bold=False, center=False):
    for row in range(row_start, row_end + 1):
        for column in range(column_start, column_end + 1):
            cell = ws.cell(row=row, column=column)
            cell.border = TABLE_BORDER
            cell.alignment = CENTER_ALIGNMENT if center else BODY_ALIGNMENT
            cell.font = BOLD_FONT if bold else BODY_FONT


def _text_chunks(value, width, *, max_lines=24):
    """按列宽保守估算换行，拆分长文本以免超过 Excel 行高上限。"""
    if not isinstance(value, str):
        return [(value, 1)]
    capacity = max(4, int(width) - 3)
    chunks, current, used, count = [], "", 0, 1
    for char in value:
        size = 2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
        if char == "\n" or used + size > capacity:
            if count >= max_lines:
                chunks.append((current, count))
                current, count = "", 2 if char == "\n" else 1
            else:
                count += 1
            used = 0
        current += char
        if char != "\n":
            used += size
    chunks.append((current, count))
    return chunks


def _merge_text(ws, row, column_count, text, *, height=26, bold=False, center=False, font=None):
    font = font or (BOLD_FONT if bold else BODY_FONT)
    width = sum(ws.column_dimensions[get_column_letter(column)].width for column in range(1, column_count + 1))
    leading = float(font.sz) * 1.4
    for chunk, count in _text_chunks(text, width * 10 / float(font.sz), max_lines=int(380 / leading)):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_count)
        ws.cell(row=row, column=1, value=chunk)
        _style_range(ws, row, row, 1, column_count, bold=bold, center=center)
        ws.cell(row=row, column=1).font = font
        ws.row_dimensions[row].height = max(height, count * leading + 12)
        row += 1
    return row


def _layout_config(layout=None):
    # 未使用原版模板的合同保留通用七列。
    return {
        "widths": [14, 23, 24, 12, 9, 9, 9],
        "headers": ["产品名称", "产品规格", "物料属性", "颜色", "数量", "含税单价\n（元）", "金额\n（元）"],
        "quantity_column": 5,
        "price_column": 6,
        "amount_column": 7,
    }


def _line_values(layout, line, index, order_number, actual_weight, sample_source=False):
    from .pipe_purchase import pipe_print_description
    display = purchase_sheet_line_display(line, sample_source=sample_source)
    specification, attributes, color = (display[key] for key in ("specification", "attributes", "color"))
    if line.get("measurement_kind") == "PIPE":
        print_line = {**line, **{
            key: _number(line.get(key)) for key in ("actual_bar_count", "raw_bar_length_m", "actual_purchase_meters")
        }}
        attributes = "\n".join(value for value in (attributes, pipe_print_description(print_line)) if value)
    quantity = _number(line.get("quantity"))
    price = _number(line.get("unit_price"))
    amount = _number(line.get("amount"))
    return [
        _text(line.get("part_name")), specification or "—", attributes or "—", color or "—",
        quantity if quantity is not None else "—", price if price is not None else "—",
        "待结算" if actual_weight or amount is None else amount,
    ]


def _add_image(ws, image_path, anchor, *, width=140, height=98):
    if not image_path:
        return False
    try:
        image = ExcelImage(str(image_path))
    except (OSError, ValueError):
        return False
    scale = min(width / image.width, height / image.height)
    image.width, image.height = image.width * scale, image.height * scale
    ws.add_image(image, anchor)
    return True


def _write_detail_rows(ws, sheet, content, layout, config, start_row):
    line_images = {
        str(item.source_purchase_order_line_id): item.image_file
        for item in sheet.line_images.select_related("image_file").all()
    } if content.get("show_image", True) else {}
    lines = list(content.get("lines") or [])
    order = content.get("order") or {}
    actual_weight = content.get("settlement_mode") == "ACTUAL_WEIGHT"
    row = start_row
    if not lines:
        return lines, line_images, _merge_text(ws, row, 7, "暂无采购明细", center=True)
    for index, line in enumerate(lines, start=1):
        values = _line_values(layout, line, index, "", actual_weight, order.get("source_type") == "SAMPLE")
        parts = [_text_chunks(value, width) for value, width in zip(values, config["widths"])]
        for continuation in range(max(len(part) for part in parts)):
            height = 1
            for column, part in enumerate(parts, 1):
                if continuation < len(part):
                    value, count = part[continuation]
                    ws.cell(row=row, column=column, value=value)
                    height = max(height, count)
            _style_range(ws, row, row, 1, 7)
            for column in (4, 5, 6, 7):
                ws.cell(row=row, column=column).alignment = CENTER_ALIGNMENT
            if continuation == 0:
                for column in (5, 6, 7):
                    ws.cell(row=row, column=column).number_format = _number_format(values[column - 1])
                    if isinstance(values[column - 1], (int, float)):
                        ws.cell(row=row, column=column).alignment = NUMBER_ALIGNMENT
                unit = _text(line.get("unit"), "").replace('"', '\\"')
                if isinstance(values[4], (int, float)) and unit:
                    ws.cell(row=row, column=5).number_format += f'" {unit}"'
                if not actual_weight and all(isinstance(values[column], (int, float)) for column in (4, 5, 6)):
                    # 与保存采购明细金额时使用相同精度，再汇总。
                    places = abs(MONEY_QUANT.as_tuple().exponent)
                    ws.cell(row=row, column=7, value=f"=ROUND(E{row}*F{row},{places})")
                    ws.cell(row=row, column=7).number_format = _number_format(line.get("amount"))
            ws.row_dimensions[row].height = max(26, height * 14 + 12)
            row += 1
    return lines, line_images, row


def _write_total_and_note(ws, content, layout, config, row, lines, line_images, *, detail_start=6):
    start_row = row
    note = f"备注：\n{_text(content.get('remark'), '')}\n税率 {_text(content.get('tax_rate'), '13')}% · 运费由{_text(content.get('freight_bearer'))}承担"
    for chunk, count in _text_chunks(note, sum(config["widths"][2:5])):
        _style_range(ws, row, row, 1, 7)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        ws.cell(row=row, column=3, value=chunk)
        ws.row_dimensions[row].height = max(110 if row == start_row else 26, count * 14 + 12)
        row += 1
    ws.merge_cells(start_row=start_row, start_column=1, end_row=row - 1, end_column=2)
    for column in (6, 7):
        ws.merge_cells(start_row=start_row, start_column=column, end_row=row - 1, end_column=column)
        ws.cell(start_row, column).alignment = CENTER_ALIGNMENT
        ws.cell(start_row, column).font = BOLD_FONT
    ws.cell(start_row, 6, "合计")
    total = _number(content.get("total_amount"))
    if content.get("settlement_mode") == "ACTUAL_WEIGHT" or total is None:
        ws.cell(start_row, 7, "待结算")
    elif lines:
        ws.cell(start_row, 7, f"=SUM(G{detail_start}:G{start_row - 1})")
        ws.cell(start_row, 7).number_format = _number_format(content.get("total_amount"))
    else:
        ws.cell(start_row, 7, total)
        ws.cell(start_row, 7).number_format = _number_format(content.get("total_amount"))
    image_path = next((path for line in lines if (path := _media_image_path(
        line_images.get(str(line.get("source_purchase_order_line_id") or ""))
    ))), None) if content.get("show_image", True) else None
    if not _add_image(ws, image_path, f"A{start_row}"):
        ws.cell(start_row, 1, "暂无图片")
        ws.cell(start_row, 1).alignment = CENTER_ALIGNMENT
    return row


def _write_sections_and_signature(ws, sheet, content, start_row, column_count):
    row = start_row
    for label, value in _structured_sections(content):
        row = _merge_text(ws, row, column_count, label, bold=True)
        text = "\n".join(f"{index}. {_text(item)}" for index, item in enumerate(value, 1)) if isinstance(value, list) else _text(value)
        row = _merge_text(ws, row, column_count, text)

    for offset, height in enumerate((26, 64, 26)):
        _style_range(ws, row + offset, row + offset, 1, 7)
        ws.merge_cells(start_row=row + offset, start_column=1, end_row=row + offset, end_column=3)
        ws.merge_cells(start_row=row + offset, start_column=4, end_row=row + offset, end_column=7)
        ws.row_dimensions[row + offset].height = height
    ws.cell(row, 1, "甲方").font = BOLD_FONT
    seal = getattr(sheet, "seal_image", None)
    _add_image(ws, _media_image_path(seal.image_file if seal else None), f"A{row + 1}", width=110, height=76)
    ws.cell(row + 2, 1, "请确认无误盖章回传")
    ws.cell(row + 2, 4, f"签订日期：{_text(_contract_date(content.get('purchase_date')))}")
    ws.cell(row + 2, 4).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    return row + 2


def build_supplier_purchase_sheet_xlsx(sheet):
    from .purchase_sheet_template_edits import apply_original_template_edits
    if uses_fabric_contract_template(sheet):
        return apply_original_template_edits(build_fabric_purchase_sheet_xlsx(sheet), sheet.content_snapshot)
    if uses_original_purchase_contract_template(sheet):
        return apply_original_template_edits(build_original_purchase_contract_xlsx(sheet), sheet.content_snapshot)
    content = sheet.content_snapshot if isinstance(sheet.content_snapshot, dict) else {}
    config = _layout_config()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "采购合同"
    worksheet.sheet_view.showGridLines = False
    for index, width in enumerate(config["widths"], start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    buyer, seller, order = (content.get(key) or {} for key in ("buyer", "seller", "order"))
    order_number = _text(order.get("source_number") if order.get("source_type") == "SAMPLE" else order.get("production_task_number"))
    title = f"{_text(content.get('title'), '购销合同')}　（订单编号  {order_number}）"
    row = _merge_text(worksheet, 1, 7, title, height=38, center=True, font=TITLE_FONT)
    party = f"采购单号：{_text(sheet.sheet_number)}\n甲方（采购方）：{_text(buyer.get('name'))}\n乙方（供货方）：{_text(seller.get('name'))}"
    row = _merge_text(worksheet, row, 7, party, height=56)
    legal = "根据《中华人民共和国民法典》《中华人民共和国产品质量法》等国家和地方政府部门有关法律法规，双方经过充分友好协商，本着诚实守信的原则，就采购方向供货方购买产品一事，达成以下约定："
    row = _merge_text(worksheet, row, 7, legal, height=46)
    row = _merge_text(worksheet, row, 7, "一、采购产品名称、规格、数量、单价及金额", bold=True)
    for column, header in enumerate(config["headers"], start=1):
        worksheet.cell(row=row, column=column, value=header)
    _style_range(worksheet, row, row, 1, 7, bold=True, center=True)
    worksheet.row_dimensions[row].height = 40
    worksheet.print_title_rows = f"{row}:{row}"
    detail_start = row + 1
    lines, images, total_row = _write_detail_rows(worksheet, sheet, content, "CONTRACT", config, detail_start)
    section_start = _write_total_and_note(worksheet, content, "CONTRACT", config, total_row, lines, images, detail_start=detail_start)
    final_row = _write_sections_and_signature(worksheet, sheet, content, section_start, 7)

    worksheet.print_area = f"A1:G{final_row}"
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins = PageMargins(left=0.28, right=0.28, top=0.2, bottom=0.28, header=0, footer=0)
    workbook.calculation.fullCalcOnLoad = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
