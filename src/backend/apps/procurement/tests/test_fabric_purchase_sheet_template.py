from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.test import SimpleTestCase
from openpyxl import load_workbook

from ..fabric_purchase_sheet_template import (
    DRAWING_PATH, MAIN, NS, SHEET_PATH, STRINGS_PATH, STYLES_PATH, TEMPLATE_PATH,
    FabricPurchaseSheetTemplateError, build_fabric_purchase_sheet_xlsx, uses_fabric_contract_template,
)


class FabricPurchaseSheetTemplateTests(SimpleTestCase):
    def setUp(self):
        self.line = {
            "part_name": "椅背布", "specification": "600D", "material": "涤纶", "unit": "米",
            "quantity": "3.0000", "unit_price": "0.3333", "amount": "0.9999", "delivery_date": "2026-10-30",
            "material_attributes_snapshot": {"fields": [
                {"key": "fabric_width", "label": "门幅", "value": "1.5000"},
                {"key": "fabric_color_name", "label": "颜色", "value": "宝蓝"},
                {"key": "fabric_color_code", "label": "色号", "value": "B-01"},
                {"key": "fabric_color_card", "label": "色卡", "value": "卡A"},
                {"key": "fabric_coating", "label": "涂层", "value": "PVC"},
            ]},
        }
        self.content = {
            "order": {"source_type": "PRODUCTION", "production_task_number": "SC2610001", "sales_order_number": "HT-NOT-TASK"},
            "buyer": {"name": "浙江嗨趣休闲用品有限公司"}, "seller": {"name": "本单面料供应商"},
            "purchase_date": "2026-10-26", "tax_rate": "9.0000", "freight_bearer": "甲方", "remark": "按本单确认样供货",
            "lines": [deepcopy(self.line)],
        }

    def sheet(self, content=None, code="FABRIC"):
        return SimpleNamespace(content_snapshot=deepcopy(content if content is not None else self.content),
                               template_version=SimpleNamespace(template=SimpleNamespace(template_code=code)))

    def test_template_selection_uses_exact_code_and_excludes_sample(self):
        self.assertTrue(uses_fabric_contract_template(self.sheet()))
        for code in ("FABRIC_OLD", "fabric", "PLASTIC", None):
            self.assertFalse(uses_fabric_contract_template(self.sheet(code=code)))
        content = deepcopy(self.content)
        content["order"]["source_type"] = "SAMPLE"
        self.assertFalse(uses_fabric_contract_template(self.sheet(content)))

    def test_original_layout_resources_title_runs_and_fixed_terms_are_preserved(self):
        output = build_fabric_purchase_sheet_xlsx(self.sheet())
        with ZipFile(TEMPLATE_PATH) as original, ZipFile(BytesIO(output)) as result:
            changed = {SHEET_PATH, STRINGS_PATH, STYLES_PATH, DRAWING_PATH}
            for item in original.namelist():
                if item not in changed:
                    self.assertEqual(result.read(item), original.read(item), item)
            source_sheet, new_sheet = (ET.fromstring(z.read(SHEET_PATH)) for z in (original, result))
            for tag in ("cols", "pageSetup", "pageMargins", "printOptions"):
                self.assertEqual(ET.tostring(source_sheet.find(f"m:{tag}", NS)), ET.tostring(new_sheet.find(f"m:{tag}", NS)), tag)
            source_styles, new_styles = (ET.fromstring(z.read(STYLES_PATH)) for z in (original, result))
            for tag in ("fonts", "fills", "borders", "cellStyleXfs"):
                self.assertEqual(ET.tostring(source_styles.find(f"m:{tag}", NS)), ET.tostring(new_styles.find(f"m:{tag}", NS)), tag)
            old_xfs, new_xfs = (root.find("m:cellXfs", NS) for root in (source_styles, new_styles))
            for index, style in enumerate(old_xfs):
                self.assertEqual(ET.tostring(style), ET.tostring(new_xfs[index]))
            old_strings, new_strings = (ET.fromstring(z.read(STRINGS_PATH)) for z in (original, result))
            def title_runs(sheet, strings):
                cell = sheet.find(".//m:c[@r='A1']", NS)
                return strings[int(cell.find("m:v", NS).text)].findall(".//m:rPr", NS)
            self.assertEqual([ET.tostring(run) for run in title_runs(source_sheet, old_strings)],
                             [ET.tostring(run) for run in title_runs(new_sheet, new_strings)])
            self.assertEqual(source_sheet.find(".//m:c[@r='A1']", NS).get("s"), new_sheet.find(".//m:c[@r='A1']", NS).get("s"))
        original_values = load_workbook(TEMPLATE_PATH, data_only=True).active
        generated = load_workbook(BytesIO(output), data_only=True).active
        for row in range(10, 19):
            self.assertEqual(generated[f"A{row - 1}"].value, original_values[f"A{row}"].value)

    def test_zero_one_two_and_many_lines_resize_formulas_merges_and_seal(self):
        for count in (0, 1, 2, 5):
            with self.subTest(count=count):
                content = deepcopy(self.content)
                content["lines"] = [deepcopy(self.line) for _ in range(count)]
                output = build_fabric_purchase_sheet_xlsx(self.sheet(content))
                formula = load_workbook(BytesIO(output), data_only=False).active
                cached = load_workbook(BytesIO(output), data_only=True).active
                delta, note = count - 2, count + 6
                self.assertEqual(cached["G5"].value, "订单编号")
                self.assertIn("SC2610001", cached["A1"].value)
                self.assertNotIn("HT-NOT-TASK", cached["A1"].value)
                self.assertEqual(cached[f"F{note}"].value, float(Decimal("0.9999") * count))
                self.assertEqual(formula[f"F{note}"].value, f"=SUM(F6:F{5 + count})" if count else "=0")
                for row in range(6, 6 + count):
                    self.assertEqual(formula[f"F{row}"].value, f"=ROUND(D{row}*E{row},4)")
                    self.assertEqual(cached[f"F{row}"].value, 0.9999)
                    self.assertEqual(formula[f"D{row}"].number_format, "0")
                if count:
                    self.assertEqual(cached["G6"].value, "SC2610001")
                merges = {str(item) for item in formula.merged_cells.ranges}
                self.assertEqual(f"G6:G{5 + count}" in merges, count > 1)
                self.assertIn(f"A{note}:D{note}", merges)
                self.assertEqual(cached[f"A{19 + delta}"].value, "签订日期：   2026.10.26")
                with ZipFile(TEMPLATE_PATH) as original, ZipFile(BytesIO(output)) as result:
                    source_drawing, drawing = (ET.fromstring(z.read(DRAWING_PATH)) for z in (original, result))
                    self.assertEqual(drawing.find(".//xdr:from/xdr:row", NS).text, str(17 + delta))
                    self.assertEqual(ET.tostring(source_drawing.find(".//xdr:ext", NS)), ET.tostring(drawing.find(".//xdr:ext", NS)))
                    old_sheet, new_sheet = (ET.fromstring(z.read(SHEET_PATH)) for z in (original, result))
                    def height(root, before):
                        return sum(Decimal(row.get("ht", "17.25")) for row in root.findall("m:sheetData/m:row", NS) if int(row.get("r")) < before)
                    expected_y = int(source_drawing.find(".//a:off", NS).get("y")) + int((height(new_sheet, 18 + delta) - height(old_sheet, 18)) * 12700)
                    self.assertEqual(int(drawing.find(".//a:off", NS).get("y")), expected_y)

    def test_fabric_fields_tax_freight_and_example_cleanup(self):
        output = build_fabric_purchase_sheet_xlsx(self.sheet())
        worksheet = load_workbook(BytesIO(output), data_only=True).active
        self.assertIn("品名：椅背布；规格：600D；材质：涤纶", worksheet["A6"].value)
        self.assertIn("色卡：卡A", worksheet["A6"].value)
        self.assertIn("涂层：PVC", worksheet["A6"].value)
        self.assertNotIn("门幅", worksheet["A6"].value)
        self.assertEqual(worksheet["B6"].value, "1.5")
        self.assertEqual(worksheet["C6"].value, "宝蓝 / 色号：B-01")
        self.assertIn("9%增值税", worksheet["A4"].value)
        self.assertIn("运费由甲方承担", worksheet["A4"].value)
        self.assertIn("运输费用由甲方负责", worksheet["A8"].value)
        self.assertEqual(worksheet["A7"].value, "备注：按本单确认样供货")
        with ZipFile(BytesIO(output)) as result:
            strings = result.read(STRINGS_PATH).decode("utf-8")
            for example in ("HT26148", "黄秀", "400*300D/PVC", "灰色#19"):
                self.assertNotIn(example, strings)

    def test_missing_top_level_fields_fall_back_to_snapshot_but_explicit_blank_does_not(self):
        content = deepcopy(self.content)
        line = content["lines"][0]
        for key, label, value in (("part_name", "品名", "旧品名"), ("specification", "规格", "旧规格"), ("material", "材质", "旧材质")):
            del line[key]
            line["material_attributes_snapshot"]["fields"].append({"key": key, "label": label, "value": value})
        worksheet = load_workbook(BytesIO(build_fabric_purchase_sheet_xlsx(self.sheet(content))), data_only=True).active
        self.assertIn("品名：旧品名；规格：旧规格；材质：旧材质", worksheet["A6"].value)
        line.update(specification="", material="")
        worksheet = load_workbook(BytesIO(build_fabric_purchase_sheet_xlsx(self.sheet(content))), data_only=True).active
        self.assertNotIn("旧规格", worksheet["A6"].value)
        self.assertNotIn("旧材质", worksheet["A6"].value)

    def test_missing_width_is_unknown_and_mixed_deliveries_keep_line_relationships(self):
        content = deepcopy(self.content)
        content["lines"] = [deepcopy(self.line) for _ in range(3)]
        content["lines"][0]["material_attributes_snapshot"] = {}
        content["lines"][1]["delivery_date"] = "2026-11-02"
        content["lines"][2]["delivery_date"] = ""
        worksheet = load_workbook(BytesIO(build_fabric_purchase_sheet_xlsx(self.sheet(content))), data_only=True).active
        self.assertEqual(worksheet["B6"].value, "—")
        for row, delivery in ((6, "2026-10-30"), (7, "2026-11-02"), (8, "未填写")):
            self.assertIn(f"交期：{delivery}", worksheet[f"A{row}"].value)
        self.assertIn("2026年10月30日、2026年11月2日", worksheet["A10"].value)

    def test_long_note_grows_without_clipping_and_excessive_text_is_rejected(self):
        content = deepcopy(self.content)
        content["remark"] = "本单备注" * 75
        worksheet = load_workbook(BytesIO(build_fabric_purchase_sheet_xlsx(self.sheet(content))), data_only=True).active
        self.assertEqual(worksheet["A7"].value, f"备注：{content['remark']}")
        self.assertGreater(worksheet.row_dimensions[7].height, 36)
        content["remark"] = "长" * 10000
        with self.assertRaisesRegex(FabricPurchaseSheetTemplateError, "合同备注文字过长"):
            build_fabric_purchase_sheet_xlsx(self.sheet(content))

    def test_unsupported_or_missing_business_data_is_rejected_without_source_mutation(self):
        for field, value in (("unit", "码"), ("quantity", "NaN"), ("unit_price", None), ("delivery_date", "不是日期")):
            content = deepcopy(self.content)
            content["lines"][0][field] = value
            with self.subTest(field=field), self.assertRaises(FabricPurchaseSheetTemplateError):
                build_fabric_purchase_sheet_xlsx(self.sheet(content))
        for order in ({"source_type": "PRODUCTION", "sales_order_number": "HT123"}, {"source_type": "SAMPLE", "production_task_number": "SC123"}):
            content = deepcopy(self.content)
            content["order"] = order
            with self.assertRaisesRegex(FabricPurchaseSheetTemplateError, "生产任务单号"):
                build_fabric_purchase_sheet_xlsx(self.sheet(content))
        sheet = self.sheet()
        before = deepcopy(sheet.content_snapshot)
        build_fabric_purchase_sheet_xlsx(sheet)
        self.assertEqual(sheet.content_snapshot, before)
