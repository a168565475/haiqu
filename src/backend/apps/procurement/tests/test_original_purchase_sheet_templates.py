from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.test import SimpleTestCase
from openpyxl import load_workbook

from ..original_purchase_sheet_templates import (
    TEMPLATE_PATHS, OriginalPurchaseSheetTemplateError,
    build_original_purchase_contract_xlsx, uses_original_purchase_contract_template,
)
from ..fabric_purchase_sheet_template import (
    DRAWING_PATH, NS, SHEET_PATH, STRINGS_PATH, STYLES_PATH, FabricPurchaseSheetTemplateError,
)


class OriginalPurchaseSheetTemplateTests(SimpleTestCase):
    def setUp(self):
        self.line = {
            "source_purchase_order_line_id": "line-1", "product_name": "本单椅子", "part_name": "本单零件",
            "specification": "19×0.8", "material": "PE", "unit": "个", "quantity": "12.0000",
            "unit_price": "0.3500", "amount": "4.2000", "delivery_date": "2026-11-01", "remark": "按本单要求",
            "material_attributes_snapshot": {"fields": [
                {"key": "accessory_color", "label": "颜色", "value": "蓝色"},
                {"key": "weight", "label": "克重", "value": "120G"},
                {"key": "tube_coating_color", "label": "喷塑颜色", "value": "不得导出"},
            ]},
        }
        self.content = {
            "order": {"source_type": "PRODUCTION", "production_task_number": "SC2610999", "sales_order_number": "HT-NOT-TASK"},
            "buyer": {"name": "本单采购方"}, "seller": {"name": "本单供应商"},
            "settlement_mode": "NORMAL", "purchase_date": "2026-10-26", "tax_rate": "9.0000",
            "freight_bearer": "甲方", "remark": "本单备注", "lines": [deepcopy(self.line)], "total_amount": "4.2",
        }

    def sheet(self, code, content=None):
        return SimpleNamespace(content_snapshot=deepcopy(content or self.content),
                               template_version=SimpleNamespace(template=SimpleNamespace(template_code=code)))

    def export(self, code, content=None):
        output = build_original_purchase_contract_xlsx(self.sheet(code, content))
        return output, load_workbook(BytesIO(output), data_only=True).active

    def test_exact_routing_excludes_samples_and_fabric(self):
        for code in TEMPLATE_PATHS:
            self.assertTrue(uses_original_purchase_contract_template(self.sheet(code)))
        for code in ("FABRIC", "FOAM", "AUXILIARY_OLD", None):
            self.assertFalse(uses_original_purchase_contract_template(self.sheet(code)))
        content = deepcopy(self.content)
        content["order"]["source_type"] = "SAMPLE"
        self.assertFalse(uses_original_purchase_contract_template(self.sheet("AUXILIARY", content)))

    def test_foam_keeps_original_style_images_and_dynamic_count_unit(self):
        output, ws = self.export("FOAM_CUTTING")
        self.assertEqual(ws["A6"].value, "SC2610999")
        self.assertEqual(ws["B6"].value, "本单零件")
        self.assertIn("克重：120G", ws["C6"].value)
        self.assertIn("颜色：蓝色", ws["C6"].value)
        self.assertEqual(ws["D5"].value, "数量（个）")
        self.assertEqual(ws["E5"].value, "含税单价（元/个）")
        self.assertEqual(ws["D6"].number_format, "0")
        self.assertEqual(ws["F6"].value, 4.2)
        self.assertIn("参考图", ws["A7"].value)
        self.assertIn("以本单文字规格", ws["A7"].value)
        self.assertNotIn("不得导出", " ".join(str(cell.value or "") for row in ws for cell in row))
        self.assert_preserved_resources("FOAM_CUTTING", output)

    def assert_preserved_resources(self, code, output):
        with ZipFile(TEMPLATE_PATHS[code]) as original, ZipFile(BytesIO(output)) as result:
            for name in original.namelist():
                if name.startswith("xl/media/"):
                    self.assertEqual(original.read(name), result.read(name), name)
            old_sheet, new_sheet = (ET.fromstring(z.read(SHEET_PATH)) for z in (original, result))
            for tag in ("pageMargins", "pageSetup"):
                self.assertEqual(ET.tostring(old_sheet.find(f"m:{tag}", NS)), ET.tostring(new_sheet.find(f"m:{tag}", NS)))
            old_styles, new_styles = (ET.fromstring(z.read(STYLES_PATH)) for z in (original, result))
            for tag in ("fonts", "fills", "borders"):
                self.assertEqual(ET.tostring(old_styles.find(f"m:{tag}", NS)), ET.tostring(new_styles.find(f"m:{tag}", NS)))
            old_drawing, new_drawing = (ET.fromstring(z.read(DRAWING_PATH)) for z in (original, result))
            self.assertEqual(len(old_drawing), len(new_drawing))
            for old, new in zip(old_drawing, new_drawing):
                self.assertEqual(old.find(".//a:xfrm/a:ext", NS).attrib, new.find(".//a:xfrm/a:ext", NS).attrib)

    def test_auxiliary_has_no_sequence_full_attributes_and_pure_color(self):
        output, ws = self.export("AUXILIARY")
        self.assertEqual([ws.cell(5, c).value for c in range(1, 8)],
                         ["产品名称", "产品规格", "物料属性", "颜色", "数量（个）", "含税单价（元/个）", "总金额（元）"])
        self.assertEqual(ws["A6"].value, "本单零件")
        self.assertEqual(ws["B6"].value, "19×0.8")
        self.assertIn("材质：PE", ws["C6"].value)
        self.assertIn("克重：120G", ws["C6"].value)
        self.assertEqual(ws["D6"].value, "蓝色")
        self.assertIn("甲方（需方）：本单采购方", ws["A2"].value)
        self.assertIn("乙方（供方）：本单供应商", ws["A2"].value)
        self.assertIn("开9%发票", ws["A8"].value)
        self.assertEqual(ws["E16"].value, "乙方")
        self.assert_preserved_resources("AUXILIARY", output)

    def test_counts_resize_footer_print_area_and_remove_old_samples(self):
        for code, original_count, last_row in (("FOAM_CUTTING", 1, 18), ("AUXILIARY", 2, 19), ("PLASTIC_PART", 7, 24)):
            for count in (1, 3, 9):
                content = deepcopy(self.content)
                content["lines"] = [{**deepcopy(self.line), "source_purchase_order_line_id": f"line-{i}"} for i in range(count)]
                content["total_amount"] = str(Decimal("4.2") * count)
                with self.subTest(code=code, count=count):
                    output, ws = self.export(code, content)
                    note = 6 + count
                    amount_col = {"FOAM_CUTTING": "F", "AUXILIARY": "G", "PLASTIC_PART": "I"}[code]
                    self.assertEqual(ws[f"{amount_col}{note}"].value, float(Decimal("4.2") * count))
                    self.assertIn("2026.10.26", ws[f"A{last_row + count - original_count}"].value)
                    self.assertIn("SC2610999", ws["A1"].value)
                    self.assertIn("9%增值税", ws["A4"].value)
                    self.assertIn("运输费用由甲方负责", ws[f"A{note + 1}"].value)
                    with ZipFile(BytesIO(output)) as z:
                        text = z.read(STRINGS_PATH).decode("utf-8")
                        for old in ("HT26148", "HT25148", "永康市恒天", "瑞安市塑料三厂", "武义诚喆", "EVA直立板", "19管夹棉扶手椅"):
                            self.assertNotIn(old, text)
                    self.assert_preserved_resources(code, output)
                    if code == "AUXILIARY":
                        self.assertIn(f"$G${last_row + count - original_count}", str(ws.print_area))

    def test_plastic_normal_uses_piece_amount_and_keeps_reference_picture(self):
        content = deepcopy(self.content)
        content["lines"].append({**deepcopy(self.line), "source_purchase_order_line_id": "line-2"})
        content["total_amount"] = "8.4"
        output, ws = self.export("PLASTIC_PART", content)
        self.assertEqual(ws["C5"].value, "颜色")
        self.assertEqual(ws["C6"].value, "蓝色")
        self.assertEqual(ws["F5"].value, "单位")
        self.assertEqual(ws["F6"].value, "个")
        self.assertEqual(ws["G6"].value, 12)
        self.assertEqual(ws["H6"].value, 0.35)
        self.assertEqual(ws["I6"].value, 4.2)
        self.assertEqual(ws["I7"].value, 4.2)
        self.assertNotIn("G6:G7", {str(m) for m in ws.merged_cells.ranges})
        self.assertIn("参考图", ws["A6"].value)
        formulas = load_workbook(BytesIO(output), data_only=False).active
        self.assertEqual(formulas["I6"].value, "=ROUND(G6*H6,4)")

    def test_plastic_unmerged_columns_have_real_borders_on_every_side(self):
        content = deepcopy(self.content)
        content["lines"] = [deepcopy(self.line) for _ in range(3)]
        content["total_amount"] = "12.6"
        output, ws = self.export("PLASTIC_PART", content)
        with ZipFile(TEMPLATE_PATHS["PLASTIC_PART"]) as source, ZipFile(BytesIO(output)) as result:
            old_styles, new_styles = (ET.fromstring(z.read(STYLES_PATH)).find("m:cellXfs", NS) for z in (source, result))
            old_sheet, new_sheet = (ET.fromstring(z.read(SHEET_PATH)) for z in (source, result))
            for row, source_row in ((6, 6), (7, 7), (8, 12)):
                for column in "BGHI":
                    target = ws[f"{column}{row}"]
                    for side in ("top", "bottom", "left", "right"):
                        self.assertEqual(getattr(target.border, side).style, "thin", f"{column}{row} {side}")
                    old_cell = old_sheet.find(f".//m:c[@r='{column}{source_row}']", NS)
                    new_cell = new_sheet.find(f".//m:c[@r='{column}{row}']", NS)
                    old_xf, new_xf = old_styles[int(old_cell.get("s"))], new_styles[int(new_cell.get("s"))]
                    for attribute in ("fontId", "fillId", "xfId"):
                        self.assertEqual(old_xf.get(attribute), new_xf.get(attribute))
        self.assertEqual(ws["G6"].number_format, "0")
        self.assertEqual(ws["H6"].number_format, "0.##")

    def test_webbing_uses_existing_unit_without_assuming_kilograms(self):
        for unit in ("米", "kg"):
            content = deepcopy(self.content)
            content["lines"][0]["unit"] = unit
            with self.subTest(unit=unit):
                output, ws = self.export("WEBBING_ACTUAL_WEIGHT", content)
                self.assertEqual(ws["B5"].value, "物料属性")
                self.assertIn("材质：PE", ws["B6"].value)
                self.assertEqual(ws["D6"].value, "19×0.8")
                self.assertEqual(ws["F5"].value, f"数量（{unit}）")
                self.assertEqual(ws["G5"].value, f"含税单价（元/{unit}）")
                self.assertEqual(ws["F6"].value, 12)
                self.assertEqual(ws["G6"].value, 0.35)
                self.assertEqual(ws["H6"].value, 4.2)
                self.assertEqual(ws["H7"].value, 4.2)
                self.assert_preserved_resources("WEBBING_ACTUAL_WEIGHT", output)

    def test_existing_actual_weight_mode_stays_pending_without_new_pricing_fields(self):
        content = deepcopy(self.content)
        content["settlement_mode"] = "ACTUAL_WEIGHT"
        content["lines"][0].update(unit="米", quantity="2000", unit_price="13", amount="")
        content["total_amount"] = None
        output, ws = self.export("WEBBING_ACTUAL_WEIGHT", content)
        self.assertEqual(ws["F6"].value, 2000)
        self.assertEqual(ws["G6"].value, 13)
        self.assertEqual(ws["H6"].value, "待结算")
        self.assertEqual(ws["H7"].value, "按实际送货重量结算")
        formula = load_workbook(BytesIO(output), data_only=False).active
        self.assertEqual(formula["H6"].value, "待结算")

    def test_missing_task_and_inconsistent_money_fail_explicitly(self):
        for mutation in (lambda c: c["order"].pop("production_task_number"),
                         lambda c: c["lines"][0].update(amount="999"),
                         lambda c: c["lines"][0].update(quantity="NaN"),
                         lambda c: c.update(settlement_mode="UNSUPPORTED")):
            content = deepcopy(self.content)
            mutation(content)
            with self.assertRaises(OriginalPurchaseSheetTemplateError):
                self.export("PLASTIC_PART", content)

    def test_source_snapshot_is_not_modified(self):
        sheet = self.sheet("AUXILIARY")
        before = deepcopy(sheet.content_snapshot)
        build_original_purchase_contract_xlsx(sheet)
        self.assertEqual(sheet.content_snapshot, before)

    def test_reference_pictures_fit_their_regions_and_seals_follow_footer(self):
        for code, end in (("FOAM_CUTTING", 7), ("AUXILIARY", 8), ("PLASTIC_PART", 13)):
            for count in (1, 4):
                content = deepcopy(self.content)
                content["lines"] = [deepcopy(self.line) for _ in range(count)]
                content["total_amount"] = str(Decimal("4.2") * count)
                with self.subTest(code=code, count=count):
                    output, ws = self.export(code, content)
                    with ZipFile(BytesIO(output)) as z, ZipFile(TEMPLATE_PATHS[code]) as original:
                        drawing = ET.fromstring(z.read(DRAWING_PATH))
                        source_drawing = ET.fromstring(original.read(DRAWING_PATH))
                        old_stamp_row = int(source_drawing[0].find("xdr:from/xdr:row", NS).text)
                        self.assertEqual(int(drawing[0].find("xdr:from/xdr:row", NS).text), old_stamp_row + count - (end - 6))
                        image = drawing[1]
                        self.assertTrue(image.tag.endswith("oneCellAnchor"))
                        image_start = int(image.find("xdr:from/xdr:rowOff", NS).text)
                        image_height = int(image.find("xdr:ext", NS).get("cy"))
                        if code == "PLASTIC_PART":
                            available = sum(ws.row_dimensions[row].height for row in range(6, 6 + count))
                        else:
                            available = ws.row_dimensions[6 + count].height
                        self.assertLessEqual(image_start + image_height, round(available * 12700))
                        self.assertNotIn(b"referenceTextHeight", z.read(SHEET_PATH))

    def test_mixed_units_keep_numeric_cells_and_each_unit_visible(self):
        for code, quantity_col, price_col in (("FOAM_CUTTING", "D", "E"), ("AUXILIARY", "E", "F"), ("WEBBING_ACTUAL_WEIGHT", "F", "G")):
            content = deepcopy(self.content)
            content["lines"].append({**deepcopy(self.line), "unit": "米"})
            content["total_amount"] = "8.4"
            with self.subTest(code=code):
                _, ws = self.export(code, content)
                self.assertEqual(ws[f"{quantity_col}5"].value, "数量")
                self.assertEqual(ws[f"{quantity_col}6"].value, 12)
                self.assertIn("个", ws[f"{quantity_col}6"].number_format)
                self.assertIn("米", ws[f"{quantity_col}7"].number_format)
                self.assertIn("元/米", ws[f"{price_col}7"].number_format)

    def test_mixed_deliveries_and_excessive_note_are_explicit(self):
        content = deepcopy(self.content)
        content["lines"].append({**deepcopy(self.line), "delivery_date": "2026-11-02"})
        content["total_amount"] = "8.4"
        _, ws = self.export("AUXILIARY", content)
        self.assertIn("交期：2026-11-01", ws["C6"].value)
        self.assertIn("交期：2026-11-02", ws["C7"].value)
        content["remark"] = "长备注" * 1500
        with self.assertRaisesRegex(FabricPurchaseSheetTemplateError, "合同备注文字过长"):
            self.export("AUXILIARY", content)

    def test_corrupt_original_image_structure_returns_template_error(self):
        mutations = (
            ("missing_image", None, None, None),
            ("missing_from", "xdr:from", None, None),
            ("missing_row", "xdr:from/xdr:row", None, None),
            ("invalid_row", "xdr:from/xdr:row", "text", "NaN"),
            ("negative_column", "xdr:from/xdr:col", "text", "-1"),
            ("invalid_row_offset", "xdr:from/xdr:rowOff", "text", "1.5"),
            ("missing_transform", ".//a:xfrm", None, None),
            ("missing_position", ".//a:xfrm/a:off", None, None),
            ("invalid_position", ".//a:xfrm/a:off", "x", "invalid"),
            ("missing_size", ".//a:xfrm/a:ext", None, None),
            ("zero_size", ".//a:xfrm/a:ext", "cy", "0"),
            ("missing_size_value", ".//a:xfrm/a:ext", "cx", None),
        )
        for code in TEMPLATE_PATHS:
            with ZipFile(TEMPLATE_PATHS[code]) as original:
                entries = {item.filename: original.read(item.filename) for item in original.infolist()}
            for label, path, attribute, value in mutations:
                with self.subTest(code=code, corruption=label):
                    drawing = ET.fromstring(entries[DRAWING_PATH])
                    if path is None:
                        drawing.remove(drawing[-1])
                    else:
                        node = drawing[0].find(path, NS)
                        if attribute is None:
                            parent = next(parent for parent in drawing[0].iter() if node in list(parent))
                            parent.remove(node)
                        elif attribute == "text":
                            node.text = value
                        elif value is None:
                            node.attrib.pop(attribute)
                        else:
                            node.set(attribute, value)
                    damaged = BytesIO()
                    with ZipFile(damaged, "w") as result:
                        for name, data in entries.items():
                            result.writestr(name, ET.tostring(drawing) if name == DRAWING_PATH else data)
                    damaged.seek(0)
                    with patch.dict(TEMPLATE_PATHS, {code: damaged}):
                        with self.assertRaisesRegex(OriginalPurchaseSheetTemplateError, "恢复.*原件"):
                            self.export(code)
