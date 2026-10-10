from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import base64
import tempfile
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import IntegrityError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient, APITestCase
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.boms.models import Bom, BomCreationMethod, BomLine, BomLineImage, BomSection, BomSectionType, BomStatus
from apps.customers.models import CustomerAccount, CustomerAssignment
from apps.files.models import BusinessFile, FileScanStatus
from apps.organizations.models import Employee
from apps.products.models import ProductInformation, RawMaterialInformation
from apps.production_tasks.models import (
    DeliveryMethod,
    OrderSource,
    ProductionTask,
    ProductionTaskProduct,
    ProductionTaskStatus,
)
from apps.reference_data.models import Currency, ReferenceStatus

from ..models import PurchaseOrder, Supplier, SupplierPurchaseSheet
from ..api.views import is_purchase_task_uniqueness_error


class PurchaseOrderApiTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="purchase_user")
        self.employee = Employee.objects.create(
            user=self.user,
            employee_code="PURCHASE_USER",
            display_name="采购测试员工",
        )
        self.user.user_permissions.set(
            Permission.objects.filter(content_type__app_label="procurement")
        )
        self.customer = CustomerAccount.objects.create(
            account_code="PURCHASE_CUSTOMER",
            account_name="采购测试客户",
            primary_country_code="CN",
        )
        CustomerAssignment.objects.create(
            customer_account=self.customer,
            employee=self.employee,
            assignment_role="PRIMARY_OWNER",
            effective_from=timezone.now(),
        )
        self.product = (
            ProductInformation.objects.filter(status="ACTIVE", archived_at__isnull=True)
            .exclude(product_code="")
            .select_related("unit")
            .first()
        )
        self.assertIsNotNone(self.product)
        self.inventory_material = (
            RawMaterialInformation.objects.filter(
                archived_at__isnull=True,
                unit__unit_name="PCS",
            )
            .select_related("unit")
            .first()
        )
        self.assertIsNotNone(self.inventory_material)
        self.cny = Currency.objects.filter(
            code="CNY",
            status=ReferenceStatus.ACTIVE,
            archived_at__isnull=True,
        ).first()
        self.assertIsNotNone(self.cny)
        self.sales_currency = Currency.objects.filter(
            code="USD",
            status=ReferenceStatus.ACTIVE,
            archived_at__isnull=True,
        ).first() or self.cny
        self.supplier = Supplier.objects.create(
            supplier_number="GYS90001",
            supplier_type="MATERIAL",
            supplier_categories=["TUBE"],
            supplier_name="启用管材供应商",
            payment_cycle="MONTHLY",
            phone="13800000001",
            is_disabled=False,
        )
        self.disabled_supplier = Supplier.objects.create(
            supplier_number="GYS90002",
            supplier_type="MATERIAL",
            supplier_categories=["TUBE"],
            supplier_name="已禁用管材供应商",
            payment_cycle="MONTHLY",
            phone="13800000002",
            is_disabled=True,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def make_task(self, suffix, *, status=ProductionTaskStatus.READY):
        today = timezone.localdate()
        return ProductionTask.objects.create(
            task_number=f"HT-PO-{suffix}",
            contract_number=f"HTC-PO-{suffix}",
            order_source=OrderSource.NEW,
            order_number=f"ORDER-PO-{suffix}",
            customer_account=self.customer,
            order_date=today,
            delivery_date=today + timedelta(days=30),
            salesperson=self.employee,
            currency=self.sales_currency,
            includes_tax=True,
            order_tolerance_percent=Decimal("3"),
            is_repeat_order=False,
            preproduction_sample_required=False,
            delivery_method=DeliveryMethod.SPLIT,
            status=status,
            created_by_employee=self.employee,
        )

    def test_list_orders_by_update_time_instead_of_creation_time(self):
        now = timezone.now()
        orders = []
        for index in range(2):
            task = self.make_task(f"UPDATED-{index}")
            order = PurchaseOrder.objects.create(
                order_number=f"CGDD20260910{index + 1:02d}",
                production_task=task,
                production_task_number_snapshot=task.task_number,
                sales_order_number_snapshot=task.order_number,
                customer_name_snapshot=self.customer.account_name,
                source_revision="updated-order-test",
                created_by_employee=self.employee,
            )
            PurchaseOrder.objects.filter(pk=order.pk).update(
                created_at=now - timedelta(days=2 - index),
                updated_at=now - timedelta(hours=index),
            )
            orders.append(str(order.pk))

        for ordering, expected in [("-updated_at", orders), ("updated_at", list(reversed(orders)))]:
            with self.subTest(ordering=ordering):
                response = self.client.get("/api/v1/purchase-orders/", {"ordering": ordering})
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual([row["id"] for row in response.data["data"]], expected)

        invalid = self.client.get("/api/v1/purchase-orders/", {"ordering": "updated_by_employee"})
        self.assertEqual(invalid.status_code, 400)

    def make_product_line(self, task, *, line_number=1, quantity=10):
        return ProductionTaskProduct.objects.create(
            production_task=task,
            line_number=line_number,
            product=self.product,
            product_number_snapshot=self.product.product_number,
            product_code_snapshot=self.product.product_code,
            item_number_snapshot=self.product.product_item.item_number,
            product_name_snapshot=self.product.product_name,
            specification_snapshot=self.product.specification,
            quantity=quantity,
            unit=self.product.unit,
            currency=self.sales_currency,
            unit_price=Decimal("1"),
            total_amount=Decimal(quantity),
            exchange_rate=Decimal("1"),
        )

    def make_bom(
        self,
        product_line,
        *,
        version=1,
        status=BomStatus.VERIFIED,
        section_type=BomSectionType.PROFILE,
        empty=False,
        **line_values,
    ):
        bom = Bom.objects.create(
            bom_number=(
                f"BOM-{product_line.production_task.task_number.replace('-', '')}"
                f"-P{product_line.line_number}-V{version}"
            ),
            production_task=product_line.production_task,
            product_line=product_line,
            version_number=version,
            creation_method=BomCreationMethod.BLANK,
            status=status,
            product_number_snapshot=product_line.product_number_snapshot,
            product_code_snapshot=product_line.product_code_snapshot,
            product_name_snapshot=product_line.product_name_snapshot,
            specification_snapshot=product_line.specification_snapshot,
            order_quantity=product_line.quantity,
            quantity_difference=Decimal("0"),
            section_count=1,
            line_count=0 if empty else 1,
            created_by_employee=self.employee,
        )
        section = BomSection.objects.create(
            bom=bom,
            line_number=1,
            section_type=section_type,
            title={
                BomSectionType.PROFILE: "管材类",
                BomSectionType.ACCESSORY: "配件类",
                BomSectionType.FABRIC: "面辅料类",
            }[section_type],
        )
        line = None
        if not empty:
            defaults = {
                "category": "五金件" if section_type == BomSectionType.ACCESSORY else "型材",
                "material": "铁",
                "part_name": "采购测试物料",
                "specification": "测试规格",
                "unit": "PCS",
                "branch_count": "12",
            }
            defaults.update(line_values)
            line = BomLine.objects.create(
                section=section,
                line_number=1,
                **defaults,
            )
            line.raw_values = {
                "source_inventory_material_id": str(self.inventory_material.pk)
            }
            line.save(update_fields=["raw_values", "updated_at"])
        return bom, line

    def get_source(self, task, bom=None):
        params = {"production_task_id": str(task.pk)}
        if bom is not None:
            params["bom_id"] = str(bom.pk)
        return self.client.get(
            "/api/v1/purchase-orders/source/",
            params,
        )

    @staticmethod
    def create_payload(task, source_data, **line_overrides):
        lines = []
        for source_line in source_data["lines"]:
            item = {"bom_line_id": source_line["bom_line_id"]}
            item.update(line_overrides)
            lines.append(item)
        return {
            "production_task_id": str(task.pk),
            "bom_id": source_data["selected_bom_id"],
            "source_revision": source_data["source_revision"],
            "lines": lines,
        }

    def create_generated_supplier_sheet(self, suffix):
        task = self.make_task(suffix)
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.supplier.pk),
            delivery_date=(timezone.localdate() + timedelta(days=10)).isoformat(),
            actual_purchase_quantity="12.5000",
            unit_price="2.0000",
        )
        payload["currency_id"] = str(self.cny.pk)
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        order_id = created.data["data"]["id"]
        confirmed = self.client.post(f"/api/v1/purchase-orders/{order_id}/confirm/", {}, format="json")
        self.assertEqual(confirmed.status_code, 200, confirmed.data)
        options = self.client.get(f"/api/v1/purchase-orders/{order_id}/sheet-options/")
        self.assertEqual(options.status_code, 200, options.data)
        generated = self.client.post(
            f"/api/v1/purchase-orders/{order_id}/generate-sheets/",
            {
                "assignments": [{
                    "supplier_id": str(self.supplier.pk),
                    "template_version_id": options.data["data"]["templates"][0]["version_id"],
                }]
            },
            format="json",
        )
        self.assertEqual(generated.status_code, 201, generated.data)
        return SupplierPurchaseSheet.objects.select_related("purchase_order", "template_version").get(
            pk=generated.data["data"][0]["id"]
        )

    def test_source_and_options_return_contract_and_cny_default(self):
        task = self.make_task("SOURCE")
        product_line = self.make_product_line(task)
        self.make_bom(product_line, branch_count="1,234.5")

        source = self.get_source(task)
        options = self.client.get("/api/v1/purchase-orders/options/")

        self.assertEqual(source.status_code, 200, source.data)
        source_data = source.data["data"]
        self.assertTrue(source_data["can_create"])
        self.assertFalse(source_data["requires_bom_selection"])
        self.assertTrue(source_data["has_creatable_bom"])
        self.assertEqual(source_data["selected_bom_id"], source_data["bom_options"][0]["bom_id"])
        self.assertRegex(source_data["source_revision"], r"^[0-9a-f]{64}$")
        self.assertEqual(source_data["preview_lines"], [])
        self.assertEqual(source_data["lines"][0]["required_quantity"], "1234.5000")
        self.assertEqual(source_data["lines"][0]["bom_status"], BomStatus.VERIFIED)
        self.assertEqual(source_data["lines"][0]["source_quantity_field"], "branch_count")
        self.assertEqual(source_data["lines"][0]["source_unit_text"], "PCS")
        self.assertEqual(source_data["lines"][0]["quantity_parse_status"], "READY")
        self.assertEqual(source_data["lines"][0]["suggested_supplier_category"], "TUBE")
        self.assertEqual(source_data["lines"][0]["category_source"], "SECTION")
        source_product = source_data["production_task"]["products"][0]
        self.assertEqual(source_product["product_number"], product_line.product_number_snapshot)
        self.assertEqual(source_product["product_code"], product_line.product_code_snapshot)
        self.assertEqual(source_product["item_number"], product_line.item_number_snapshot)
        self.assertEqual(source_product["product_name"], product_line.product_name_snapshot)
        self.assertEqual(source_product["specification"], product_line.specification_snapshot)
        self.assertEqual(source_product["color_card"], product_line.color_card_snapshot)
        self.assertEqual(source_product["color_name"], product_line.color_name_snapshot)
        self.assertEqual(source_product["color_code"], product_line.color_code_snapshot)
        self.assertEqual(source_product["order_quantity"], "10")
        self.assertEqual(source_product["unit_name"], product_line.unit.unit_name)

        self.assertEqual(options.status_code, 200, options.data)
        options_data = options.data["data"]
        self.assertEqual(options_data["default_currency_id"], str(self.cny.pk))
        self.assertEqual(
            [item["value"] for item in options_data["statuses"]],
            ["DRAFT", "SUBMITTED", "PURCHASING", "COMPLETED"],
        )
        self.assertEqual(len(options_data["created_date_presets"]), 3)
        task_option = next(item for item in options_data["production_tasks"] if item["id"] == str(task.pk))
        self.assertTrue(task_option["can_create"])
        self.assertNotIn("products", task_option)
        self.assertFalse(any(item["id"] == str(self.disabled_supplier.pk) for item in options_data["suppliers"]))

    def test_purchase_source_excludes_cutting_detail_section(self):
        task = self.make_task("NO-CUTTING")
        product_line = self.make_product_line(task)
        bom, purchasable_line = self.make_bom(product_line)
        cutting_section = BomSection.objects.create(
            bom=bom,
            line_number=2,
            section_type=BomSectionType.FABRIC,
            title="裁剪明细",
        )
        cutting_line = BomLine.objects.create(
            section=cutting_section,
            line_number=1,
            category="裁剪",
            material="面料",
            part_name="裁剪片",
            specification="测试裁剪尺寸",
            unit="PCS",
            branch_count="20",
        )

        first_source = self.get_source(task)

        self.assertEqual(first_source.status_code, 200, first_source.data)
        first_data = first_source.data["data"]
        self.assertTrue(first_data["can_create"])
        self.assertEqual(first_data["bom_options"][0]["line_count"], 1)
        self.assertEqual(
            [line["bom_line_id"] for line in first_data["lines"]],
            [str(purchasable_line.pk)],
        )
        self.assertNotIn(str(cutting_line.pk), [line["bom_line_id"] for line in first_data["lines"]])

        original_revision = first_data["source_revision"]
        cutting_line.part_name = "修改后的裁剪片"
        cutting_line.save(update_fields=["part_name", "updated_at"])
        second_source = self.get_source(task)

        self.assertEqual(second_source.status_code, 200, second_source.data)
        self.assertEqual(second_source.data["data"]["source_revision"], original_revision)

    def test_integrity_error_mapping_only_accepts_production_task_uniqueness(self):
        task_unique = IntegrityError(
            "UNIQUE constraint failed: procurement_purchase_order.production_task_id"
        )
        number_unique = IntegrityError(
            "UNIQUE constraint failed: procurement_purchase_order.order_number"
        )

        self.assertTrue(is_purchase_task_uniqueness_error(task_unique))
        self.assertFalse(is_purchase_task_uniqueness_error(number_unique))

    def test_create_recalculates_amount_and_list_filters_and_stats(self):
        task = self.make_task("CREATE")
        product_line = self.make_product_line(task)
        product_line.color_card_snapshot = "志龙"
        product_line.color_name_snapshot = "橙色"
        product_line.color_code_snapshot = "ORANGE-01"
        product_line.save(
            update_fields=["color_card_snapshot", "color_name_snapshot", "color_code_snapshot"]
        )
        self.make_bom(product_line, branch_count="5")
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.supplier.pk),
            delivery_date=(timezone.localdate() + timedelta(days=10)).isoformat(),
            actual_purchase_quantity="3.2500",
            unit_price="2.4000",
            amount="999999.0000",
        )
        payload.update({"currency_id": str(self.cny.pk), "is_mailed": True, "total_amount": "1"})

        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")

        self.assertEqual(created.status_code, 201, created.data)
        data = created.data["data"]
        self.assertRegex(data["order_number"], r"^CGDD[0-9]{8}[0-9]{2}$")
        self.assertEqual(data["status"], "DRAFT")
        self.assertEqual(Decimal(data["total_amount"]), Decimal("7.8000"))
        self.assertEqual(Decimal(data["lines"][0]["amount"]), Decimal("7.8000"))
        self.assertEqual(data["lines"][0]["supplier_name"], self.supplier.supplier_name)
        self.assertTrue(
            AuditLog.objects.filter(
                object_type="PURCHASE_ORDER",
                object_id=data["id"],
                action="CREATED",
            ).exists()
        )

        listed = self.client.get(
            "/api/v1/purchase-orders/",
            {
                "supplier_id": str(self.supplier.pk),
                "created_after": timezone.localdate().isoformat(),
            },
        )
        stats = self.client.get("/api/v1/purchase-orders/stats/")
        self.assertEqual(listed.status_code, 200, listed.data)
        self.assertEqual(listed.data["meta"]["total"], 1)
        self.assertEqual(listed.data["data"][0]["supplier_assigned_line_count"], 1)
        self.assertEqual(listed.data["data"][0]["supplier_count"], 1)
        self.assertEqual(listed.data["data"][0]["order_date"], task.order_date.isoformat())
        self.assertEqual(listed.data["data"][0]["delivery_date"], task.delivery_date.isoformat())
        self.assertEqual(
            listed.data["data"][0]["product_summary"],
            [
                {
                    "product_line_id": str(product_line.pk),
                    "product_name": product_line.product_name_snapshot,
                    "specification": product_line.specification_snapshot,
                    "color_card": "志龙",
                    "color_name": "橙色",
                    "color_code": "ORANGE-01",
                }
            ],
        )
        self.assertEqual(stats.data["data"]["total"], 1)
        self.assertEqual(stats.data["data"]["amounts"][0]["currency_code"], "CNY")
        self.assertEqual(Decimal(stats.data["data"]["amounts"][0]["total_amount"]), Decimal("7.8000"))

        bad_supplier = self.client.get("/api/v1/purchase-orders/", {"supplier_id": "bad"})
        bad_date = self.client.get("/api/v1/purchase-orders/", {"created_after": "2026-9-5"})
        impossible_date = self.client.get(
            "/api/v1/purchase-orders/",
            {"created_after": "2026-02-30"},
        )
        self.assertEqual(bad_supplier.status_code, 400)
        self.assertEqual(bad_date.status_code, 400)
        self.assertEqual(impossible_date.status_code, 400)

    def test_purchase_order_number_uses_date_and_two_digit_daily_sequence(self):
        created_numbers = []
        for suffix in ("NUMBER-1", "NUMBER-2"):
            task = self.make_task(suffix)
            product_line = self.make_product_line(task)
            self.make_bom(product_line, branch_count="1")
            source_data = self.get_source(task).data["data"]

            response = self.client.post(
                "/api/v1/purchase-orders/",
                self.create_payload(task, source_data),
                format="json",
            )

            self.assertEqual(response.status_code, 201, response.data)
            created_numbers.append(response.data["data"]["order_number"])

        today_prefix = f"CGDD{timezone.localdate():%Y%m%d}"
        self.assertEqual(created_numbers, [f"{today_prefix}01", f"{today_prefix}02"])

    def test_currency_is_required_before_header_total_can_be_calculated(self):
        task = self.make_task("NO-CURRENCY")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            actual_purchase_quantity="2",
            unit_price="3",
        )

        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")

        self.assertEqual(created.status_code, 201, created.data)
        self.assertIsNone(created.data["data"]["total_amount"])
        self.assertEqual(Decimal(created.data["data"]["lines"][0]["amount"]), Decimal("6"))

    def test_disabled_supplier_and_overflow_amount_are_rejected(self):
        task = self.make_task("DISABLED")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        disabled_payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.disabled_supplier.pk),
        )

        disabled = self.client.post("/api/v1/purchase-orders/", disabled_payload, format="json")

        self.assertEqual(disabled.status_code, 400, disabled.data)
        self.assertEqual(PurchaseOrder.objects.count(), 0)

        overflow_payload = self.create_payload(
            task,
            source_data,
            actual_purchase_quantity="99999999999999.9999",
            unit_price="99999999999999.9999",
        )
        overflow = self.client.post("/api/v1/purchase-orders/", overflow_payload, format="json")
        self.assertEqual(overflow.status_code, 400, overflow.data)
        self.assertEqual(PurchaseOrder.objects.count(), 0)

    def test_product_without_bom_does_not_block_selected_valid_bom(self):
        task = self.make_task("MISSING")
        first_product = self.make_product_line(task, line_number=1)
        first_bom, first_line = self.make_bom(first_product)
        self.make_product_line(task, line_number=2)

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        self.assertTrue(source.data["data"]["can_create"])
        self.assertEqual(source.data["data"]["selected_bom_id"], str(first_bom.pk))
        self.assertEqual(len(source.data["data"]["lines"]), 1)
        self.assertEqual(source.data["data"]["lines"][0]["bom_status"], BomStatus.VERIFIED)
        self.assertEqual(
            [item["line_number"] for item in source.data["data"]["production_task"]["products"]],
            [1, 2],
        )
        self.assertNotIn("BOM_MISSING", [item["code"] for item in source.data["data"]["blockers"]])
        missing_option = next(item for item in source.data["data"]["bom_options"] if item["bom_id"] is None)
        self.assertIn("BOM_MISSING", [item["code"] for item in missing_option["blockers"]])

        payload = self.create_payload(task, source.data["data"])
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["data"]["lines"][0]["bom_line_id"], str(first_line.pk))

    def test_multiple_products_require_bom_selection_and_selected_valid_bom_is_editable(self):
        task = self.make_task("MULTI")
        pending_product = self.make_product_line(task, line_number=1)
        pending_bom, _ = self.make_bom(
            pending_product,
            status=BomStatus.PENDING_CONFIRMATION,
            empty=True,
        )
        valid_product = self.make_product_line(task, line_number=2)
        valid_bom, valid_line = self.make_bom(valid_product, branch_count="8")

        unselected = self.get_source(task)

        self.assertEqual(unselected.status_code, 200, unselected.data)
        self.assertFalse(unselected.data["data"]["can_create"])
        self.assertTrue(unselected.data["data"]["requires_bom_selection"])
        self.assertTrue(unselected.data["data"]["has_creatable_bom"])
        self.assertIsNone(unselected.data["data"]["selected_bom_id"])
        self.assertEqual(len(unselected.data["data"]["bom_options"]), 2)

        selected = self.get_source(task, valid_bom)

        self.assertEqual(selected.status_code, 200, selected.data)
        selected_data = selected.data["data"]
        self.assertTrue(selected_data["can_create"])
        self.assertEqual(selected_data["selected_bom_id"], str(valid_bom.pk))
        self.assertEqual([item["bom_line_id"] for item in selected_data["lines"]], [str(valid_line.pk)])
        self.assertNotIn("BOM_NOT_VERIFIED", [item["code"] for item in selected_data["blockers"]])

        pending = self.get_source(task, pending_bom)
        self.assertFalse(pending.data["data"]["can_create"])
        self.assertIn("BOM_NOT_VERIFIED", [item["code"] for item in pending.data["data"]["blockers"]])
        self.assertIn("BOM_EMPTY", [item["code"] for item in pending.data["data"]["blockers"]])

        created = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(
                task,
                selected_data,
                supplier_id=str(self.supplier.pk),
                delivery_date=(timezone.localdate() + timedelta(days=5)).isoformat(),
                actual_purchase_quantity="8",
                unit_price="2.5",
            ),
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["data"]["lines"][0]["supplier_name"], self.supplier.supplier_name)

    def test_latest_unverified_bom_never_falls_back_to_old_verified_version(self):
        task = self.make_task("LATEST")
        product_line = self.make_product_line(task)
        old_bom, old_line = self.make_bom(product_line, version=1, status=BomStatus.VERIFIED)
        latest_bom, latest_line = self.make_bom(
            product_line,
            version=2,
            status=BomStatus.PENDING_CONFIRMATION,
            branch_count="99",
        )

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        self.assertFalse(source.data["data"]["can_create"])
        self.assertEqual(source.data["data"]["lines"], [])
        self.assertEqual(len(source.data["data"]["preview_lines"]), 1)
        self.assertEqual(source.data["data"]["preview_lines"][0]["bom_line_id"], str(latest_line.pk))
        self.assertEqual(
            source.data["data"]["preview_lines"][0]["bom_status"],
            BomStatus.PENDING_CONFIRMATION,
        )
        blocker = next(
            item for item in source.data["data"]["blockers"] if item["code"] == "BOM_NOT_VERIFIED"
        )
        self.assertEqual(blocker["bom_id"], str(latest_bom.pk))
        self.assertNotEqual(blocker["bom_id"], str(old_bom.pk))

        old_version = self.get_source(task, old_bom)
        self.assertEqual(old_version.status_code, 400, old_version.data)

        payload = {
            "production_task_id": str(task.pk),
            "bom_id": str(latest_bom.pk),
            "source_revision": source.data["data"]["source_revision"],
            "lines": [{"bom_line_id": str(latest_line.pk)}],
        }
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 409, created.data)
        self.assertFalse(PurchaseOrder.objects.filter(production_task=task).exists())
        self.assertTrue(BomLine.objects.filter(pk=old_line.pk).exists())

    def test_pending_validation_task_cannot_prepare_purchase_order(self):
        task = self.make_task("PENDING-DRAFT", status=ProductionTaskStatus.PENDING_VALIDATION)
        product_line = self.make_product_line(task)
        self.make_bom(product_line)

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        self.assertFalse(source.data["data"]["can_create"])
        self.assertIn("TASK_NOT_READY", [item["code"] for item in source.data["data"]["blockers"]])
        blocked_source = {**source.data["data"], "lines": source.data["data"]["preview_lines"]}
        payload = self.create_payload(task, blocked_source)
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 409, created.data)
        self.assertIn("PMC", str(created.data))
        self.assertFalse(PurchaseOrder.objects.filter(production_task=task).exists())

    def test_unresolved_material_identity_blocks_purchase_order_creation(self):
        task = self.make_task("NO-MATERIAL-ID")
        product_line = self.make_product_line(task)
        _, bom_line = self.make_bom(product_line)
        bom_line.raw_values = {}
        bom_line.save(update_fields=["raw_values", "updated_at"])

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        self.assertFalse(source.data["data"]["can_create"])
        blocker = next(
            item
            for item in source.data["data"]["blockers"]
            if item["code"] == "BOM_MATERIAL_IDENTITY_INVALID"
        )
        self.assertEqual(blocker["invalid_line_count"], 1)
        self.assertEqual(
            blocker["invalid_lines"][0]["reason"],
            "MATERIAL_IDENTITY_UNRESOLVED",
        )

    def test_invalid_quantity_is_visible_but_does_not_block_manual_purchase(self):
        task = self.make_task("INVALID-QTY")
        product_line = self.make_product_line(task)
        self.make_bom(product_line, branch_count="1,2,3")

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        self.assertTrue(source.data["data"]["can_create"])
        line = source.data["data"]["lines"][0]
        self.assertIsNone(line["required_quantity"])
        self.assertEqual(line["source_quantity_text"], "1,2,3")
        self.assertEqual(line["quantity_parse_status"], "INVALID")

    def test_derived_quantity_multiplies_full_precision_and_unknown_suggestion_is_null(self):
        task = self.make_task("DERIVED")
        product_line = self.make_product_line(task, quantity=10000)
        self.make_bom(
            product_line,
            section_type=BomSectionType.ACCESSORY,
            category="",
            material="",
            part_name="未知采购物料",
            raw_material="",
            branch_count="",
            budget_quantity="",
            total_consumption="",
            unit_consumption="0.00006",
        )

        source = self.get_source(task)

        self.assertEqual(source.status_code, 200, source.data)
        line = source.data["data"]["lines"][0]
        self.assertEqual(line["required_quantity"], "0.6000")
        self.assertEqual(line["source_quantity_field"], "unit_consumption*order_quantity")
        self.assertEqual(line["source_quantity_text"], "0.00006 × 10000.0000")
        self.assertIsNone(line["suggested_supplier_category"])
        self.assertEqual(line["category_source"], "UNKNOWN")

    def test_long_invalid_derived_quantity_keeps_full_audit_text_and_can_be_saved(self):
        task = self.make_task("LONG-SOURCE")
        product_line = self.make_product_line(task)
        invalid_source_text = "x" * 100
        self.make_bom(
            product_line,
            section_type=BomSectionType.ACCESSORY,
            branch_count="",
            budget_quantity="",
            total_consumption="",
            unit_consumption=invalid_source_text,
        )
        source_data = self.get_source(task).data["data"]

        line = source_data["lines"][0]
        self.assertEqual(line["quantity_parse_status"], "INVALID")
        self.assertEqual(line["source_quantity_text"], f"{invalid_source_text} × 10.0000")

        created = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(task, source_data),
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(
            created.data["data"]["lines"][0]["source_quantity_text"],
            f"{invalid_source_text} × 10.0000",
        )

    def test_source_revision_detects_line_snapshot_change(self):
        task = self.make_task("REVISION")
        product_line = self.make_product_line(task)
        _, bom_line = self.make_bom(product_line)
        old_source = self.get_source(task).data["data"]

        # QuerySet.update 不触发 auto_now；来源修订仍应基于真实行快照发现变化。
        BomLine.objects.filter(pk=bom_line.pk).update(part_name="已变更采购物料")
        new_source = self.get_source(task).data["data"]

        self.assertNotEqual(old_source["source_revision"], new_source["source_revision"])
        stale = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(task, old_source),
            format="json",
        )
        self.assertEqual(stale.status_code, 409, stale.data)
        self.assertEqual(stale.data["error"]["code"], "PURCHASE_SOURCE_CHANGED")
        self.assertEqual(
            stale.data["error"]["meta"]["current_source_revision"],
            new_source["source_revision"],
        )

    def test_archived_purchase_order_still_blocks_a_second_order(self):
        task = self.make_task("ARCHIVED")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        created = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(task, source_data),
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        PurchaseOrder.objects.filter(pk=created.data["data"]["id"]).update(
            archived_at=timezone.now(),
            archive_reason="测试归档",
        )

        source = self.get_source(task)

        self.assertFalse(source.data["data"]["can_create"])
        self.assertIn(
            "PURCHASE_ORDER_EXISTS",
            [item["code"] for item in source.data["data"]["blockers"]],
        )
        recreated = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(task, source_data),
            format="json",
        )
        self.assertEqual(recreated.status_code, 409, recreated.data)
        self.assertEqual(PurchaseOrder.objects.filter(production_task=task).count(), 1)

    def test_purchase_endpoints_hide_out_of_scope_production_tasks(self):
        task = self.make_task("SCOPE")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]

        outsider_user = get_user_model().objects.create_user(username="purchase_outsider")
        Employee.objects.create(
            user=outsider_user,
            employee_code="PURCHASE_OUTSIDER",
            display_name="采购范围外员工",
        )
        outsider_user.user_permissions.set(
            Permission.objects.filter(content_type__app_label="procurement")
        )
        self.client.force_authenticate(outsider_user)

        source = self.get_source(task)
        created = self.client.post(
            "/api/v1/purchase-orders/",
            self.create_payload(task, source_data),
            format="json",
        )
        listed = self.client.get("/api/v1/purchase-orders/")
        stats = self.client.get("/api/v1/purchase-orders/stats/")
        options = self.client.get("/api/v1/purchase-orders/options/")

        self.assertEqual(source.status_code, 404, source.data)
        self.assertEqual(created.status_code, 404, created.data)
        self.assertEqual(listed.data["meta"]["total"], 0)
        self.assertEqual(stats.data["data"]["total"], 0)
        self.assertFalse(
            any(item["id"] == str(task.pk) for item in options.data["data"]["production_tasks"])
        )

    def test_submit_purchase_order_enables_material_receipt_gate(self):
        task = self.make_task("SUBMIT")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.supplier.pk),
            delivery_date=(timezone.localdate() + timedelta(days=7)).isoformat(),
            actual_purchase_quantity="12.5000",
            unit_price="2.0000",
        )
        payload["currency_id"] = str(self.cny.pk)
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        order_id = created.data["data"]["id"]

        submitted = self.client.post(f"/api/v1/purchase-orders/{order_id}/confirm/", {}, format="json")
        self.assertEqual(submitted.status_code, 200, submitted.data)
        self.assertEqual(submitted.data["data"]["status"], "SUBMITTED")
        self.assertEqual(submitted.data["data"]["status_display"], "已确认")
        repeated = self.client.post(f"/api/v1/purchase-orders/{order_id}/confirm/", {}, format="json")
        self.assertEqual(repeated.status_code, 409, repeated.data)

    def test_delete_purchase_order_without_downstream_records_and_recreate(self):
        task = self.make_task("DELETE")
        product_line = self.make_product_line(task)
        self.make_bom(product_line)
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.supplier.pk),
            delivery_date=(timezone.localdate() + timedelta(days=7)).isoformat(),
            actual_purchase_quantity="12.5000",
            unit_price="2.0000",
        )
        payload["currency_id"] = str(self.cny.pk)
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        order_id = created.data["data"]["id"]
        confirmed = self.client.post(f"/api/v1/purchase-orders/{order_id}/confirm/", {}, format="json")
        self.assertEqual(confirmed.status_code, 200, confirmed.data)

        deleted = self.client.delete(f"/api/v1/purchase-orders/{order_id}/")

        self.assertEqual(deleted.status_code, 204, getattr(deleted, "data", None))
        self.assertFalse(PurchaseOrder.objects.filter(pk=order_id).exists())
        self.assertTrue(
            AuditLog.objects.filter(
                object_type="PURCHASE_ORDER",
                object_id=order_id,
                action="DELETED",
            ).exists()
        )
        refreshed_source = self.get_source(task)
        self.assertEqual(refreshed_source.status_code, 200, refreshed_source.data)
        self.assertTrue(refreshed_source.data["data"]["can_create"])

    def test_confirm_generate_edit_preview_publish_and_version_purchase_sheet(self):
        task = self.make_task("SHEET")
        product_line = self.make_product_line(task)
        _, bom_line = self.make_bom(product_line)
        bom_image_file = BusinessFile.objects.create(
            uploaded_by_employee=self.employee,
            original_name="bom-source.png",
            storage_key="bom-line-images/bom-source.png",
            content_type="image/png",
            size_bytes=1,
            content_hash="bom-source-image",
            scan_status=FileScanStatus.CLEAN,
            scanned_at=timezone.now(),
        )
        BomLineImage.objects.create(
            bom_line=bom_line,
            business_file=bom_image_file,
            position=1,
        )
        source_data = self.get_source(task).data["data"]
        payload = self.create_payload(
            task,
            source_data,
            supplier_id=str(self.supplier.pk),
            delivery_date=(timezone.localdate() + timedelta(days=9)).isoformat(),
            actual_purchase_quantity="12.5000",
            unit_price="2.0000",
        )
        payload["currency_id"] = str(self.cny.pk)
        created = self.client.post("/api/v1/purchase-orders/", payload, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        order_id = created.data["data"]["id"]

        confirmed = self.client.post(f"/api/v1/purchase-orders/{order_id}/confirm/", {}, format="json")
        self.assertEqual(confirmed.status_code, 200, confirmed.data)
        confirmed_line = confirmed.data["data"]["lines"][0]
        edited_quantity = "13.0000"
        edited_price = "2.5000"
        expected_delivery_date = timezone.localdate() + timedelta(days=10)
        order_update_payload = {
            "is_mailed": True,
            "lines": [{
                "id": confirmed_line["id"],
                "supplier_id": str(self.supplier.pk),
                "delivery_date": expected_delivery_date.isoformat(),
                "actual_purchase_quantity": edited_quantity,
                "unit_price": edited_price,
            }],
        }
        edited_order = self.client.patch(
            f"/api/v1/purchase-orders/{order_id}/",
            order_update_payload,
            format="json",
            HTTP_IF_MATCH=confirmed.headers["ETag"],
        )
        self.assertEqual(edited_order.status_code, 200, edited_order.data)
        self.assertEqual(edited_order.data["data"]["status"], "SUBMITTED")
        self.assertEqual(edited_order.data["data"]["lines"][0]["actual_purchase_quantity"], edited_quantity)
        self.assertEqual(edited_order.data["data"]["lines"][0]["amount"], "32.5000")
        self.assertTrue(edited_order.data["data"]["is_mailed"])
        options = self.client.get(f"/api/v1/purchase-orders/{order_id}/sheet-options/")
        self.assertEqual(options.status_code, 200, options.data)
        self.assertEqual(len(options.data["data"]["suppliers"]), 1)
        template_version_id = options.data["data"]["templates"][0]["version_id"]

        generated = self.client.post(
            f"/api/v1/purchase-orders/{order_id}/generate-sheets/",
            {
                "assignments": [
                    {
                        "supplier_id": str(self.supplier.pk),
                        "template_version_id": template_version_id,
                    }
                ]
            },
            format="json",
        )
        self.assertEqual(generated.status_code, 201, generated.data)
        self.assertEqual(len(generated.data["data"]), 1)
        sheet_id = generated.data["data"][0]["id"]
        self.assertRegex(generated.data["data"][0]["sheet_number"], r"^CGDD[0-9]{10}$")
        self.assertEqual(len(generated.data["data"][0]["line_images"]), 1)
        self.assertEqual(
            generated.data["data"][0]["line_images"][0]["image_file_id"],
            str(bom_image_file.pk),
        )
        supplier_filter_options = self.client.get("/api/v1/supplier-purchase-sheets/options/")
        self.assertEqual(supplier_filter_options.status_code, 200, supplier_filter_options.data)
        self.assertEqual(supplier_filter_options.data["data"]["suppliers"][0]["id"], str(self.supplier.pk))
        filtered_sheets = self.client.get(
            f"/api/v1/supplier-purchase-sheets/?supplier_id={self.supplier.pk}"
        )
        self.assertEqual(filtered_sheets.status_code, 200, filtered_sheets.data)
        self.assertEqual(len(filtered_sheets.data["data"]), 1)

        blocked_order_update = self.client.patch(
            f"/api/v1/purchase-orders/{order_id}/",
            order_update_payload,
            format="json",
            HTTP_IF_MATCH=edited_order.headers["ETag"],
        )
        self.assertEqual(blocked_order_update.status_code, 409, blocked_order_update.data)

        detail = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        content = detail.data["data"]["content_snapshot"]
        self.assertIn("payment_and_delivery", content)
        self.assertIn("quality_standards", content)
        self.assertIn("delivery_handover", content)
        self.assertIn("special_instructions", content)
        self.assertIn("dispute_resolution", content)
        self.assertIn("effectiveness", content)
        self.assertNotIn("{{采购交期}}", content["payment_and_delivery"])
        self.assertIn(
            f"{expected_delivery_date.year}年{expected_delivery_date.month}月{expected_delivery_date.day}日",
            content["payment_and_delivery"],
        )
        self.assertNotIn("terms", content)
        self.assertGreaterEqual(len(content["quality_standards"]), 1)
        content["remark"] = "采购表编辑测试"
        updated = self.client.patch(
            f"/api/v1/supplier-purchase-sheets/{sheet_id}/",
            {"content": content},
            format="json",
            HTTP_IF_MATCH=detail.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["data"]["content_snapshot"]["remark"], "采购表编辑测试")

        xlsx_response = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/download-xlsx/")
        self.assertEqual(xlsx_response.status_code, 200)
        self.assertEqual(
            xlsx_response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(xlsx_response.content), data_only=False)
        worksheet = workbook.active
        self.assertEqual(worksheet.page_setup.orientation, "portrait")
        self.assertIn("订单编号", worksheet["A1"].value)
        self.assertEqual(worksheet["A5"].value, "序号")
        self.assertEqual(worksheet["A6"].value, 1)
        section_values = [str(worksheet.cell(row=row, column=1).value or "") for row in range(1, worksheet.max_row + 1)]
        self.assertTrue(any(value.startswith("二、支付方式和交期") for value in section_values))
        self.assertTrue(any(value.startswith("三、产品质量标准") for value in section_values))
        self.assertTrue(any(value.startswith("七、合同生效") for value in section_values))

        preview = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/preview-pdf/")
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview["Content-Type"], "application/pdf")

        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            image_path = Path(media_root) / "purchase-sheet-images" / "test.png"
            image_path.parent.mkdir(parents=True, exist_ok=True)
            image_path.write_bytes(base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            ))
            image_file = BusinessFile.objects.create(
                uploaded_by_employee=self.employee,
                original_name="test.png",
                storage_key="purchase-sheet-images/test.png",
                content_type="image/png",
                size_bytes=image_path.stat().st_size,
                content_hash="test-image",
                scan_status=FileScanStatus.CLEAN,
                scanned_at=timezone.now(),
            )
            with patch(
                "apps.procurement.purchase_sheet_services.save_scanned_image",
                return_value=image_file,
            ):
                image_upload = self.client.post(
                    f"/api/v1/supplier-purchase-sheets/{sheet_id}/line-images/{content['lines'][0]['source_purchase_order_line_id']}/",
                    {"file": SimpleUploadedFile("test.png", b"unused", content_type="image/png")},
                    format="multipart",
                    HTTP_IF_MATCH=updated.headers["ETag"],
                )
            self.assertEqual(image_upload.status_code, 200, image_upload.data)
            self.assertEqual(len(image_upload.data["data"]["line_images"]), 1)
            image_response = self.client.get(
                f"/api/v1/supplier-purchase-sheets/{sheet_id}/line-images/{content['lines'][0]['source_purchase_order_line_id']}/"
            )
            self.assertEqual(image_response.status_code, 200)
            self.assertEqual(image_response["Content-Type"], "image/png")
            with patch(
                "apps.procurement.purchase_sheet_services.save_scanned_image",
                return_value=image_file,
            ):
                seal_upload = self.client.post(
                    f"/api/v1/supplier-purchase-sheets/{sheet_id}/seal-image/",
                    {"file": SimpleUploadedFile("seal.png", b"unused", content_type="image/png")},
                    format="multipart",
                    HTTP_IF_MATCH=image_upload.headers["ETag"],
                )
            self.assertEqual(seal_upload.status_code, 200, seal_upload.data)
            self.assertEqual(seal_upload.data["data"]["seal_image"]["image_file_id"], str(image_file.pk))
            seal_response = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/seal-image/")
            self.assertEqual(seal_response.status_code, 200)
            self.assertEqual(seal_response["Content-Type"], "image/png")
            published = self.client.post(
                f"/api/v1/supplier-purchase-sheets/{sheet_id}/publish/",
                {},
                format="json",
                HTTP_IF_MATCH=seal_upload.headers["ETag"],
            )
            self.assertEqual(published.status_code, 200, published.data)
            self.assertEqual(published.data["data"]["status"], "PUBLISHED")
            self.assertTrue(published.data["data"]["pdf_ready"])
            downloaded = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/download-pdf/")
            self.assertEqual(downloaded.status_code, 200)

        new_version = self.client.post(
            f"/api/v1/supplier-purchase-sheets/{sheet_id}/new-version/",
            {},
            format="json",
            HTTP_IF_MATCH=published.headers["ETag"],
        )
        self.assertEqual(new_version.status_code, 201, new_version.data)
        self.assertEqual(new_version.data["data"]["version_number"], 2)
        self.assertEqual(new_version.data["data"]["status"], "DRAFT")
        self.assertEqual(len(new_version.data["data"]["line_images"]), 1)
        self.assertEqual(new_version.data["data"]["seal_image"]["image_file_id"], str(image_file.pk))
        self.assertEqual(
            SupplierPurchaseSheet.objects.filter(series_id=new_version.data["data"]["series_id"]).count(),
            2,
        )
        historical = self.client.get(f"/api/v1/supplier-purchase-sheets/{sheet_id}/")
        self.assertEqual(historical.status_code, 200, historical.data)
        self.assertEqual(historical.data["data"]["status"], "PUBLISHED")

        sheet_ids = list(
            SupplierPurchaseSheet.objects.filter(purchase_order_id=order_id).values_list("id", flat=True)
        )
        business_file_ids = set(
            SupplierPurchaseSheet.objects.filter(pk__in=sheet_ids, pdf_file_id__isnull=False)
            .values_list("pdf_file_id", flat=True)
        )
        business_file_ids.update(
            image_file_id
            for image_file_id in SupplierPurchaseSheet.objects.filter(pk__in=sheet_ids)
            .values_list("line_images__image_file_id", flat=True)
            if image_file_id is not None
        )
        business_file_ids.update(
            image_file_id
            for image_file_id in SupplierPurchaseSheet.objects.filter(pk__in=sheet_ids)
            .values_list("seal_image__image_file_id", flat=True)
            if image_file_id is not None
        )

        deleted = self.client.delete(f"/api/v1/purchase-orders/{order_id}/")
        self.assertEqual(deleted.status_code, 204, getattr(deleted, "data", None))
        self.assertFalse(PurchaseOrder.objects.filter(pk=order_id).exists())
        self.assertFalse(SupplierPurchaseSheet.objects.filter(pk__in=sheet_ids).exists())
        self.assertFalse(BusinessFile.objects.filter(pk__in=business_file_ids).exists())
        self.assertEqual(
            AuditLog.objects.filter(
                object_type="SUPPLIER_PURCHASE_SHEET",
                object_id__in=[str(sheet_id_value) for sheet_id_value in sheet_ids],
                action="DELETED",
            ).count(),
            2,
        )

    def test_batch_publish_is_atomic_and_rejects_duplicate_selection(self):
        first = self.create_generated_supplier_sheet("BATCH-PUBLISH")
        second_supplier = Supplier.objects.create(
            supplier_number="GYS90003",
            supplier_type="MATERIAL",
            supplier_categories=["TUBE"],
            supplier_name="第二家批量发布供应商",
            payment_cycle="MONTHLY",
            phone="13800000003",
            is_disabled=False,
        )
        second_content = deepcopy(first.content_snapshot)
        second_content["lines"][0]["unit_price"] = ""
        second = SupplierPurchaseSheet.objects.create(
            purchase_order=first.purchase_order,
            supplier=second_supplier,
            supplier_number_snapshot=second_supplier.supplier_number,
            supplier_name_snapshot=second_supplier.supplier_name,
            supplier_phone_snapshot=second_supplier.phone,
            sheet_number="CGDD2099123102",
            version_number=1,
            is_current=True,
            template_version=first.template_version,
            status="DRAFT",
            content_snapshot=second_content,
            total_amount=first.total_amount,
            created_by_employee=self.employee,
            updated_by_employee=self.employee,
        )

        duplicate = self.client.post(
            "/api/v1/supplier-purchase-sheets/batch-publish/",
            {"ids": [str(first.pk), str(first.pk)]},
            format="json",
        )
        self.assertEqual(duplicate.status_code, 400, duplicate.data)

        invalid = self.client.post(
            "/api/v1/supplier-purchase-sheets/batch-publish/",
            {"ids": [str(first.pk), str(second.pk)]},
            format="json",
        )
        self.assertEqual(invalid.status_code, 400, invalid.data)
        self.assertIn(second.sheet_number, str(invalid.data))
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(first.status, "DRAFT")
        self.assertEqual(second.status, "DRAFT")
        self.assertIsNone(first.pdf_file_id)
        self.assertIsNone(second.pdf_file_id)

        second_content["lines"][0]["unit_price"] = "2.0000"
        second.content_snapshot = second_content
        second.save(update_fields=["content_snapshot", "updated_at"])
        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            published = self.client.post(
                "/api/v1/supplier-purchase-sheets/batch-publish/",
                {"ids": [str(second.pk), str(first.pk)]},
                format="json",
            )

        self.assertEqual(published.status_code, 200, published.data)
        self.assertEqual(
            [item["id"] for item in published.data["data"]],
            [str(second.pk), str(first.pk)],
        )
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(first.status, "PUBLISHED")
        self.assertEqual(second.status, "PUBLISHED")
        self.assertIsNotNone(first.pdf_file_id)
        self.assertIsNotNone(second.pdf_file_id)
        self.assertEqual(
            AuditLog.objects.filter(
                object_type="SUPPLIER_PURCHASE_SHEET",
                object_id__in=[str(first.pk), str(second.pk)],
                action="PUBLISHED",
            ).count(),
            2,
        )

    def test_purchase_viewer_can_create_and_version_contract_template(self):
        view_permission = Permission.objects.get(
            content_type__app_label="procurement",
            codename="view_purchaseorder",
        )
        self.user.user_permissions.set([view_permission])
        self.user = get_user_model().objects.get(pk=self.user.pk)
        self.client.force_authenticate(self.user)

        created = self.client.post(
            "/api/v1/purchase-contract-templates/",
            {
                "template_name": "QA 通用采购合同",
                "version_name": "QA 初始版",
                "supplier_categories": [],
                "status": "ACTIVE",
                "content_config": {
                    "title": "采购合同",
                    "buyer": {"name": "浙江嗨趣休闲用品有限公司"},
                    "settlement_mode": "NORMAL",
                    "tax_rate": "13",
                    "freight_bearer": "乙方",
                    "show_image": True,
                    "terms": ["按约定交期送货。"],
                },
                "change_summary": "建立 QA 通用模板",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertRegex(created.data["data"]["template_code"], r"^TPL_[A-F0-9]{12}$")
        self.assertEqual(created.data["data"]["current_version"]["version_name"], "QA 初始版")
        template_id = created.data["data"]["id"]
        updated = self.client.patch(
            f"/api/v1/purchase-contract-templates/{template_id}/",
            {
                "template_name": "QA 通用采购合同",
                "version_name": "QA 当前版",
                "content_config": {
                    **created.data["data"]["current_version"]["content_config"],
                    "freight_bearer": "甲方",
                    "payment_and_delivery": "月结，按采购交期送达。",
                    "quality_standards": ["符合确认样。", "外观无明显瑕疵。"],
                    "delivery_handover": "送达指定仓库并完成签收。",
                    "special_instructions": "未经确认不得换料。",
                    "dispute_resolution": "协商不成由甲方所在地法院管辖。",
                    "effectiveness": "双方签章后生效。",
                },
                "change_summary": "更新当前模板付款与质量内容",
                "save_mode": "CURRENT",
            },
            format="json",
            HTTP_IF_MATCH=created.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["data"]["current_version_number"], 1)
        self.assertEqual(updated.data["data"]["current_version"]["version_name"], "QA 当前版")
        self.assertEqual(len(updated.data["data"]["versions"]), 1)
        self.assertEqual(updated.data["data"]["current_version"]["content_config"]["quality_standards"][1], "外观无明显瑕疵。")
        self.assertTrue(
            AuditLog.objects.filter(
                object_type="PURCHASE_CONTRACT_TEMPLATE",
                object_id=template_id,
                action="CURRENT_VERSION_UPDATED",
            ).exists()
        )

        new_version = self.client.patch(
            f"/api/v1/purchase-contract-templates/{template_id}/",
            {
                "template_name": "QA 通用采购合同 V2",
                "version_name": "包装标签版",
                "content_config": {
                    **updated.data["data"]["current_version"]["content_config"],
                    "special_instructions": "新版本增加包装标签要求。",
                },
                "change_summary": "建立第二版",
                "save_mode": "NEW_VERSION",
            },
            format="json",
            HTTP_IF_MATCH=updated.headers["ETag"],
        )
        self.assertEqual(new_version.status_code, 200, new_version.data)
        self.assertEqual(new_version.data["data"]["current_version_number"], 2)
        self.assertEqual(new_version.data["data"]["current_version"]["version_name"], "包装标签版")
        self.assertEqual(len(new_version.data["data"]["versions"]), 2)
        self.assertEqual(new_version.data["data"]["versions"][0]["change_summary"], "建立第二版")
        self.assertTrue(
            AuditLog.objects.filter(
                object_type="PURCHASE_CONTRACT_TEMPLATE",
                object_id=template_id,
                action="VERSION_CREATED",
            ).exists()
        )
