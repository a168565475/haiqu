import uuid
from copy import deepcopy
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import IntegrityError, transaction
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.inventory.models import MaterialInbound, MaterialInboundLine
from apps.inventory.tests import test_inventory_api as inventory_tests
from apps.organizations.models import Employee
from apps.reference_data.models import Currency
from apps.samples.models import SampleMaterialProcurementRequest, SampleMaterialRequest, SampleMaterialPurchaseSource, SampleMaterialArrival, SampleRequest
from apps.samples.procurement_status import sample_status_display
from apps.samples.tests import test_material_arrivals as sample_tests
from apps.procurement.models import PurchaseOrder, PurchaseOrderLine


class SamplePurchaseApiTests(APITestCase):
    def test_sample_product_attributes_reach_purchase_without_overwriting_materials(self):
        from apps.samples.models import SampleProduct, SampleItem
        product = SampleProduct.objects.create(sample_request=self.sample, line_number=1,
            product_name='打样椅', specification='60×80', material_snapshot='铝制成品')
        SampleItem.objects.create(sample_request=self.sample, sample_product=product, line_number=1,
            product_description_snapshot='打样椅', specification_snapshot='60×80',
            color_name_snapshot='成品蓝', requested_quantity=2, sample_quantity=2, unit_code='PCS')
        preview = self.client.get(self.root + 'source/', {'source_type': 'SAMPLE', 'sample_procurement_request_id': str(self.procurement.pk)})
        self.assertEqual(preview.status_code, 200, preview.data)
        summary = preview.data['data']['sample_request']
        self.assertIn('材质：铝制成品', summary['summary'])
        self.assertEqual(summary['products'][0]['color_name'], '成品蓝')
        created = self.create_order()
        self.assertEqual(created.data['data']['product_summary'][0]['specification'], '60×80')
        detail = self.client.get(self.root + self.sample_order_id + '/')
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertEqual(detail.data['data']['product_summary'][0]['material'], '铝制成品')
        self.assertEqual(self.purchase_line.color_snapshot, self.demand['color'])
        self.assertEqual(self.purchase_line.specification_snapshot, self.demand['specification'])

    payload = inventory_tests.MaterialInventoryApiTests.payload
    confirm_quality = inventory_tests.MaterialInventoryApiTests.confirm_quality

    def setUp(self):
        sample_tests.MaterialArrivalTests.setUp(self)
        self.currency = Currency.objects.get(code='CNY')
        SampleMaterialPurchaseSource.objects.filter(procurement_request=self.procurement).delete()
        self.sample.requested_by_employee = self.employee
        self.sample.sample_due_date = timezone.localdate()
        self.sample.save()
        self.user.user_permissions.add(Permission.objects.get(content_type__app_label='samples', codename='view_sample_procurement'))
        self.client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
        self.root = '/api/v1/purchase-orders/'
        self.demand.update(rawMaterialId=str(self.raw_material.pk), material=self.raw_material.material)
        self.procurement.lines = [self.demand]
        self.procurement.save()
        record = self.procurement.material_request
        record.draft['lines'] = [self.demand]
        record.save()

    def purchase_payload(self):
        result = self.client.get(self.root + 'source/', {'sample_procurement_request_id': str(self.procurement.pk)})
        self.assertEqual(result.status_code, 200, result.data)
        source = result.data['data']
        self.assertEqual(source['bom_options'], [])
        self.assertNotIn('production_task', source)
        return {'source_type': 'SAMPLE', 'sample_procurement_request_id': str(self.procurement.pk),
                'source_revision': source['source_revision'], 'currency_id': str(self.currency.pk), 'is_mailed': False,
                'lines': [{'sample_line_key': self.key, 'supplier_id': str(self.supplier.pk),
                           'delivery_date': timezone.localdate().isoformat(), 'actual_purchase_quantity': '4', 'unit_price': '2'}]}

    def create_order(self):
        result = self.client.post(self.root, self.purchase_payload(), format='json')
        self.assertEqual(result.status_code, 201, result.data)
        self.sample_order_id = result.data['data']['id']
        self.purchase_line = PurchaseOrderLine.objects.get(purchase_order_id=self.sample_order_id)
        return result

    def receive_all(self):
        order_id = str(self.purchase_line.purchase_order_id)
        current = self.client.get('/api/v1/material-inventory/eligible-lines/', {'purchase_order_id': order_id})
        self.assertEqual(current.status_code, 200, current.data)
        return self.client.post('/api/v1/material-inventory/sample-receivings/', {
            'purchase_order_id': order_id, 'client_request_id': str(uuid.uuid4()),
            'items': [{'purchase_order_line_id': str(self.purchase_line.pk), 'zone_id': self.zone_id}],
        }, format='json', HTTP_IF_MATCH=current['ETag'])

    def test_create_has_real_sample_source_and_no_bom(self):
        options = self.client.get(self.root + 'options/').data['data']
        self.assertEqual(options['sample_requests'][0]['source_type'], 'SAMPLE')
        sources = self.client.get(self.root + 'options/', {'scope': 'sources'})
        self.assertEqual(sources.status_code, 200, sources.data)
        source = next(row for row in sources.data['data']['sample_requests'] if row['id'] == str(self.procurement.pk))
        self.assertEqual(source['source_type'], 'SAMPLE')
        self.assertTrue(source['can_create'])
        created = self.create_order()
        order = PurchaseOrder.objects.get(pk=self.sample_order_id)
        self.assertIsNone(order.production_task_id)
        self.assertIsNone(self.purchase_line.bom_id)
        self.assertIsNone(self.purchase_line.source_bom_line_id_snapshot)
        self.assertIsNone(self.purchase_line.product_line_id_snapshot)
        self.assertEqual(created.data['data']['source_number'], self.sample.request_number)
        self.assertFalse(created.data['data']['is_historical_unlinked'])
        self.assertEqual(self.purchase_line.sample_material_source.line_key, uuid.UUID(self.key))
        self.assertEqual(sample_status_display(SampleRequest.objects.get(pk=self.sample.pk)), '采购中')
        self.assertEqual(self.client.get(self.root + self.sample_order_id + '/').status_code, 200)
        self.assertEqual(self.client.get(self.root).status_code, 200)
        self.assertEqual(self.client.get(self.root + 'stats/').status_code, 200)

    def test_duplicate_and_invalid_source_do_not_create_partial_orders(self):
        payload = self.purchase_payload()
        before = PurchaseOrder.objects.count()
        for field, value in [('production_task_id', str(self.task.pk)), ('bom_id', str(self.bom.pk))]:
            invalid = {**payload, field: value}
            self.assertEqual(self.client.post(self.root, invalid, format='json').status_code, 400)
        invalid = deepcopy(payload)
        invalid['lines'][0]['sample_line_key'] = str(uuid.uuid4())
        self.assertEqual(self.client.post(self.root, invalid, format='json').status_code, 400)
        invalid = deepcopy(payload)
        invalid['source_revision'] = '0' * 64
        self.assertEqual(self.client.post(self.root, invalid, format='json').status_code, 409)
        self.assertEqual(PurchaseOrder.objects.count(), before)
        self.assertEqual(self.client.post(self.root, payload, format='json').status_code, 201)
        self.assertEqual(self.client.post(self.root, payload, format='json').status_code, 409)
        self.assertEqual(PurchaseOrder.objects.count(), before + 1)

    def test_database_rejects_ambiguous_source(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            PurchaseOrder.objects.create(order_number='CGDD2026091499', source_type='SAMPLE', production_task=self.task)

    def test_delete_unreceived_order_releases_sample_source(self):
        self.create_order()
        result = self.client.delete(self.root + self.sample_order_id + '/')
        self.assertEqual(result.status_code, 204, getattr(result, 'data', None))
        self.assertFalse(SampleMaterialPurchaseSource.objects.filter(procurement_request=self.procurement).exists())
        from apps.audit.models import AuditLog
        audit = AuditLog.objects.get(object_type='PURCHASE_ORDER', object_id=self.sample_order_id, action='DELETED')
        self.assertEqual(audit.before_data['lines'][0]['sample_line_key'], self.key)
        self.assertEqual(audit.before_data['lines'][0]['purchase_order_line_id'], str(self.purchase_line.pk))
        self.assertIsNone(audit.before_data['lines'][0]['source_bom_line_id'])
        self.assertEqual(sample_status_display(SampleRequest.objects.get(pk=self.sample.pk)), '等待采购')
        self.create_order()

    @override_settings(ACCESS_MANAGEMENT_ENABLED=True, MODULE_ACCESS_ENABLED=True, ACCESS_RESOURCE_SCOPES_ENABLED=True)
    def test_update_confirm_generate_and_direct_full_arrival(self):
        created = self.create_order()
        changed = self.client.patch(self.root + self.sample_order_id + '/', {
            'is_mailed': False, 'lines': [{'id': str(self.purchase_line.pk), 'supplier_id': str(self.supplier.pk),
            'delivery_date': timezone.localdate().isoformat(), 'actual_purchase_quantity': '4', 'unit_price': '3'}],
        }, format='json', HTTP_IF_MATCH=created['ETag'])
        self.assertEqual(changed.status_code, 200, changed.data)
        stale = self.client.patch(self.root + self.sample_order_id + '/', {'is_mailed': True, 'lines': []}, format='json', HTTP_IF_MATCH=created['ETag'])
        self.assertIn(stale.status_code, [400, 409, 412])
        confirmed = self.client.post(self.root + self.sample_order_id + '/confirm/', {}, format='json')
        self.assertEqual(confirmed.status_code, 200, confirmed.data)
        sheet_options = self.client.get(self.root + self.sample_order_id + '/sheet-options/')
        self.assertEqual(sheet_options.status_code, 200, sheet_options.data)
        template = sheet_options.data['data']['templates'][0]
        generated = self.client.post(self.root + self.sample_order_id + '/generate-sheets/', {'assignments': [
            {'supplier_id': str(self.supplier.pk), 'template_version_id': str(template['version_id'])},
        ]}, format='json')
        self.assertEqual(generated.status_code, 201, generated.data)
        sheet_id = generated.data['data'][0]['id']
        detail = self.client.get('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/')
        content = detail.data['data']['content_snapshot']
        self.assertEqual(content['order']['source_number'], self.sample.request_number)
        self.assertEqual(content['lines'][0]['material'], self.demand['material'])
        self.assertEqual(content['lines'][0]['remark'], '黑色')
        saved = self.client.patch('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/', {'content': content}, format='json', HTTP_IF_MATCH=detail['ETag'])
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(saved.data['data']['content_snapshot']['lines'][0]['remark'], '黑色')
        from io import BytesIO
        from openpyxl import load_workbook
        exported = self.client.get('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/download-xlsx/')
        self.assertEqual(exported.status_code, 200)
        binary = b''.join(exported.streaming_content) if exported.streaming else exported.content
        cells = '\n'.join(str(cell.value or '') for row in load_workbook(BytesIO(binary)).active for cell in row)
        self.assertIn(self.sample.request_number, cells)
        self.assertIn(self.demand['material'], cells)
        self.assertIn('黑色', cells)
        self.assertEqual(self.client.get('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/').status_code, 200)
        preview = self.client.get('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/preview-pdf/')
        self.assertEqual(preview.status_code, 404)
        published = self.client.post('/api/v1/supplier-purchase-sheets/' + str(sheet_id) + '/publish/', {}, format='json', HTTP_IF_MATCH=saved['ETag'])
        self.assertEqual(published.status_code, 200, published.data)
        eligible = self.client.get('/api/v1/material-inventory/eligible-purchase-orders/')
        self.assertEqual(eligible.status_code, 200, eligible.data)
        self.assertIn(self.sample_order_id, [str(row['id']) for row in eligible.data['data']])
        # 仓库人员无需品质或样品核对权限，直接全部入库并回写样品来源。
        self.user.user_permissions.remove(*Permission.objects.filter(content_type__app_label='samples'))
        self.user.user_permissions.remove(Permission.objects.get(codename='quality_confirm_material'))
        self.client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
        # 模块门禁开启时，旧直授不等于仓管MODULE范围。先证明越权被拒，再给真实岗位动作。
        denied = self.client.get('/api/v1/material-inventory/eligible-lines/', {'purchase_order_id': str(self.purchase_line.purchase_order_id)})
        self.assertEqual(denied.status_code, 404, denied.data)
        from django.contrib.auth.models import Group
        from apps.access_control.models import RoleActionGrant
        from apps.organizations.models import RoleProfile, EmployeeRoleAssignment
        receiver_role = RoleProfile.objects.create(auth_group=Group.objects.create(name='SAMPLE_DIRECT_RECEIVER'),
            role_code='SAMPLE_DIRECT_RECEIVER', display_name='样品收货测试仓管')
        EmployeeRoleAssignment.objects.create(employee=self.employee, role_profile=receiver_role, effective_from=timezone.now())
        RoleActionGrant.objects.create(role=receiver_role, scope='MODULE', permission=Permission.objects.get(
            content_type__app_label='inventory', codename='receive_material'))
        self.client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
        received = self.receive_all()
        self.assertEqual(received.status_code, 201, received.data)
        self.assertEqual(SampleMaterialArrival.objects.filter(procurement_request=self.procurement).count(), 1)
        self.assertEqual(SampleMaterialArrival.objects.get(procurement_request=self.procurement).accepted_quantity, 4)
        self.assertEqual(received.data['data']['status'], 'CONFIRMED')
        from apps.wecom.models import NotificationEvent
        ready_event = NotificationEvent.objects.get(event_code='SAMPLE_MATERIALS_AVAILABLE', object_id=str(self.sample.pk))
        self.assertEqual(ready_event.context['document_no'], self.sample.request_number)
        self.assertNotIn('product_name', ready_event.context)  # 此历史夹具无产品，不用单号冒充产品名。
        self.assertFalse(NotificationEvent.objects.filter(event_code='MATERIAL_RECEIPT_RECORDED').exists())
        self.assertIsNone(received.data['data']['quality_confirmed_at'])
        self.sample.refresh_from_db()
        # 到货不会跳过样品领料确认。
        self.assertEqual(self.sample.status, 'MATERIAL_ISSUE_PENDING')
        self.assertEqual(self.receive_all().status_code, 409)
        self.assertEqual(self.client.delete(self.root + self.sample_order_id + '/').status_code, 409)

    def test_same_material_different_colors_survive_sheet_edit_and_exports(self):
        from io import BytesIO
        from openpyxl import load_workbook
        material = 'Q235材质' * 30
        self.raw_material.material = material
        self.raw_material.save(update_fields=['material'])
        demands = [{**self.demand, 'material': material}, {**self.demand, 'key': str(uuid.uuid4()), 'material': material, 'color': '白色'}]
        self.procurement.lines = demands
        self.procurement.save()
        payload = self.purchase_payload()
        payload['lines'] = [{**payload['lines'][0], 'sample_line_key': row['key']} for row in demands]
        created = self.client.post(self.root, payload, format='json')
        self.assertEqual(created.status_code, 201, created.data)
        order_id = created.data['data']['id']
        self.assertEqual(SampleMaterialPurchaseSource.objects.filter(procurement_request=self.procurement).count(), 2)
        self.assertEqual(self.client.post(self.root + order_id + '/confirm/', {}, format='json').status_code, 200)
        template = self.client.get(self.root + order_id + '/sheet-options/').data['data']['templates'][0]
        generated = self.client.post(self.root + order_id + '/generate-sheets/', {'assignments': [
            {'supplier_id': str(self.supplier.pk), 'template_version_id': str(template['version_id'])},
        ]}, format='json')
        self.assertEqual(generated.status_code, 201, generated.data)
        url = '/api/v1/supplier-purchase-sheets/' + generated.data['data'][0]['id'] + '/'
        detail = self.client.get(url)
        content = detail.data['data']['content_snapshot']
        self.assertEqual([row['remark'] for row in content['lines']], ['黑色', '白色'])
        saved = self.client.patch(url, {'content': content}, format='json', HTTP_IF_MATCH=detail['ETag'])
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual([row['material'] for row in saved.data['data']['content_snapshot']['lines']], [material, material])
        exported = self.client.get(url + 'download-xlsx/')
        self.assertEqual(exported.status_code, 200)
        binary = b''.join(exported.streaming_content) if exported.streaming else exported.content
        cells = '\n'.join(str(cell.value or '') for row in load_workbook(BytesIO(binary)).active for cell in row)
        for text in [material, '黑色', '白色', self.sample.request_number]:
            self.assertIn(text, cells)
        # 可选导出验收产物，仅在隔离测试运行器显式指定目录时落盘。
        import os
        if os.environ.get('PURCHASE_SOURCE_QA_OUTPUT'):
            from pathlib import Path
            output = Path(os.environ['PURCHASE_SOURCE_QA_OUTPUT'])
            output.mkdir(parents=True, exist_ok=True)
            (output / 'sample-colors.xlsx').write_bytes(binary)

    def test_historical_pending_full_return_and_failed_arrival_do_not_partially_advance(self):
        from unittest.mock import patch
        from rest_framework.exceptions import ValidationError
        self.create_order()
        self.assertEqual(self.client.post(self.root + self.sample_order_id + '/confirm/', {}, format='json').status_code, 200)
        # 本用例验证入库回写，准备已正式发布的采购状态；正式发布另有整链路测试。
        PurchaseOrder.objects.filter(pk=self.sample_order_id).update(status='PURCHASING')
        from apps.inventory.tests.test_sample_receiving import legacy_sample_receipt
        received = legacy_sample_receipt(self, '4')
        with patch('apps.samples.material_arrival_services.receive_material_arrival', side_effect=ValidationError('测试回写失败')):
            failed = self.confirm_quality(received, return_quantity='1')
        self.assertEqual(failed.status_code, 400, failed.data)
        inbound = MaterialInbound.objects.get(pk=received.data['data']['id'])
        self.assertNotEqual(inbound.status, 'CONFIRMED')
        self.assertFalse(SampleMaterialArrival.objects.filter(procurement_request=self.procurement).exists())
        from apps.procurement.models import PurchaseReturn
        self.assertFalse(PurchaseReturn.objects.filter(purchase_order_id=self.sample_order_id).exists())
        quality = self.confirm_quality(received, return_quantity='4')
        self.assertEqual(quality.status_code, 200, quality.data)
        self.assertFalse(SampleMaterialArrival.objects.filter(procurement_request=self.procurement).exists())
        self.sample.refresh_from_db()
        # 回写失败或全量退货没有合格到货，仍等待采购到货。
        self.assertEqual(self.sample.status, 'PROCUREMENT_PENDING')

    def test_unauthorized_source_cannot_be_read_or_received(self):
        self.create_order()
        order = PurchaseOrder.objects.get(pk=self.sample_order_id)
        order.status = 'PURCHASING'
        order.save()
        received = self.receive_all()
        self.assertEqual(received.status_code, 201, received.data)
        outsider = get_user_model().objects.create_user(username='outside_sample_purchase')
        Employee.objects.create(user=outsider, employee_code='OUTSIDE_PUR', display_name='范围外员工')
        outsider.user_permissions.set(Permission.objects.filter(content_type__app_label__in=['procurement', 'inventory']))
        self.client.force_authenticate(outsider)
        self.assertEqual(self.client.get(self.root + self.sample_order_id + '/').status_code, 404)
        self.assertEqual(self.client.get(self.root + 'source/', {'sample_procurement_request_id': str(self.procurement.pk)}).status_code, 404)
        batches = self.client.get('/api/v1/material-inventory/quality-batches/')
        self.assertEqual(batches.status_code, 200, batches.data)
        self.assertEqual(batches.data['data'], [])
        self.assertEqual(self.confirm_quality(received).status_code, 404)
        before = MaterialInbound.objects.count()
        denied = self.client.post('/api/v1/material-inventory/receipts/', self.payload('2', '2'), format='json')
        self.assertEqual(denied.status_code, 400, denied.data)
        self.assertEqual(MaterialInbound.objects.count(), before)

    def test_returns_completion_and_replenishment_do_not_block_sample(self):
        from apps.inventory.tests.test_sample_receiving import legacy_sample_receipt
        self.create_order()
        self.assertEqual(self.client.post(self.root + self.sample_order_id + '/confirm/', {}, format='json').status_code, 200)
        # 本用例验证入库回写，准备已正式发布的采购状态；正式发布另有整链路测试。
        PurchaseOrder.objects.filter(pk=self.sample_order_id).update(status='PURCHASING')
        received = legacy_sample_receipt(self, '4')
        self.assertEqual(received.status_code, 201, received.data)
        quality = self.confirm_quality(received, return_quantity='2')
        self.assertEqual(quality.status_code, 200, quality.data)
        self.sample.refresh_from_db()
        # 退货后的合格到货仍允许先领已有材料。
        self.assertEqual(self.sample.status, 'MATERIAL_ISSUE_PENDING')
        from apps.procurement.models import PurchaseReturn
        returned = PurchaseReturn.objects.get(purchase_order_id=self.sample_order_id)
        return_url = '/api/v1/purchase-returns/' + str(returned.pk) + '/'
        detail = self.client.get(return_url)
        result = self.client.post(return_url + 'confirm/', {}, format='json', HTTP_IF_MATCH=detail['ETag'])
        self.assertEqual(result.status_code, 200, result.data)
        result = self.client.post(return_url + 'complete/', {}, format='json', HTTP_IF_MATCH=result['ETag'])
        self.assertEqual(result.status_code, 200, result.data)
        detail = self.client.get(self.root + self.sample_order_id + '/')
        result = self.client.post(self.root + self.sample_order_id + '/complete/', {'remark': '退货完成，补货继续跟进'}, format='json', HTTP_IF_MATCH=detail['ETag'])
        self.assertEqual(result.status_code, 200, result.data)
        received = self.receive_all()
        self.assertEqual(received.status_code, 201, received.data)
        self.sample.refresh_from_db()
        # 到货不会跳过样品领料确认。
        self.assertEqual(self.sample.status, 'MATERIAL_ISSUE_PENDING')
        self.assertEqual(PurchaseOrder.objects.get(pk=self.sample_order_id).status, 'COMPLETED')


from django.test import TransactionTestCase


class SamplePurchaseConcurrencyTests(TransactionTestCase):
    setUp = SamplePurchaseApiTests.setUp
    purchase_payload = SamplePurchaseApiTests.purchase_payload

    def test_sample_order_and_supplier_sheet_do_not_reverse_supplier_and_counter_locks(self):
        from concurrent.futures import ThreadPoolExecutor
        from datetime import date
        from threading import Event
        from unittest.mock import patch
        from django.db import close_old_connections, connection, connections
        from rest_framework.test import APIClient
        from apps.procurement.models import PurchaseContractTemplate, PurchaseContractTemplateVersion
        from apps.procurement.purchase_sheet_services import generate_supplier_purchase_sheets
        from apps.procurement.services import next_purchase_order_number

        payload = self.purchase_payload()
        self.order.status = 'SUBMITTED'
        self.order.save(update_fields=['status'])
        template = PurchaseContractTemplate.objects.create(template_code='SAMPLE-LOCK', template_name='并发采购模板')
        version = PurchaseContractTemplateVersion.objects.create(template=template, version_number=1)
        counter_locked, supplier_locked = Event(), Event()

        def sheet_number():
            number = next_purchase_order_number(date(2099, 12, 20))
            counter_locked.set()
            self.assertTrue(supplier_locked.wait(timeout=10), '样品采购未取得供应商锁')
            return number

        def sample_number():
            supplier_locked.set()
            return next_purchase_order_number(date(2099, 12, 20))

        def create(kind):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '5s'")
                    cursor.execute("SET statement_timeout = '10s'")
                if kind == 'sheet':
                    sheets = generate_supplier_purchase_sheets(
                        order_id=self.order.pk,
                        assignments=[{'supplier_id': self.supplier.pk, 'template_version_id': version.pk}],
                        allowed_task_ids=[self.task.pk], actor_employee=self.employee,
                    )
                    return sheets[0].sheet_number
                self.assertTrue(counter_locked.wait(timeout=10), '采购表未取得共用流水锁')
                client = APIClient()
                client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
                response = client.post(self.root, payload, format='json')
                self.assertEqual(response.status_code, 201, response.data)
                return response.data['data']['order_number']
            finally:
                connections.close_all()

        with patch('apps.procurement.purchase_sheet_services.next_purchase_order_number', side_effect=sheet_number), \
                patch('apps.procurement.services.next_purchase_order_number', side_effect=sample_number):
            with ThreadPoolExecutor(max_workers=2) as pool:
                sheet, sample = [pool.submit(create, kind) for kind in ('sheet', 'sample')]
                numbers = [sheet.result(timeout=20), sample.result(timeout=20)]
        self.assertEqual(numbers, ['CGDD991220001', 'CGDD991220002'])

    def test_same_sample_concurrent_creation_is_serialized(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from django.db import close_old_connections, connections
        from rest_framework.test import APIClient
        payload = self.purchase_payload()
        barrier = Barrier(2)
        def submit(_):
            close_old_connections()
            try:
                client = APIClient()
                client.force_authenticate(get_user_model().objects.get(pk=self.user.pk))
                barrier.wait(timeout=20)
                result = client.post(self.root, payload, format='json')
                return result.status_code
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            statuses = list(pool.map(submit, range(2)))
        self.assertEqual(sorted(statuses), [201, 409])
        self.assertEqual(PurchaseOrder.objects.filter(sample_procurement_request=self.procurement).count(), 1)
        self.assertEqual(SampleMaterialPurchaseSource.objects.filter(procurement_request=self.procurement).count(), 1)
