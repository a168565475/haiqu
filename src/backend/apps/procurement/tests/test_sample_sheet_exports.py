from types import SimpleNamespace
from unittest.mock import Mock
from django.test import SimpleTestCase
from openpyxl import Workbook
from ..purchase_sheet_xlsx import _layout_config, _write_detail_rows


class SampleSheetExportTests(SimpleTestCase):
    def test_long_sample_requirements_continue_without_duplicate_quantities(self):
        material = '材质' * 250
        color = '颜色' * 250
        sheet = SimpleNamespace(line_images=Mock())
        sheet.line_images.select_related.return_value.all.return_value = []
        for layout in ('AUXILIARY', 'PLASTIC', 'FABRIC', 'WEBBING', 'FOAM'):
            with self.subTest(layout=layout):
                ws = Workbook().active
                config = _layout_config(layout)
                content = {'order': {'source_type': 'SAMPLE', 'source_number': 'DY-001'}, 'lines': [
                    {'part_name': '测试材料', 'specification': '规格', 'material': material, 'remark': color,
                     'quantity': '4', 'unit_price': '2', 'amount': '8'},
                ]}
                _, _, end = _write_detail_rows(ws, sheet, content, layout, config, 6)
                self.assertGreater(end, 7)
                material_text = ''.join(str(ws.cell(row, 3).value or '') for row in range(6, end))
                color_text = ''.join(str(ws.cell(row, 4).value or '') for row in range(6, end))
                self.assertIn(material, material_text)
                self.assertIn(color, color_text)
                self.assertTrue(all(26 <= ws.row_dimensions[row].height <= 348 for row in range(6, end)))
                quantities = [ws.cell(row, config['quantity_column']).value for row in range(6, end)]
                self.assertEqual(quantities[0], 4)
                self.assertTrue(all(value is None for value in quantities[1:]))
                self.assertEqual(ws.cell(6, config['amount_column']).value, '=ROUND(E6*F6,4)')
