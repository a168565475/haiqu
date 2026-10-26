from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from ..purchase_sheet_services import _default_content, normalize_sheet_content
from ..purchase_sheet_xlsx import _layout_config, _line_values, _write_detail_rows, _write_total_and_note, build_supplier_purchase_sheet_xlsx


class SupplierPurchaseSheetBomAttributeTests(SimpleTestCase):
    def setUp(self):
        self.attributes = {
            "schema_version": 1,
            "section_type": "HARDWARE",
            "section_title": "五金件",
            "fields": [
                {"key": "part_name", "label": "品名", "value": "连接件"},
                {"key": "material", "label": "材质", "value": "钢"},
                {"key": "specification", "label": "规格", "value": "M8"},
                {"key": "tube_coating_color", "label": "喷塑颜色", "value": "黑色"},
            ],
        }
        self.line = SimpleNamespace(
            pk="line-1",
            section_type_snapshot="HARDWARE",
            section_title_snapshot="五金件",
            product_name_snapshot="成品",
            product_specification_snapshot="成品规格",
            part_name_snapshot="连接件",
            specification_snapshot="M8",
            material_snapshot="钢",
            raw_material_snapshot="钢材",
            material_attributes_snapshot=deepcopy(self.attributes),
            actual_purchase_quantity=Decimal("2"),
            source_unit_text="PCS",
            unit_price=Decimal("3.5"),
            amount=Decimal("7"),
            delivery_date=date(2026, 10, 20),
            color_snapshot="",
            raw_bar_length_m=None,
        )
        self.order = SimpleNamespace(
            source_type="PRODUCTION",
            sales_order_number_snapshot="SO-1",
            production_task_number_snapshot="PT-1",
            sales_order_number="SO-1",
            sample_number_snapshot="",
        )
        self.supplier = SimpleNamespace(supplier_name="供应商", phone="")
        self.template_version = SimpleNamespace(content_config={})

    def test_default_content_copies_order_bom_partition_and_attributes(self):
        content = _default_content(self.order, self.supplier, [self.line], self.template_version, "CGDD2026101201")

        row = content["lines"][0]
        self.assertEqual(row["section_type"], "HARDWARE")
        self.assertEqual(row["section_title"], "五金件")
        self.assertEqual(row["material_attributes_snapshot"], self.attributes)

    def test_update_keeps_order_source_attributes_when_submitted_content_changes(self):
        current = _default_content(self.order, self.supplier, [self.line], self.template_version, "CGDD2026101201")
        sheet = SimpleNamespace(content_snapshot=current, sheet_number="CGDD2026101201")
        submitted = deepcopy(current)
        submitted["lines"][0]["section_title"] = "被修改分区"
        submitted["lines"][0]["material_attributes_snapshot"] = {
            "fields": [{"label": "规格", "value": "被修改规格"}],
        }

        normalized, _ = normalize_sheet_content(sheet, submitted)

        row = normalized["lines"][0]
        self.assertEqual(row["section_title"], "五金件")
        self.assertEqual(row["material_attributes_snapshot"], self.attributes)

    def test_xlsx_line_splits_specification_and_requirements_without_powder_color(self):
        values = _line_values(
            "DEFAULT",
            {
                "part_name": "连接件",
                "specification": "M8",
                "section_title": "五金件",
                "material_attributes_snapshot": self.attributes,
                "material": "钢",
            },
            1,
            "SO-1",
            False,
        )

        self.assertEqual(values[:3], ["连接件", "M8", "材质：钢"])
        self.assertEqual(values[3], "—")
        self.assertNotIn("喷塑颜色", values[2])
        self.assertNotIn("黑色", values[2])

    def test_all_xlsx_categories_use_the_same_seven_columns(self):
        line = {
            "part_name": "半通", "specification": "16", "material": "PE", "remark": "按样",
            "material_attributes_snapshot": {"fields": [
                {"key": "plastic_color", "label": "颜色", "value": "黑色"},
            ]},
        }
        for layout in ("AUXILIARY", "FABRIC", "WEBBING", "PLASTIC", "FOAM"):
            with self.subTest(layout=layout):
                values = _line_values(layout, line, 1, "SO-1", False)
                self.assertEqual(len(values), 7)
                self.assertEqual(values[:4], ["半通", "16", "材质：PE；要求：按样", "黑色"])
                self.assertEqual(_layout_config(layout)["widths"], [14, 23, 24, 12, 9, 9, 9])

    def test_default_xlsx_columns_formulas_and_number_formats_match_seven_columns(self):
        config = _layout_config("AUXILIARY")
        self.assertEqual(config["headers"][:4], ["产品名称", "产品规格", "物料属性", "颜色"])
        self.assertEqual(len(config["headers"]), 7)
        sheet = SimpleNamespace(line_images=Mock())
        sheet.line_images.select_related.return_value.all.return_value = []
        content = {"total_amount": "21.60", "lines": [{
            "part_name": "半通", "specification": "16", "material": "PE", "quantity": "48.0000",
            "unit_price": "0.4500", "amount": "21.60",
        }]}
        worksheet = Workbook().active
        lines, images, end = _write_detail_rows(worksheet, sheet, content, "AUXILIARY", config, 6)
        _write_total_and_note(worksheet, content, "AUXILIARY", config, end, lines, images)
        self.assertEqual([worksheet.cell(6, index).value for index in range(1, 4)], ["半通", "16", "材质：PE"])
        self.assertEqual(worksheet["E6"].number_format, "0")
        self.assertEqual(worksheet["F6"].number_format, "0.##")
        self.assertEqual(worksheet["G6"].value, "=ROUND(E6*F6,4)")
        self.assertEqual(worksheet["G6"].number_format, "0.#")
        self.assertEqual(worksheet["G7"].value, "=SUM(G6:G6)")
        self.assertEqual(worksheet["G7"].number_format, "0.#")

    def test_long_production_attributes_continue_without_losing_text_or_repeating_quantity(self):
        material = "物料属性" * 300
        sheet = SimpleNamespace(line_images=Mock())
        sheet.line_images.select_related.return_value.all.return_value = []
        content = {"order": {"source_type": "PRODUCTION"}, "lines": [{
            "part_name": "采购材料", "specification": "规格", "material": material,
            "quantity": "48", "unit_price": "0.45", "amount": "21.6",
        }]}
        worksheet = Workbook().active
        _, _, end = _write_detail_rows(worksheet, sheet, content, "AUXILIARY", _layout_config("AUXILIARY"), 6)
        self.assertGreater(end, 7)
        self.assertIn(material, "".join(str(worksheet.cell(row, 3).value or "") for row in range(6, end)))
        self.assertEqual(worksheet["E6"].value, 48)
        self.assertTrue(all(worksheet.cell(row, 5).value is None for row in range(7, end)))

    def test_pipe_description_removes_trailing_zeroes_without_changing_saved_line(self):
        line = {
            "part_name": "管材", "specification": "16", "material": "铁", "measurement_kind": "PIPE",
            "actual_bar_count": "2.0000", "raw_bar_length_m": "6.5000", "actual_purchase_meters": "13.0000",
        }
        original = deepcopy(line)
        values = _line_values("AUXILIARY", line, 1, "SO-1", False)
        self.assertIn("2 根 × 6.5 米/根 = 13 米；单价：元/米", values[2])
        self.assertEqual(line, original)

    def test_complete_workbook_matches_pdf_content_order_units_and_page_setup(self):
        content = {
            "title": "塑料件采购合同", "order": {"source_type": "PRODUCTION", "sales_order_number": "SALE-OTHER", "production_task_number": "HT-001"},
            "buyer": {"name": "甲方公司"}, "seller": {"name": "乙方工厂"}, "purchase_date": "2026-10-26",
            "show_image": False, "total_amount": "21.6000", "tax_rate": "13", "freight_bearer": "乙方", "remark": "到厂验收",
            "lines": [{"part_name": "半通", "specification": "16", "material": "PE", "unit": "只",
                       "quantity": "48.0000", "unit_price": "0.4500", "amount": "21.6000"}],
        }
        sheet = SimpleNamespace(content_snapshot=content, sheet_number="CG-001", line_images=Mock(), seal_image=None)
        workbook = load_workbook(BytesIO(build_supplier_purchase_sheet_xlsx(sheet)))
        worksheet = workbook.active
        self.assertEqual(worksheet["A1"].value, "塑料件采购合同　（订单编号  HT-001）")
        self.assertEqual(worksheet["A2"].value, "采购单号：CG-001\n甲方（采购方）：甲方公司\n乙方（供货方）：乙方工厂")
        self.assertIn("中华人民共和国民法典", worksheet["A3"].value)
        self.assertEqual(worksheet["A4"].value, "一、采购产品名称、规格、数量、单价及金额")
        self.assertEqual([cell.value for cell in worksheet[5]], _layout_config()["headers"])
        self.assertEqual(worksheet["E6"].value, 48)
        self.assertEqual(worksheet["E6"].number_format, '0" 只"')
        self.assertEqual(worksheet["G6"].value, "=ROUND(E6*F6,4)")
        self.assertEqual(worksheet["A7"].value, "暂无图片")
        self.assertIn("备注：\n到厂验收", worksheet["C7"].value)
        self.assertIn("税率 13% · 运费由乙方承担", worksheet["C7"].value)
        self.assertEqual(worksheet["F7"].value, "合计")
        self.assertEqual(worksheet["G7"].value, "=SUM(G6:G6)")
        self.assertEqual(worksheet["A8"].value, "二、支付方式和交期")
        self.assertEqual(worksheet.cell(worksheet.max_row, 1).value, "请确认无误盖章回传")
        self.assertEqual(worksheet.cell(worksheet.max_row, 4).value, "签订日期：2026.10.26")
        self.assertEqual(worksheet.max_column, 7)
        self.assertEqual(worksheet.page_setup.orientation, "portrait")
        self.assertEqual(str(worksheet.page_setup.paperSize), worksheet.PAPERSIZE_A4)
        self.assertEqual(worksheet.page_setup.fitToWidth, 1)
        self.assertEqual(worksheet.page_setup.fitToHeight, 0)
        self.assertTrue(all(getattr(worksheet.page_margins, side) > 0 for side in ("left", "right", "top", "bottom")))
        sheet.line_images.select_related.assert_not_called()

    def test_weight_settlement_and_missing_values_match_pdf(self):
        values = _line_values("WEBBING", {}, 1, "", True)
        self.assertEqual(values, ["—", "—", "—", "—", "—", "—", "待结算"])
        sheet = SimpleNamespace(content_snapshot={"settlement_mode": "ACTUAL_WEIGHT", "show_image": False, "lines": [{}]},
                                sheet_number="CG-W", line_images=Mock(), seal_image=None)
        worksheet = load_workbook(BytesIO(build_supplier_purchase_sheet_xlsx(sheet))).active
        self.assertEqual(worksheet["G6"].value, "待结算")
        self.assertEqual(worksheet["G7"].value, "待结算")

    def test_long_contract_notes_and_terms_are_preserved_in_bounded_rows(self):
        note, payment = "备注内容" * 500, "交期内容" * 1000
        sheet = SimpleNamespace(content_snapshot={"show_image": False, "remark": note, "payment_and_delivery": payment, "lines": []},
                                sheet_number="CG-LONG", line_images=Mock(), seal_image=None)
        worksheet = load_workbook(BytesIO(build_supplier_purchase_sheet_xlsx(sheet))).active
        notes = "".join(str(worksheet.cell(row, 3).value or "") for row in range(1, worksheet.max_row + 1))
        terms = "".join(str(worksheet.cell(row, 1).value or "") for row in range(1, worksheet.max_row + 1))
        self.assertIn(note, notes)
        self.assertIn(payment, terms)
        self.assertTrue(all(dimension.height <= 409 for dimension in worksheet.row_dimensions.values()))

    def test_only_first_available_image_in_detail_order_is_used(self):
        sheet = SimpleNamespace(line_images=Mock())
        first, second = object(), object()
        sheet.line_images.select_related.return_value.all.return_value = [
            SimpleNamespace(source_purchase_order_line_id="second", image_file=second),
            SimpleNamespace(source_purchase_order_line_id="first", image_file=first),
        ]
        content = {"lines": [{"source_purchase_order_line_id": "first"}, {"source_purchase_order_line_id": "second"}]}
        worksheet, config = Workbook().active, _layout_config()
        lines, images, row = _write_detail_rows(worksheet, sheet, content, "CONTRACT", config, 6)
        with patch("apps.procurement.purchase_sheet_xlsx._media_image_path", side_effect=lambda item: "first.png" if item is first else "second.png"), patch(
            "apps.procurement.purchase_sheet_xlsx._add_image", return_value=True,
        ) as add_image:
            _write_total_and_note(worksheet, content, "CONTRACT", config, row, lines, images)
        self.assertEqual(add_image.call_args.args[1], "first.png")

    def test_amount_formulas_use_the_same_per_line_precision_as_saved_pdf_amounts(self):
        sheet = SimpleNamespace(line_images=Mock())
        sheet.line_images.select_related.return_value.all.return_value = []
        content = {"lines": [{"quantity": "0.3333", "unit_price": "0.3333", "amount": "0.1111"} for _ in range(5)]}
        worksheet = Workbook().active
        _write_detail_rows(worksheet, sheet, content, "CONTRACT", _layout_config(), 6)
        self.assertEqual([worksheet.cell(row, 7).value for row in range(6, 11)], [f"=ROUND(E{row}*F{row},4)" for row in range(6, 11)])
