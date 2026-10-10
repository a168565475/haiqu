import re
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import Supplier, SupplierCategory, SupplierPaymentCycle, SupplierType


SUPPLIER_NUMBER_PATTERN = re.compile(r"^GYS([0-9]{5})$")
REQUIRED_HEADERS = {
    "供应商编号",
    "供应商类型",
    "供应商类别",
    "供应商名称",
    "付款周期",
    "备注",
    "联系电话",
    "是否禁用",
    "提交人",
    "提交时间",
    "更新时间",
}
TYPE_MAP = {label: value for value, label in SupplierType.choices}
CATEGORY_MAP = {label: value for value, label in SupplierCategory.choices}
PAYMENT_CYCLE_MAP = {label: value for value, label in SupplierPaymentCycle.choices}
DISABLED_MAP = {"": False, "正常": False, "否": False, "禁用": True, "是": True}


class SupplierWorkbookError(ValueError):
    pass


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date_time(value, *, row_number, field_name):
    if value in {None, ""}:
        return None
    parsed = value
    if isinstance(value, date) and not isinstance(value, datetime):
        parsed = datetime.combine(value, time.min)
    elif not isinstance(value, datetime):
        try:
            parsed = datetime.fromisoformat(_text(value))
        except ValueError as exc:
            raise SupplierWorkbookError(
                f"第 {row_number} 行“{field_name}”不是有效日期时间。"
            ) from exc
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def _mapped_value(mapping, raw_value, *, row_number, field_name, allow_blank=False):
    label = _text(raw_value)
    if allow_blank and not label:
        return ""
    try:
        return mapping[label]
    except KeyError as exc:
        raise SupplierWorkbookError(
            f"第 {row_number} 行“{field_name}”包含不支持的值：{label or '空'}。"
        ) from exc


def _categories(row, *, row_number):
    raw_value = row.get("供应商类别-文本") or row.get("供应商类别")
    labels = [item.strip() for item in re.split(r"[,，、;；]", _text(raw_value)) if item.strip()]
    if not labels:
        raise SupplierWorkbookError(f"第 {row_number} 行“供应商类别”不能为空。")
    if len(labels) != len(set(labels)):
        raise SupplierWorkbookError(f"第 {row_number} 行“供应商类别”存在重复选项。")
    try:
        return [CATEGORY_MAP[label] for label in labels]
    except KeyError as exc:
        raise SupplierWorkbookError(
            f"第 {row_number} 行“供应商类别”包含不支持的值：{exc.args[0]}。"
        ) from exc


def read_supplier_workbook(workbook_path):
    path = Path(workbook_path)
    if not path.is_file():
        raise SupplierWorkbookError(f"找不到供应商表格：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "供应商信息" not in workbook.sheetnames:
            raise SupplierWorkbookError("表格缺少“供应商信息”工作表。")
        sheet = workbook["供应商信息"]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [_text(value) for value in next(rows)]
        except StopIteration as exc:
            raise SupplierWorkbookError("“供应商信息”工作表为空。") from exc
        missing_headers = sorted(REQUIRED_HEADERS.difference(headers))
        if missing_headers:
            raise SupplierWorkbookError(f"表格缺少列：{'、'.join(missing_headers)}。")

        imported_rows = []
        seen_numbers = set()
        for row_number, values in enumerate(rows, start=2):
            raw_row = dict(zip(headers, values))
            if not any(value is not None and _text(value) for value in values):
                continue
            supplier_number = _text(raw_row.get("供应商编号"))
            if not SUPPLIER_NUMBER_PATTERN.fullmatch(supplier_number):
                raise SupplierWorkbookError(
                    f"第 {row_number} 行供应商编号“{supplier_number or '空'}”格式不正确。"
                )
            if supplier_number in seen_numbers:
                raise SupplierWorkbookError(f"供应商编号“{supplier_number}”在表格中重复。")
            seen_numbers.add(supplier_number)

            supplier_name = _text(raw_row.get("供应商名称"))
            if not supplier_name:
                raise SupplierWorkbookError(f"第 {row_number} 行“供应商名称”不能为空。")
            disabled_label = _text(raw_row.get("是否禁用"))
            if disabled_label not in DISABLED_MAP:
                raise SupplierWorkbookError(
                    f"第 {row_number} 行“是否禁用”包含不支持的值：{disabled_label}。"
                )

            imported_row = {
                    "supplier_number": supplier_number,
                    "supplier_type": _mapped_value(
                        TYPE_MAP,
                        raw_row.get("供应商类型"),
                        row_number=row_number,
                        field_name="供应商类型",
                    ),
                    "supplier_categories": _categories(raw_row, row_number=row_number),
                    "supplier_name": supplier_name,
                    "payment_cycle": _mapped_value(
                        PAYMENT_CYCLE_MAP,
                        raw_row.get("付款周期"),
                        row_number=row_number,
                        field_name="付款周期",
                        allow_blank=True,
                    ),
                    "remark": _text(raw_row.get("备注")),
                    "phone": _text(raw_row.get("联系电话")),
                    "is_disabled": DISABLED_MAP[disabled_label],
                    "submitted_by_name": _text(raw_row.get("提交人")),
                    "submitted_at": _date_time(
                        raw_row.get("提交时间"), row_number=row_number, field_name="提交时间"
                    ),
                    "source_updated_at": _date_time(
                        raw_row.get("更新时间"), row_number=row_number, field_name="更新时间"
                    ),
                }
            model_values = {
                key: value for key, value in imported_row.items() if key != "source_updated_at"
            }
            try:
                Supplier(**model_values).full_clean(
                    validate_unique=False,
                    validate_constraints=False,
                )
            except ValidationError as exc:
                raise SupplierWorkbookError(
                    f"第 {row_number} 行数据不符合供应商规则：{'；'.join(exc.messages)}"
                ) from exc
            imported_rows.append(imported_row)
    finally:
        workbook.close()

    if not imported_rows:
        raise SupplierWorkbookError("表格中没有可导入的供应商数据。")
    return imported_rows


def supplier_workbook_summary(rows):
    number_values = [int(SUPPLIER_NUMBER_PATTERN.fullmatch(row["supplier_number"]).group(1)) for row in rows]
    type_labels = dict(SupplierType.choices)
    category_labels = dict(SupplierCategory.choices)
    type_counts = Counter(type_labels[row["supplier_type"]] for row in rows)
    category_counts = Counter(
        category_labels[category]
        for row in rows
        for category in row["supplier_categories"]
    )
    return {
        "rows": len(rows),
        "minimum_number": f"GYS{min(number_values):05d}",
        "maximum_number": f"GYS{max(number_values):05d}",
        "next_number": f"GYS{max(number_values) + 1:05d}",
        "type_counts": dict(type_counts),
        "category_counts": dict(category_counts),
        "monthly_count": sum(row["payment_cycle"] == SupplierPaymentCycle.MONTHLY for row in rows),
        "disabled_count": sum(row["is_disabled"] for row in rows),
    }


@transaction.atomic
def import_supplier_rows(rows, *, expected_existing):
    existing_count = Supplier.objects.count()
    if existing_count != expected_existing:
        raise SupplierWorkbookError(
            f"数据库现有供应商为 {existing_count} 条，与确认值 {expected_existing} 不一致，已停止导入。"
        )

    created_count = 0
    updated_count = 0
    for row in rows:
        supplier = Supplier.objects.filter(supplier_number=row["supplier_number"]).first()
        if supplier is None:
            supplier = Supplier(supplier_number=row["supplier_number"])
            created_count += 1
        else:
            updated_count += 1
        for field_name in (
            "supplier_type",
            "supplier_categories",
            "supplier_name",
            "payment_cycle",
            "remark",
            "phone",
            "is_disabled",
            "submitted_by_name",
            "submitted_at",
        ):
            setattr(supplier, field_name, row[field_name])
        supplier.full_clean()
        supplier.save()
        source_created_at = row["submitted_at"] or row["source_updated_at"] or timezone.now()
        source_updated_at = row["source_updated_at"] or source_created_at
        Supplier.objects.filter(pk=supplier.pk).update(
            created_at=source_created_at,
            updated_at=source_updated_at,
        )

    maximum_number = Supplier.objects.order_by("-supplier_number").values_list(
        "supplier_number", flat=True
    ).first()
    maximum_sequence = int(SUPPLIER_NUMBER_PATTERN.fullmatch(maximum_number).group(1))
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT setval(%s::regclass, %s, true)",
            ["haiqu_supplier_number_seq", maximum_sequence],
        )
    return {"created": created_count, "updated": updated_count, "total": len(rows)}
