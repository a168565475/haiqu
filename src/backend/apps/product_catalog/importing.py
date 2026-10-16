"""产品来源文件预检、受控导入与逐行核验。"""

from collections import defaultdict
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from django.contrib.auth import get_user_model
from django.db import connection, transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import PermissionDenied, ValidationError

from apps.access_control.sales_policy import has_permission
from apps.core.audit import write_audit_log
from apps.core.auth import active_employee_for

from .models import ProductCatalog, ProductCatalogCategory, ProductCatalogImport, ProductCatalogName
from .services import create_product_record


HEADERS = ("类型", "产品编码", "类别", "品名", "规格", "单位", "货号", "仓库号", "仓库名称", "规模")
BUSINESS_FIELDS = {
    "product_code": "产品编码", "item_number": "货号", "specification": "规格",
    "unit": "单位", "scale": "规模",
}


def preflight_product_catalog(source_bytes, *, filename, sheet="产品信息"):
    if not isinstance(filename, str) or not filename or len(filename) > 255 or "/" in filename or "\\" in filename or not filename.lower().endswith(".xlsx"):
        raise ValidationError({"filename": "来源必须使用不含路径、长度不超过 255 的 .xlsx 文件名。"})
    if not isinstance(sheet, str) or not sheet or len(sheet) > 255:
        raise ValidationError({"sheet": "必须指定有效工作表名称。"})
    if not isinstance(source_bytes, bytes) or not source_bytes:
        raise ValidationError({"source": "来源文件内容不能为空，且必须是原始字节。"})
    rows = []
    try:
        workbook = load_workbook(BytesIO(source_bytes), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, KeyError, ValueError, ParseError) as exc:
        raise ValidationError({"source": "无法读取有效的 xlsx 来源文件，请检查文件格式。"}) from exc
    try:
        if sheet not in workbook.sheetnames:
            raise ValidationError({"sheet": "来源文件不存在指定工作表。"})
        try:
            sheet_rows = list(workbook[sheet].iter_rows())
        except (BadZipFile, OSError, KeyError, ValueError, ParseError) as exc:
            raise ValidationError({"source": "无法读取工作表中的行，请检查来源文件。"}) from exc
        headers = sheet_rows[0] if sheet_rows else ()
        if tuple(cell.value for cell in headers) != HEADERS:
            raise ValidationError({"headers": "表头必须依次为：" + "、".join(HEADERS)})
        # 部分导出文件没有 dimension；实际迭代保留 Excel 行号，不依赖 max_row。
        for row_number, cells in enumerate(sheet_rows[1:], start=2):
            values = [cell.value for cell in cells]
            if all(value is None for value in values):
                continue
            if any(value is not None for value in values[len(HEADERS):]):
                raise ValidationError({"source": f"第 {row_number} 行存在表头以外的数据。"})
            for cell in cells:
                if cell.data_type in ("f", "e"):
                    raise ValidationError({"source": f"第 {row_number} 行存在公式或错误单元格，不能自动导入。"})
                if cell.value is not None and not isinstance(cell.value, str):
                    raise ValidationError({"source": f"第 {row_number} 行存在非文本值，不能自动转换。"})
            values = values[:len(HEADERS)] + [None] * max(0, len(HEADERS) - len(values))
            raw = dict(zip(HEADERS, values))
            category, name = raw["类别"] or "", raw["品名"] or ""
            if not category.strip() or not name.strip():
                raise ValidationError({"source": f"第 {row_number} 行必须填写类别和完整品名。"})
            if raw["类型"] != "成品" or raw["规模"] not in ("常规", "定制"):
                raise ValidationError({"source": f"第 {row_number} 行的类型或规模不是本次确认的原值。"})
            data = {field: raw[header] if raw[header] is not None else "" for field, header in BUSINESS_FIELDS.items()}
            for field in ("specification", "unit"):
                if data[field] == "/":
                    data[field] = ""
            limits = [("类别", category, ProductCatalogCategory._meta.get_field("name").max_length),
                      ("品名", name, ProductCatalogName._meta.get_field("name").max_length)]
            limits.extend((BUSINESS_FIELDS[field], value, ProductCatalog._meta.get_field(field).max_length) for field, value in data.items())
            for label, value, limit in limits:
                if len(value) > limit:
                    raise ValidationError({"source": f"第 {row_number} 行的{label}超过 {limit} 字符，不能截断导入。"})
            rows.append({"row": row_number, "category": category, "name": name, "data": data, "raw_fields": raw})
    finally:
        workbook.close()
    if not rows:
        raise ValidationError({"source": "工作表没有可导入的产品行。"})
    return {"filename": filename, "sha256": sha256(source_bytes).hexdigest(), "sheet": sheet, "rows": rows}


def _summary(preflight, *, mode, created=0, existing=0, batch=None):
    rows = preflight["rows"]
    return {
        "mode": mode, "filename": preflight["filename"], "sha256": preflight["sha256"], "sheet": preflight["sheet"],
        "total_rows": len(rows), "category_count": len({row["category"] for row in rows}),
        "name_count": len({(row["category"], row["name"]) for row in rows}),
        "blank_codes": sum(row["data"]["product_code"] == "" for row in rows),
        "created": created, "existing": existing, "pending": len(rows) - created - existing,
        "batch_id": str(batch.pk) if batch else None,
    }


def _source_state(preflight, source_bytes):
    filename, digest, sheet = (preflight[key] for key in ("filename", "sha256", "sheet"))
    conflicts = set(ProductCatalogImport.objects.filter(filename=filename, sheet=sheet).exclude(sha256=digest).values_list("sha256", flat=True))
    conflicts.update(ProductCatalog.objects.filter(source_filename=filename, source_sheet=sheet).exclude(source_sha256=digest).exclude(source_sha256="").values_list("source_sha256", flat=True))
    if conflicts:
        raise ValidationError({"source": "同文件名、同工作表已有不同内容的导入，禁止整批重复建档。", "conflicting_sha256": sorted(conflicts)})
    batch = ProductCatalogImport.objects.filter(filename=filename, sha256=digest, sheet=sheet).first()
    existing = {product.source_row: product for product in ProductCatalog.objects.filter(source_filename=filename, source_sha256=digest, source_sheet=sheet)}
    expected = {row["row"]: row for row in preflight["rows"]}
    if batch is not None:
        if bytes(batch.source_bytes) != source_bytes or batch.row_count != len(expected):
            raise ValidationError({"source": "已保存批次的原始文件或行数不一致，请先核验。"})
        if set(existing) != set(expected):
            raise ValidationError({"source": "已保存批次的来源行不完整，请先核验，不能自动补写。"})
    elif existing:
        raise ValidationError({"source": "已有来源产品缺少完整导入批次，请先核验。"})
    for number, product in existing.items():
        if product.source_import_id != batch.pk or product.source_raw_fields != expected[number]["raw_fields"]:
            raise ValidationError({"source": f"已保存第 {number} 行的原始来源不一致，请先核验。"})
    return batch, existing


def _actor_request(actor):
    user = get_user_model().objects.filter(pk=getattr(actor, "pk", None), is_active=True).first()
    employee = active_employee_for(user)
    if user is None or employee is None:
        raise PermissionDenied("写入必须指定关联有效员工的现有启用用户。")
    request = SimpleNamespace(user=user, employee=employee, META={"HTTP_USER_AGENT": "manage.py import_product_catalog"}, headers={})
    if not has_permission(request, "product_catalog.add_productcatalog"):
        raise PermissionDenied("指定用户没有新增产品权限。")
    return request


def import_product_catalog(source_bytes, *, filename, sheet="产品信息", apply=False, actor=None):
    preflight = preflight_product_catalog(source_bytes, filename=filename, sheet=sheet)
    if not apply:
        batch, existing = _source_state(preflight, source_bytes)
        return _summary(preflight, mode="preview", existing=len(existing), batch=batch)
    request = _actor_request(actor)
    with transaction.atomic():
        # 首次导入尚无批次可行锁，按文件名和工作表锁住整个检查、建档事务。
        lock_key = int.from_bytes(sha256((filename + "\0" + sheet).encode("utf-8")).digest()[:8], "big", signed=True)
        with connection.cursor() as cursor:
            cursor.execute("SELECT pg_advisory_xact_lock(%s)", [lock_key])
        batch, existing = _source_state(preflight, source_bytes)
        if existing:
            return _summary(preflight, mode="applied", existing=len(existing), batch=batch)
        batch = ProductCatalogImport.objects.create(filename=filename, sha256=preflight["sha256"], sheet=sheet, source_bytes=source_bytes, row_count=len(preflight["rows"]))
        for row in preflight["rows"]:
            category, _ = ProductCatalogCategory.objects.get_or_create(name=row["category"])
            name, _ = ProductCatalogName.objects.get_or_create(category=category, name=row["name"])
            if category.archived_at is not None or name.archived_at is not None:
                raise ValidationError({"source": f"第 {row['row']} 行对应类别或品名已归档，不能自动恢复。"})
            create_product_record(
                data={"name_id": name.pk, **row["data"]},
                source={"filename": filename, "sha256": preflight["sha256"], "sheet": sheet, "row": row["row"], "raw_fields": row["raw_fields"], "import_id": batch.pk},
                request=request,
            )
        result = _summary(preflight, mode="applied", created=len(preflight["rows"]), batch=batch)
        write_audit_log(object_type="PRODUCT_CATALOG_IMPORT", object_id=batch.pk, action="CREATED", actor_employee=request.employee, after_data=result, request=request)
        return result


def verify_product_catalog_import(*, batch):
    result = {"filename": batch.filename, "sha256": batch.sha256, "sheet": batch.sheet, "batch_id": str(batch.pk),
              "valid": False, "total_rows": batch.row_count, "matched_rows": 0, "batch_errors": [], "differences": []}
    source_bytes = bytes(batch.source_bytes)
    if sha256(source_bytes).hexdigest() != batch.sha256:
        result["batch_errors"].append("source_sha256")
        return result
    try:
        preflight = preflight_product_catalog(source_bytes, filename=batch.filename, sheet=batch.sheet)
    except ValidationError:
        result["batch_errors"].append("source_bytes")
        return result
    if batch.row_count != len(preflight["rows"]):
        result["batch_errors"].append("row_count")
    result["total_rows"] = len(preflight["rows"])
    products = defaultdict(list)
    for product in batch.products.select_related("name__category").prefetch_related("colors"):
        products[product.source_row].append(product)
    expected_rows = {row["row"] for row in preflight["rows"]}
    for row in preflight["rows"]:
        matches = products.get(row["row"], [])
        if len(matches) != 1:
            result["differences"].append({"row": row["row"], "fields": ["missing_product" if not matches else "duplicate_source_row"]})
            continue
        product, = matches
        fields = []
        if product.name.category.name != row["category"]:
            fields.append("category")
        if product.name.name != row["name"]:
            fields.append("name")
        fields.extend(field for field, value in row["data"].items() if getattr(product, field) != value)
        for field, expected in (("source_filename", batch.filename), ("source_sha256", batch.sha256), ("source_sheet", batch.sheet), ("source_raw_fields", row["raw_fields"])):
            if getattr(product, field) != expected:
                fields.append(field)
        if product.archived_at is not None:
            fields.append("archived_at")
        if [color.value for color in product.colors.all() if color.archived_at is None] != ["无颜色"]:
            fields.append("colors")
        if fields:
            result["differences"].append({"row": row["row"], "fields": fields})
        else:
            result["matched_rows"] += 1
    for row in products:
        if row not in expected_rows:
            result["differences"].append({"row": row, "fields": ["unexpected_source_row"]})
    result["valid"] = not result["batch_errors"] and not result["differences"]
    return result
