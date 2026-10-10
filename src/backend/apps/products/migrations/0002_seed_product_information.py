import uuid
from pathlib import Path

from django.db import migrations
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("88d23ee6-032c-45f8-9ab8-03effeb29771")
RESOURCE_PATH = Path(__file__).resolve().parent.parent / "resources" / "product_information.xlsx"
EXPECTED_HEADERS = (
    "标题",
    "品号",
    "类型",
    "产品编码",
    "类别",
    "品名",
    "规格",
    "单位",
    "产品状态",
    "货号",
    "仓库号",
    "仓库名称",
    "规模",
    "唯一校验",
)
STATUS_MAP = {"正常": "ACTIVE", "废弃": "INACTIVE"}
SCALE_MAP = {"定制": "CUSTOM", "常规": "REGULAR"}
SUPPLEMENT_ITEM_NUMBER = "XY-205-1"
SUPPLEMENT_ITEM_CATEGORY = "外袋"
SUPPLEMENT_ITEM_DESCRIPTION = "来源：产品信息_20260831115757.xlsx（补充缺失货号）"
EXPECTED_HISTORICAL_MISMATCH = ("1004200001", "三折床-不带棚", "XY-212")


def stable_id(kind, value):
    return uuid.uuid5(SEED_NAMESPACE, f"{kind}:{value}")


def clean_cell(value):
    return "" if value is None else str(value).strip()


def load_rows():
    workbook = load_workbook(RESOURCE_PATH, read_only=True, data_only=True)
    try:
        worksheet = workbook["产品信息"]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    headers = tuple(clean_cell(value) for value in rows[0])
    if headers != EXPECTED_HEADERS:
        raise RuntimeError(f"产品信息表头不符合预期：{headers}")
    data_rows = [tuple(clean_cell(value) for value in row) for row in rows[1:] if any(value is not None for value in row)]
    if len(data_rows) != 185:
        raise RuntimeError(f"产品信息数据行数应为 185，实际为 {len(data_rows)}。")
    return data_rows


def validate_rows(rows):
    product_numbers = [row[1] for row in rows]
    product_codes = [row[3] for row in rows if row[3]]
    unique_checks = [row[13] for row in rows]
    if len(product_numbers) != len(set(product_numbers)) or any(not value for value in product_numbers):
        raise RuntimeError("产品信息品号为空或重复。")
    if len(product_codes) != len(set(product_codes)):
        raise RuntimeError("产品信息非空产品编码存在重复。")
    if len(unique_checks) != len(set(unique_checks)) or any(not value for value in unique_checks):
        raise RuntimeError("产品信息唯一校验列为空或重复。")
    if {row[2] for row in rows} != {"成品"}:
        raise RuntimeError("产品信息类型目前只允许成品。")
    if not {row[8] for row in rows}.issubset(STATUS_MAP):
        raise RuntimeError("产品信息存在未知状态。")
    if not {row[12] for row in rows}.issubset(SCALE_MAP):
        raise RuntimeError("产品信息存在未知规模。")
    required_indexes = (1, 2, 4, 5, 6, 7, 8, 9, 12, 13)
    if any(not row[index] for row in rows for index in required_indexes):
        raise RuntimeError("产品信息必填列存在空值。")


def seed_product_information(apps, schema_editor):
    ProductInformation = apps.get_model("products", "ProductInformation")
    ProductCategory = apps.get_model("reference_data", "ProductCategory")
    ProductionUnit = apps.get_model("reference_data", "ProductionUnit")
    ProductItem = apps.get_model("reference_data", "ProductItem")
    db_alias = schema_editor.connection.alias

    rows = load_rows()
    validate_rows(rows)
    category_by_name = {
        item.category_name: item
        for item in ProductCategory.objects.using(db_alias).filter(archived_at__isnull=True)
    }
    unit_by_name = {
        item.unit_name: item
        for item in ProductionUnit.objects.using(db_alias).filter(archived_at__isnull=True)
    }
    if SUPPLEMENT_ITEM_CATEGORY not in category_by_name:
        raise RuntimeError(f"补充货号引用的产品类别不存在：{SUPPLEMENT_ITEM_CATEGORY}")

    supplement_item = ProductItem.objects.using(db_alias).filter(item_number=SUPPLEMENT_ITEM_NUMBER).first()
    if supplement_item is None:
        supplement_item = ProductItem.objects.using(db_alias).create(
            id=stable_id("PRODUCT_ITEM", SUPPLEMENT_ITEM_NUMBER),
            item_number=SUPPLEMENT_ITEM_NUMBER,
            category_id=category_by_name[SUPPLEMENT_ITEM_CATEGORY].pk,
            status="ACTIVE",
            sort_order=36,
            description=SUPPLEMENT_ITEM_DESCRIPTION,
        )

    item_by_number = {
        item.item_number: item
        for item in ProductItem.objects.using(db_alias).filter(archived_at__isnull=True)
    }
    missing_categories = sorted({row[4] for row in rows} - set(category_by_name))
    missing_units = sorted({row[7] for row in rows} - set(unit_by_name))
    missing_items = sorted({row[9] for row in rows} - set(item_by_number))
    if missing_categories or missing_units or missing_items:
        raise RuntimeError(
            f"产品信息引用缺失：类别={missing_categories}，单位={missing_units}，货号={missing_items}"
        )

    mismatches = []
    for row in rows:
        product_number = row[1]
        category = category_by_name[row[4]]
        product_item = item_by_number[row[9]]
        if product_item.category_id != category.pk:
            mismatches.append((product_number, row[4], row[9]))
        ProductInformation.objects.using(db_alias).update_or_create(
            id=stable_id("PRODUCT_INFORMATION", product_number),
            defaults={
                "product_number": product_number,
                "product_type": "FINISHED",
                "product_code": row[3],
                "category_id": category.pk,
                "product_name": row[5],
                "specification": row[6],
                "unit_id": unit_by_name[row[7]].pk,
                "product_item_id": product_item.pk,
                "warehouse_number": row[10],
                "warehouse_name": row[11],
                "scale": SCALE_MAP[row[12]],
                "status": STATUS_MAP[row[8]],
            },
        )

    if mismatches != [EXPECTED_HISTORICAL_MISMATCH]:
        raise RuntimeError(f"产品信息历史类别/货号异常与预期不符：{mismatches}")


def remove_seeded_product_information(apps, schema_editor):
    ProductInformation = apps.get_model("products", "ProductInformation")
    ProductItem = apps.get_model("reference_data", "ProductItem")
    db_alias = schema_editor.connection.alias
    rows = load_rows()
    ids = [stable_id("PRODUCT_INFORMATION", row[1]) for row in rows]
    ProductInformation.objects.using(db_alias).filter(pk__in=ids).delete()
    ProductItem.objects.using(db_alias).filter(
        item_number=SUPPLEMENT_ITEM_NUMBER,
        description=SUPPLEMENT_ITEM_DESCRIPTION,
    ).delete()


class Migration(migrations.Migration):
    dependencies = [
        ("products", "0001_initial"),
    ]

    operations = [
        migrations.RunPython(seed_product_information, remove_seeded_product_information),
    ]

