import uuid
from collections import Counter, defaultdict
from pathlib import Path

from django.db import migrations
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("fa9d56d6-8fe1-41d4-bfb4-31ea456cf983")
RESOURCE_PATH = Path(__file__).resolve().parent.parent / "resources" / "raw_material_information.xlsx"
EXPECTED_HEADERS = (
    "品号",
    "类型",
    "类别",
    "品名",
    "是否包材辅料",
    "规格",
    "单位",
    "颜色",
    "材质",
    "箱子形状",
    "备注",
    "直径",
    "长度",
    "宽度（门幅）",
    "厚度",
    "型号",
    "仓库号",
    "仓库名称",
    "重量（kg）",
    "流转单位",
    "数字",
    "唯一校验",
)
ITEM_TYPE_MAP = {
    "材料": "MATERIAL",
    "半成品": "SEMI_FINISHED",
}
BOX_SHAPE_MAP = {
    "": "",
    "半盖": "HALF_COVER",
    "全盖": "FULL_COVER",
    "天地盖": "LID_AND_BASE",
    "展示盒": "DISPLAY_BOX",
    "其他": "OTHER",
}
EXPECTED_WAREHOUSES = {
    ("CK001", "原料布面仓"),
    ("CK002", "裁剪布面仓"),
    ("CK003", "原料管材仓"),
    ("CK004", "管材喷塑仓"),
    ("CK005", "五金仓"),
    ("CK006", "裁剪布面仓"),
    ("CK008", "包材仓"),
    ("CK010", "辅料仓"),
}


def stable_id(row_number, material_number):
    return uuid.uuid5(SEED_NAMESPACE, f"RAW_MATERIAL_INFORMATION:{row_number}:{material_number}")


def clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_rows():
    workbook = load_workbook(RESOURCE_PATH, read_only=True, data_only=True)
    try:
        worksheet = workbook["原料信息"]
        headers = tuple(clean_cell(cell.value) for cell in worksheet[1])
        rows = [
            (row_number, tuple(clean_cell(value) for value in values))
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True),
                start=2,
            )
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()

    if headers != EXPECTED_HEADERS:
        raise RuntimeError(f"原料信息表头不符合预期：{headers}")
    if len(rows) != 4897:
        raise RuntimeError(f"原料信息数据行数应为 4897，实际为 {len(rows)}。")
    return rows


def validate_rows(rows):
    values = [row for _row_number, row in rows]
    if {row[1] for row in values} != set(ITEM_TYPE_MAP):
        raise RuntimeError("原料信息类型必须且只能包含材料、半成品。")
    if not {row[4] for row in values}.issubset({"", "是", "否"}):
        raise RuntimeError("是否包材辅料存在未知值。")
    if not {row[9] for row in values}.issubset(BOX_SHAPE_MAP):
        raise RuntimeError("箱子形状存在未知值。")
    required_indexes = (0, 1, 2, 3, 5, 6, 16, 17)
    if any(not row[index] for row in values for index in required_indexes):
        raise RuntimeError("原料信息必填列存在空值。")

    number_counts = Counter(row[0] for row in values)
    duplicate_numbers = {key: count for key, count in number_counts.items() if count > 1}
    if len(duplicate_numbers) != 26 or sum(count - 1 for count in duplicate_numbers.values()) != 26:
        raise RuntimeError(f"原料信息历史重复品号与预期不符：{duplicate_numbers}")
    warehouses = {(row[16], row[17]) for row in values}
    if warehouses != EXPECTED_WAREHOUSES:
        raise RuntimeError(f"原料信息仓库组合与预期不符：{warehouses}")


def seed_raw_material_information(apps, schema_editor):
    RawMaterialInformation = apps.get_model("products", "RawMaterialInformation")
    ProductCategory = apps.get_model("reference_data", "ProductCategory")
    ProductionUnit = apps.get_model("reference_data", "ProductionUnit")
    ProductColor = apps.get_model("reference_data", "ProductColor")
    db_alias = schema_editor.connection.alias

    rows = load_rows()
    validate_rows(rows)
    category_by_key = {
        (item.item_type, item.category_name): item
        for item in ProductCategory.objects.using(db_alias).filter(archived_at__isnull=True)
    }
    unit_by_name = {
        item.unit_name: item
        for item in ProductionUnit.objects.using(db_alias).filter(archived_at__isnull=True)
    }
    colors_by_name = defaultdict(list)
    for color in ProductColor.objects.using(db_alias).filter(archived_at__isnull=True):
        colors_by_name[color.color_name].append(color)

    missing_categories = sorted(
        {
            (ITEM_TYPE_MAP[row[1]], row[2])
            for _row_number, row in rows
        }
        - set(category_by_key)
    )
    missing_units = sorted({row[6] for _row_number, row in rows} - set(unit_by_name))
    if missing_categories or missing_units:
        raise RuntimeError(f"原料信息引用缺失：类别={missing_categories}，单位={missing_units}")

    instances = []
    for row_number, row in rows:
        item_type = ITEM_TYPE_MAP[row[1]]
        color_matches = colors_by_name.get(row[7], [])
        color_id = color_matches[0].pk if len(color_matches) == 1 else None
        instances.append(
            RawMaterialInformation(
                id=stable_id(row_number, row[0]),
                material_number=row[0],
                item_type=item_type,
                category_id=category_by_key[(item_type, row[2])].pk,
                material_name=row[3],
                is_packaging_auxiliary=row[4] == "是",
                specification=row[5],
                unit_id=unit_by_name[row[6]].pk,
                color_id=color_id,
                color_snapshot=row[7],
                material=row[8],
                box_shape=BOX_SHAPE_MAP[row[9]],
                notes=row[10],
                diameter=row[11],
                length=row[12],
                width=row[13],
                thickness=row[14],
                model_number=row[15],
                warehouse_number=row[16],
                warehouse_name=row[17],
                weight_kg=row[18],
                flow_unit=row[19],
                number_value=row[20],
                source_unique_check=row[21],
                source_row_number=row_number,
            )
        )
    RawMaterialInformation.objects.using(db_alias).bulk_create(instances, batch_size=500)


def remove_seeded_raw_material_information(apps, schema_editor):
    RawMaterialInformation = apps.get_model("products", "RawMaterialInformation")
    db_alias = schema_editor.connection.alias
    rows = load_rows()
    ids = [stable_id(row_number, row[0]) for row_number, row in rows]
    RawMaterialInformation.objects.using(db_alias).filter(pk__in=ids).delete()


class Migration(migrations.Migration):
    dependencies = [
        ("products", "0003_raw_material_information"),
    ]

    operations = [
        migrations.RunPython(
            seed_raw_material_information,
            remove_seeded_raw_material_information,
        ),
    ]
