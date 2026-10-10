from datetime import timedelta
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient, APITestCase

from apps.audit.models import AuditLog
from apps.crm.models import CustomerTimelineEvent
from apps.customers.models import (
    CustomerAccount,
    CustomerAssignment,
    CustomerContact,
    CustomerEntityAddress,
    CustomerTransactionEntity,
)
from apps.email_accounts.models import EmailAccount, SmtpSecurity
from apps.email_accounts.services import EmailDeliveryUncertain, encrypt_smtp_password
from apps.inquiries.models import Inquiry, InquiryItem, InquiryStatus
from apps.organizations.models import Employee
from apps.workflow.models import ApprovalStatus, BusinessTask

from ..models import (
    Quotation,
    QuotationContainerConfig,
    QuotationDispatchRecord,
    QuotationEmailDispatchAttempt,
    QuotationEmailDispatchAttemptStatus,
    QuotationTemplate,
    QuotationVersion,
    QuotationVersionStatus,
)


class QuotationApiTests(APITestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="quotation_api_user")
        self.employee = Employee.objects.create(
            user=self.user,
            employee_code="QUOTATION_API_USER",
            display_name="报价接口测试员工",
        )
        self.approver_user = user_model.objects.create_user(username="quotation_approver")
        self.approver = Employee.objects.create(
            user=self.approver_user,
            employee_code="QUOTATION_APPROVER",
            display_name="报价审批员工",
        )
        self.other_user = user_model.objects.create_user(username="quotation_other")
        self.other_employee = Employee.objects.create(
            user=self.other_user,
            employee_code="QUOTATION_OTHER",
            display_name="其他报价员工",
        )
        permissions = Permission.objects.filter(content_type__app_label__in=["quotations", "workflow"])
        self.user.user_permissions.set(permissions)
        self.user.user_permissions.add(*Permission.objects.filter(
            content_type__app_label="inquiries", codename="change_inquiry"
        ))
        self.user.user_permissions.add(*Permission.objects.filter(
            content_type__app_label="customers", codename="add_customertransactionentity"
        ))
        self.approver_user.user_permissions.set(permissions)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.customer = self.create_customer("QUOTE_CUSTOMER", "报价客户", self.employee)
        self.other_customer = self.create_customer(
            "OTHER_QUOTE_CUSTOMER",
            "其他报价客户",
            self.other_employee,
        )
        self.entity = CustomerTransactionEntity.objects.create(
            customer_account=self.customer,
            entity_code="QUOTE_ENTITY",
            legal_name="报价交易主体",
            registration_country_code="CN",
        )
        self.address = CustomerEntityAddress.objects.create(
            transaction_entity=self.entity,
            address_code="QUOTE_ADDRESS",
            country_code="CN",
            city="深圳",
            address_line1="测试路 1 号",
        )
        self.contact = CustomerContact.objects.create(
            customer_account=self.customer,
            transaction_entity=self.entity,
            display_name="报价联系人",
        )
        self.inquiry = Inquiry.objects.create(
            customer_account=self.customer,
            transaction_entity=self.entity,
            customer_contact=self.contact,
            owner_employee=self.employee,
            inquiry_number="QUOTE_SOURCE_INQUIRY",
            subject="报价来源询价",
            received_at=timezone.now(),
        )
        InquiryItem.objects.create(
            inquiry=self.inquiry, line_number=1, product_description="询盘产品", specification="规格 A",
            quantity="12.0000", unit_code="PCS", target_unit_price="9.0000", target_currency_code="USD",
        )

    def test_create_quote_from_inquiry_creates_real_draft_task_and_requirement_snapshot(self):
        self.inquiry.status = "READY_FOR_QUOTE"
        self.inquiry.save(update_fields=["status", "updated_at"])
        context = self.client.get(f"/api/v1/quotations/inquiry-context/?inquiry_id={self.inquiry.pk}")
        self.assertEqual(context.status_code, 200, context.data)
        response = self.client.post(
            "/api/v1/quotations/create-from-inquiry/",
            {
                "inquiry_id": str(self.inquiry.pk), "transaction_entity_id": str(self.entity.pk),
                "subject": "询盘转报价", "currency_code": "USD", "due_at": (timezone.now() + timedelta(days=1)).isoformat(),
                "idempotency_key": "inq-quote-test-001",
            }, format="json", HTTP_IF_MATCH=context["ETag"],
        )
        self.assertEqual(response.status_code, 201, response.data)
        quotation = Quotation.objects.get(pk=response.data["data"]["quotation"]["id"])
        version = quotation.versions.get(version_number=1)
        self.inquiry.refresh_from_db()
        self.assertEqual(self.inquiry.status, "QUOTING")
        self.assertEqual(version.lines.count(), 1)
        self.assertIsNone(version.lines.first().unit_price)
        self.assertEqual(version.customer_requirement_snapshot["items"][0]["target_unit_price"], "9.0000")
        self.assertEqual(version.quotation_template.template_code, "HAIQU_FIXED_BILINGUAL")
        self.assertTrue(BusinessTask.objects.filter(title__contains=quotation.quotation_number).exists())

    def test_create_quote_from_inquiry_copies_independent_color_snapshot(self):
        item = self.inquiry.items.first()
        item.color = "哑光黑"
        item.save(update_fields=["color", "updated_at"])
        self.inquiry.status = "READY_FOR_QUOTE"
        self.inquiry.save(update_fields=["status", "updated_at"])

        context = self.client.get(f"/api/v1/quotations/inquiry-context/?inquiry_id={self.inquiry.pk}")
        response = self.client.post(
            "/api/v1/quotations/create-from-inquiry/",
            {
                "inquiry_id": str(self.inquiry.pk),
                "transaction_entity_id": str(self.entity.pk),
                "subject": "颜色快照测试",
                "currency_code": "USD",
                "due_at": (timezone.now() + timedelta(days=1)).isoformat(),
                "idempotency_key": "inq-quote-color-test-001",
            },
            format="json",
            HTTP_IF_MATCH=context["ETag"],
        )
        self.assertEqual(response.status_code, 201, response.data)

        quotation = Quotation.objects.get(pk=response.data["data"]["quotation"]["id"])
        line = quotation.versions.get(version_number=1).lines.get()
        self.assertEqual(line.color_snapshot, "哑光黑")
        self.assertEqual(line.specification_snapshot, "规格 A；颜色：哑光黑")

        from apps.quotations.api.serializers import QuotationLineSerializer

        serialized = QuotationLineSerializer(line).data
        self.assertEqual(serialized["color_snapshot"], "哑光黑")

    def test_create_quote_from_inquiry_requires_explicit_choice_when_active_quote_exists(self):
        self.inquiry.status = "READY_FOR_QUOTE"
        self.inquiry.save(update_fields=["status", "updated_at"])
        Quotation.objects.create(
            customer_account=self.customer, transaction_entity=self.entity, source_inquiry=self.inquiry,
            owner_employee=self.employee, quotation_number="QUOTE_EXISTING", subject="已有报价",
        )
        context = self.client.get(f"/api/v1/quotations/inquiry-context/?inquiry_id={self.inquiry.pk}")
        response = self.client.post(
            "/api/v1/quotations/create-from-inquiry/",
            {"inquiry_id": str(self.inquiry.pk), "transaction_entity_id": str(self.entity.pk), "subject": "第二份报价", "currency_code": "USD", "due_at": (timezone.now() + timedelta(days=1)).isoformat(), "idempotency_key": "inq-quote-test-002"},
            format="json", HTTP_IF_MATCH=context["ETag"],
        )
        self.assertEqual(response.status_code, 409, response.data)

    def create_customer(self, code, name, owner):
        customer = CustomerAccount.objects.create(
            account_code=code,
            account_name=name,
            primary_country_code="CN",
        )
        CustomerAssignment.objects.create(
            customer_account=customer,
            employee=owner,
            assignment_role="PRIMARY_OWNER",
            effective_from=timezone.now(),
        )
        return customer

    def create_quotation(self):
        response = self.client.post(
            "/api/v1/quotations/",
            {
                "customer_account_id": str(self.customer.pk),
                "transaction_entity_id": str(self.entity.pk),
                "source_inquiry_id": str(self.inquiry.pk),
                "owner_employee_id": str(self.employee.pk),
                "subject": "标准报价",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        return response

    def create_version(self, quotation_id, quotation_etag, template=None, from_version=None):
        payload = {"currency_code": "USD", "language_code": "zh-hans"}
        if template:
            payload["quotation_template_id"] = str(template.pk)
        if from_version:
            payload = {"from_version_id": str(from_version)}
        response = self.client.post(
            f"/api/v1/quotations/{quotation_id}/versions/",
            payload,
            format="json",
            HTTP_IF_MATCH=quotation_etag,
        )
        self.assertEqual(response.status_code, 201, response.data)
        return response

    def test_seller_profile_updates_only_controlled_template_fields_with_etag_and_audit(self):
        template = QuotationTemplate.objects.create(
            template_code="SELLER_PROFILE_TEST",
            template_name="卖方资料测试模板",
            template_config={"brand_name": "海曲", "payment_terms": "保留原有配置"},
        )
        profile = self.client.get(f"/api/v1/quotation-templates/{template.pk}/seller-profile/")
        self.assertEqual(profile.status_code, 200, profile.data)
        self.assertIn("ETag", profile)

        updated = self.client.patch(
            f"/api/v1/quotation-templates/{template.pk}/seller-profile/",
            {
                "company_name": "JINHUA SUPERGROUP LEISURE PRODUCTS CO.,LTD",
                "contact_person": "张三",
                "phone": "+86 755 12345678",
                "fax": "+86 755 87654321",
                "skype": "supergroup.sales",
                "qq": "12345678",
                "email": "sales@example.com",
                "website": "https://example.com",
            },
            format="json",
            HTTP_IF_MATCH=profile["ETag"],
        )

        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertIn("ETag", updated)
        template.refresh_from_db()
        self.assertEqual(template.template_version, 2)
        self.assertEqual(template.template_config["seller_company_name"], "JINHUA SUPERGROUP LEISURE PRODUCTS CO.,LTD")
        self.assertEqual(template.template_config["seller_fax"], "+86 755 87654321")
        self.assertEqual(template.template_config["seller_skype"], "supergroup.sales")
        self.assertEqual(template.template_config["seller_qq"], "12345678")
        self.assertEqual(template.template_config["brand_name"], "海曲")
        self.assertEqual(template.template_config["payment_terms"], "保留原有配置")
        self.assertEqual(template.template_config["seller_contact_person"], "张三")
        self.assertEqual(template.template_config["seller_email"], "sales@example.com")
        self.assertTrue(AuditLog.objects.filter(
            object_type="QUOTATION_TEMPLATE",
            object_id=template.pk,
            action="SELLER_PROFILE_UPDATED",
        ).exists())

        stale = self.client.patch(
            f"/api/v1/quotation-templates/{template.pk}/seller-profile/",
            {"phone": "+86 755 87654321"},
            format="json",
            HTTP_IF_MATCH=profile["ETag"],
        )
        invalid_email = self.client.patch(
            f"/api/v1/quotation-templates/{template.pk}/seller-profile/",
            {"email": "not-an-email"},
            format="json",
            HTTP_IF_MATCH=updated["ETag"],
        )
        self.assertEqual(stale.status_code, 412, stale.data)
        self.assertEqual(invalid_email.status_code, 400, invalid_email.data)

    @patch("apps.files.uploads._clamav_scan", return_value="ClamAV test clean")
    def test_product_image_upload_preview_and_remove_reuses_scanned_business_file(self, _scan):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        line = self.add_line(version.data["data"]["id"], version.headers["ETag"])
        line_id = line.data["data"]["id"]
        with TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            uploaded = self.client.post(
                f"/api/v1/quotation-lines/{line_id}/product-image/",
                {"file": SimpleUploadedFile("product.jpg", b"\xff\xd8\xff\xdb\x00\x04", content_type="image/jpeg")},
                format="multipart",
                HTTP_IF_MATCH=line.headers["ETag"],
            )
            self.assertEqual(uploaded.status_code, 200, uploaded.data)
            self.assertIn("ETag", uploaded)
            self.assertIsNotNone(uploaded.data["data"]["product_image"])

            preview = self.client.get(f"/api/v1/quotation-lines/{line_id}/product-image/")
            removed = self.client.delete(
                f"/api/v1/quotation-lines/{line_id}/product-image/",
                HTTP_IF_MATCH=uploaded.headers["ETag"],
            )
            rejected_gif = self.client.post(
                f"/api/v1/quotation-lines/{line_id}/product-image/",
                {"file": SimpleUploadedFile("product.gif", b"GIF89a\x00\x00\x00", content_type="image/gif")},
                format="multipart",
                HTTP_IF_MATCH=removed.headers["ETag"],
            )

            self.assertEqual(preview.status_code, 200)
            self.assertEqual(preview["Content-Type"], "image/jpeg")
            self.assertEqual(removed.status_code, 200, removed.data)
            self.assertIsNone(removed.data["data"]["product_image"])
            self.assertEqual(rejected_gif.status_code, 400, rejected_gif.data)

    @patch("apps.files.uploads._clamav_scan", return_value="ClamAV test clean")
    def test_formal_pdf_embeds_clean_product_image_without_changing_source_file(self, _scan):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        line = self.add_line(version.data["data"]["id"], version.headers["ETag"])
        line_id = line.data["data"]["id"]
        image_bytes = BytesIO()
        Image.new("RGB", (8, 6), color=(32, 86, 154)).save(image_bytes, format="PNG")

        with TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            uploaded = self.client.post(
                f"/api/v1/quotation-lines/{line_id}/product-image/",
                {"file": SimpleUploadedFile("product.png", image_bytes.getvalue(), content_type="image/png")},
                format="multipart",
                HTTP_IF_MATCH=line.headers["ETag"],
            )
            matrix = self.client.put(
                f"/api/v1/quotation-versions/{version.data['data']['id']}/price-matrix/",
                {"price_bases": [{
                    "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                    "currency_code": "USD", "is_default": True,
                    "line_prices": [{"quotation_line_id": line_id, "unit_price": "20.0000"}],
                }]},
                format="json",
                HTTP_IF_MATCH=uploaded.headers["ETag"],
            )
            locked = self.lock_version(version.data["data"]["id"], matrix.headers["ETag"])
            issued = self.client.post(
                f"/api/v1/quotation-versions/{version.data['data']['id']}/issue/",
                {}, format="json", HTTP_IF_MATCH=locked.headers["ETag"],
            )
            generated = self.client.post(
                f"/api/v1/quotation-versions/{version.data['data']['id']}/generate-formal-pdf/",
                {}, format="json", HTTP_IF_MATCH=issued.headers["ETag"],
            )
            scan_count_before_locked_upload = _scan.call_count
            rejected_locked_upload = self.client.post(
                f"/api/v1/quotation-lines/{line_id}/product-image/",
                {"file": SimpleUploadedFile("replacement.png", image_bytes.getvalue(), content_type="image/png")},
                format="multipart",
                HTTP_IF_MATCH=generated.headers["ETag"],
            )
            downloaded = self.client.get(
                f"/api/v1/quotation-versions/{version.data['data']['id']}/formal-document-download/"
            )
            pdf_content = b"".join(downloaded.streaming_content)

        self.assertEqual(uploaded.status_code, 200, uploaded.data)
        self.assertEqual(generated.status_code, 200, generated.data)
        self.assertEqual(rejected_locked_upload.status_code, 409, rejected_locked_upload.data)
        self.assertEqual(_scan.call_count, scan_count_before_locked_upload)
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn(b"/Image", pdf_content)

    def test_draft_terms_are_saved_with_etag_and_return_server_payment_breakdown(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation["ETag"])
        line = self.add_line(version.data["data"]["id"], version["ETag"])
        current = self.client.get(f"/api/v1/quotation-versions/{version.data['data']['id']}/")

        updated = self.client.patch(
            f"/api/v1/quotation-versions/{version.data['data']['id']}/",
            {
                "terms_snapshot": {
                    "payment_terms": "30% 订金 / 70% 余款",
                    "deposit_percent": "30",
                    "lead_time": "收到订金后 35 天",
                    "incoterm_code": "FOB",
                    "incoterm_label": "FOB 宁波",
                }
            },
            format="json",
            HTTP_IF_MATCH=current["ETag"],
        )

        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["data"]["terms_snapshot"]["deposit_percent"], "30")
        self.assertEqual(updated.data["data"]["payment_breakdown"]["deposit_amount"], "54.0000")
        self.assertEqual(updated.data["data"]["payment_breakdown"]["balance_amount"], "126.0000")
        refreshed = self.client.get(f"/api/v1/quotation-versions/{version.data['data']['id']}/")
        self.assertEqual(refreshed.data["data"]["terms_snapshot"]["lead_time"], "收到订金后 35 天")
        self.assertEqual(refreshed.data["data"]["payment_breakdown"], updated.data["data"]["payment_breakdown"])

        rejected = self.client.patch(
            f"/api/v1/quotation-versions/{version.data['data']['id']}/",
            {"terms_snapshot": {"deposit_percent": "101"}},
            format="json",
            HTTP_IF_MATCH=updated["ETag"],
        )
        self.assertEqual(rejected.status_code, 400, rejected.data)

        locked = self.lock_version(version.data["data"]["id"], updated["ETag"])
        self.assertEqual(locked.status_code, 200, locked.data)
        self.assertEqual(locked.data["data"]["terms_snapshot"]["payment_terms"], "30% 订金 / 70% 余款")
        self.assertEqual(locked.data["data"]["terms_snapshot"]["deposit_percent"], "30")
        self.assertEqual(locked.data["data"]["terms_snapshot"]["lead_time"], "收到订金后 35 天")
        self.assertEqual(locked.data["data"]["terms_snapshot"]["payment"], "T/T 30 days")
        self.assertEqual(locked.data["data"]["terms_snapshot"]["tax_rate"], "0.100000")

    def add_line(self, version_id, etag, description="测试产品", specification="规格 A"):
        response = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/lines/",
            {
                "product_description_snapshot": description,
                "specification_snapshot": specification,
                "quantity": "10.0000",
                "unit_code": "PCS",
                "unit_price": "20.0000",
                "discount_rate": "0.100000",
            },
            format="json",
            HTTP_IF_MATCH=etag,
        )
        self.assertEqual(response.status_code, 201, response.data)
        return response

    def configure_export_packaging(self, line_id, etag, *, container_type="40HQ", unit_price="18.7500"):
        packaging = self.client.put(
            f"/api/v1/quotation-lines/{line_id}/packaging/",
            {
                "length_cm": "78.0000",
                "width_cm": "58.0000",
                "height_cm": "61.0000",
                "pieces_per_carton": 4,
                "carton_count": 1,
                "gross_weight_kg": "24.5000",
                "net_weight_kg": "23.0000",
            },
            format="json",
            HTTP_IF_MATCH=etag,
        )
        self.assertEqual(packaging.status_code, 200, packaging.data)
        tier = self.client.post(
            f"/api/v1/quotation-lines/{line_id}/container-tiers/",
            {
                "container_type_code": container_type,
                "effective_volume_cbm": "68.0000",
                "load_factor": "0.900000",
                "unit_price": unit_price,
            },
            format="json",
            HTTP_IF_MATCH=packaging["ETag"],
        )
        self.assertEqual(tier.status_code, 201, tier.data)
        return tier

    def test_line_mutation_recalculates_version_amounts(self):
        quotation_response = self.create_quotation()
        version_response = self.create_version(
            quotation_response.data["data"]["id"],
            quotation_response["ETag"],
        )
        line_response = self.add_line(version_response.data["data"]["id"], version_response["ETag"])
        version_response = self.client.get(
            f"/api/v1/quotation-versions/{version_response.data['data']['id']}/"
        )
        self.assertEqual(version_response.status_code, 200, version_response.data)
        self.assertEqual(version_response.data["data"]["subtotal_amount"], "200.0000")
        self.assertEqual(version_response.data["data"]["discount_amount"], "20.0000")
        self.assertEqual(version_response.data["data"]["total_amount"], "180.0000")
        update_response = self.client.patch(
            f"/api/v1/quotation-lines/{line_response.data['data']['id']}/",
            {"unit_price": "30.0000"},
            format="json",
            HTTP_IF_MATCH=version_response["ETag"],
        )
        self.assertEqual(update_response.status_code, 200, update_response.data)
        self.assertEqual(update_response.data["data"]["line_total"], "270.0000")
        version_response = self.client.get(
            f"/api/v1/quotation-versions/{version_response.data['data']['id']}/"
        )
        self.assertEqual(version_response.data["data"]["total_amount"], "270.0000")

    def test_packaging_and_container_tier_are_etag_protected_and_calculated(self):
        quotation_response = self.create_quotation()
        version_response = self.create_version(
            quotation_response.data["data"]["id"],
            quotation_response["ETag"],
        )
        line_response = self.add_line(version_response.data["data"]["id"], version_response["ETag"])
        line_id = line_response.data["data"]["id"]
        missing_etag = self.client.put(
            f"/api/v1/quotation-lines/{line_id}/packaging/",
            {"length_cm": "50", "width_cm": "40", "height_cm": "30", "pieces_per_carton": 10, "carton_count": 5},
            format="json",
        )
        self.assertEqual(missing_etag.status_code, 428, missing_etag.data)
        packaging_response = self.client.put(
            f"/api/v1/quotation-lines/{line_id}/packaging/",
            {"length_cm": "50", "width_cm": "40", "height_cm": "30", "pieces_per_carton": 10, "carton_count": 5},
            format="json",
            HTTP_IF_MATCH=line_response["ETag"],
        )
        self.assertEqual(packaging_response.status_code, 200, packaging_response.data)
        self.assertEqual(packaging_response.data["data"]["carton_cbm"], "0.060")
        self.assertEqual(packaging_response.data["data"]["total_cbm"], "0.300")
        config = QuotationContainerConfig.objects.create(
            container_type_code="TEST_CONTAINER",
            effective_volume_cbm="30.0000",
            default_load_factor="0.800000",
        )
        tier_response = self.client.post(
            f"/api/v1/quotation-lines/{line_id}/container-tiers/",
            {"container_config_id": str(config.pk), "unit_price": "12.5000"},
            format="json",
            HTTP_IF_MATCH=packaging_response["ETag"],
        )
        self.assertEqual(tier_response.status_code, 201, tier_response.data)
        self.assertEqual(tier_response.data["data"]["auto_carton_count"], 400)
        self.assertEqual(tier_response.data["data"]["auto_quantity"], "4000.0000")
        self.assertEqual(tier_response.data["data"]["unit_price"], "12.5000")
        delete_response = self.client.delete(
            f"/api/v1/quotation-container-tiers/{tier_response.data['data']['id']}/",
            HTTP_IF_MATCH=tier_response["ETag"],
        )
        self.assertEqual(delete_response.status_code, 204, delete_response.data)

    def test_quote_scoped_container_tier_does_not_create_global_configuration(self):
        quotation_response = self.create_quotation()
        version_response = self.create_version(
            quotation_response.data["data"]["id"],
            quotation_response["ETag"],
        )
        line_response = self.add_line(version_response.data["data"]["id"], version_response["ETag"])
        line_id = line_response.data["data"]["id"]
        packaging_response = self.client.put(
            f"/api/v1/quotation-lines/{line_id}/packaging/",
            {"length_cm": "50", "width_cm": "40", "height_cm": "30", "pieces_per_carton": 10, "carton_count": 5},
            format="json",
            HTTP_IF_MATCH=line_response["ETag"],
        )
        self.assertEqual(packaging_response.status_code, 200, packaging_response.data)
        initial_global_count = QuotationContainerConfig.objects.count()

        tier_response = self.client.post(
            f"/api/v1/quotation-lines/{line_id}/container-tiers/",
            {"container_type_code": "40HQ", "effective_volume_cbm": "30.0000", "load_factor": "0.800000", "unit_price": "18.7500"},
            format="json",
            HTTP_IF_MATCH=packaging_response["ETag"],
        )

        self.assertEqual(tier_response.status_code, 201, tier_response.data)
        self.assertIsNone(tier_response.data["data"]["container_config"])
        self.assertEqual(tier_response.data["data"]["container_type_code"], "40HQ")
        self.assertEqual(tier_response.data["data"]["auto_carton_count"], 400)
        self.assertEqual(tier_response.data["data"]["final_quantity"], "4000.0000")
        self.assertEqual(tier_response.data["data"]["unit_price"], "18.7500")
        self.assertEqual(QuotationContainerConfig.objects.count(), initial_global_count)

        list_response = self.client.get(f"/api/v1/quotation-lines/{line_id}/container-tiers/")
        self.assertEqual(list_response.status_code, 200, list_response.data)
        self.assertEqual(list_response.data["data"][0]["container_type_code"], "40HQ")
        self.assertIsNone(list_response.data["data"][0]["container_config"])

        duplicate_response = self.client.post(
            f"/api/v1/quotation-lines/{line_id}/container-tiers/",
            {"container_type_code": "40HQ", "effective_volume_cbm": "30.0000", "load_factor": "0.800000", "unit_price": "18.7500"},
            format="json",
            HTTP_IF_MATCH=tier_response["ETag"],
        )
        self.assertEqual(duplicate_response.status_code, 400, duplicate_response.data)

    def test_packaging_cbm_rounds_half_up_to_three_decimal_places(self):
        quotation_response = self.create_quotation()
        version_response = self.create_version(
            quotation_response.data["data"]["id"],
            quotation_response["ETag"],
        )
        line_response = self.add_line(version_response.data["data"]["id"], version_response["ETag"])

        response = self.client.put(
            f"/api/v1/quotation-lines/{line_response.data['data']['id']}/packaging/",
            {"length_cm": "10", "width_cm": "10", "height_cm": "5", "pieces_per_carton": 1, "carton_count": 1},
            format="json",
            HTTP_IF_MATCH=line_response["ETag"],
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["carton_cbm"], "0.001")
        self.assertEqual(response.data["data"]["total_cbm"], "0.001")

    def test_packaging_dimensions_that_round_to_zero_return_validation_error(self):
        quotation_response = self.create_quotation()
        version_response = self.create_version(
            quotation_response.data["data"]["id"],
            quotation_response["ETag"],
        )
        line_response = self.add_line(version_response.data["data"]["id"], version_response["ETag"])

        response = self.client.put(
            f"/api/v1/quotation-lines/{line_response.data['data']['id']}/packaging/",
            {"length_cm": "1", "width_cm": "1", "height_cm": "1", "pieces_per_carton": 1, "carton_count": 1},
            format="json",
            HTTP_IF_MATCH=line_response["ETag"],
        )
        after = self.client.get(f"/api/v1/quotation-lines/{line_response.data['data']['id']}/packaging/")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("0.001 m³", str(response.data))
        self.assertEqual(after.status_code, 200, after.data)
        self.assertIsNone(after.data["data"])

    def lock_version(self, version_id, etag):
        return self.client.post(
            f"/api/v1/quotation-versions/{version_id}/lock/",
            {
                "address_id": str(self.address.pk),
                "contact_id": str(self.contact.pk),
                "tax_rate": "0.100000",
                "terms": {"payment": "T/T 30 days"},
            },
            format="json",
            HTTP_IF_MATCH=etag,
        )

    def test_create_quotation_generates_number_audit_and_timeline(self):
        response = self.create_quotation()
        quotation = Quotation.objects.get(pk=response.data["data"]["id"])

        self.assertRegex(quotation.quotation_number, r"^QUO-\d{6}-\d{6,}$")
        self.assertEqual(quotation.outcome_status, "OPEN")
        self.assertTrue(
            AuditLog.objects.filter(object_type="QUOTATION", object_id=quotation.pk, action="CREATED").exists()
        )
        self.assertTrue(
            CustomerTimelineEvent.objects.filter(
                customer_account=self.customer,
                source_id=quotation.pk,
                event_type="QUOTATION_CREATED",
            ).exists()
        )

    def test_creation_options_only_returns_customers_with_create_scope_and_entities(self):
        response = self.client.get("/api/v1/quotations/creation-options/")
        self.assertEqual(response.status_code, 200, response.data)
        customers = response.data["data"]
        self.assertEqual([item["id"] for item in customers], [str(self.customer.pk)])
        self.assertEqual(customers[0]["entities"], [{
            "id": str(self.entity.pk),
            "entity_code": "QUOTE_ENTITY",
            "legal_name": "报价交易主体",
        }])

    def test_create_draft_creates_quotation_and_first_version_in_one_request(self):
        response = self.client.post(
            "/api/v1/quotations/create-draft/",
            {
                "customer_account_id": str(self.customer.pk),
                "transaction_entity_id": str(self.entity.pk),
                "owner_employee_id": str(self.employee.pk),
                "subject": "事务草稿报价",
                "currency_code": "USD",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        quotation = Quotation.objects.get(pk=response.data["data"]["id"])
        self.assertEqual(quotation.versions.count(), 1)
        self.assertEqual(quotation.versions.get().status, QuotationVersionStatus.DRAFT)
        self.assertIn("Quotation-Version-ETag", response)

    def test_list_does_not_leak_unassigned_customer_quotations(self):
        visible_response = self.create_quotation()
        hidden_entity = CustomerTransactionEntity.objects.create(
            customer_account=self.other_customer,
            entity_code="HIDDEN_QUOTE_ENTITY",
            legal_name="隐藏主体",
        )
        hidden = Quotation.objects.create(
            customer_account=self.other_customer,
            transaction_entity=hidden_entity,
            owner_employee=self.other_employee,
            quotation_number="HIDDEN_QUOTATION",
            subject="不可见报价",
        )

        response = self.client.get("/api/v1/quotations/")
        returned_ids = {item["id"] for item in response.data["data"]}

        self.assertIn(visible_response.data["data"]["id"], returned_ids)
        self.assertNotIn(str(hidden.pk), returned_ids)

    def test_line_total_tiers_lock_and_amounts_are_server_calculated(self):
        quotation = self.create_quotation()
        quotation_id = quotation.data["data"]["id"]
        version = self.create_version(quotation_id, quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        line_id = line.data["data"]["id"]
        tier = self.client.post(
            f"/api/v1/quotation-lines/{line_id}/quantity-tiers/",
            {"min_quantity": "1.0000", "max_quantity": "100.0000", "unit_price": "19.0000"},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, tier.headers["ETag"])
        locked_line = self.client.patch(
            f"/api/v1/quotation-lines/{line_id}/",
            {"unit_price": "1.0000"},
            format="json",
            HTTP_IF_MATCH=locked.headers["ETag"],
        )

        self.assertEqual(tier.status_code, 201, tier.data)
        self.assertEqual(locked.status_code, 200, locked.data)
        self.assertEqual(locked.data["data"]["subtotal_amount"], "180.0000")
        self.assertEqual(locked.data["data"]["discount_amount"], "20.0000")
        self.assertEqual(locked.data["data"]["tax_amount"], "18.0000")
        self.assertEqual(locked.data["data"]["total_amount"], "198.0000")
        self.assertEqual(locked_line.status_code, 409, locked_line.data)

    def test_overlapping_quantity_tier_is_rejected(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        line = self.add_line(version.data["data"]["id"], version.headers["ETag"])
        url = f"/api/v1/quotation-lines/{line.data['data']['id']}/quantity-tiers/"
        first = self.client.post(
            url,
            {"min_quantity": "1", "max_quantity": "100", "unit_price": "19"},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        overlap = self.client.post(
            url,
            {"min_quantity": "50", "max_quantity": "150", "unit_price": "18"},
            format="json",
            HTTP_IF_MATCH=first.headers["ETag"],
        )

        self.assertEqual(first.status_code, 201, first.data)
        self.assertEqual(overlap.status_code, 409, overlap.data)

    def test_copy_version_copies_lines_and_tiers_with_new_ids(self):
        quotation = self.create_quotation()
        quotation_id = quotation.data["data"]["id"]
        first = self.create_version(quotation_id, quotation.headers["ETag"])
        line = self.add_line(first.data["data"]["id"], first.headers["ETag"])
        tier = self.client.post(
            f"/api/v1/quotation-lines/{line.data['data']['id']}/quantity-tiers/",
            {"min_quantity": "1", "unit_price": "18"},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        first_version = Quotation.objects.get(pk=quotation_id).versions.get(pk=first.data["data"]["id"])
        first_version.address_snapshot = {"address_line1": "California"}
        first_version.contact_snapshot = {"display_name": "Buyer A"}
        first_version.save(update_fields=["address_snapshot", "contact_snapshot", "updated_at"])
        quotation_etag = self.client.get(f"/api/v1/quotations/{quotation_id}/").headers["ETag"]
        first_detail = self.client.get(f"/api/v1/quotation-versions/{first.data['data']['id']}/")
        copied = self.create_version(
            quotation_id,
            quotation_etag,
            from_version=first.data["data"]["id"],
        )

        self.assertEqual(copied.data["data"]["version_number"], 2)
        self.assertEqual(str(copied.data["data"]["parent_version"]), first.data["data"]["id"])
        self.assertEqual(copied.data["data"]["total_amount"], first_detail.data["data"]["total_amount"])
        self.assertEqual(len(copied.data["data"]["lines"]), 1)
        self.assertNotEqual(copied.data["data"]["lines"][0]["id"], line.data["data"]["id"])
        self.assertEqual(str(copied.data["data"]["lines"][0]["lineage_id"]), str(line.data["data"]["lineage_id"]))
        self.assertEqual(len(copied.data["data"]["lines"][0]["quantity_tiers"]), 1)
        self.assertEqual(copied.data["data"]["address_snapshot"]["address_line1"], "California")
        self.assertEqual(copied.data["data"]["contact_snapshot"]["display_name"], "Buyer A")
        diff = self.client.get(
            f"/api/v1/quotation-versions/{copied.data['data']['id']}/diff/",
            {"base_version_id": first.data["data"]["id"]},
        )
        self.assertEqual(diff.status_code, 200, diff.data)
        self.assertEqual(diff.data["data"]["changes"], [])

    def test_version_diff_includes_terms_price_basis_and_real_line_price_changes(self):
        quotation = self.create_quotation()
        quotation_id = quotation.data["data"]["id"]
        first = self.create_version(quotation_id, quotation.headers["ETag"])
        line = self.add_line(first.data["data"]["id"], first.headers["ETag"])
        first_matrix = self.client.put(
            f"/api/v1/quotation-versions/{first.data['data']['id']}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        self.assertEqual(first_matrix.status_code, 200, first_matrix.data)

        quotation_etag = self.client.get(f"/api/v1/quotations/{quotation_id}/").headers["ETag"]
        copied = self.create_version(
            quotation_id,
            quotation_etag,
            from_version=first.data["data"]["id"],
        )
        copied_id = copied.data["data"]["id"]
        copied_line_id = copied.data["data"]["lines"][0]["id"]
        updated = self.client.patch(
            f"/api/v1/quotation-versions/{copied_id}/",
            {
                "valid_until": (timezone.localdate() + timedelta(days=30)).isoformat(),
                "terms_snapshot": {"payment_terms": "30% 订金 / 70% 余款", "lead_time": "35 天"},
            },
            format="json",
            HTTP_IF_MATCH=copied.headers["ETag"],
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        copied_matrix = self.client.put(
            f"/api/v1/quotation-versions/{copied_id}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": copied_line_id, "unit_price": "19.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=updated.headers["ETag"],
        )
        self.assertEqual(copied_matrix.status_code, 200, copied_matrix.data)

        diff = self.client.get(
            f"/api/v1/quotation-versions/{copied_id}/diff/",
            {"base_version_id": first.data["data"]["id"]},
        )
        fields = {item["field"] for item in diff.data["data"]["changes"]}
        self.assertEqual(diff.status_code, 200, diff.data)
        self.assertIn("付款条件", fields)
        self.assertIn("交期", fields)
        self.assertIn("测试产品 · FOB 宁波", fields)
        self.assertIn("测试产品 单价", fields)
        self.assertIn("报价合计", fields)

    def test_required_approval_blocks_issue_until_approver_accepts(self):
        template = QuotationTemplate.objects.create(
            template_code="APPROVAL_TEMPLATE",
            template_name="需审批模板",
            requires_approval=True,
        )
        quotation = self.create_quotation()
        version = self.create_version(
            quotation.data["data"]["id"], quotation.headers["ETag"], template=template
        )
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version_id}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, matrix.headers["ETag"])
        submitted = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/submit-approval/",
            {"approver_employee_id": str(self.approver.pk), "reason": "请审批价格"},
            format="json",
            HTTP_IF_MATCH=locked.headers["ETag"],
        )
        blocked = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {},
            format="json",
            HTTP_IF_MATCH=submitted.headers["Quotation-Version-ETag"],
        )
        self.client.force_authenticate(user=self.approver_user)
        approved = self.client.post(
            f"/api/v1/approvals/{submitted.data['data']['id']}/approve/",
            {"note": "价格合理"},
            format="json",
            HTTP_IF_MATCH=submitted.headers["ETag"],
        )
        self.client.force_authenticate(user=self.user)
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {},
            format="json",
            HTTP_IF_MATCH=approved.headers["Quotation-Version-ETag"],
        )

        self.assertEqual(submitted.status_code, 200, submitted.data)
        self.assertEqual(blocked.status_code, 409, blocked.data)
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertEqual(approved.data["data"]["status"], ApprovalStatus.APPROVED)
        self.assertEqual(issued.status_code, 200, issued.data)
        self.assertEqual(issued.data["data"]["status"], QuotationVersionStatus.ISSUED)

    def test_document_dispatch_and_accept_flow(self):
        quotation = self.create_quotation()
        quotation_id = quotation.data["data"]["id"]
        version = self.create_version(quotation_id, quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version_id}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, matrix.headers["ETag"])
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {},
            format="json",
            HTTP_IF_MATCH=locked.headers["ETag"],
        )
        fernet_key = "MDEyMzQ1Njc4OWFiY2RlZjAxMjM0NTY3ODlhYmNkZWY="
        with TemporaryDirectory() as media_root, self.settings(
            MEDIA_ROOT=media_root,
            EMAIL_ACCOUNT_FERNET_KEYS=[fernet_key],
        ), patch(
            "apps.quotations.uploads._clamav_scan", return_value="ClamAV test clean"
        ), patch("apps.quotations.services.send_via_email_account") as send_email:
            EmailAccount.objects.create(
                employee=self.employee,
                display_name="报价测试邮箱",
                sender_name="报价测试",
                sender_email="quotes@example.test",
                smtp_host="smtp.example.test",
                smtp_port=587,
                smtp_username="quotes@example.test",
                smtp_password_ciphertext=encrypt_smtp_password("test-app-password"),
                smtp_security=SmtpSecurity.STARTTLS,
                is_default=True,
            )
            document = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/formal-document/",
                {"file": SimpleUploadedFile("quotation.pdf", b"%PDF-1.4\n%%EOF", content_type="application/pdf")},
                format="multipart",
                HTTP_IF_MATCH=issued.headers["ETag"],
            )
            emailed = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                {
                    "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000001",
                    "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                    "recipient": "buyer@example.test",
                    "subject": "Formal quotation",
                    "body": "Please find our formal quotation attached.",
                },
                format="json",
                HTTP_IF_MATCH=document.headers["ETag"],
            )
            replayed = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                {
                    "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000001",
                    "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                    "recipient": "buyer@example.test",
                    "subject": "Formal quotation",
                    "body": "Please find our formal quotation attached.",
                },
                format="json",
                HTTP_IF_MATCH=document.headers["ETag"],
            )
            resent = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                {
                    "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000002",
                    "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                    "recipient": "buyer-again@example.test",
                    "subject": "Formal quotation follow-up",
                    "body": "Please find the quotation attached again.",
                },
                format="json",
                HTTP_IF_MATCH=emailed.headers["ETag"],
            )
            self.assertEqual(send_email.call_count, 2)
            EmailAccount.objects.all().delete()
            failed_before_smtp = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                {
                    "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000004",
                    "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                    "recipient": "buyer@example.test",
                    "subject": "Formal quotation without account",
                    "body": "This request must fail before SMTP.",
                },
                format="json",
                HTTP_IF_MATCH=resent.headers["ETag"],
            )
            self.assertEqual(send_email.call_count, 2)
        dispatched = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/manual-dispatch/",
            {
                "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000003",
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "recipient_snapshot": {"name": "报价联系人", "email": "buyer@example.com"},
                "evidence_note": "业务员确认已通过客户指定渠道发送。",
            },
            format="json",
            HTTP_IF_MATCH=resent.headers["ETag"],
        )
        quote_etag = self.client.get(f"/api/v1/quotations/{quotation_id}/").headers["ETag"]
        accepted = self.client.post(
            f"/api/v1/quotations/{quotation_id}/accept/",
            {
                "accepted_version_id": version_id,
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "customer_contact_id": str(self.contact.pk),
                "accepted_at": timezone.now().isoformat(),
                "source": "EMAIL",
                "note": "客户接受",
            },
            format="json",
            HTTP_IF_MATCH=quote_etag,
        )
        blocked_manual_resend = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/manual-dispatch/",
            {
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "recipient_snapshot": {"name": "报价联系人"},
                "evidence_note": "报价已接受后不应再产生发送记录。",
            },
            format="json",
            HTTP_IF_MATCH=dispatched.headers["ETag"],
        )
        blocked_generic_resend = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/dispatch/",
            {
                "channel": "MANUAL",
                "recipient_snapshot": {"name": "报价联系人"},
                "external_reference": "SHOULD-NOT-BE-CREATED",
            },
            format="json",
            HTTP_IF_MATCH=dispatched.headers["ETag"],
        )

        self.assertEqual(issued.status_code, 200, issued.data)
        self.assertEqual(document.status_code, 200, document.data)
        self.assertEqual(len(document.data["data"]["documents"]), 1)
        self.assertEqual(emailed.status_code, 200, emailed.data)
        self.assertEqual(replayed.status_code, 200, replayed.data)
        self.assertEqual(replayed.data["email_attempt"]["status"], QuotationEmailDispatchAttemptStatus.ACCEPTED)
        self.assertEqual(emailed.data["data"]["status"], QuotationVersionStatus.SENT)
        self.assertFalse(emailed.data["data"]["dispatches"][0]["is_simulated"])
        self.assertEqual(resent.status_code, 200, resent.data)
        self.assertEqual(failed_before_smtp.status_code, 409, failed_before_smtp.data)
        self.assertEqual(
            QuotationEmailDispatchAttempt.objects.get(
                quotation_version_id=version_id,
                idempotency_key="b010f3b8-7fe1-4b56-a111-000000000004",
            ).status,
            QuotationEmailDispatchAttemptStatus.FAILED,
        )
        email_dispatches = [item for item in resent.data["data"]["dispatches"] if item["channel"] == "EMAIL"]
        self.assertEqual(len(email_dispatches), 2)
        self.assertEqual(
            {item["recipient_snapshot"]["email"] for item in email_dispatches},
            {"buyer@example.test", "buyer-again@example.test"},
        )
        self.assertEqual(dispatched.status_code, 200, dispatched.data)
        self.assertEqual(dispatched.data["data"]["status"], QuotationVersionStatus.SENT)
        self.assertEqual(accepted.status_code, 200, accepted.data)
        self.assertEqual(accepted.data["data"]["outcome_status"], "WON")
        self.assertEqual(str(accepted.data["data"]["accepted_version"]), version_id)
        self.assertEqual(blocked_manual_resend.status_code, 409, blocked_manual_resend.data)
        self.assertEqual(blocked_generic_resend.status_code, 409, blocked_generic_resend.data)
        self.assertIn("不能再次发送", str(blocked_manual_resend.data))
        self.assertEqual(
            QuotationDispatchRecord.objects.filter(quotation_version_id=version_id).count(),
            3,
        )
        attempt = QuotationEmailDispatchAttempt.objects.get(
            quotation_version_id=version_id,
            idempotency_key="b010f3b8-7fe1-4b56-a111-000000000001",
        )
        self.assertEqual(attempt.body_html, "Please find our formal quotation attached.")
        self.assertEqual(attempt.attachment_snapshot["content_type"], "application/pdf")
        summary = self.client.get("/api/v1/quotations/summary/")
        self.assertEqual(summary.status_code, 200, summary.data)
        self.assertEqual(summary.data["data"]["accepted"], 1)
        self.assertEqual(summary.data["data"]["sent"], 0)

        attempts_url = f"/api/v1/quotation-versions/{version_id}/email-attempts/"
        self.client.force_authenticate(user=self.other_user)
        forbidden_attempts = self.client.get(attempts_url)
        self.assertEqual(forbidden_attempts.status_code, 403, forbidden_attempts.data)

        dispatch_permission = Permission.objects.get(
            content_type__app_label="quotations",
            codename="dispatch_quotation",
        )
        self.other_user.user_permissions.add(dispatch_permission)
        self.other_user = get_user_model().objects.get(pk=self.other_user.pk)
        self.client.force_authenticate(user=self.other_user)
        out_of_scope_attempts = self.client.get(attempts_url)
        self.assertEqual(out_of_scope_attempts.status_code, 404, out_of_scope_attempts.data)

    def test_manual_dispatch_uses_evidence_without_requiring_formal_document(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version_id}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, matrix.headers["ETag"])
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {},
            format="json",
            HTTP_IF_MATCH=locked.headers["ETag"],
        )

        email_without_document = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
            {
                "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000004",
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "recipient": "buyer@example.test",
                "subject": "Formal quotation",
                "body": "Please find our formal quotation attached.",
            },
            format="json",
            HTTP_IF_MATCH=issued.headers["ETag"],
        )
        manual = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/manual-dispatch/",
            {
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "recipient_snapshot": {"recipient": "WhatsApp buyer"},
                "evidence_note": "业务员确认已通过 WhatsApp 发送。",
            },
            format="json",
            HTTP_IF_MATCH=issued.headers["ETag"],
        )

        self.assertEqual(email_without_document.status_code, 409, email_without_document.data)
        self.assertEqual(manual.status_code, 200, manual.data)
        self.assertEqual(manual.data["data"]["status"], QuotationVersionStatus.SENT)
        self.assertFalse(manual.data["data"]["documents"])

    def test_uncertain_email_is_recovered_by_manual_resolution_without_resending(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version_id}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, matrix.headers["ETag"])
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {}, format="json", HTTP_IF_MATCH=locked.headers["ETag"],
        )
        fernet_key = "MDEyMzQ1Njc4OWFiY2RlZjAxMjM0NTY3ODlhYmNkZWY="
        idempotency_key = "b010f3b8-7fe1-4b56-a111-000000000005"
        with TemporaryDirectory() as media_root, self.settings(
            MEDIA_ROOT=media_root,
            EMAIL_ACCOUNT_FERNET_KEYS=[fernet_key],
        ):
            EmailAccount.objects.create(
                employee=self.employee,
                display_name="结果待核对测试邮箱",
                sender_email="quotes@example.test",
                smtp_host="smtp.example.test",
                smtp_port=587,
                smtp_username="quotes@example.test",
                smtp_password_ciphertext=encrypt_smtp_password("test-app-password"),
                smtp_security=SmtpSecurity.STARTTLS,
                is_default=True,
            )
            generated = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/generate-formal-pdf/",
                {}, format="json", HTTP_IF_MATCH=issued.headers["ETag"],
            )
            payload = {
                "idempotency_key": idempotency_key,
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "recipient": "buyer@example.test",
                "cc": ["sales@example.test"],
                "subject": "Formal quotation",
                "body": "<p>Please check the attached quotation.</p>",
            }
            with patch(
                "apps.quotations.services.send_via_email_account",
                side_effect=EmailDeliveryUncertain("socket closed after DATA"),
            ) as send_email:
                uncertain = self.client.post(
                    f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                    payload,
                    format="json",
                    HTTP_IF_MATCH=generated.headers["ETag"],
                )
                same_key = self.client.post(
                    f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                    payload,
                    format="json",
                    HTTP_IF_MATCH=generated.headers["ETag"],
                )
                self.assertEqual(send_email.call_count, 1)

            attempts = self.client.get(
                f"/api/v1/quotation-versions/{version_id}/email-attempts/",
                {"idempotency_key": idempotency_key},
            )
            attempt_id = attempts.data["data"][0]["id"]
            resolved = self.client.post(
                f"/api/v1/quotation-email-attempts/{attempt_id}/resolve/",
                {"result": "ACCEPTED", "note": "已核对发件箱，确认邮件已投递。"},
                format="json",
                HTTP_IF_MATCH=generated.headers["ETag"],
            )
            not_sent_payload = {
                **payload,
                "idempotency_key": "b010f3b8-7fe1-4b56-a111-000000000006",
                "subject": "Second uncertain attempt",
            }
            with patch(
                "apps.quotations.services.send_via_email_account",
                side_effect=EmailDeliveryUncertain("connection lost after DATA"),
            ):
                second_uncertain = self.client.post(
                    f"/api/v1/quotation-versions/{version_id}/email-dispatch/",
                    not_sent_payload,
                    format="json",
                    HTTP_IF_MATCH=resolved.headers["ETag"],
                )
            second_attempt = QuotationEmailDispatchAttempt.objects.get(
                quotation_version_id=version_id,
                idempotency_key=not_sent_payload["idempotency_key"],
            )
            marked_not_sent = self.client.post(
                f"/api/v1/quotation-email-attempts/{second_attempt.pk}/resolve/",
                {"result": "NOT_SENT", "note": "已核对发件箱，确认没有这封邮件。"},
                format="json",
                HTTP_IF_MATCH=resolved.headers["ETag"],
            )

        self.assertEqual(uncertain.status_code, 409, uncertain.data)
        self.assertEqual(same_key.status_code, 409, same_key.data)
        self.assertEqual(attempts.status_code, 200, attempts.data)
        self.assertEqual(attempts.data["data"][0]["status"], QuotationEmailDispatchAttemptStatus.UNCERTAIN)
        self.assertEqual(attempts.data["data"][0]["body_html"], payload["body"])
        self.assertEqual(resolved.status_code, 200, resolved.data)
        self.assertEqual(resolved.data["email_attempt"]["status"], QuotationEmailDispatchAttemptStatus.ACCEPTED)
        self.assertEqual(resolved.data["data"]["status"], QuotationVersionStatus.SENT)
        self.assertEqual(second_uncertain.status_code, 409, second_uncertain.data)
        self.assertEqual(marked_not_sent.status_code, 200, marked_not_sent.data)
        self.assertEqual(marked_not_sent.data["email_attempt"]["status"], QuotationEmailDispatchAttemptStatus.FAILED)
        self.assertEqual(
            QuotationDispatchRecord.objects.filter(quotation_version_id=version_id, channel="EMAIL").count(),
            1,
        )

    def test_fixed_template_pdf_is_generated_registered_and_downloadable(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        line = self.add_line(version_id, version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version_id}/price-matrix/",
            {"price_bases": [
                {
                    "label": "EXW Shenzhen", "incoterm_code": "EXW", "named_place": "Shenzhen",
                    "currency_code": "USD", "is_default": False,
                    "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "18.0000"}],
                },
                {
                    "label": "FOB Ningbo", "incoterm_code": "FOB", "named_place": "Ningbo",
                    "currency_code": "USD", "is_default": True,
                    "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.5000"}],
                },
            ]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        locked = self.lock_version(version_id, matrix.headers["ETag"])
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {}, format="json", HTTP_IF_MATCH=locked.headers["ETag"],
        )

        with TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            generated = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/generate-formal-pdf/",
                {}, format="json", HTTP_IF_MATCH=issued.headers["ETag"],
            )
            regenerated = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/generate-formal-pdf/",
                {}, format="json", HTTP_IF_MATCH=generated.headers["ETag"],
            )
            dispatched = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/manual-dispatch/",
                {
                    "price_basis_id": regenerated.data["data"]["price_bases"][0]["id"],
                    "recipient_snapshot": {"recipient": "WhatsApp buyer"},
                    "evidence_note": "测试人工发送证据。",
                },
                format="json",
                HTTP_IF_MATCH=regenerated.headers["ETag"],
            )
            blocked_regeneration = self.client.post(
                f"/api/v1/quotation-versions/{version_id}/generate-formal-pdf/",
                {}, format="json", HTTP_IF_MATCH=dispatched.headers["ETag"],
            )
            downloaded = self.client.get(
                f"/api/v1/quotation-versions/{version_id}/formal-document-download/"
            )
            pdf_content = b"".join(downloaded.streaming_content)

        self.assertEqual(generated.status_code, 200, generated.data)
        self.assertEqual(generated.data["data"]["status"], QuotationVersionStatus.ISSUED)
        self.assertEqual(len(generated.data["data"]["documents"]), 1)
        self.assertTrue(generated.data["data"]["documents"][0]["is_formal"])
        self.assertEqual(regenerated.status_code, 200, regenerated.data)
        self.assertEqual(len(regenerated.data["data"]["documents"]), 2)
        self.assertEqual(sum(item["is_formal"] for item in regenerated.data["data"]["documents"]), 1)
        self.assertEqual(dispatched.status_code, 200, dispatched.data)
        self.assertEqual(blocked_regeneration.status_code, 409, blocked_regeneration.data)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded["Content-Type"], "application/pdf")
        self.assertTrue(pdf_content.startswith(b"%PDF"))

    def test_stale_etag_and_cross_customer_entity_are_rejected(self):
        foreign_entity = CustomerTransactionEntity.objects.create(
            customer_account=self.other_customer,
            entity_code="FOREIGN_QUOTE_ENTITY",
            legal_name="其他客户主体",
        )
        rejected = self.client.post(
            "/api/v1/quotations/",
            {
                "customer_account_id": str(self.customer.pk),
                "transaction_entity_id": str(foreign_entity.pk),
                "owner_employee_id": str(self.employee.pk),
                "subject": "跨客户报价",
            },
            format="json",
        )
        quotation = self.create_quotation()
        url = f"/api/v1/quotations/{quotation.data['data']['id']}/"
        changed = self.client.patch(
            url,
            {"subject": "已修改"},
            format="json",
            HTTP_IF_MATCH=quotation.headers["ETag"],
        )
        stale = self.client.patch(
            url,
            {"subject": "旧版本修改"},
            format="json",
            HTTP_IF_MATCH=quotation.headers["ETag"],
        )

        self.assertEqual(rejected.status_code, 400, rejected.data)
        self.assertIn("transaction_entity_id", rejected.data["error"]["field_errors"])
        self.assertEqual(changed.status_code, 200, changed.data)
        self.assertEqual(stale.status_code, 412, stale.data)

    def test_list_summary_and_options_return_quotation_list_contract(self):
        quotation = self.create_quotation()
        quotation_id = quotation.data["data"]["id"]
        self.create_version(quotation_id, quotation.headers["ETag"])

        listing = self.client.get("/api/v1/quotations/")
        summary = self.client.get("/api/v1/quotations/summary/")
        options = self.client.get("/api/v1/quotations/options/")
        exported = self.client.get("/api/v1/quotations/export/")

        self.assertEqual(listing.status_code, 200, listing.data)
        self.assertEqual(summary.status_code, 200, summary.data)
        self.assertEqual(options.status_code, 200, options.data)
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(
            exported["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(listing.data["data"][0]["display_status"], "DRAFT")
        self.assertEqual(listing.data["data"][0]["current_version"]["version_number"], 1)
        self.assertEqual(summary.data["data"]["all"], 1)
        self.assertEqual(summary.data["data"]["draft"], 1)
        self.assertEqual(options.data["data"]["currencies"][0]["value"], "USD")

    def test_price_history_only_returns_visible_real_quotation_snapshots(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        self.add_line(version.data["data"]["id"], version.headers["ETag"])
        now = timezone.now()
        Quotation.objects.get(pk=quotation.data["data"]["id"]).versions.filter(
            pk=version.data["data"]["id"]
        ).update(
            status="SENT",
            locked_at=now,
            issued_at=now,
            transaction_entity_snapshot={"legal_name": "测试交易主体"},
            address_snapshot={"address_line1": "测试地址"},
            contact_snapshot={"display_name": "测试联系人"},
        )

        response = self.client.get(
            "/api/v1/quotations/price-history/",
            {"customer_account_id": str(self.customer.pk), "product": "测试产品", "period_months": 12},
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["source_label"], "真实报价快照，不含订单成交、成本或市场价格")
        self.assertEqual(response.data["data"]["calculation_warning"], "")
        self.assertEqual(response.data["data"]["records"][0]["status"], "SENT")
        self.assertEqual(len(response.data["data"]["records"]), 1)
        self.assertEqual(response.data["data"]["records"][0]["unit_price"], "20.0000")
        self.assertEqual(response.data["data"]["metrics"]["latest_quote_price"], "20.0000")

    def test_price_history_blocks_average_when_units_are_not_comparable(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        first_line = self.add_line(version.data["data"]["id"], version.headers["ETag"])
        second_line = self.client.post(
            f"/api/v1/quotation-versions/{version.data['data']['id']}/lines/",
            {
                "product_description_snapshot": "测试产品",
                "specification_snapshot": "整箱规格",
                "quantity": "2.0000",
                "unit_code": "CTN",
                "unit_price": "200.0000",
                "discount_rate": "0",
            },
            format="json",
            HTTP_IF_MATCH=first_line.headers["ETag"],
        )
        self.assertEqual(second_line.status_code, 201, second_line.data)
        now = timezone.now()
        Quotation.objects.get(pk=quotation.data["data"]["id"]).versions.filter(
            pk=version.data["data"]["id"]
        ).update(
            status="SENT",
            locked_at=now,
            issued_at=now,
            transaction_entity_snapshot={"legal_name": "测试交易主体"},
            address_snapshot={"address_line1": "测试地址"},
            contact_snapshot={"display_name": "测试联系人"},
        )

        response = self.client.get(
            "/api/v1/quotations/price-history/",
            {"customer_account_id": str(self.customer.pk), "product": "测试产品", "currency_code": "USD"},
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["data"]["unit_codes"], ["CTN", "PCS"])
        self.assertIn("多个计量单位", response.data["data"]["calculation_warning"])
        self.assertIsNone(response.data["data"]["metrics"]["average_quote_price"])
        self.assertEqual(response.data["data"]["series"], [])

    def test_formal_xlsx_export_uses_frozen_quote_data_and_adds_product_blocks(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        first_line = self.add_line(version_id, version.headers["ETag"], "露营床 A", "牛津布 / 蓝色")
        first_tier = self.configure_export_packaging(first_line.data["data"]["id"], first_line["ETag"], unit_price="18.7500")
        second_line = self.add_line(version_id, first_tier["ETag"], "露营床 B", "铝合金 / 灰色")
        second_tier = self.configure_export_packaging(second_line.data["data"]["id"], second_line["ETag"], container_type="20GP", unit_price="22.5000")
        third_line = self.add_line(version_id, second_tier["ETag"], "露营床 C", "钢管 / 绿色")
        third_tier = self.configure_export_packaging(third_line.data["data"]["id"], third_line["ETag"], container_type="40H1", unit_price="25.0000")

        draft_download = self.client.get(f"/api/v1/quotation-versions/{version_id}/formal-xlsx-download/")
        self.assertEqual(draft_download.status_code, 400, draft_download.data)

        version_model = QuotationVersion.objects.select_related("quotation_template").get(pk=version_id)
        template = version_model.quotation_template
        template.template_config = {
            **template.template_config,
            "seller_company_name": "JINHUA SUPERGROUP LEISURE PRODUCTS CO.,LTD",
            "seller_fax": "+86 579 12345678",
            "seller_skype": "supergroup.sales",
            "seller_qq": "12345678",
        }
        template.save(update_fields=["template_config", "updated_at"])

        locked = self.lock_version(version_id, third_tier["ETag"])
        self.assertEqual(locked.status_code, 200, locked.data)
        issued = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/issue/",
            {},
            format="json",
            HTTP_IF_MATCH=locked["ETag"],
        )
        self.assertEqual(issued.status_code, 200, issued.data)
        exported = self.client.get(f"/api/v1/quotation-versions/{version_id}/formal-xlsx-download/")
        workbook = load_workbook(BytesIO(exported.content), data_only=False)
        worksheet = workbook.active

        self.assertEqual(exported.status_code, 200, exported.content)
        self.assertEqual(
            exported["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        quotation_number = Quotation.objects.get(pk=quotation.data["data"]["id"]).quotation_number
        self.assertEqual(worksheet["A3"].value, f"Quotation: {quotation_number}")
        self.assertTrue(str(worksheet["A6"].value).startswith("ITEM-"))
        self.assertEqual(worksheet["B4"].value, "露营床 A")
        self.assertEqual(worksheet["B12"].value, "露营床 B")
        self.assertEqual(worksheet["B20"].value, "露营床 C")
        self.assertEqual(worksheet["C6"].value, "牛津布 / 蓝色")
        self.assertEqual(worksheet["K6"].value, "40HQ")
        self.assertEqual(worksheet["L6"].value, 884)
        self.assertEqual(worksheet["M6"].value, 18.75)
        self.assertEqual(worksheet["K14"].value, "20GP")
        self.assertEqual(worksheet["M14"].value, 22.5)
        self.assertEqual(worksheet["K22"].value, "40H1")
        self.assertEqual(worksheet["M22"].value, 25)
        self.assertEqual(worksheet["B28"].value, "REMARKS:")
        self.assertNotIn("Fanny", " ".join(str(cell.value or "") for row in worksheet.iter_rows() for cell in row))
        repeated = self.client.get(f"/api/v1/quotation-versions/{version_id}/formal-xlsx-download/")
        repeated_workbook = load_workbook(BytesIO(repeated.content), data_only=False)
        self.assertEqual(repeated_workbook.active["A6"].value, worksheet["A6"].value)
        self.assertEqual(QuotationVersion.objects.get(pk=version_id).status, QuotationVersionStatus.ISSUED)
        with ZipFile(BytesIO(exported.content)) as exported_zip:
            drawing_xml = exported_zip.read("xl/drawings/drawing1.xml").decode("utf-8")
            self.assertIn("JINHUA SUPERGROUP LEISURE PRODUCTS CO.,LTD", drawing_xml)

    def test_simulated_dispatch_is_audited_without_progressing_quote_or_inquiry(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        version_id = version.data["data"]["id"]
        current = self.client.get(f"/api/v1/quotation-versions/{version_id}/")
        self.inquiry.status = "QUOTING"
        self.inquiry.save(update_fields=["status", "updated_at"])

        simulated = self.client.post(
            f"/api/v1/quotation-versions/{version_id}/simulate-dispatch/",
            {"recipient": "demo-recipient@example.test", "notes": "仅供演示"},
            format="json",
            HTTP_IF_MATCH=current.headers["ETag"],
        )
        self.inquiry.refresh_from_db()

        self.assertEqual(simulated.status_code, 200, simulated.data)
        self.assertEqual(simulated.data["data"]["status"], QuotationVersionStatus.DRAFT)
        self.assertTrue(simulated.data["data"]["dispatches"][0]["is_simulated"])
        self.assertEqual(self.inquiry.status, "QUOTING")

    def test_price_matrix_keeps_multiple_bases_and_syncs_default_price(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        line = self.add_line(version.data["data"]["id"], version.headers["ETag"])

        response = self.client.put(
            f"/api/v1/quotation-versions/{version.data['data']['id']}/price-matrix/",
            {
                "price_bases": [
                    {
                        "label": "EXW 宁波", "incoterm_code": "EXW", "named_place": "宁波",
                        "currency_code": "USD", "is_default": False,
                        "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "18.0000"}],
                    },
                    {
                        "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                        "currency_code": "USD", "is_default": True,
                        "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.5000"}],
                    },
                ]
            },
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(len(response.data["data"]["price_bases"]), 2)
        self.assertEqual(response.data["data"]["lines"][0]["unit_price"], "20.5000")
        self.assertEqual(response.data["data"]["total_amount"], "184.5000")

    def test_accept_records_version_price_basis_contact_and_source(self):
        quotation = self.create_quotation()
        version = self.create_version(quotation.data["data"]["id"], quotation.headers["ETag"])
        line = self.add_line(version.data["data"]["id"], version.headers["ETag"])
        matrix = self.client.put(
            f"/api/v1/quotation-versions/{version.data['data']['id']}/price-matrix/",
            {"price_bases": [{
                "label": "FOB 宁波", "incoterm_code": "FOB", "named_place": "宁波",
                "currency_code": "USD", "is_default": True,
                "line_prices": [{"quotation_line_id": line.data["data"]["id"], "unit_price": "20.0000"}],
            }]},
            format="json",
            HTTP_IF_MATCH=line.headers["ETag"],
        )
        version_object = self.customer.quotations.get(pk=quotation.data["data"]["id"]).versions.get()
        version_object.status = QuotationVersionStatus.SENT
        version_object.locked_at = timezone.now()
        version_object.transaction_entity_snapshot = {"legal_name": self.entity.legal_name}
        version_object.address_snapshot = {"address_line1": self.address.address_line1}
        version_object.contact_snapshot = {"display_name": self.contact.display_name}
        version_object.save(update_fields=[
            "status", "locked_at", "transaction_entity_snapshot", "address_snapshot",
            "contact_snapshot", "updated_at",
        ])
        quote_detail = self.client.get(f"/api/v1/quotations/{quotation.data['data']['id']}/")

        accepted = self.client.post(
            f"/api/v1/quotations/{quotation.data['data']['id']}/accept/",
            {
                "accepted_version_id": version.data["data"]["id"],
                "price_basis_id": matrix.data["data"]["price_bases"][0]["id"],
                "customer_contact_id": str(self.contact.pk),
                "accepted_at": timezone.now().isoformat(),
                "source": "EMAIL",
                "note": "客户邮件确认接受。",
            },
            format="json",
            HTTP_IF_MATCH=quote_detail.headers["ETag"],
        )

        self.assertEqual(accepted.status_code, 200, accepted.data)
        self.assertEqual(accepted.data["data"]["outcome_status"], "WON")
        self.assertEqual(accepted.data["data"]["acceptance"]["source"], "EMAIL")
        self.assertEqual(accepted.data["data"]["acceptance"]["customer_contact_id"], str(self.contact.pk))
        self.inquiry.refresh_from_db()
        self.assertEqual(self.inquiry.status, InquiryStatus.QUOTE_ACCEPTED)
