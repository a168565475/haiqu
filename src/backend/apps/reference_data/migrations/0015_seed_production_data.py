import hashlib
import shutil
import uuid
from pathlib import Path

from django.conf import settings
from django.db import migrations
from django.utils import timezone
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("04724775-ed7a-4c53-bca4-7e27f2ffaf42")
RESOURCE_ROOT = Path(__file__).resolve().parent.parent / "resources" / "production_data"
IMAGE_ROOT = RESOURCE_ROOT / "images"
ITEM_TYPE_MAP = {
    "成品": "FINISHED",
    "半成品": "SEMI_FINISHED",
    "材料": "MATERIAL",
}


def stable_id(kind, value):
    return uuid.uuid5(SEED_NAMESPACE, f"{kind}:{value}")


def clean_cell(value):
    return "" if value is None else str(value).strip()


def load_rows(filename, expected_headers, expected_count):
    workbook = load_workbook(RESOURCE_ROOT / filename, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = tuple(clean_cell(value) for value in rows[0])
    if headers != expected_headers:
        raise RuntimeError(f"{filename} 表头不符合预期：{headers}")
    data_rows = [tuple(clean_cell(value) for value in row) for row in rows[1:] if any(value is not None for value in row)]
    if len(data_rows) != expected_count:
        raise RuntimeError(f"{filename} 数据行数应为 {expected_count}，实际为 {len(data_rows)}。")
    return data_rows


def load_color_rows():
    workbook = load_workbook(RESOURCE_ROOT / "product_colors.xlsx", read_only=False, data_only=True)
    try:
        worksheet = workbook.active
        headers = tuple(clean_cell(cell.value) for cell in worksheet[1])
        expected_headers = ("色卡", "色号", "颜色名称", "图片", "唯一校验")
        if headers != expected_headers:
            raise RuntimeError(f"product_colors.xlsx 表头不符合预期：{headers}")
        rows = []
        for row_number in range(2, worksheet.max_row + 1):
            cells = worksheet[row_number]
            values = tuple(clean_cell(cell.value) for cell in cells)
            if not any(values):
                continue
            image_target = cells[3].hyperlink.target if cells[3].hyperlink else ""
            if image_target:
                image_name = Path(image_target.replace("\\", "/")).name
                if image_name == "image":
                    image_name = "combined-pantone-7771c-2324c.png"
            else:
                image_name = ""
            rows.append((*values[:3], image_name, values[4]))
    finally:
        workbook.close()
    if len(rows) != 324:
        raise RuntimeError(f"product_colors.xlsx 数据行数应为 324，实际为 {len(rows)}。")
    return rows


def validate_seed_data(categories, units, colors, items):
    category_names = [row[1] for row in categories]
    if len(category_names) != len(set(category_names)) or any(row[0] not in ITEM_TYPE_MAP for row in categories):
        raise RuntimeError("类别数据存在重复名称或未知品号类型。")
    unit_codes = [row[0] for row in units]
    unit_names = [row[1] for row in units]
    if len(unit_codes) != len(set(unit_codes)) or len(unit_names) != len(set(unit_names)):
        raise RuntimeError("单位编号或单位名称存在重复。")
    color_keys = [(row[0], row[1], row[2]) for row in colors]
    unique_checks = [row[4] for row in colors]
    if len(color_keys) != len(set(color_keys)) or len(unique_checks) != len(set(unique_checks)):
        raise RuntimeError("颜色组合键或唯一校验列存在重复。")
    item_ids = [row[0] for row in items]
    item_numbers = [row[2] for row in items]
    if len(item_ids) != len(set(item_ids)) or len(item_numbers) != len(set(item_numbers)):
        raise RuntimeError("货号 data_id 或货号存在重复。")
    missing_categories = sorted({row[1] for row in items} - set(category_names))
    if missing_categories:
        raise RuntimeError(f"货号引用了不存在的产品类别：{missing_categories}")
    image_rows = [row for row in colors if row[3]]
    if len(image_rows) != 23:
        raise RuntimeError(f"颜色图片记录应为 23 条，实际为 {len(image_rows)}。")
    for row in image_rows:
        image_path = IMAGE_ROOT / row[3]
        if not image_path.is_file():
            raise RuntimeError(f"颜色图片资源不存在：{row[3]}")
        content = image_path.read_bytes()
        if not content.startswith(b"\x89PNG\r\n\x1a\n") or not 0 < len(content) <= 5 * 1024 * 1024:
            raise RuntimeError(f"颜色图片不是有效的 5MB 以内 PNG：{row[3]}")


def seed_production_data(apps, schema_editor):
    ProductCategory = apps.get_model("reference_data", "ProductCategory")
    ProductionUnit = apps.get_model("reference_data", "ProductionUnit")
    ProductColor = apps.get_model("reference_data", "ProductColor")
    ProductItem = apps.get_model("reference_data", "ProductItem")
    BusinessFile = apps.get_model("files", "BusinessFile")
    db_alias = schema_editor.connection.alias

    categories = load_rows(
        "product_categories.xlsx",
        ("品号类型", "品号类别"),
        49,
    )
    units = load_rows(
        "production_units.xlsx",
        ("单位编号", "单位名称"),
        12,
    )
    colors = load_color_rows()
    items = load_rows(
        "product_items.xlsx",
        ("data_id", "产品类别", "货号"),
        35,
    )
    validate_seed_data(categories, units, colors, items)

    category_by_name = {}
    for sort_order, (source_type, category_name) in enumerate(categories, start=1):
        category_id = stable_id("PRODUCT_CATEGORY", category_name)
        category, _ = ProductCategory.objects.using(db_alias).update_or_create(
            id=category_id,
            defaults={
                "item_type": ITEM_TYPE_MAP[source_type],
                "category_name": category_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：类别维护.xlsx",
            },
        )
        category_by_name[category_name] = category

    for sort_order, (unit_code, unit_name) in enumerate(units, start=1):
        ProductionUnit.objects.using(db_alias).update_or_create(
            id=stable_id("PRODUCTION_UNIT", unit_code),
            defaults={
                "unit_code": unit_code,
                "unit_name": unit_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：单位维护.xlsx",
            },
        )

    media_root = Path(settings.MEDIA_ROOT)
    for sort_order, (color_card, color_code, color_name, image_name, _unique_check) in enumerate(colors, start=1):
        identity = f"{color_card}\u241f{color_code}\u241f{color_name}"
        image_file = None
        if image_name:
            image_path = IMAGE_ROOT / image_name
            content = image_path.read_bytes()
            digest = hashlib.sha256(content).hexdigest()
            business_file_id = stable_id("PRODUCT_COLOR_IMAGE", identity)
            storage_key = f"production-colors/seed/{business_file_id}.png"
            destination = media_root / storage_key
            destination.parent.mkdir(parents=True, exist_ok=True)
            if not destination.exists() or hashlib.sha256(destination.read_bytes()).hexdigest() != digest:
                shutil.copyfile(image_path, destination)
            image_file, _ = BusinessFile.objects.using(db_alias).update_or_create(
                id=business_file_id,
                defaults={
                    "uploaded_by_employee_id": None,
                    "original_name": image_name,
                    "storage_key": storage_key,
                    "content_type": "image/png",
                    "size_bytes": len(content),
                    "content_hash": digest,
                    "scan_status": "CLEAN",
                    "scanned_at": timezone.now(),
                    "scan_message": "项目内置且已校验签名的颜色图片。",
                },
            )
        ProductColor.objects.using(db_alias).update_or_create(
            id=stable_id("PRODUCT_COLOR", identity),
            defaults={
                "color_card": color_card,
                "color_code": color_code,
                "color_name": color_name,
                "image_file_id": image_file.pk if image_file else None,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：颜色维护.xlsx",
            },
        )

    for sort_order, (legacy_data_id, category_name, item_number) in enumerate(items, start=1):
        ProductItem.objects.using(db_alias).update_or_create(
            id=stable_id("PRODUCT_ITEM", legacy_data_id),
            defaults={
                "legacy_data_id": legacy_data_id,
                "category_id": category_by_name[category_name].pk,
                "item_number": item_number,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：货号维护.xlsx",
            },
        )


def remove_seeded_production_data(apps, schema_editor):
    ProductCategory = apps.get_model("reference_data", "ProductCategory")
    ProductionUnit = apps.get_model("reference_data", "ProductionUnit")
    ProductColor = apps.get_model("reference_data", "ProductColor")
    ProductItem = apps.get_model("reference_data", "ProductItem")
    BusinessFile = apps.get_model("files", "BusinessFile")
    db_alias = schema_editor.connection.alias

    business_files = list(
        BusinessFile.objects.using(db_alias).filter(storage_key__startswith="production-colors/seed/")
    )
    ProductItem.objects.using(db_alias).filter(description="来源：货号维护.xlsx").delete()
    ProductColor.objects.using(db_alias).filter(description="来源：颜色维护.xlsx").delete()
    ProductionUnit.objects.using(db_alias).filter(description="来源：单位维护.xlsx").delete()
    ProductCategory.objects.using(db_alias).filter(description="来源：类别维护.xlsx").delete()
    BusinessFile.objects.using(db_alias).filter(pk__in=[item.pk for item in business_files]).delete()
    media_root = Path(settings.MEDIA_ROOT).resolve()
    for business_file in business_files:
        file_path = (media_root / business_file.storage_key).resolve()
        try:
            file_path.relative_to(media_root)
        except ValueError:
            continue
        file_path.unlink(missing_ok=True)


class Migration(migrations.Migration):
    dependencies = [
        ("reference_data", "0014_production_data"),
    ]

    operations = [
        migrations.RunPython(seed_production_data, remove_seeded_production_data),
    ]
