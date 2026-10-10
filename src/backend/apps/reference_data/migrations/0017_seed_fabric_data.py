import uuid
from pathlib import Path

from django.db import migrations
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("49e52e78-b156-4f5c-a83b-5fd9d0148acb")
RESOURCE_ROOT = Path(__file__).resolve().parent.parent / "resources" / "fabric_data"


def stable_id(kind, value):
    return uuid.uuid5(SEED_NAMESPACE, f"{kind}:{value}")


def clean_cell(value):
    return "" if value is None else str(value).strip()


def load_values(filename, expected_header, expected_count):
    workbook = load_workbook(RESOURCE_ROOT / filename, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header = clean_cell(worksheet["A1"].value)
        values = [
            clean_cell(row[0])
            for row in worksheet.iter_rows(min_row=2, max_col=1, values_only=True)
        ]
    finally:
        workbook.close()
    if header != expected_header:
        raise RuntimeError(f"{filename} 表头应为 {expected_header}，实际为 {header}。")
    if len(values) != expected_count:
        raise RuntimeError(f"{filename} 数据行数应为 {expected_count}，实际为 {len(values)}。")
    if any(not value for value in values):
        raise RuntimeError(f"{filename} 存在空白数据。")
    if len(values) != len(set(values)):
        raise RuntimeError(f"{filename} 存在重复数据。")
    return values


def seed_fabric_data(apps, schema_editor):
    FabricMaterial = apps.get_model("reference_data", "FabricMaterial")
    FabricSpecification = apps.get_model("reference_data", "FabricSpecification")
    FabricWidth = apps.get_model("reference_data", "FabricWidth")
    db_alias = schema_editor.connection.alias

    materials = load_values("fabric_materials.xlsx", "材质", 7)
    specifications = load_values("fabric_specifications.xlsx", "规格", 16)
    widths = load_values("fabric_widths.xlsx", "门幅", 8)

    for sort_order, material_name in enumerate(materials, start=1):
        FabricMaterial.objects.using(db_alias).update_or_create(
            id=stable_id("FABRIC_MATERIAL", material_name),
            defaults={
                "material_name": material_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：面料材质维护_20260831105317.xlsx",
            },
        )

    for sort_order, specification_name in enumerate(specifications, start=1):
        FabricSpecification.objects.using(db_alias).update_or_create(
            id=stable_id("FABRIC_SPECIFICATION", specification_name),
            defaults={
                "specification_name": specification_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：面料规格维护_20260831104955.xlsx",
            },
        )

    for sort_order, width_value in enumerate(widths, start=1):
        FabricWidth.objects.using(db_alias).update_or_create(
            id=stable_id("FABRIC_WIDTH", width_value),
            defaults={
                "width_value": width_value,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": "来源：面料门幅维护_20260831105335.xlsx",
            },
        )


def remove_seeded_fabric_data(apps, schema_editor):
    FabricMaterial = apps.get_model("reference_data", "FabricMaterial")
    FabricSpecification = apps.get_model("reference_data", "FabricSpecification")
    FabricWidth = apps.get_model("reference_data", "FabricWidth")
    db_alias = schema_editor.connection.alias
    FabricMaterial.objects.using(db_alias).filter(
        description="来源：面料材质维护_20260831105317.xlsx"
    ).delete()
    FabricSpecification.objects.using(db_alias).filter(
        description="来源：面料规格维护_20260831104955.xlsx"
    ).delete()
    FabricWidth.objects.using(db_alias).filter(
        description="来源：面料门幅维护_20260831105335.xlsx"
    ).delete()


class Migration(migrations.Migration):
    dependencies = [
        ("reference_data", "0016_fabric_data"),
    ]

    operations = [
        migrations.RunPython(seed_fabric_data, remove_seeded_fabric_data),
    ]
