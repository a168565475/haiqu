import uuid
from pathlib import Path

from django.db import migrations
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("230d436b-b5dc-4f12-af46-dd1544ba9f34")
RESOURCE_ROOT = Path(__file__).resolve().parent.parent / "resources" / "tube_data"


SOURCE_SPECS = (
    ("TubeProductName", "product_name", "tube_product_names.xlsx", "品名", 79, "管材品名维护_20260831103828.xlsx"),
    ("TubeThickness", "thickness_value", "tube_thicknesses.xlsx", "厚度", 14, "管材厚度维护_20260831104010.xlsx"),
    ("TubeDiameter", "diameter_value", "tube_diameters.xlsx", "直径", 24, "管材直径维护_20260831103901.xlsx"),
    ("TubeMaterial", "material_name", "tube_materials.xlsx", "材质", 3, "管材材质维护_20260831104317.xlsx"),
    ("TubeShape", "shape_name", "tube_shapes.xlsx", "形状", 7, "管材形状维护_20260831104409.xlsx"),
    ("TubeCoatingColor", "color_name", "tube_coating_colors.xlsx", "喷塑颜色", 29, "管材喷塑颜色维护_20260831104441.xlsx"),
)


def stable_id(model_name, value):
    return uuid.uuid5(SEED_NAMESPACE, f"{model_name}:{value}")


def clean_cell(value):
    return "" if value is None else str(value).strip()


def load_values(filename, expected_header, expected_count):
    workbook = load_workbook(RESOURCE_ROOT / filename, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header = clean_cell(worksheet["A1"].value)
        values = [
            clean_cell(row[0])
            for row in worksheet.iter_rows(min_row=2, max_col=1, values_only=True)
        ]
    finally:
        workbook.close()
    if header != expected_header:
        raise RuntimeError(f"{filename} 表头应为 {expected_header}，实际为 {header}。")
    if len(values) != expected_count:
        raise RuntimeError(f"{filename} 数据行数应为 {expected_count}，实际为 {len(values)}。")
    if any(not value for value in values):
        raise RuntimeError(f"{filename} 存在空白数据。")
    if len(values) != len(set(values)):
        raise RuntimeError(f"{filename} 存在重复数据。")
    return values


def seed_tube_data(apps, schema_editor):
    db_alias = schema_editor.connection.alias
    validated_sources = []
    for model_name, field_name, filename, header, count, source_name in SOURCE_SPECS:
        validated_sources.append(
            (model_name, field_name, source_name, load_values(filename, header, count))
        )

    for model_name, field_name, source_name, values in validated_sources:
        model = apps.get_model("reference_data", model_name)
        for sort_order, value in enumerate(values, start=1):
            model.objects.using(db_alias).update_or_create(
                id=stable_id(model_name, value),
                defaults={
                    field_name: value,
                    "status": "ACTIVE",
                    "sort_order": sort_order,
                    "description": f"来源：{source_name}",
                },
            )


def remove_seeded_tube_data(apps, schema_editor):
    db_alias = schema_editor.connection.alias
    for model_name, _field_name, _filename, _header, _count, source_name in SOURCE_SPECS:
        model = apps.get_model("reference_data", model_name)
        model.objects.using(db_alias).filter(description=f"来源：{source_name}").delete()


class Migration(migrations.Migration):
    dependencies = [
        ("reference_data", "0018_tube_data"),
    ]

    operations = [
        migrations.RunPython(seed_tube_data, remove_seeded_tube_data),
    ]
