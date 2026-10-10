import uuid
from pathlib import Path

from django.db import migrations
from openpyxl import load_workbook


SEED_NAMESPACE = uuid.UUID("e39a6213-f64a-46b0-9c5a-d68255a8eb28")
RESOURCE_ROOT = Path(__file__).resolve().parent.parent / "resources" / "accessory_data"
DIRECT_NAME_SOURCE = "来源：辅料品名维护_20260831105443.xlsx"
SUPPLEMENTAL_NAME_SOURCE = "由辅料规格/颜色表关联补充；单位待维护。"
SPECIFICATION_SOURCE = "来源：辅料规格维护_20260831105206.xlsx"
COLOR_SOURCE = "来源：辅料颜色维护_20260831105533.xlsx"
MATERIAL_SOURCE = "来源：辅料材质维护_20260831105554.xlsx"


def stable_id(kind, *values):
    return uuid.uuid5(SEED_NAMESPACE, f"{kind}:" + "\0".join(values))


def clean_cell(value):
    return "" if value is None else str(value).strip()


def load_rows(filename, expected_headers, expected_count):
    workbook = load_workbook(RESOURCE_ROOT / filename, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        width = len(expected_headers)
        headers = tuple(clean_cell(cell.value) for cell in worksheet[1][:width])
        rows = [
            tuple(clean_cell(value) for value in row)
            for row in worksheet.iter_rows(min_row=2, max_col=width, values_only=True)
        ]
    finally:
        workbook.close()
    if headers != expected_headers:
        raise RuntimeError(f"{filename} 表头应为 {expected_headers}，实际为 {headers}。")
    if len(rows) != expected_count:
        raise RuntimeError(f"{filename} 数据行数应为 {expected_count}，实际为 {len(rows)}。")
    return rows


def validate_unique(values, label):
    if len(values) != len(set(values)):
        raise RuntimeError(f"{label} 存在重复数据。")


def load_and_validate_sources():
    names = load_rows(
        "accessory_product_names.xlsx",
        ("品名", "单位"),
        112,
    )
    specifications = load_rows(
        "accessory_specifications.xlsx",
        ("品名", "规格", "克重", "唯一校验"),
        349,
    )
    colors = load_rows(
        "accessory_colors.xlsx",
        ("品名", "颜色", "唯一校验"),
        284,
    )
    materials = load_rows(
        "accessory_materials.xlsx",
        ("材质",),
        27,
    )

    if any(not name or not unit for name, unit in names):
        raise RuntimeError("辅料品名或单位存在空白数据。")
    validate_unique([name for name, _unit in names], "辅料品名")

    if any(not name or not value or not validation for name, value, _weight, validation in specifications):
        raise RuntimeError("辅料规格存在空白的品名、规格或唯一校验值。")
    if any(validation != f"{name}{value}" for name, value, _weight, validation in specifications):
        raise RuntimeError("辅料规格的唯一校验值与品名＋规格不一致。")
    validate_unique([(name, value) for name, value, _weight, _validation in specifications], "辅料品名＋规格")

    if any(not name or not color or not validation for name, color, validation in colors):
        raise RuntimeError("辅料颜色存在空白的品名、颜色或唯一校验值。")
    if any(validation != f"{name}{color}" for name, color, validation in colors):
        raise RuntimeError("辅料颜色的唯一校验值与品名＋颜色不一致。")
    validate_unique([(name, color) for name, color, _validation in colors], "辅料品名＋颜色")

    if any(not material[0] for material in materials):
        raise RuntimeError("辅料材质存在空白数据。")
    validate_unique([material[0] for material in materials], "辅料材质")

    direct_names = {name for name, _unit in names}
    supplemental_names = []
    seen_names = set(direct_names)
    for name in [row[0] for row in specifications] + [row[0] for row in colors]:
        if name not in seen_names:
            supplemental_names.append(name)
            seen_names.add(name)
    if len(supplemental_names) != 36:
        raise RuntimeError(f"规格/颜色表应补充 36 个品名，实际为 {len(supplemental_names)}。")
    return names, supplemental_names, specifications, colors, materials


def seed_accessory_data(apps, schema_editor):
    AccessoryProductName = apps.get_model("reference_data", "AccessoryProductName")
    AccessorySpecification = apps.get_model("reference_data", "AccessorySpecification")
    AccessoryColor = apps.get_model("reference_data", "AccessoryColor")
    AccessoryMaterial = apps.get_model("reference_data", "AccessoryMaterial")
    db_alias = schema_editor.connection.alias

    names, supplemental_names, specifications, colors, materials = load_and_validate_sources()

    for sort_order, (product_name, unit_name) in enumerate(names, start=1):
        AccessoryProductName.objects.using(db_alias).update_or_create(
            id=stable_id("ACCESSORY_PRODUCT_NAME", product_name),
            defaults={
                "product_name": product_name,
                "unit_name": unit_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": DIRECT_NAME_SOURCE,
            },
        )
    for sort_order, product_name in enumerate(supplemental_names, start=len(names) + 1):
        AccessoryProductName.objects.using(db_alias).update_or_create(
            id=stable_id("ACCESSORY_PRODUCT_NAME", product_name),
            defaults={
                "product_name": product_name,
                "unit_name": "",
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": SUPPLEMENTAL_NAME_SOURCE,
            },
        )

    for sort_order, (product_name, specification_value, weight, _validation) in enumerate(specifications, start=1):
        AccessorySpecification.objects.using(db_alias).update_or_create(
            id=stable_id("ACCESSORY_SPECIFICATION", product_name, specification_value),
            defaults={
                "product_name_id": stable_id("ACCESSORY_PRODUCT_NAME", product_name),
                "specification_value": specification_value,
                "weight": weight,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": SPECIFICATION_SOURCE,
            },
        )

    for sort_order, (product_name, color_name, _validation) in enumerate(colors, start=1):
        AccessoryColor.objects.using(db_alias).update_or_create(
            id=stable_id("ACCESSORY_COLOR", product_name, color_name),
            defaults={
                "product_name_id": stable_id("ACCESSORY_PRODUCT_NAME", product_name),
                "color_name": color_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": COLOR_SOURCE,
            },
        )

    for sort_order, (material_name,) in enumerate(materials, start=1):
        AccessoryMaterial.objects.using(db_alias).update_or_create(
            id=stable_id("ACCESSORY_MATERIAL", material_name),
            defaults={
                "material_name": material_name,
                "status": "ACTIVE",
                "sort_order": sort_order,
                "description": MATERIAL_SOURCE,
            },
        )


def remove_seeded_accessory_data(apps, schema_editor):
    AccessoryProductName = apps.get_model("reference_data", "AccessoryProductName")
    AccessorySpecification = apps.get_model("reference_data", "AccessorySpecification")
    AccessoryColor = apps.get_model("reference_data", "AccessoryColor")
    AccessoryMaterial = apps.get_model("reference_data", "AccessoryMaterial")
    db_alias = schema_editor.connection.alias
    AccessorySpecification.objects.using(db_alias).filter(description=SPECIFICATION_SOURCE).delete()
    AccessoryColor.objects.using(db_alias).filter(description=COLOR_SOURCE).delete()
    AccessoryMaterial.objects.using(db_alias).filter(description=MATERIAL_SOURCE).delete()
    AccessoryProductName.objects.using(db_alias).filter(
        description__in=[DIRECT_NAME_SOURCE, SUPPLEMENTAL_NAME_SOURCE]
    ).delete()


class Migration(migrations.Migration):
    dependencies = [
        ("reference_data", "0020_accessory_data"),
    ]

    operations = [
        migrations.RunPython(seed_accessory_data, remove_seeded_accessory_data),
    ]
