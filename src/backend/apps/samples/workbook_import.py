import hashlib
import json
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.apps import apps
from django.conf import settings
from django.db import connection, transaction
from django.db.models import Q
from django.db.models.deletion import ProtectedError
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apps.customers.models import CustomerAccount, CustomerLifecycleStage
from apps.files.models import BusinessFile
from apps.organizations.models import Employee
from apps.reference_data.models import (
    ProductCategory,
    ProductColor,
    ProductItemType,
    ReferenceStatus,
)
from apps.workflow.models import BusinessTask

from .models import (
    FeeApprovalStatus,
    FeeVerificationStatus,
    SampleDeliveryBatch,
    SampleFeeDecision,
    SampleFeeRecord,
    SampleFileLink,
    SampleItem,
    SampleProduct,
    SampleProductSamplingDetail,
    SampleRequest,
    SampleRequestStatus,
    SampleShippingMethod,
    SampleTaskLink,
    SampleTolerance,
    SampleType,
    SampleWorkflowRecord,
    SampleWorkflowStage,
)


SHEET_NAME = "打样单"
IMPORT_EMPLOYEE_CODE = "SYSTEM_ENTE_IMPORT"
PLACEHOLDER_CUSTOMER_CODE = "LEGACY_SAMPLE_HISTORY"
PLACEHOLDER_CUSTOMER_SOURCE = "SAMPLE_IMPORT"
PLACEHOLDER_CUSTOMER_EXTERNAL_ID = "SAMPLE_WORKBOOK_HISTORY"
MISSING_CATEGORY_NAME = "矮脚斜躺沙滩椅"

SHIPPING_METHODS = {
    "邮寄": SampleShippingMethod.MAIL,
    "其他": SampleShippingMethod.OTHER,
}
DELIVERY_BATCHES = {
    "一次交货": SampleDeliveryBatch.SINGLE,
    "分批交货": SampleDeliveryBatch.PARTIAL,
    "分配交货": SampleDeliveryBatch.PARTIAL,
}
SAMPLE_TYPES = {
    "是": SampleType.PRE_PRODUCTION,
    "否": SampleType.PRE_SALES,
}
TOLERANCES = {
    "": SampleTolerance.NONE,
    "无": SampleTolerance.NONE,
    "±0.5cm": SampleTolerance.PLUS_MINUS_0_5_CM,
    "±1cm": SampleTolerance.PLUS_MINUS_1_CM,
}

COLUMN_NUMBER = {
    "A": 1,
    "B": 2,
    "C": 3,
    "D": 4,
    "E": 5,
    "F": 6,
    "G": 7,
    "H": 8,
    "I": 9,
    "J": 10,
    "K": 11,
    "L": 12,
    "M": 13,
    "N": 14,
    "O": 15,
    "P": 16,
    "Q": 17,
    "R": 18,
    "S": 19,
    "T": 20,
    "U": 21,
    "V": 22,
    "W": 23,
    "X": 24,
    "Y": 25,
    "Z": 26,
    "AA": 27,
    "AB": 28,
    "AC": 29,
    "AD": 30,
    "AE": 31,
    "AF": 32,
    "AG": 33,
    "AH": 34,
    "AI": 35,
    "AJ": 36,
    "AK": 37,
    "AL": 38,
    "AM": 39,
    "AN": 40,
    "AO": 41,
    "AP": 42,
    "AQ": 43,
    "AR": 44,
    "AS": 45,
    "AT": 46,
    "AU": 47,
    "AV": 48,
    "AW": 49,
    "AX": 50,
    "AY": 51,
    "AZ": 52,
    "BA": 53,
    "BB": 54,
    "BC": 55,
    "BD": 56,
    "BE": 57,
    "BF": 58,
    "BG": 59,
    "BH": 60,
    "BI": 61,
    "BJ": 62,
    "BK": 63,
    "BL": 64,
    "BM": 65,
    "BN": 66,
    "BO": 67,
    "BP": 68,
    "BQ": 69,
    "BR": 70,
    "BS": 71,
    "BT": 72,
    "BU": 73,
    "BV": 74,
}


class SampleWorkbookImportError(ValueError):
    pass


@dataclass
class SampleWorkbookImportPlan:
    source_path: Path
    source_hash: str
    samples: list[dict]
    corrections: list[dict] = field(default_factory=list)
    missing_attachments: list[dict] = field(default_factory=list)
    source_audit: list[dict] = field(default_factory=list)
    categories_to_create: list[str] = field(default_factory=list)


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _cell(row, column):
    return row[COLUMN_NUMBER[column] - 1]


def _required_text(value, *, row_number, label):
    text = _text(value)
    if not text:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不能为空。")
    return text


def _parse_yes_no(value, *, row_number, label, allow_blank=False):
    text = _text(value)
    if allow_blank and not text:
        return False
    if text == "是":
        return True
    if text == "否":
        return False
    raise SampleWorkbookImportError(f"第 {row_number} 行{label}只能填写“是”或“否”。")


def _parse_positive_integer(value, *, row_number, label):
    try:
        number = Decimal(_text(value))
    except InvalidOperation as exc:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效数字。") from exc
    if number <= 0 or number != number.to_integral_value():
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}必须是正整数。")
    return int(number)


def _parse_nonnegative_integer(value, *, row_number, label):
    if value is None or _text(value) == "":
        return 0
    try:
        number = Decimal(_text(value))
    except InvalidOperation as exc:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效数字。") from exc
    if number < 0 or number != number.to_integral_value():
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}必须是非负整数。")
    return int(number)


def _parse_optional_nonnegative_decimal(value, *, row_number, label):
    if value is None or _text(value) == "":
        return None
    try:
        number = Decimal(_text(value))
    except InvalidOperation as exc:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效金额。") from exc
    if number < 0 or number.as_tuple().exponent < -4:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}必须是非负数且最多四位小数。")
    return number


def _parse_dimension(value, *, row_number, label):
    if value is None or _text(value) == "":
        return None
    try:
        number = Decimal(_text(value))
    except InvalidOperation as exc:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效数字。") from exc
    if number <= 0 or number.as_tuple().exponent < -2:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}必须大于0且最多两位小数。")
    return number


def _parse_date(value, *, workbook_epoch, row_number, label, required=True):
    if value is None or _text(value) == "":
        if required:
            raise SampleWorkbookImportError(f"第 {row_number} 行{label}不能为空。")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        parsed = from_excel(value, workbook_epoch)
        return parsed.date() if isinstance(parsed, datetime) else parsed
    try:
        return date.fromisoformat(_text(value))
    except ValueError as exc:
        raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效日期。") from exc


def _parse_datetime(value, *, workbook_epoch, row_number, label, fallback_date=None):
    if value is None or _text(value) == "":
        if fallback_date is None:
            return None
        parsed = datetime.combine(fallback_date, time.min)
    elif isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, (int, float)):
        parsed = from_excel(value, workbook_epoch)
        if isinstance(parsed, date) and not isinstance(parsed, datetime):
            parsed = datetime.combine(parsed, time.min)
    else:
        try:
            parsed = datetime.fromisoformat(_text(value).replace("Z", "+00:00"))
        except ValueError as exc:
            raise SampleWorkbookImportError(f"第 {row_number} 行{label}不是有效日期时间。") from exc
    return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed


def _split_file_names(value):
    text = _text(value)
    if not text:
        return []
    for separator in ("，", "；", ";", "\n"):
        text = text.replace(separator, ",")
    return [item.strip() for item in text.split(",") if item.strip()]


def _record_missing_attachments(target, *, request_number, source_row, field_name, value):
    for file_name in _split_file_names(value):
        target.append(
            {
                "request_number": request_number,
                "source_row": source_row,
                "field": field_name,
                "file_name": file_name,
            }
        )


def _normalize_request_numbers(manual_value, generated_value, *, row_number, corrections):
    raw_manual = _required_text(manual_value, row_number=row_number, label="打样单号-手填").upper()
    manual = raw_manual[4:] if raw_manual.startswith("HTS-") else raw_manual
    expected = f"HTS-{manual}"
    raw_generated = _required_text(generated_value, row_number=row_number, label="生成打样单号").upper()
    if raw_manual.startswith("HTS-") and raw_generated == f"HTS-{raw_manual}":
        corrections.append(
            {
                "source_row": row_number,
                "field": "生成打样单号",
                "from": raw_generated,
                "to": expected,
                "reason": "去除重复 HTS- 前缀",
            }
        )
        raw_generated = expected
    if raw_generated != expected:
        raise SampleWorkbookImportError(
            f"第 {row_number} 行生成打样单号 {raw_generated} 与手填号 {manual} 不一致。"
        )
    if len(manual) > 46 or len(raw_generated) > 50:
        raise SampleWorkbookImportError(f"第 {row_number} 行打样单号超过系统长度限制。")
    return manual, raw_generated


def _status_for_row(row, *, row_number):
    process_status = _text(_cell(row, "BT"))
    current_node = _text(_cell(row, "BU"))
    if process_status == "流转完成":
        return SampleRequestStatus.FINISHED
    if process_status == "进行中" and current_node == "金工上传参数":
        return SampleRequestStatus.METALWORK_PENDING
    if process_status == "进行中" and current_node == "裁剪上传参数":
        return SampleRequestStatus.CUTTING_PENDING
    if not process_status and not current_node and (_cell(row, "AY") or _cell(row, "BC")):
        return SampleRequestStatus.FINISHED
    raise SampleWorkbookImportError(
        f"第 {row_number} 行流程状态无法映射：{process_status or '空'} / {current_node or '空'}。"
    )


def _sampling_selection(raw_category):
    category = _text(raw_category)
    if category == "是":
        return "", True
    if category == "否" or not category:
        return "", False
    return category, True


def _validate_headers(sheet):
    expected = {
        "A1": "标题",
        "B1": "打样单编号",
        "C1": "打样单号-手填",
        "D1": "生成打样单号",
        "X1": "样品明细",
        "AL1": "打样信息明细",
        "AP1": "特殊说明",
        "BT1": "流程状态",
        "BU1": "当前节点",
        "X2": "产品类别",
        "Y2": "产品名称",
        "Z2": "规格",
        "AA2": "色卡",
        "AB2": "颜色",
        "AC2": "色号",
        "AD2": "打样数量",
        "AK2": "样品明细UUID",
        "AL2": "类型",
        "AM2": "类别",
        "AN2": "描述",
        "AO2": "图片",
    }
    mismatches = [
        f"{cell} 应为“{label}”"
        for cell, label in expected.items()
        if _text(sheet[cell].value) != label
    ]
    if mismatches:
        raise SampleWorkbookImportError("表格结构与已确认模板不一致：" + "；".join(mismatches))


def _source_audit_row(row, *, row_number, request_number):
    fields = {
        "业务员": _text(_cell(row, "E")),
        "制单人": _text(_cell(row, "L")),
        "跟单员": _text(_cell(row, "AR")),
        "审批人": _text(_cell(row, "AT")),
        "金工上传人": _text(_cell(row, "AX")),
        "裁剪上传人": _text(_cell(row, "BB")),
        "提交人": _text(_cell(row, "BQ")),
        "当前负责人": _text(_cell(row, "BV")),
        "客户编号": _text(_cell(row, "BD")),
        "客户代码": _text(_cell(row, "BK")),
        "客户名称": _text(_cell(row, "BM")),
    }
    return {
        "request_number": request_number,
        "source_row": row_number,
        **fields,
    }


def _finalize_sample(sample):
    if not sample["items"]:
        raise SampleWorkbookImportError(f"第 {sample['source_row']} 行样品没有样品明细。")
    item_total = sum(item["sample_quantity"] for item in sample["items"])
    if item_total != sample["total_quantity"]:
        raise SampleWorkbookImportError(
            f"{sample['request_number']} 总打样数量为 {sample['total_quantity']}，明细合计为 {item_total}。"
        )
    for item in sample["items"]:
        if item["cutting_reported_quantity"] > item["sample_quantity"]:
            raise SampleWorkbookImportError(f"第 {item['source_row']} 行裁剪报工数量超过打样数量。")
        if item["metalwork_reported_quantity"] > item["sample_quantity"]:
            raise SampleWorkbookImportError(f"第 {item['source_row']} 行金工报工数量超过打样数量。")


def build_sample_workbook_import_plan(source_path):
    source_path = Path(source_path)
    if not source_path.is_file():
        raise SampleWorkbookImportError(f"源文件不存在：{source_path}")
    if source_path.suffix.lower() != ".xlsx":
        raise SampleWorkbookImportError("样品导入只支持 XLSX 文件。")
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()

    import_employee = Employee.objects.filter(employee_code=IMPORT_EMPLOYEE_CODE).first()
    if import_employee is None:
        raise SampleWorkbookImportError(f"缺少导入服务员工：{IMPORT_EMPLOYEE_CODE}")

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SampleWorkbookImportError(f"缺少工作表：{SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        _validate_headers(sheet)
        workbook_epoch = workbook.epoch
        rows = list(sheet.iter_rows(min_row=3, max_col=COLUMN_NUMBER["BV"], values_only=True))
    finally:
        workbook.close()

    categories = {item.category_name: item for item in ProductCategory.objects.all()}
    colors = {
        (item.color_card, item.color_name, item.color_code): item
        for item in ProductColor.objects.filter(status=ReferenceStatus.ACTIVE)
    }
    corrections = []
    missing_attachments = []
    source_audit = []
    categories_to_create = set()
    samples = []
    current = None
    seen_request_numbers = set()
    seen_manual_numbers = set()
    seen_item_ids = set()

    for offset, row in enumerate(rows, start=3):
        if _text(_cell(row, "A")):
            if current is not None:
                _finalize_sample(current)
                samples.append(current)

            manual_number, request_number = _normalize_request_numbers(
                _cell(row, "C"),
                _cell(row, "D"),
                row_number=offset,
                corrections=corrections,
            )
            if request_number in seen_request_numbers:
                raise SampleWorkbookImportError(f"正式编号重复：{request_number}")
            if manual_number in seen_manual_numbers:
                raise SampleWorkbookImportError(f"手填编号重复：{manual_number}")
            seen_request_numbers.add(request_number)
            seen_manual_numbers.add(manual_number)

            category_name = _required_text(_cell(row, "N"), row_number=offset, label="产品类别")
            if category_name not in categories:
                if category_name == MISSING_CATEGORY_NAME:
                    categories_to_create.add(category_name)
                else:
                    raise SampleWorkbookImportError(f"第 {offset} 行产品类别未维护：{category_name}")

            shipping_text = _required_text(_cell(row, "F"), row_number=offset, label="寄送方式")
            delivery_text = _required_text(_cell(row, "I"), row_number=offset, label="交货批次")
            sample_type_text = _required_text(_cell(row, "J"), row_number=offset, label="产前确认样")
            if shipping_text not in SHIPPING_METHODS:
                raise SampleWorkbookImportError(f"第 {offset} 行寄送方式无法映射：{shipping_text}")
            if delivery_text not in DELIVERY_BATCHES:
                raise SampleWorkbookImportError(f"第 {offset} 行交货批次无法映射：{delivery_text}")
            if sample_type_text not in SAMPLE_TYPES:
                raise SampleWorkbookImportError(f"第 {offset} 行产前确认样无法映射：{sample_type_text}")

            submitted_on = _parse_date(
                _cell(row, "M"),
                workbook_epoch=workbook_epoch,
                row_number=offset,
                label="提交日期",
            )
            requested_at = _parse_datetime(
                _cell(row, "BR"),
                workbook_epoch=workbook_epoch,
                row_number=offset,
                label="提交时间",
                fallback_date=submitted_on,
            )
            updated_at = _parse_datetime(
                _cell(row, "BS"),
                workbook_epoch=workbook_epoch,
                row_number=offset,
                label="更新时间",
                fallback_date=submitted_on,
            )
            sample_due_date = _parse_date(
                _cell(row, "H"),
                workbook_epoch=workbook_epoch,
                row_number=offset,
                label="样品交期",
            )
            current = {
                "source_row": offset,
                "manual_request_number": manual_number,
                "request_number": request_number,
                "shipping_method": SHIPPING_METHODS[shipping_text],
                "order_date": _parse_date(
                    _cell(row, "G"),
                    workbook_epoch=workbook_epoch,
                    row_number=offset,
                    label="下单日期",
                ),
                "sample_due_date": sample_due_date,
                "delivery_batch": DELIVERY_BATCHES[delivery_text],
                "sample_type": SAMPLE_TYPES[sample_type_text],
                "is_repeat_order": _parse_yes_no(_cell(row, "K"), row_number=offset, label="翻单"),
                "submitted_on": submitted_on,
                "purpose": _text(_cell(row, "AP")),
                "notes": _text(_cell(row, "AS")),
                "requested_at": requested_at,
                "requested_due_at": _parse_datetime(
                    None,
                    workbook_epoch=workbook_epoch,
                    row_number=offset,
                    label="样品交期",
                    fallback_date=sample_due_date,
                ),
                "updated_at": updated_at,
                "status": _status_for_row(row, row_number=offset),
                "approval_text": _text(_cell(row, "AQ")),
                "approval_at": _parse_datetime(
                    _cell(row, "AU"),
                    workbook_epoch=workbook_epoch,
                    row_number=offset,
                    label="审批时间",
                    fallback_date=submitted_on,
                ),
                "metalwork_at": _parse_datetime(
                    _cell(row, "AY"),
                    workbook_epoch=workbook_epoch,
                    row_number=offset,
                    label="金工上传时间",
                ),
                "cutting_at": _parse_datetime(
                    _cell(row, "BC"),
                    workbook_epoch=workbook_epoch,
                    row_number=offset,
                    label="裁剪上传时间",
                ),
                "metalwork_files": _split_file_names(_cell(row, "AV")),
                "cutting_files": _split_file_names(_cell(row, "AZ")),
                "fee_amount": _parse_optional_nonnegative_decimal(
                    _cell(row, "BJ"), row_number=offset, label="打样费"
                ),
                "category_name": category_name,
                "product_name": _required_text(_cell(row, "O"), row_number=offset, label="产品名称"),
                "specification": _required_text(_cell(row, "P"), row_number=offset, label="规格"),
                "length_cm": _parse_dimension(_cell(row, "S"), row_number=offset, label="长"),
                "armrest_width_cm": _parse_dimension(_cell(row, "T"), row_number=offset, label="扶手宽"),
                "seat_width_cm": _parse_dimension(_cell(row, "U"), row_number=offset, label="坐宽"),
                "back_height_cm": _parse_dimension(_cell(row, "V"), row_number=offset, label="背高"),
                "seat_height_cm": _parse_dimension(_cell(row, "W"), row_number=offset, label="坐高"),
                "total_quantity": _parse_positive_integer(_cell(row, "R"), row_number=offset, label="总打样数量"),
                "items": [],
                "sampling_details": [],
                "source_people": _source_audit_row(row, row_number=offset, request_number=request_number),
            }
            source_audit.append(current["source_people"])
            _record_missing_attachments(
                missing_attachments,
                request_number=request_number,
                source_row=offset,
                field_name="任务单",
                value=_cell(row, "Q"),
            )
            _record_missing_attachments(
                missing_attachments,
                request_number=request_number,
                source_row=offset,
                field_name="产品及展示图",
                value=_cell(row, "BI"),
            )
            _record_missing_attachments(
                missing_attachments,
                request_number=request_number,
                source_row=offset,
                field_name="金工参数",
                value=_cell(row, "AV"),
            )
            _record_missing_attachments(
                missing_attachments,
                request_number=request_number,
                source_row=offset,
                field_name="裁剪参数",
                value=_cell(row, "AZ"),
            )

        if current is None:
            if any(_text(value) for value in row):
                raise SampleWorkbookImportError(f"第 {offset} 行出现在首条样品记录之前。")
            continue

        has_item = any(_text(_cell(row, column)) for column in ("X", "Y", "Z", "AA", "AB", "AC", "AD", "AK"))
        if has_item:
            item_id_text = _required_text(_cell(row, "AK"), row_number=offset, label="样品明细UUID")
            try:
                item_id = uuid.UUID(item_id_text)
            except ValueError as exc:
                raise SampleWorkbookImportError(f"第 {offset} 行样品明细UUID无效。") from exc
            if item_id in seen_item_ids:
                raise SampleWorkbookImportError(f"样品明细UUID重复：{item_id}")
            seen_item_ids.add(item_id)

            quantity = _parse_positive_integer(_cell(row, "AD"), row_number=offset, label="打样数量")
            cutting_quantity = _parse_nonnegative_integer(
                _cell(row, "AG"), row_number=offset, label="裁剪报工数量"
            )
            metalwork_quantity = _parse_nonnegative_integer(
                _cell(row, "AH"), row_number=offset, label="金工报工数量"
            )
            color_card = _required_text(_cell(row, "AA"), row_number=offset, label="色卡")
            color_name = _required_text(_cell(row, "AB"), row_number=offset, label="颜色")
            color_code = _required_text(_cell(row, "AC"), row_number=offset, label="色号")
            lookup_card = color_card
            if color_card == "潘通定色" and color_name == "7743C（近墨绿）" and color_code == "7743C":
                lookup_card = "潘通色卡"
                corrections.append(
                    {
                        "source_row": offset,
                        "field": "色卡",
                        "from": color_card,
                        "to": lookup_card,
                        "reason": "关联现有同色号维护记录，快照保留原文",
                    }
                )
            product_color = colors.get((lookup_card, color_name, color_code))
            if product_color is None:
                raise SampleWorkbookImportError(
                    f"第 {offset} 行颜色无法关联：{color_card} / {color_name} / {color_code}"
                )
            tolerance_text = _text(_cell(row, "AE"))
            if tolerance_text not in TOLERANCES:
                raise SampleWorkbookImportError(f"第 {offset} 行公差无法映射：{tolerance_text}")
            current["items"].append(
                {
                    "source_row": offset,
                    "id": item_id,
                    "category_name": _required_text(_cell(row, "X"), row_number=offset, label="明细产品类别"),
                    "product_name": _required_text(_cell(row, "Y"), row_number=offset, label="明细产品名称"),
                    "specification": _required_text(_cell(row, "Z"), row_number=offset, label="明细规格"),
                    "product_color_id": product_color.pk,
                    "color_card": color_card,
                    "color_name": color_name,
                    "color_code": color_code,
                    "sample_quantity": quantity,
                    "tolerance": TOLERANCES[tolerance_text],
                    "notes": _text(_cell(row, "AF")),
                    "cutting_reported_quantity": cutting_quantity,
                    "metalwork_reported_quantity": metalwork_quantity,
                    "cutting_completed": _parse_yes_no(
                        _cell(row, "AI"), row_number=offset, label="裁剪是否报工完成"
                    ),
                    "metalwork_completed": _parse_yes_no(
                        _cell(row, "AJ"), row_number=offset, label="金工是否报工完成"
                    ),
                }
            )

        detail_type = _text(_cell(row, "AL"))
        if detail_type:
            category_label, is_selected = _sampling_selection(_cell(row, "AM"))
            current["sampling_details"].append(
                {
                    "source_row": offset,
                    "detail_type": detail_type,
                    "category_label": category_label,
                    "is_selected": is_selected,
                    "description": _text(_cell(row, "AN")),
                }
            )
            _record_missing_attachments(
                missing_attachments,
                request_number=current["request_number"],
                source_row=offset,
                field_name="打样信息明细图片",
                value=_cell(row, "AO"),
            )

    if current is not None:
        _finalize_sample(current)
        samples.append(current)

    if not samples:
        raise SampleWorkbookImportError("工作表没有可导入的样品记录。")
    if len(samples) != len(seen_request_numbers) or len(samples) != len(seen_manual_numbers):
        raise SampleWorkbookImportError("样品编号唯一性校验失败。")

    fee_rows = [sample for sample in samples if sample["fee_amount"] is not None]
    if len(fee_rows) != 1 or fee_rows[0]["request_number"] != "HTS-26008" or fee_rows[0]["fee_amount"] != Decimal("100"):
        raise SampleWorkbookImportError("打样费数据与已确认口径不一致，应仅有 HTS-26008 的人民币 100 元。")

    return SampleWorkbookImportPlan(
        source_path=source_path,
        source_hash=source_hash,
        samples=samples,
        corrections=corrections,
        missing_attachments=missing_attachments,
        source_audit=source_audit,
        categories_to_create=sorted(categories_to_create),
    )


def sample_import_plan_summary(plan):
    return {
        "sample_count": len(plan.samples),
        "product_count": len(plan.samples),
        "item_count": sum(len(sample["items"]) for sample in plan.samples),
        "sampling_detail_count": sum(len(sample["sampling_details"]) for sample in plan.samples),
        "status_counts": dict(sorted(Counter(sample["status"] for sample in plan.samples).items())),
        "correction_count": len(plan.corrections),
        "missing_attachment_count": len(plan.missing_attachments),
        "categories_to_create": plan.categories_to_create,
        "source_hash": plan.source_hash,
    }


def _models_with_existing_tables():
    existing_tables = set(connection.introspection.table_names())
    return [
        model
        for model in apps.get_models()
        if model._meta.managed and model._meta.db_table in existing_tables
    ]


def _raw_delete_tree(root_model, root_ids):
    root_ids = set(root_ids)
    if not root_ids:
        return {}
    models = _models_with_existing_tables()
    collected = {root_model: root_ids}
    depths = {root_model: 0}

    changed = True
    while changed:
        changed = False
        for model in models:
            relation_query = Q()
            relation_depths = []
            for model_field in model._meta.concrete_fields:
                if not model_field.many_to_one:
                    continue
                target_model = model_field.remote_field.model
                if target_model not in collected or target_model is model:
                    continue
                relation_query |= Q(**{f"{model_field.attname}__in": collected[target_model]})
                relation_depths.append(depths[target_model] + 1)
            if not relation_depths:
                continue
            found_ids = set(model._base_manager.filter(relation_query).values_list(model._meta.pk.attname, flat=True))
            new_ids = found_ids - collected.get(model, set())
            if new_ids:
                collected.setdefault(model, set()).update(new_ids)
                depths.setdefault(model, max(relation_depths))
                changed = True

    deleted = {}
    for model in sorted(collected, key=lambda item: depths[item], reverse=True):
        count = model._base_manager.filter(pk__in=collected[model])._raw_delete(connection.alias)
        if count:
            deleted[model._meta.label] = count
    return deleted


def _sample_business_files(sample_ids):
    file_ids = set(
        SampleFileLink.objects.filter(sample_request_id__in=sample_ids).values_list("business_file_id", flat=True)
    )
    file_ids.update(
        SampleProduct.objects.filter(sample_request_id__in=sample_ids)
        .exclude(task_sheet_file_id=None)
        .values_list("task_sheet_file_id", flat=True)
    )
    file_ids.update(
        SampleProductSamplingDetail.objects.filter(sample_product__sample_request_id__in=sample_ids)
        .exclude(image_file_id=None)
        .values_list("image_file_id", flat=True)
    )
    file_ids.update(
        SampleWorkflowRecord.objects.filter(sample_request_id__in=sample_ids)
        .exclude(business_file_id=None)
        .values_list("business_file_id", flat=True)
    )
    return file_ids


def _safe_media_path(storage_key):
    media_root = Path(settings.MEDIA_ROOT).resolve()
    target = (media_root / storage_key).resolve()
    try:
        target.relative_to(media_root)
    except ValueError as exc:
        raise SampleWorkbookImportError(f"文件路径超出媒体目录，拒绝删除：{storage_key}") from exc
    return target


def _workflow_records_for_sample(sample_request, data, actor_employee):
    created_at = data["requested_at"]
    approval_at = data["approval_at"] or created_at
    updated_at = data["updated_at"] or approval_at
    records = [
        SampleWorkflowRecord(
            sample_request=sample_request,
            stage=SampleWorkflowStage.REQUEST_CREATED,
            note="历史表格导入：创建样品需求。",
            operated_by_employee=actor_employee,
            operated_at=created_at,
        ),
        SampleWorkflowRecord(
            sample_request=sample_request,
            stage=SampleWorkflowStage.REQUIREMENTS_COMPLETED,
            note="历史表格导入：样品需求已完善。",
            operated_by_employee=actor_employee,
            operated_at=created_at,
        ),
        SampleWorkflowRecord(
            sample_request=sample_request,
            stage=SampleWorkflowStage.FINANCE_CONFIRMED,
            note=f"历史表格导入：财务确认{data['approval_text'] or '同意'}。",
            operated_by_employee=actor_employee,
            operated_at=approval_at,
        ),
        SampleWorkflowRecord(
            sample_request=sample_request,
            stage=SampleWorkflowStage.PMC_APPROVED,
            note="历史表格导入：PMC 审核通过。",
            operated_by_employee=actor_employee,
            operated_at=approval_at,
        ),
    ]
    if data["status"] in {SampleRequestStatus.CUTTING_PENDING, SampleRequestStatus.FINISHED}:
        metalwork_at = data["metalwork_at"] or updated_at
        missing = "、".join(data["metalwork_files"])
        records.append(
            SampleWorkflowRecord(
                sample_request=sample_request,
                stage=SampleWorkflowStage.METALWORK_PARAMETERS,
                note=("历史表格导入：金工参数已上传。" + (f" 原附件缺失：{missing}" if missing else "")),
                operated_by_employee=actor_employee,
                operated_at=metalwork_at,
            )
        )
    if data["status"] == SampleRequestStatus.FINISHED:
        cutting_at = data["cutting_at"] or updated_at
        missing = "、".join(data["cutting_files"])
        records.extend(
            [
                SampleWorkflowRecord(
                    sample_request=sample_request,
                    stage=SampleWorkflowStage.CUTTING_PARAMETERS,
                    note=("历史表格导入：裁剪参数已上传。" + (f" 原附件缺失：{missing}" if missing else "")),
                    operated_by_employee=actor_employee,
                    operated_at=cutting_at,
                ),
                SampleWorkflowRecord(
                    sample_request=sample_request,
                    stage=SampleWorkflowStage.WORKFLOW_FINISHED,
                    note="历史表格导入：原样品流程已结束。",
                    operated_by_employee=actor_employee,
                    operated_at=updated_at,
                ),
            ]
        )
    return records


@transaction.atomic
def execute_sample_workbook_import(plan, *, expected_existing):
    locked_sample_ids = list(
        SampleRequest.objects.select_for_update().order_by("id").values_list("id", flat=True)
    )
    if len(locked_sample_ids) != expected_existing:
        raise SampleWorkbookImportError(
            f"当前样品数量为 {len(locked_sample_ids)}，与确认的 {expected_existing} 条不一致，已停止替换。"
        )

    import_employee = Employee.objects.select_for_update().filter(employee_code=IMPORT_EMPLOYEE_CODE).first()
    if import_employee is None:
        raise SampleWorkbookImportError(f"缺少导入服务员工：{IMPORT_EMPLOYEE_CODE}")

    file_ids = _sample_business_files(locked_sample_ids)
    files = list(BusinessFile.objects.filter(id__in=file_ids))
    task_ids = set(
        SampleTaskLink.objects.filter(sample_request_id__in=locked_sample_ids).values_list(
            "business_task_id", flat=True
        )
    )
    deleted = Counter()
    if task_ids:
        deleted.update(_raw_delete_tree(BusinessTask, task_ids))
    deleted.update(_raw_delete_tree(SampleRequest, locked_sample_ids))

    storage_keys = []
    preserved_file_ids = []
    for business_file in files:
        try:
            storage_key = business_file.storage_key
            business_file.delete()
            storage_keys.append(storage_key)
        except ProtectedError:
            preserved_file_ids.append(str(business_file.pk))

    placeholder_customer, customer_created = CustomerAccount.objects.get_or_create(
        account_code=PLACEHOLDER_CUSTOMER_CODE,
        defaults={
            "account_name": "历史样品客户",
            "lifecycle_stage": CustomerLifecycleStage.LOST,
            "source_system": PLACEHOLDER_CUSTOMER_SOURCE,
            "external_id": PLACEHOLDER_CUSTOMER_EXTERNAL_ID,
            "notes": "仅用于无法关联原客户的历史样品表格导入。",
            "recorded_by_employee": import_employee,
            "recorded_at": timezone.now(),
        },
    )
    if (
        placeholder_customer.source_system != PLACEHOLDER_CUSTOMER_SOURCE
        or placeholder_customer.external_id != PLACEHOLDER_CUSTOMER_EXTERNAL_ID
    ):
        raise SampleWorkbookImportError(f"客户编码 {PLACEHOLDER_CUSTOMER_CODE} 已被其他数据占用。")

    categories = {item.category_name: item for item in ProductCategory.objects.all()}
    created_categories = []
    for category_name in plan.categories_to_create:
        category, created = ProductCategory.objects.get_or_create(
            category_name=category_name,
            defaults={
                "item_type": ProductItemType.FINISHED,
                "status": ReferenceStatus.ACTIVE,
                "description": "由历史样品表格导入补充。",
            },
        )
        if category.item_type != ProductItemType.FINISHED:
            raise SampleWorkbookImportError(f"产品类别 {category_name} 不是成品类别。")
        categories[category_name] = category
        if created:
            created_categories.append(category_name)

    created_samples = []
    created_products = []
    created_items = []
    created_sampling_details = []
    created_workflow_records = []
    created_fees = []

    for data in plan.samples:
        sample_request = SampleRequest(
            customer_account=placeholder_customer,
            owner_employee=import_employee,
            request_number=data["request_number"],
            manual_request_number=data["manual_request_number"],
            requested_by_employee=import_employee,
            created_by_employee=import_employee,
            shipping_method=data["shipping_method"],
            order_date=data["order_date"],
            sample_due_date=data["sample_due_date"],
            delivery_batch=data["delivery_batch"],
            is_repeat_order=data["is_repeat_order"],
            submitted_on=data["submitted_on"],
            creation_idempotency_key=f"sample-xlsx-{plan.source_hash[:24]}-{data['source_row']}",
            status=data["status"],
            sample_type=data["sample_type"],
            purpose=data["purpose"],
            requested_at=data["requested_at"],
            requested_due_at=data["requested_due_at"],
            notes=data["notes"],
        )
        created_samples.append(sample_request)
    SampleRequest.objects.bulk_create(created_samples)

    for sample_request, data in zip(created_samples, plan.samples, strict=True):
        product = SampleProduct(
            sample_request=sample_request,
            line_number=1,
            category=categories[data["category_name"]],
            category_name_snapshot=data["category_name"],
            product_name=data["product_name"],
            specification=data["specification"],
            length_cm=data["length_cm"],
            armrest_width_cm=data["armrest_width_cm"],
            seat_width_cm=data["seat_width_cm"],
            back_height_cm=data["back_height_cm"],
            seat_height_cm=data["seat_height_cm"],
        )
        created_products.append(product)
        for line_number, item in enumerate(data["items"], start=1):
            created_items.append(
                SampleItem(
                    id=item["id"],
                    sample_request=sample_request,
                    sample_product=product,
                    line_number=line_number,
                    product_description_snapshot=item["product_name"],
                    specification_snapshot=item["specification"],
                    requested_quantity=Decimal(item["sample_quantity"]),
                    unit_code="PCS",
                    product_color_id=item["product_color_id"],
                    color_card_snapshot=item["color_card"],
                    color_name_snapshot=item["color_name"],
                    color_code_snapshot=item["color_code"],
                    tolerance=item["tolerance"],
                    sample_quantity=item["sample_quantity"],
                    cutting_reported_quantity=item["cutting_reported_quantity"],
                    metalwork_reported_quantity=item["metalwork_reported_quantity"],
                    cutting_completed=item["cutting_completed"],
                    metalwork_completed=item["metalwork_completed"],
                    notes=item["notes"],
                )
            )
        for line_number, detail in enumerate(data["sampling_details"], start=1):
            created_sampling_details.append(
                SampleProductSamplingDetail(
                    sample_product=product,
                    line_number=line_number,
                    detail_type=detail["detail_type"],
                    category_label=detail["category_label"],
                    is_selected=detail["is_selected"],
                    description=detail["description"],
                )
            )
        created_workflow_records.extend(
            _workflow_records_for_sample(sample_request, data, import_employee)
        )
        if data["fee_amount"] is not None:
            created_fees.append(
                SampleFeeRecord(
                    sample_request=sample_request,
                    fee_decision=SampleFeeDecision.PAID,
                    amount=data["fee_amount"],
                    currency_code="CNY",
                    requested_decision=SampleFeeDecision.PAID,
                    sample_fee_amount=data["fee_amount"],
                    shipping_fee_amount=0,
                    discount_amount=0,
                    payable_amount=data["fee_amount"],
                    approval_status=FeeApprovalStatus.NOT_REQUIRED,
                    verification_status=FeeVerificationStatus.MANUALLY_CONFIRMED,
                    decision_reason="历史表格导入：人民币打样费已人工确认。",
                    decided_by_employee=import_employee,
                    decided_at=data["approval_at"] or data["requested_at"],
                )
            )

    SampleProduct.objects.bulk_create(created_products)
    SampleItem.objects.bulk_create(created_items)
    SampleProductSamplingDetail.objects.bulk_create(created_sampling_details)
    SampleWorkflowRecord.objects.bulk_create(created_workflow_records)
    SampleFeeRecord.objects.bulk_create(created_fees)

    cleanup_results = []

    def cleanup_storage():
        for storage_key in storage_keys:
            target = _safe_media_path(storage_key)
            try:
                if target.exists():
                    target.unlink()
                cleanup_results.append({"storage_key": storage_key, "status": "deleted"})
            except OSError as exc:
                cleanup_results.append(
                    {"storage_key": storage_key, "status": "failed", "error": str(exc)}
                )

    transaction.on_commit(cleanup_storage)
    return {
        "deleted": dict(sorted(deleted.items())),
        "deleted_sample_count": len(locked_sample_ids),
        "deleted_file_record_count": len(storage_keys),
        "preserved_shared_file_ids": preserved_file_ids,
        "storage_cleanup": cleanup_results,
        "placeholder_customer_id": str(placeholder_customer.pk),
        "placeholder_customer_created": customer_created,
        "created_categories": created_categories,
        "sample_count": len(created_samples),
        "product_count": len(created_products),
        "item_count": len(created_items),
        "sampling_detail_count": len(created_sampling_details),
        "workflow_record_count": len(created_workflow_records),
        "fee_record_count": len(created_fees),
    }


def json_safe(value):
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (Decimal, uuid.UUID, Path)):
        return str(value)
    return value


def build_sample_import_report(plan, *, mode, result=None):
    return json_safe(
        {
            "mode": mode,
            "source_file": str(plan.source_path),
            "summary": sample_import_plan_summary(plan),
            "result": result or {},
            "corrections": plan.corrections,
            "missing_attachments": plan.missing_attachments,
            "source_audit": plan.source_audit,
        }
    )
