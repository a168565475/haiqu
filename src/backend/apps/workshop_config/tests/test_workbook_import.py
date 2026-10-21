import io
import json
import tempfile
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from ..models import Craft, Operation, PlanCategory, Workshop, WorkshopEmployee
from ..workbook_import import (
    OPERATION_CORRECTION_HEADERS,
    analyze_craft_workbook,
    analyze_employee_workbook,
    analyze_operation_correction_workbook,
    analyze_workbooks,
)


class WorkshopWorkbookImportTests(TestCase):
    def _workbook(self, path, sheet_name, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()

    def _paths(self, directory, *, duplicate_operation=False):
        employee = directory / "employees.xlsx"
        craft = directory / "crafts.xlsx"
        operation = directory / "operations.xlsx"
        self._workbook(employee, "车间员工信息", ["车间名称", "岗位", "姓名", "状态"], [["装配车间", "组长", "员工甲", "在职"]])
        self._workbook(craft, "工艺信息", ["类型", "车间", "工艺", "唯一校验"], [["厂内", "装配车间", "装配", ""]])
        rows = [["GX001", "厂内", "装配车间", "测试品", "装配", "打包", "", ""]]
        if duplicate_operation:
            rows.append(["GX001", "厂内", "装配车间", "测试品", "装配", "打包", "", ""])
        self._workbook(operation, "工序信息", ["工序编号", "类型", "车间", "品名", "工艺", "工序", "品名-工艺-工序", "提交校验"], rows)
        return employee, craft, operation

    def _operation_correction_workbook(self, path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工序待修正"
        for _ in range(7):
            sheet.append([])
        sheet.append(OPERATION_CORRECTION_HEADERS)
        sheet.append(
            ["", "厂内", "装配车间", "测试品", "装配", "打包", "", "", 2, "缺工序编号", "", "", "", "GX000001", "测试品", "不适用", "不适用", "待复核"]
        )
        sheet.append(
            ["", "厂内", "装配车间", "测试品", "装配", "打包", "", "", 3, "重复业务组合", "", "", "业务组1", "GX000002", "测试品", "删除", "不适用", "待确认重复记录"]
        )
        sheet.append(
            ["", "厂内", "装配车间", "", "装配", "检查", "", "", 4, "品名为空", "", "", "", "GX000003", "", "不适用", "不适用", "待补品名"]
        )
        sheet.append(
            ["", "厂内", "装配车间", "-", "装配", "检查", "", "", 5, "品名为“-”", "", "", "", "GX000004", "-", "不适用", "删除", "待确认占位品名"]
        )
        workbook.save(path)
        workbook.close()

    def _metal_operation_correction_workbook(self, path, *, operation_number="GX-METAL-001", operation_name="弯管"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工序待修正"
        for _ in range(7):
            sheet.append([])
        sheet.append(OPERATION_CORRECTION_HEADERS)
        sheet.append([
            "", "厂内", "金工车间", "-", "管材加工", operation_name, "", "", 1,
            "品名为“-”", "", "", "", operation_number, "-", "不适用", "删除", "待确认占位品名",
        ])
        workbook.save(path)
        workbook.close()

    def test_valid_small_workbooks_pass_read_only_analysis(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = self._paths(Path(temp_dir))
            plan = analyze_workbooks(
                employee_path=paths[0],
                craft_path=paths[1],
                operation_path=paths[2],
                expected_employees=1,
                expected_crafts=1,
                expected_operations=1,
            )
        self.assertTrue(plan.report["can_apply"])
        self.assertEqual(plan.report["counts"]["operations_with_number"], 1)

    def test_employee_only_dry_run_accepts_fixed_workshop(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = self._paths(Path(temp_dir))
            plan = analyze_employee_workbook(employee_path=paths[0], expected_employees=1)
            output = io.StringIO()
            call_command(
                "import_workshop_employees",
                employees=str(paths[0]),
                expect_employees=1,
                stdout=output,
            )
        self.assertTrue(plan.report["can_apply"])
        self.assertEqual(plan.report["workshops"], ["装配车间"])
        self.assertIn("员工 dry-run 完成：未写入数据库", output.getvalue())
        self.assertEqual(WorkshopEmployee.objects.count(), 0)

    def test_employee_only_apply_cannot_create_unlinked_identity(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            paths = self._paths(directory)
            supervisor_file = directory / "supervisors.json"
            additional_file = directory / "additional.json"
            supervisor_file.write_text(
                json.dumps(
                    [
                        {"车间": "装配车间", "岗位": "组长", "姓名": "员工甲"},
                        {"车间": "金工车间", "岗位": "车间主管", "姓名": "员工乙"},
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            additional_file.write_text(
                json.dumps(
                    [
                        {
                            "车间": "金工车间",
                            "岗位": "车间主管",
                            "姓名": "员工乙",
                            "状态": "在职",
                            "是否为主管": True,
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(CommandError, "统一员工主档"):
                call_command(
                    "import_workshop_employees",
                    employees=str(paths[0]),
                    expect_employees=1,
                    apply=True,
                    supervisor_file=str(supervisor_file),
                    additional_employees_file=str(additional_file),
                    expect_additional_employees=1,
                    expect_existing_employees=0,
                    stdout=io.StringIO(),
                )
        self.assertEqual(WorkshopEmployee.objects.count(), 0)

    def test_combined_apply_cannot_create_unlinked_identity_or_partial_crafts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = self._paths(Path(temp_dir))
            with self.assertRaisesRegex(CommandError, "统一员工主档"):
                call_command(
                    "import_workshop_workbooks", employees=str(paths[0]), crafts=str(paths[1]),
                    operations=str(paths[2]), expect_employees=1, expect_crafts=1, expect_operations=1,
                    apply=True, stdout=io.StringIO(),
                )
        self.assertEqual(WorkshopEmployee.objects.count(), 0)
        self.assertEqual(Craft.objects.count(), 0)
        self.assertEqual(Operation.objects.count(), 0)

    def test_craft_only_dry_run_excludes_exact_confirmed_row(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = self._paths(Path(temp_dir))
            self._workbook(
                paths[1],
                "工艺信息",
                ["类型", "车间", "工艺", "唯一校验"],
                [
                    ["厂内", "车缝车间", "裁件", ""],
                    ["厂内", "车缝车间", "裁剪", ""],
                    ["厂外", "", "裁剪", ""],
                ],
            )
            plan = analyze_craft_workbook(
                craft_path=paths[1],
                expected_crafts=3,
                excluded_craft_keys={("厂内", "车缝车间", "裁件")},
                expected_import_crafts=2,
            )
        self.assertTrue(plan.report["can_apply"])
        self.assertEqual(plan.report["counts"]["excluded_crafts"], 1)
        self.assertEqual(
            {(row["类型"], row["车间"], row["工艺"]) for row in plan.crafts},
            {("厂内", "车缝车间", "裁剪"), ("厂外", "", "裁剪")},
        )

    def test_craft_only_apply_writes_only_remaining_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            craft_path = directory / "crafts.xlsx"
            self._workbook(
                craft_path,
                "工艺信息",
                ["类型", "车间", "工艺", "唯一校验"],
                [
                    ["厂内", "车缝车间", "裁件", ""],
                    ["厂内", "车缝车间", "裁剪", ""],
                ],
            )
            call_command(
                "import_workshop_crafts",
                crafts=str(craft_path),
                expect_crafts=2,
                expect_import_crafts=1,
                exclude_craft=["厂内|车缝车间|裁件"],
                apply=True,
                expect_existing_crafts=0,
                stdout=io.StringIO(),
            )
        self.assertEqual(Craft.objects.count(), 1)
        self.assertTrue(
            Craft.objects.filter(
                craft_type="INTERNAL",
                workshop__name="车缝车间",
                name="裁剪",
                status="ACTIVE",
            ).exists()
        )

    def test_operation_correction_dry_run_excludes_problem_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            operation_path = Path(temp_dir) / "operation-correction.xlsx"
            self._operation_correction_workbook(operation_path)
            plan = analyze_operation_correction_workbook(
                operation_path=operation_path,
                expected_source_operations=4,
                expected_import_operations=1,
                expected_duplicate_business_exclusions=1,
                expected_missing_product_exclusions=1,
                expected_placeholder_product_exclusions=1,
            )
        self.assertTrue(plan.report["can_apply"])
        self.assertEqual(plan.operations[0]["最终工序编号"], "GX000001")
        self.assertEqual(len(plan.excluded_operations), 3)
        self.assertEqual(Operation.objects.count(), 0)

    def test_operation_correction_apply_writes_only_clean_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            operation_path = Path(temp_dir) / "operation-correction.xlsx"
            self._operation_correction_workbook(operation_path)
            workshop = Workshop.objects.get(name="装配车间")
            Craft.objects.create(craft_type="INTERNAL", workshop=workshop, name="装配")
            output = io.StringIO()
            call_command(
                "import_workshop_operations",
                operations=str(operation_path),
                expect_source_operations=4,
                expect_import_operations=1,
                expect_duplicate_business_exclusions=1,
                expect_missing_product_exclusions=1,
                expect_placeholder_product_exclusions=1,
                expect_existing_crafts=1,
                expect_existing_operations=0,
                apply=True,
                stdout=output,
            )
        self.assertIn("工序正式导入完成：1 条，排除 3 条问题行", output.getvalue())
        self.assertEqual(Operation.objects.count(), 1)
        self.assertTrue(
            Operation.objects.filter(
                operation_number="GX000001",
                product_name="测试品",
                craft__name="装配",
                name="打包",
                status="ACTIVE",
            ).exists()
        )

    def test_metal_operation_import_accepts_placeholder_product_without_unit(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            operation_path = Path(temp_dir) / "metal-operation-correction.xlsx"
            self._metal_operation_correction_workbook(operation_path)
            workshop, _ = Workshop.objects.get_or_create(name="金工车间")
            craft = Craft.objects.create(
                craft_type="INTERNAL", workshop=workshop, name="管材加工",
                plan_category=PlanCategory.METAL,
            )
            output = io.StringIO()

            call_command(
                "import_metal_operations",
                operations=str(operation_path),
                expect_source_operations=1,
                expect_import_operations=0,
                expect_duplicate_business_exclusions=0,
                expect_missing_product_exclusions=0,
                expect_placeholder_product_exclusions=1,
                expect_metal_operations=1,
                expect_existing_operations=0,
                expect_existing_metal_operations=0,
                apply=True,
                stdout=output,
            )

            repeat_output = io.StringIO()
            call_command(
                "import_metal_operations",
                operations=str(operation_path),
                expect_source_operations=1,
                expect_import_operations=0,
                expect_duplicate_business_exclusions=0,
                expect_missing_product_exclusions=0,
                expect_placeholder_product_exclusions=1,
                expect_metal_operations=1,
                expect_existing_operations=0,
                expect_existing_metal_operations=0,
                apply=True,
                stdout=repeat_output,
            )

        self.assertIn("金工工序正式导入完成：1 条", output.getvalue())
        operation = Operation.objects.get(operation_number="GX-METAL-001")
        self.assertEqual(operation.craft, craft)
        self.assertEqual(operation.product_name, "")
        self.assertIsNone(operation.reporting_unit_id)

        self.assertIn("已存在且内容一致：1 条，未重复写入", repeat_output.getvalue())
        self.assertEqual(Operation.objects.count(), 1)

    def test_metal_operation_import_blocks_existing_business_key(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            operation_path = Path(temp_dir) / "metal-operation-correction.xlsx"
            self._metal_operation_correction_workbook(operation_path)
            workshop, _ = Workshop.objects.get_or_create(name="金工车间")
            craft = Craft.objects.create(
                craft_type="INTERNAL", workshop=workshop, name="管材加工",
                plan_category=PlanCategory.METAL,
            )
            Operation.objects.create(
                operation_number="EXISTING-001",
                product_name="",
                craft=craft,
                name="弯管",
            )
            output = io.StringIO()

            call_command(
                "import_metal_operations",
                operations=str(operation_path),
                expect_source_operations=1,
                expect_import_operations=0,
                expect_duplicate_business_exclusions=0,
                expect_missing_product_exclusions=0,
                expect_placeholder_product_exclusions=1,
                expect_metal_operations=1,
                expect_existing_operations=1,
                expect_existing_metal_operations=1,
                stdout=output,
            )

        self.assertIn("METAL_OPERATION_BUSINESS_COLLISION", output.getvalue())
        self.assertEqual(Operation.objects.count(), 1)

    def test_dry_run_reports_duplicates_and_writes_nothing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            paths = self._paths(Path(temp_dir), duplicate_operation=True)
            output = io.StringIO()
            call_command(
                "import_workshop_workbooks",
                employees=str(paths[0]),
                crafts=str(paths[1]),
                operations=str(paths[2]),
                expect_employees=1,
                expect_crafts=1,
                expect_operations=2,
                stdout=output,
            )
        self.assertIn("DUPLICATE_OPERATION_NUMBER", output.getvalue())
        self.assertIn("dry-run 完成：未写入数据库", output.getvalue())
        self.assertEqual(Workshop.objects.count(), 4)
        self.assertEqual(WorkshopEmployee.objects.count(), 0)
        self.assertEqual(Craft.objects.count(), 0)
        self.assertEqual(Operation.objects.count(), 0)
