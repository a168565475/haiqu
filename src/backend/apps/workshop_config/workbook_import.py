from collections import Counter
from dataclasses import dataclass
import json
from pathlib import Path

from openpyxl import load_workbook

from .constants import FIXED_WORKSHOP_NAMES


EMPLOYEE_HEADERS = ["车间名称", "岗位", "姓名", "状态"]
CRAFT_HEADERS = ["类型", "车间", "工艺", "唯一校验"]
OPERATION_HEADERS = ["工序编号", "类型", "车间", "品名", "工艺", "工序", "品名-工艺-工序", "提交校验"]
OPERATION_CORRECTION_HEADERS = [
    "工序编号",
    "类型",
    "车间",
    "品名",
    "工艺",
    "工序",
    "品名-工艺-工序",
    "提交校验",
    "源行号",
    "问题类型",
    "处理要求",
    "重复编号组",
    "重复业务组",
    "最终工序编号",
    "最终品名",
    "重复业务处理",
    "“-”品名确认",
    "修正状态",
]


@dataclass(frozen=True)
class WorkbookPlan:
    employees: list[dict]
    crafts: list[dict]
    operations: list[dict]
    report: dict


@dataclass(frozen=True)
class EmployeeWorkbookPlan:
    employees: list[dict]
    report: dict


@dataclass(frozen=True)
class CraftWorkbookPlan:
    crafts: list[dict]
    excluded_crafts: list[dict]
    report: dict


@dataclass(frozen=True)
class OperationCorrectionPlan:
    operations: list[dict]
    excluded_operations: list[dict]
    report: dict


def _text(value):
    return str(value or "").strip()


def _read_sheet(path, expected_sheet, expected_headers):
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise ValueError(f"工作簿不存在：{workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if expected_sheet not in workbook.sheetnames:
            raise ValueError(f"{workbook_path.name} 缺少工作表：{expected_sheet}")
        sheet = workbook[expected_sheet]
        headers = [_text(cell.value) for cell in sheet[1]]
        if headers[: len(expected_headers)] != expected_headers:
            raise ValueError(
                f"{workbook_path.name} 表头不正确，预期 {expected_headers}，实际 {headers[:len(expected_headers)]}"
            )
        rows = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = [_text(value) for value in row[: len(expected_headers)]]
            values.extend([""] * (len(expected_headers) - len(values)))
            if not any(values):
                continue
            rows.append({"row": row_number, **dict(zip(expected_headers, values, strict=True))})
        return rows
    finally:
        workbook.close()


def _read_sheet_at_header(path, expected_sheet, expected_headers, header_row):
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise ValueError(f"工作簿不存在：{workbook_path}")
    # 修正表由表格工具生成，可能不含可供只读模式使用的尺寸元数据。
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        if expected_sheet not in workbook.sheetnames:
            raise ValueError(f"{workbook_path.name} 缺少工作表：{expected_sheet}")
        sheet = workbook[expected_sheet]
        headers = [_text(sheet.cell(header_row, column).value) for column in range(1, len(expected_headers) + 1)]
        if headers != expected_headers:
            raise ValueError(f"{workbook_path.name} 第 {header_row} 行表头不正确，预期 {expected_headers}，实际 {headers}")
        rows = []
        for row_number in range(header_row + 1, sheet.max_row + 1):
            values = [
                _text(sheet.cell(row_number, column).value)
                for column in range(1, len(expected_headers) + 1)
            ]
            if not any(values):
                continue
            rows.append({"row": row_number, **dict(zip(expected_headers, values, strict=True))})
        return rows
    finally:
        workbook.close()


def _issue(code, message, count, sample_rows=None):
    issue = {"code": code, "message": message, "count": count}
    if sample_rows:
        issue["sample_rows"] = list(sample_rows)[:20]
    return issue


def _read_and_validate_employees(employee_path, expected_employees):
    employees = _read_sheet(employee_path, "车间员工信息", EMPLOYEE_HEADERS)
    for row in employees:
        row["车间"] = row.pop("车间名称")

    blockers = []
    if len(employees) != expected_employees:
        blockers.append(
            _issue("UNEXPECTED_ROW_COUNT", "车间员工行数与预期不一致。", abs(len(employees) - expected_employees))
        )

    valid_employee_statuses = {"在职", "离职"}
    invalid_employee_rows = [row["row"] for row in employees if row["状态"] not in valid_employee_statuses]
    missing_employee_fields = [
        row["row"] for row in employees if not row["车间"] or not row["岗位"] or not row["姓名"]
    ]
    employee_keys = [(row["车间"], row["岗位"], row["姓名"]) for row in employees]
    employee_key_counts = Counter(employee_keys)
    duplicate_employee_rows = [
        row["row"]
        for row in employees
        if employee_key_counts[(row["车间"], row["岗位"], row["姓名"])] > 1
    ]
    unknown_workshop_rows = [row["row"] for row in employees if row["车间"] not in FIXED_WORKSHOP_NAMES]
    if invalid_employee_rows:
        blockers.append(
            _issue(
                "INVALID_EMPLOYEE_STATUS",
                "员工状态只能是“在职”或“离职”。",
                len(invalid_employee_rows),
                invalid_employee_rows,
            )
        )
    if missing_employee_fields:
        blockers.append(
            _issue(
                "MISSING_EMPLOYEE_FIELD",
                "员工的车间、岗位和姓名不能为空。",
                len(missing_employee_fields),
                missing_employee_fields,
            )
        )
    if duplicate_employee_rows:
        blockers.append(
            _issue(
                "DUPLICATE_EMPLOYEE",
                "存在重复的“车间＋岗位＋姓名”。",
                len(duplicate_employee_rows),
                duplicate_employee_rows,
            )
        )
    if unknown_workshop_rows:
        blockers.append(
            _issue(
                "UNKNOWN_EMPLOYEE_WORKSHOP",
                f"员工车间必须是固定值：{'、'.join(FIXED_WORKSHOP_NAMES)}。",
                len(unknown_workshop_rows),
                unknown_workshop_rows,
            )
        )
    return employees, blockers


def analyze_employee_workbook(*, employee_path, expected_employees=81):
    employees, blockers = _read_and_validate_employees(employee_path, expected_employees)
    workshop_names = sorted({row["车间"] for row in employees if row["车间"]})
    report = {
        "mode": "dry-run",
        "can_apply": not blockers,
        "counts": {
            "employees": len(employees),
            "active_employees": sum(row["状态"] == "在职" for row in employees),
            "left_employees": sum(row["状态"] == "离职" for row in employees),
            "workshops": len(workshop_names),
        },
        "workshops": workshop_names,
        "blockers": blockers,
        "warnings": [],
    }
    return EmployeeWorkbookPlan(employees=employees, report=report)


def analyze_craft_workbook(
    *,
    craft_path,
    expected_crafts=11,
    excluded_craft_keys=None,
    expected_import_crafts=None,
):
    crafts = _read_sheet(craft_path, "工艺信息", CRAFT_HEADERS)
    blockers = []
    excluded_craft_keys = set(excluded_craft_keys or [])

    if len(crafts) != expected_crafts:
        blockers.append(
            _issue("UNEXPECTED_ROW_COUNT", "工艺行数与预期不一致。", abs(len(crafts) - expected_crafts))
        )

    invalid_craft_rows = [
        row["row"]
        for row in crafts
        if row["类型"] not in {"厂内", "厂外"} or not row["工艺"]
    ]
    invalid_craft_workshops = [
        row["row"]
        for row in crafts
        if (row["类型"] == "厂内" and row["车间"] not in FIXED_WORKSHOP_NAMES)
        or (row["类型"] == "厂外" and row["车间"])
    ]
    craft_business_keys = [(row["类型"], row["工艺"]) for row in crafts]
    craft_business_key_counts = Counter(craft_business_keys)
    duplicate_craft_rows = [
        row["row"]
        for row in crafts
        if craft_business_key_counts[(row["类型"], row["工艺"])] > 1
    ]
    if invalid_craft_rows:
        blockers.append(
            _issue(
                "INVALID_CRAFT",
                "工艺类型或工艺名称不正确。",
                len(invalid_craft_rows),
                invalid_craft_rows,
            )
        )
    if invalid_craft_workshops:
        blockers.append(
            _issue(
                "INVALID_CRAFT_WORKSHOP",
                "厂内工艺须关联固定车间，厂外工艺须留空车间。",
                len(invalid_craft_workshops),
                invalid_craft_workshops,
            )
        )
    if duplicate_craft_rows:
        blockers.append(
            _issue(
                "DUPLICATE_CRAFT",
                "存在重复的“类型＋工艺名称”。",
                len(duplicate_craft_rows),
                duplicate_craft_rows,
            )
        )

    rows_by_exact_key = {
        (row["类型"], row["车间"], row["工艺"]): row
        for row in crafts
    }
    missing_exclusions = sorted(excluded_craft_keys - set(rows_by_exact_key))
    if missing_exclusions:
        blockers.append(
            _issue(
                "EXCLUDED_CRAFT_NOT_FOUND",
                "要求排除的工艺未在工作簿中精确匹配。",
                len(missing_exclusions),
            )
        )

    excluded_crafts = [
        row
        for row in crafts
        if (row["类型"], row["车间"], row["工艺"]) in excluded_craft_keys
    ]
    import_crafts = [
        row
        for row in crafts
        if (row["类型"], row["车间"], row["工艺"]) not in excluded_craft_keys
    ]
    if expected_import_crafts is not None and len(import_crafts) != expected_import_crafts:
        blockers.append(
            _issue(
                "UNEXPECTED_IMPORT_ROW_COUNT",
                "排除规则处理后的工艺行数与预期不一致。",
                abs(len(import_crafts) - expected_import_crafts),
            )
        )

    report = {
        "mode": "dry-run",
        "can_apply": not blockers,
        "counts": {
            "source_crafts": len(crafts),
            "excluded_crafts": len(excluded_crafts),
            "crafts_to_import": len(import_crafts),
            "internal_crafts": sum(row["类型"] == "厂内" for row in import_crafts),
            "external_crafts": sum(row["类型"] == "厂外" for row in import_crafts),
        },
        "excluded": [
            {
                "row": row["row"],
                "类型": row["类型"],
                "车间": row["车间"],
                "工艺": row["工艺"],
            }
            for row in excluded_crafts
        ],
        "blockers": blockers,
        "warnings": [],
        "ignored_columns": ["唯一校验"],
    }
    return CraftWorkbookPlan(
        crafts=import_crafts,
        excluded_crafts=excluded_crafts,
        report=report,
    )


def analyze_operation_correction_workbook(
    *,
    operation_path,
    expected_source_operations=947,
    expected_import_operations=776,
    expected_duplicate_business_exclusions=117,
    expected_missing_product_exclusions=21,
    expected_placeholder_product_exclusions=33,
):
    operations = _read_sheet_at_header(
        operation_path,
        "工序待修正",
        OPERATION_CORRECTION_HEADERS,
        header_row=8,
    )
    blockers = []
    status_counts = Counter(row["修正状态"] for row in operations)
    allowed_statuses = {"待复核", "待确认重复记录", "待补品名", "待确认占位品名"}

    if len(operations) != expected_source_operations:
        blockers.append(
            _issue(
                "UNEXPECTED_SOURCE_ROW_COUNT",
                "工序修正表的数据行数与预期不一致。",
                abs(len(operations) - expected_source_operations),
            )
        )

    unexpected_status_rows = [row["row"] for row in operations if row["修正状态"] not in allowed_statuses]
    if unexpected_status_rows:
        blockers.append(
            _issue(
                "UNEXPECTED_CORRECTION_STATUS",
                "工序修正状态为空或不在允许范围内。",
                len(unexpected_status_rows),
                unexpected_status_rows,
            )
        )

    expected_status_counts = {
        "待复核": expected_import_operations,
        "待确认重复记录": expected_duplicate_business_exclusions,
        "待补品名": expected_missing_product_exclusions,
        "待确认占位品名": expected_placeholder_product_exclusions,
    }
    for status, expected_count in expected_status_counts.items():
        actual_count = status_counts[status]
        if actual_count != expected_count:
            blockers.append(
                _issue(
                    "UNEXPECTED_STATUS_COUNT",
                    f"状态“{status}”的行数与预期不一致，预期 {expected_count}，实际 {actual_count}。",
                    abs(actual_count - expected_count),
                )
            )

    import_operations = [row for row in operations if row["修正状态"] == "待复核"]
    excluded_operations = [row for row in operations if row["修正状态"] != "待复核"]

    required_field_rows = [
        row["row"]
        for row in import_operations
        if not row["最终工序编号"]
        or not row["最终品名"]
        or not row["类型"]
        or not row["工艺"]
        or not row["工序"]
        or not row["源行号"]
    ]
    if required_field_rows:
        blockers.append(
            _issue(
                "MISSING_IMPORT_FIELD",
                "待导入行的最终工序编号、最终品名、类型、工艺、工序和源行号均不能为空。",
                len(required_field_rows),
                required_field_rows,
            )
        )

    invalid_scope_rows = [
        row["row"]
        for row in import_operations
        if row["类型"] not in {"厂内", "厂外"}
        or (row["类型"] == "厂内" and row["车间"] not in FIXED_WORKSHOP_NAMES)
        or (row["类型"] == "厂外" and bool(row["车间"]))
    ]
    if invalid_scope_rows:
        blockers.append(
            _issue(
                "INVALID_OPERATION_SCOPE",
                "待导入工序的类型或车间范围不正确。",
                len(invalid_scope_rows),
                invalid_scope_rows,
            )
        )

    oversized_rows = [
        row["row"]
        for row in import_operations
        if len(row["最终工序编号"]) > 50
        or len(row["最终品名"]) > 200
        or len(row["工序"]) > 150
    ]
    if oversized_rows:
        blockers.append(
            _issue(
                "FIELD_TOO_LONG",
                "待导入工序存在超出数据库字段长度的内容。",
                len(oversized_rows),
                oversized_rows,
            )
        )

    number_counts = Counter(row["最终工序编号"].upper() for row in import_operations if row["最终工序编号"])
    duplicate_number_rows = [
        row["row"]
        for row in import_operations
        if row["最终工序编号"] and number_counts[row["最终工序编号"].upper()] > 1
    ]
    if duplicate_number_rows:
        blockers.append(
            _issue(
                "DUPLICATE_FINAL_OPERATION_NUMBER",
                "待导入的最终工序编号存在重复。",
                len(duplicate_number_rows),
                duplicate_number_rows,
            )
        )

    business_keys = [
        (row["最终品名"], row["类型"], row["工艺"], row["工序"])
        for row in import_operations
    ]
    business_key_counts = Counter(business_keys)
    duplicate_business_rows = [
        row["row"]
        for row in import_operations
        if business_key_counts[(row["最终品名"], row["类型"], row["工艺"], row["工序"])] > 1
    ]
    if duplicate_business_rows:
        blockers.append(
            _issue(
                "DUPLICATE_IMPORT_BUSINESS_KEY",
                "删除问题行后仍存在重复的“品名＋类型＋工艺＋工序”。",
                len(duplicate_business_rows),
                duplicate_business_rows,
            )
        )

    source_rows = [row["源行号"] for row in operations]
    invalid_source_rows = [row["row"] for row in operations if not row["源行号"].isdigit()]
    duplicate_source_rows = [
        row["row"]
        for row in operations
        if row["源行号"] and source_rows.count(row["源行号"]) > 1
    ]
    if invalid_source_rows:
        blockers.append(
            _issue(
                "INVALID_SOURCE_ROW",
                "源行号必须为整数。",
                len(invalid_source_rows),
                invalid_source_rows,
            )
        )
    if duplicate_source_rows:
        blockers.append(
            _issue(
                "DUPLICATE_SOURCE_ROW",
                "源行号存在重复。",
                len(duplicate_source_rows),
                duplicate_source_rows,
            )
        )

    report = {
        "mode": "dry-run",
        "can_apply": not blockers,
        "counts": {
            "source_operations": len(operations),
            "operations_to_import": len(import_operations),
            "excluded_operations": len(excluded_operations),
            "excluded_duplicate_business": status_counts["待确认重复记录"],
            "excluded_missing_product": status_counts["待补品名"],
            "excluded_placeholder_product": status_counts["待确认占位品名"],
        },
        "blockers": blockers,
        "warnings": [
            _issue(
                "ROWS_EXCLUDED_BY_USER_DECISION",
                "按已确认方案，不导入重复业务、空品名和占位品名问题行。",
                len(excluded_operations),
            )
        ],
        "ignored_columns": [
            "工序编号",
            "品名",
            "品名-工艺-工序",
            "提交校验",
            "问题类型",
            "处理要求",
            "重复编号组",
            "重复业务组",
            "重复业务处理",
            "“-”品名确认",
        ],
    }
    return OperationCorrectionPlan(
        operations=import_operations,
        excluded_operations=excluded_operations,
        report=report,
    )


def read_supervisor_keys(path, employee_rows):
    if not path:
        raise ValueError("正式导入必须提供主管名单，不能根据岗位推断主管。")
    try:
        with open(path, "r", encoding="utf-8") as file:
            payload = json.load(file)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"主管名单无法读取：{exc}") from exc
    if not isinstance(payload, list):
        raise ValueError("主管名单必须是 JSON 数组。")

    employee_keys = {(row["车间"], row["岗位"], row["姓名"]) for row in employee_rows}
    supervisor_keys = set()
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"主管名单第 {index} 项不是对象。")
        key = tuple(str(item.get(field, "")).strip() for field in ("车间", "岗位", "姓名"))
        if not all(key) or key not in employee_keys:
            raise ValueError(f"主管名单第 {index} 项无法匹配员工工作簿。")
        if key in supervisor_keys:
            raise ValueError(f"主管名单第 {index} 项重复。")
        supervisor_keys.add(key)
    return supervisor_keys


def read_additional_employees(path, existing_employee_rows):
    if not path:
        return []
    try:
        with open(path, "r", encoding="utf-8") as file:
            payload = json.load(file)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"新增员工名单无法读取：{exc}") from exc
    if not isinstance(payload, list):
        raise ValueError("新增员工名单必须是 JSON 数组。")

    existing_keys = {
        (row["车间"], row["岗位"], row["姓名"])
        for row in existing_employee_rows
    }
    additional_keys = set()
    additional_rows = []
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"新增员工名单第 {index} 项不是对象。")
        row = {
            "row": f"新增-{index}",
            "车间": str(item.get("车间", "")).strip(),
            "岗位": str(item.get("岗位", "")).strip(),
            "姓名": str(item.get("姓名", "")).strip(),
            "状态": str(item.get("状态", "")).strip(),
            "是否为主管": item.get("是否为主管"),
        }
        key = (row["车间"], row["岗位"], row["姓名"])
        if not all(key):
            raise ValueError(f"新增员工名单第 {index} 项的车间、岗位和姓名不能为空。")
        if row["车间"] not in FIXED_WORKSHOP_NAMES:
            raise ValueError(f"新增员工名单第 {index} 项不是固定车间。")
        if row["状态"] not in {"在职", "离职"}:
            raise ValueError(f"新增员工名单第 {index} 项状态只能是“在职”或“离职”。")
        if not isinstance(row["是否为主管"], bool):
            raise ValueError(f"新增员工名单第 {index} 项必须明确“是否为主管”为 true 或 false。")
        if key in existing_keys or key in additional_keys:
            raise ValueError(f"新增员工名单第 {index} 项与员工工作簿或新增名单重复。")
        additional_keys.add(key)
        additional_rows.append(row)
    return additional_rows


def analyze_workbooks(
    *,
    employee_path,
    craft_path,
    operation_path,
    expected_employees=81,
    expected_crafts=11,
    expected_operations=947,
):
    employees, employee_blockers = _read_and_validate_employees(employee_path, expected_employees)
    crafts = _read_sheet(craft_path, "工艺信息", CRAFT_HEADERS)
    operations = _read_sheet(operation_path, "工序信息", OPERATION_HEADERS)
    blockers = []
    warnings = []

    for label, rows, expected in [
        ("工艺", crafts, expected_crafts),
        ("工序", operations, expected_operations),
    ]:
        if len(rows) != expected:
            blockers.append(_issue("UNEXPECTED_ROW_COUNT", f"{label}行数与预期不一致。", abs(len(rows) - expected)))

    blockers.extend(employee_blockers)

    employee_workshop_names = {row["车间"] for row in employees if row["车间"]}
    craft_workshop_names = {row["车间"] for row in crafts if row["类型"] == "厂内" and row["车间"]}
    workshop_names = employee_workshop_names | craft_workshop_names
    valid_craft_types = {"厂内", "厂外"}
    invalid_craft_rows = []
    invalid_craft_workshops = []
    for row in crafts:
        if row["类型"] not in valid_craft_types or not row["工艺"]:
            invalid_craft_rows.append(row["row"])
        if row["类型"] == "厂内" and row["车间"] not in FIXED_WORKSHOP_NAMES:
            invalid_craft_workshops.append(row["row"])
        if row["类型"] == "厂外" and row["车间"]:
            invalid_craft_workshops.append(row["row"])
    craft_keys = [(row["类型"], row["工艺"]) for row in crafts]
    craft_key_counts = Counter(craft_keys)
    duplicate_craft_rows = [row["row"] for row in crafts if craft_key_counts[(row["类型"], row["工艺"])] > 1]
    if invalid_craft_rows:
        blockers.append(_issue("INVALID_CRAFT", "工艺类型或工艺名称不正确。", len(invalid_craft_rows), invalid_craft_rows))
    if invalid_craft_workshops:
        blockers.append(_issue("INVALID_CRAFT_WORKSHOP", "厂内工艺须关联现有车间，厂外工艺须留空车间。", len(invalid_craft_workshops), invalid_craft_workshops))
    if duplicate_craft_rows:
        blockers.append(_issue("DUPLICATE_CRAFT", "存在重复的“类型＋工艺名称”。", len(duplicate_craft_rows), duplicate_craft_rows))

    known_crafts = {(row["类型"], row["工艺"]): row for row in crafts}
    missing_number_rows = [row["row"] for row in operations if not row["工序编号"]]
    number_counts = Counter(row["工序编号"] for row in operations if row["工序编号"])
    duplicate_number_rows = [
        row["row"] for row in operations if row["工序编号"] and number_counts[row["工序编号"]] > 1
    ]
    missing_product_rows = [row["row"] for row in operations if not row["品名"]]
    operation_keys = [(row["品名"], row["类型"], row["工艺"], row["工序"]) for row in operations]
    operation_key_counts = Counter(operation_keys)
    duplicate_business_rows = [
        row["row"]
        for row in operations
        if operation_key_counts[(row["品名"], row["类型"], row["工艺"], row["工序"])] > 1
    ]
    invalid_operation_craft_rows = []
    for row in operations:
        craft = known_crafts.get((row["类型"], row["工艺"]))
        if craft is None:
            invalid_operation_craft_rows.append(row["row"])
            continue
        expected_workshop = craft["车间"] if row["类型"] == "厂内" else ""
        if row["车间"] != expected_workshop:
            invalid_operation_craft_rows.append(row["row"])
    if missing_number_rows:
        blockers.append(_issue("MISSING_OPERATION_NUMBER", "工序编号不能为空，且须由管理员补充唯一编号。", len(missing_number_rows), missing_number_rows))
    if duplicate_number_rows:
        blockers.append(_issue("DUPLICATE_OPERATION_NUMBER", "工序编号存在重复。", len(duplicate_number_rows), duplicate_number_rows))
    if missing_product_rows:
        blockers.append(_issue("MISSING_PRODUCT_NAME", "工序品名不能为空。", len(missing_product_rows), missing_product_rows))
    if duplicate_business_rows:
        duplicate_groups = sum(1 for count in operation_key_counts.values() if count > 1)
        duplicate_excess = sum(count - 1 for count in operation_key_counts.values() if count > 1)
        issue = _issue(
            "DUPLICATE_OPERATION_BUSINESS_KEY",
            f"存在 {duplicate_groups} 组重复的“品名＋工艺＋工序”。",
            duplicate_excess,
            duplicate_business_rows,
        )
        issue["group_count"] = duplicate_groups
        blockers.append(issue)
    if invalid_operation_craft_rows:
        blockers.append(_issue("INVALID_OPERATION_CRAFT", "工序的类型、车间与工艺资料不一致。", len(invalid_operation_craft_rows), invalid_operation_craft_rows))

    craft_names = {row["工艺"] for row in crafts}
    if {"裁件", "裁剪"}.issubset(craft_names):
        blockers.append(_issue("CRAFT_NAME_CONFLICT", "同时存在“裁件”和“裁剪”，须确认是否为同一工艺。", 2))

    dash_product_rows = [row["row"] for row in operations if row["品名"] == "-"]
    if dash_product_rows:
        warnings.append(_issue("PLACEHOLDER_PRODUCT_NAME", "品名使用“-”，请确认是否应保留为正式文本。", len(dash_product_rows), dash_product_rows))

    report = {
        "mode": "dry-run",
        "can_apply": not blockers,
        "counts": {
            "employees": len(employees),
            "active_employees": sum(row["状态"] == "在职" for row in employees),
            "left_employees": sum(row["状态"] == "离职" for row in employees),
            "workshops": len(workshop_names),
            "crafts": len(crafts),
            "internal_crafts": sum(row["类型"] == "厂内" for row in crafts),
            "external_crafts": sum(row["类型"] == "厂外" for row in crafts),
            "operations": len(operations),
            "operations_with_number": sum(bool(row["工序编号"]) for row in operations),
        },
        "blockers": blockers,
        "warnings": warnings,
        "ignored_columns": ["品名-工艺-工序", "提交校验", "唯一校验"],
    }
    return WorkbookPlan(employees=employees, crafts=crafts, operations=operations, report=report)
